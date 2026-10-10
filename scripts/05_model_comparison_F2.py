# -*- coding: utf-8 -*-
r"""
A²/O F2 四指标六模型比较、正式验证与代表配置输出
=========================================================

本脚本在原有 F3 六模型筛选代码和 F0-F3 消融代码基础上重构，单次运行同时完成
TP、TN、NH4+-N 和 COD 四个目标在 F2 特征体系下的模型比较。

核心原则：
1. 四个目标均比较 Elastic Net、SVR、GPR、Random Forest、XGBoost 和 MLP；
2. 保留原项目固定随机种子、两阶段参数筛选、Bootstrap .632+、时间顺序验证、
   扩展窗口验证、lag1/roll3 朴素动态基线和工况分组误差分析；
3. 每类算法先独立完成超参数搜索，最终在本特征组内只保留 1 个“算法代表配置”，
   因而每个目标输出 6 个代表配置；这里不提前宣布全局最终模型；
4. 最终机器学习模型应在 F0-F3 四版全部完成后，针对每个目标在 24 个
   “特征组×算法”代表配置中统一比较；
5. 进水缺失值处理严格位于每一次训练/验证划分内部。训练集唯一原始样本可进行
   短缺口线性插值（最大 7 个原始样本跨度），剩余缺失用训练中位数兜底；
   OOB、时间验证和扩展窗口验证样本仅使用训练中位数，不使用未来观测双向插值；
6. Q 不进入任何模型特征。

默认输入：
${PROJECT_ROOT}/outputs/preprocessing/dynamic/动态可用数据与插补比较结果.xlsx

默认输出：
${PROJECT_ROOT}/outputs/model_comparison/F2

输出文件：
- TP_F2_六模型比较结果.xlsx
- TN_F2_六模型比较结果.xlsx
- NH4_N_F2_六模型比较结果.xlsx
- COD_F2_六模型比较结果.xlsx
- F2_四指标模型比较说明.txt

正式默认参数保持原方案：
- 主随机种子 20260620
- 初筛 Bootstrap 100 次
- 正式 Bootstrap .632+ 1000 次
- 时间顺序验证：前 80% 训练、后 20% 验证
- 扩展窗口：初始训练 60%，验证窗口 10，步长 10
- 正式候选：每种算法 Bootstrap 综合排名前 3 ∪ Time_R2 前 3 ∪ Time_RMSE 前 3

依赖：
pip install pandas numpy openpyxl scikit-learn xgboost

如只想快速检查代码能否运行，可临时设置环境变量 MODEL_SMOKE_TEST=1。
正式分析请不要设置该变量。
"""

from __future__ import annotations

import hashlib
import json
import math
import os
import pickle
import re
import time
import traceback
import warnings
from copy import deepcopy
from datetime import datetime
from itertools import product
from pathlib import Path
from typing import Any, Callable, Dict, Iterable, List, Mapping, Optional, Sequence, Tuple

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from sklearn.ensemble import RandomForestRegressor
from sklearn.exceptions import ConvergenceWarning
from sklearn.gaussian_process import GaussianProcessRegressor
from sklearn.gaussian_process.kernels import ConstantKernel, Matern, RBF
from sklearn.linear_model import ElasticNet
from sklearn.metrics import mean_absolute_error, mean_squared_error, r2_score
from sklearn.neural_network import MLPRegressor
from sklearn.preprocessing import StandardScaler
from sklearn.svm import SVR

try:
    from xgboost import XGBRegressor
    XGBOOST_AVAILABLE = True
except Exception:
    XGBOOST_AVAILABLE = False


# =============================================================================
# 1. 路径、目标与固定参数
# =============================================================================

PROJECT_ROOT = Path(__file__).resolve().parents[1]

FEATURE_GROUP = "F2"
DEFAULT_INPUT_FILE = Path(
    str(PROJECT_ROOT / "outputs/preprocessing/dynamic/动态可用数据与插补比较结果.xlsx")
)
DEFAULT_OUTPUT_DIR = Path(str(PROJECT_ROOT / "outputs/model_comparison/F2"))
INPUT_FILE = Path(os.environ.get("MODEL_INPUT_FILE", str(DEFAULT_INPUT_FILE)))
OUTPUT_DIR = Path(os.environ.get("MODEL_OUTPUT_DIR", str(DEFAULT_OUTPUT_DIR)))

RANDOM_SEED = 20260620
BOOTSTRAP_SCREEN_N = int(os.environ.get("MODEL_BOOTSTRAP_SCREEN_N", "100"))
BOOTSTRAP_FINAL_N = int(os.environ.get("MODEL_BOOTSTRAP_FINAL_N", "1000"))
TIME_TRAIN_RATIO = 0.80
KEEP_TOP_BOOTSTRAP_PER_MODEL = 3
KEEP_TOP_TIME_R2_PER_MODEL = 3
KEEP_TOP_TIME_RMSE_PER_MODEL = 3
EXPANDING_INITIAL_TRAIN_RATIO = 0.60
EXPANDING_VALID_WINDOW_SIZE = int(os.environ.get("MODEL_EXPANDING_WINDOW", "10"))
EXPANDING_STEP_SIZE = int(os.environ.get("MODEL_EXPANDING_STEP", "10"))

# 旧版阈值保留为诊断字段，不再作为本阶段算法代表配置的强制淘汰门槛。
LEGACY_TIME_R2_REFERENCE = 0.20
LEGACY_EXPANDING_R2_REFERENCE = 0.00

MAX_INTERPOLATION_GAP = 7
N_JOBS_TREE = int(os.environ.get("MODEL_N_JOBS", "-1"))
EPS = 1e-6
SEED_STABILITY_REPEATS = int(os.environ.get("MODEL_SEED_STABILITY_REPEATS", "5"))

ENABLE_CHECKPOINT = os.environ.get("MODEL_ENABLE_CHECKPOINT", "1").strip() != "0"
REMOVE_CHECKPOINT_AFTER_SUCCESS = os.environ.get("MODEL_REMOVE_CHECKPOINT", "1").strip() != "0"
SKIP_COMPLETED_TARGETS = os.environ.get("MODEL_SKIP_COMPLETED", "0").strip() == "1"

SMOKE_TEST = os.environ.get("MODEL_SMOKE_TEST", "0").strip() == "1"
CANDIDATE_LIMIT_PER_MODEL = int(os.environ.get("MODEL_CANDIDATES_PER_MODEL", "0"))
if SMOKE_TEST:
    BOOTSTRAP_SCREEN_N = 2
    BOOTSTRAP_FINAL_N = 2
    EXPANDING_VALID_WINDOW_SIZE = 250
    EXPANDING_STEP_SIZE = 250
    SEED_STABILITY_REPEATS = 0
    if CANDIDATE_LIMIT_PER_MODEL <= 0:
        CANDIDATE_LIMIT_PER_MODEL = 1

TEXTUAL_MISSING_VALUES = {
    "", " ", "NA", "N/A", "na", "n/a", "NaN", "nan", "None", "none",
    "NULL", "null", "-", "—", "–", "－", "/", "\\", "未测", "缺失",
}

RAW_FEATURES = [
    "Aerobic_DO",
    "Anaerobic_Feed_Fraction",
    "Inf.TP",
    "Inf.COD",
    "Inf.NH4+-N",
    "Inf.TN",
    "HRT",
    "Anaerobic_Volume_Fraction",
    "Anoxic_Volume_Fraction",
]
IMPUTE_COLUMNS = ["Inf.TP", "Inf.COD", "Inf.NH4+-N", "Inf.TN"]
OPERATION_COLUMNS = [
    "Aerobic_DO",
    "Anaerobic_Feed_Fraction",
    "HRT",
    "Anaerobic_Volume_Fraction",
    "Anoxic_Volume_Fraction",
]
MECHANISM_FEATURES = [
    "Anaerobic_HRT",
    "Anoxic_HRT",
    "Aerobic_HRT",
    "C_N_ratio",
    "N_P_ratio",
    "COD_volumetric_load",
    "TN_volumetric_load",
    "NH4_N_volumetric_load",
    "TP_volumetric_load",
]

TARGET_CONFIGS: Dict[str, Dict[str, Any]] = {
    "TP": {
        "sheet": "TP动态可用表",
        "target_col": "eff.TP",
        "history": ["eff.TP_lag1", "eff.TP_lag2", "eff.TP_roll3"],
        "file_prefix": "TP",
    },
    "TN": {
        "sheet": "TN动态可用表",
        "target_col": "eff.TN",
        "history": ["eff.TN_lag1", "eff.TN_lag2", "eff.TN_roll3"],
        "file_prefix": "TN",
    },
    "NH4+-N": {
        "sheet": "氨氮动态可用表",
        "target_col": "eff.NH4+-N",
        "history": ["eff.NH4+-N_lag1", "eff.NH4+-N_lag2", "eff.NH4+-N_roll3"],
        "file_prefix": "NH4_N",
    },
    "COD": {
        "sheet": "COD动态可用表",
        "target_col": "eff.COD",
        "history": ["eff.COD_lag1", "eff.COD_lag2", "eff.COD_roll3"],
        "file_prefix": "COD",
    },
}
TARGET_ORDER = ["TP", "TN", "NH4+-N", "COD"]
MODEL_ORDER = ["ElasticNet", "SVR", "GPR", "RandomForest", "XGBoost", "MLP"]
STOCHASTIC_MODELS = {"RandomForest", "XGBoost", "MLP"}


def feature_list_for_target(history_features: Sequence[str]) -> List[str]:
    if FEATURE_GROUP == "F0":
        features = list(RAW_FEATURES)
    elif FEATURE_GROUP == "F1":
        features = list(RAW_FEATURES) + list(MECHANISM_FEATURES)
    elif FEATURE_GROUP == "F2":
        features = list(RAW_FEATURES) + list(history_features)
    elif FEATURE_GROUP == "F3":
        features = list(RAW_FEATURES) + list(MECHANISM_FEATURES) + list(history_features)
    else:
        raise ValueError(f"未知特征组：{FEATURE_GROUP}")

    expected = {"F0": 9, "F1": 18, "F2": 12, "F3": 21}[FEATURE_GROUP]
    if len(features) != expected or len(set(features)) != expected:
        raise RuntimeError(f"{FEATURE_GROUP} 特征数量或重复检查失败：{features}")
    if "Q" in features:
        raise RuntimeError("Q 不允许进入模型特征。")
    return features


# =============================================================================
# 2. 通用工具、文件输出与检查点
# =============================================================================


def now_str() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def log(message: str) -> None:
    print(f"[{now_str()}] {message}", flush=True)


def print_section(title: str) -> None:
    print("\n" + "=" * 112)
    print(title)
    print("=" * 112)


def stable_seed_from_text(text: Any) -> int:
    digest = hashlib.md5(str(text).encode("utf-8")).hexdigest()
    return int(digest[:8], 16)


def clean_column_name(value: Any) -> str:
    text = str(value)
    text = text.replace("\ufeff", "").replace("\xa0", " ").replace("\u3000", " ")
    text = text.replace("\n", "").replace("\r", "").strip()
    return re.sub(r"\s+", "", text)


def convert_to_numeric_series(series: pd.Series) -> pd.Series:
    output = series.copy()
    if output.dtype == "object" or str(output.dtype).startswith("string"):
        output = (
            output.astype("string")
            .str.strip()
            .str.replace(",", "", regex=False)
            .str.replace("，", "", regex=False)
        )
        output = output.mask(output.isin(TEXTUAL_MISSING_VALUES), np.nan)
    return pd.to_numeric(output, errors="coerce")


def params_to_json(params: Mapping[str, Any]) -> str:
    def converter(obj: Any) -> Any:
        if isinstance(obj, tuple):
            return list(obj)
        if isinstance(obj, np.integer):
            return int(obj)
        if isinstance(obj, np.floating):
            return float(obj)
        return str(obj)
    return json.dumps(dict(params), ensure_ascii=False, default=converter, sort_keys=True)


def safe_float(value: Any) -> float:
    try:
        number = float(value)
    except Exception:
        return np.nan
    return number if np.isfinite(number) else np.nan


def sanitize_excel_value(value: Any) -> Any:
    if isinstance(value, (list, tuple, dict, set)):
        return json.dumps(value, ensure_ascii=False, default=str)
    return value


def sanitize_dataframe_for_excel(df: pd.DataFrame) -> pd.DataFrame:
    if df is None:
        return pd.DataFrame()
    output = df.copy()
    for col in output.columns:
        if output[col].dtype == "object":
            output[col] = output[col].map(sanitize_excel_value)
    return output


def format_excel_workbook(writer: pd.ExcelWriter) -> None:
    workbook = writer.book
    for sheet_name in workbook.sheetnames:
        worksheet = workbook[sheet_name]
        if worksheet.max_row >= 1:
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            for cell in worksheet[1]:
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        worksheet.sheet_view.showGridLines = False
        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="center", wrap_text=False)
                if isinstance(cell.value, datetime):
                    cell.number_format = "yyyy-mm-dd"
                elif isinstance(cell.value, float):
                    cell.number_format = "0.000000"
        for col_idx in range(1, worksheet.max_column + 1):
            letter = get_column_letter(col_idx)
            max_length = 0
            for cell in worksheet[letter][: min(worksheet.max_row, 300)]:
                value = "" if cell.value is None else str(cell.value)
                max_length = max(max_length, len(value))
            worksheet.column_dimensions[letter].width = min(max(max_length + 2, 10), 48)


def safe_write_excel(output_file: Path, tables: Mapping[str, pd.DataFrame]) -> Path:
    output_file.parent.mkdir(parents=True, exist_ok=True)

    def do_write(path: Path) -> None:
        with pd.ExcelWriter(path, engine="openpyxl", datetime_format="yyyy-mm-dd") as writer:
            for sheet_name, table in tables.items():
                safe_name = str(sheet_name)[:31]
                sanitize_dataframe_for_excel(table).to_excel(writer, sheet_name=safe_name, index=False)
            format_excel_workbook(writer)

    try:
        do_write(output_file)
        return output_file
    except PermissionError:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        new_file = output_file.with_name(f"{output_file.stem}_{timestamp}{output_file.suffix}")
        log(f"原 Excel 可能正在打开，改为输出：{new_file}")
        do_write(new_file)
        return new_file


def safe_write_text(output_file: Path, text: str) -> Path:
    output_file.parent.mkdir(parents=True, exist_ok=True)
    try:
        output_file.write_text(text, encoding="utf-8-sig")
        return output_file
    except PermissionError:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        new_file = output_file.with_name(f"{output_file.stem}_{timestamp}{output_file.suffix}")
        new_file.write_text(text, encoding="utf-8-sig")
        return new_file


def config_fingerprint(target_name: str, cfg: Mapping[str, Any], candidates: Sequence[Mapping[str, Any]]) -> str:
    payload = {
        "target": target_name,
        "feature_group": FEATURE_GROUP,
        "features": feature_list_for_target(cfg["history"]),
        "random_seed": RANDOM_SEED,
        "bootstrap_screen": BOOTSTRAP_SCREEN_N,
        "bootstrap_final": BOOTSTRAP_FINAL_N,
        "time_ratio": TIME_TRAIN_RATIO,
        "expanding_initial": EXPANDING_INITIAL_TRAIN_RATIO,
        "expanding_window": EXPANDING_VALID_WINDOW_SIZE,
        "expanding_step": EXPANDING_STEP_SIZE,
        "imputation_gap": MAX_INTERPOLATION_GAP,
        "candidate_json": [c["参数JSON"] for c in candidates],
        "candidate_models": [c["模型"] for c in candidates],
    }
    return hashlib.md5(json.dumps(payload, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest()


def checkpoint_path(target_name: str, stage_tag: str) -> Path:
    return OUTPUT_DIR / f"_checkpoint_{FEATURE_GROUP}_{target_name.replace('+', 'p')}_{stage_tag}.pkl"


def load_or_compute_checkpoint(
    path: Path,
    fingerprint: str,
    compute_func: Callable[[], Any],
) -> Any:
    if ENABLE_CHECKPOINT and path.exists():
        try:
            with path.open("rb") as handle:
                payload = pickle.load(handle)
            if isinstance(payload, dict) and payload.get("fingerprint") == fingerprint:
                log(f"读取检查点：{path.name}")
                return payload["value"]
            log(f"检查点配置已变化，重新计算：{path.name}")
        except Exception as exc:
            log(f"检查点读取失败，将重新计算：{path.name} | {exc}")
    value = compute_func()
    if ENABLE_CHECKPOINT:
        with path.open("wb") as handle:
            pickle.dump({"fingerprint": fingerprint, "value": value}, handle, protocol=pickle.HIGHEST_PROTOCOL)
    return value


def remove_target_checkpoints(target_name: str) -> None:
    if not REMOVE_CHECKPOINT_AFTER_SUCCESS:
        return
    pattern = f"_checkpoint_{FEATURE_GROUP}_{target_name.replace('+', 'p')}_*.pkl"
    for path in OUTPUT_DIR.glob(pattern):
        try:
            path.unlink()
        except Exception:
            pass


# =============================================================================
# 3. 数据读取、插补和特征构造
# =============================================================================


def read_target_table(target_name: str, cfg: Mapping[str, Any]) -> pd.DataFrame:
    if not INPUT_FILE.exists():
        raise FileNotFoundError(f"未找到输入文件：{INPUT_FILE}")
    excel = pd.ExcelFile(INPUT_FILE, engine="openpyxl")
    if cfg["sheet"] not in excel.sheet_names:
        raise ValueError(f"输入文件缺少工作表：{cfg['sheet']}。现有工作表：{excel.sheet_names}")

    df = pd.read_excel(
        INPUT_FILE,
        sheet_name=cfg["sheet"],
        engine="openpyxl",
        na_values=list(TEXTUAL_MISSING_VALUES),
        keep_default_na=True,
    )
    df.columns = [clean_column_name(c) for c in df.columns]

    required = ["样本序号", "Date"] + RAW_FEATURES + [cfg["target_col"]] + list(cfg["history"])
    missing = [col for col in required if col not in df.columns]
    if missing:
        raise ValueError("动态可用表缺少必要字段：\n" + "\n".join(missing))

    df["Date"] = pd.to_datetime(df["Date"], errors="coerce")
    numeric_cols = ["样本序号"] + RAW_FEATURES + [cfg["target_col"]] + list(cfg["history"])
    for col in numeric_cols:
        df[col] = convert_to_numeric_series(df[col])

    df = df.sort_values(["样本序号", "Date"], kind="mergesort").reset_index(drop=True)
    df["_row_id"] = np.arange(len(df), dtype=int)

    if df.empty:
        raise ValueError(f"{target_name} 动态可用表为空。")
    if df["样本序号"].isna().any() or df["Date"].isna().any():
        raise ValueError(f"{target_name} 数据中存在无效样本序号或日期。")
    if df["样本序号"].duplicated().any():
        dup = df.loc[df["样本序号"].duplicated(keep=False), "样本序号"].tolist()
        raise ValueError(f"{target_name} 样本序号重复：{dup[:20]}")
    if not df["样本序号"].is_monotonic_increasing:
        raise ValueError(f"{target_name} 数据未按样本序号递增。")

    must_complete = [cfg["target_col"]] + list(cfg["history"]) + OPERATION_COLUMNS
    missing_counts = {col: int(df[col].isna().sum()) for col in must_complete if df[col].isna().any()}
    if missing_counts:
        raise ValueError(
            f"{target_name} 的目标、历史状态或运行工况存在缺失，不应统计插补：\n"
            + json.dumps(missing_counts, ensure_ascii=False, indent=2)
        )
    if (df["HRT"] <= 0).any():
        bad = df.loc[df["HRT"] <= 0, "样本序号"].tolist()
        raise ValueError(f"{target_name} HRT 必须大于0，样本：{bad[:20]}")
    aerobic_fraction = 1.0 - df["Anaerobic_Volume_Fraction"] - df["Anoxic_Volume_Fraction"]
    if (aerobic_fraction <= 0).any():
        bad = df.loc[aerobic_fraction <= 0, "样本序号"].tolist()
        raise ValueError(f"{target_name} 好氧区容积分数必须大于0，样本：{bad[:20]}")
    return df


def interpolate_short_internal_gaps(
    series: pd.Series,
    sequence_coordinates: pd.Series,
    max_gap: int,
) -> pd.Series:
    """仅插补两侧均有训练观测、且按原始样本序号计算跨度不超过 max_gap 的内部缺口。"""
    output = convert_to_numeric_series(series).astype(float).copy()
    coordinates = convert_to_numeric_series(sequence_coordinates).astype(float).to_numpy()
    values = output.to_numpy(copy=True)
    n = len(values)
    position = 0
    while position < n:
        if np.isfinite(values[position]):
            position += 1
            continue
        start = position
        while position < n and not np.isfinite(values[position]):
            position += 1
        end = position - 1
        left = start - 1
        right = end + 1
        if (
            left >= 0 and right < n
            and np.isfinite(values[left]) and np.isfinite(values[right])
            and np.isfinite(coordinates[left]) and np.isfinite(coordinates[right])
        ):
            original_gap_span = coordinates[right] - coordinates[left] - 1.0
            coordinate_span = coordinates[right] - coordinates[left]
            if 0 <= original_gap_span <= max_gap and coordinate_span > 0:
                left_value = float(values[left])
                right_value = float(values[right])
                for index in range(start, end + 1):
                    fraction = (coordinates[index] - coordinates[left]) / coordinate_span
                    values[index] = left_value + fraction * (right_value - left_value)
    return pd.Series(values, index=output.index, name=output.name)


def fit_train_imputer_and_apply(
    train_raw: pd.DataFrame,
    valid_raw: pd.DataFrame,
) -> Tuple[pd.DataFrame, pd.DataFrame, Dict[str, Any]]:
    """
    在训练集唯一原始样本上拟合短缺口线性插值与中位数兜底；Bootstrap重复样本再映射回来。
    验证/OOB样本不做双向线性插值，只使用训练中位数补全其真实缺失。
    """
    train = train_raw.copy(deep=True)
    valid = valid_raw.copy(deep=True)
    unique_train = (
        train.drop_duplicates(subset=["_row_id"], keep="first")
        .sort_values(["_row_id"], kind="mergesort")
        .copy()
    )
    info: Dict[str, Any] = {
        "method": "training_only_short_internal_linear_interpolation_plus_training_median",
        "max_gap": MAX_INTERPOLATION_GAP,
        "train_medians": {},
        "training_missing_before": {},
        "training_linear_filled": {},
        "training_median_filled": {},
        "validation_missing_filled_by_training_median": {},
    }

    for col in IMPUTE_COLUMNS:
        original = convert_to_numeric_series(unique_train[col]).astype(float)
        median_value = original.median(skipna=True)
        if pd.isna(median_value):
            raise ValueError(f"训练数据变量 {col} 全部缺失，无法插补。")
        interpolated = interpolate_short_internal_gaps(
            original, unique_train["样本序号"], MAX_INTERPOLATION_GAP
        )
        linear_filled = int(original.isna().sum() - interpolated.isna().sum())
        median_filled = int(interpolated.isna().sum())
        unique_train[col] = interpolated.fillna(float(median_value))

        mapping = unique_train.set_index("_row_id")[col]
        train[col] = train["_row_id"].map(mapping)
        valid_series = convert_to_numeric_series(valid[col]).astype(float)
        valid_missing = int(valid_series.isna().sum())
        valid[col] = valid_series.fillna(float(median_value))

        info["train_medians"][col] = float(median_value)
        info["training_missing_before"][col] = int(original.isna().sum())
        info["training_linear_filled"][col] = linear_filled
        info["training_median_filled"][col] = median_filled
        info["validation_missing_filled_by_training_median"][col] = valid_missing

    remaining_train = {c: int(train[c].isna().sum()) for c in IMPUTE_COLUMNS if train[c].isna().any()}
    remaining_valid = {c: int(valid[c].isna().sum()) for c in IMPUTE_COLUMNS if valid[c].isna().any()}
    if remaining_train or remaining_valid:
        raise ValueError(
            "插补后仍存在缺失：" + json.dumps(
                {"train": remaining_train, "valid": remaining_valid}, ensure_ascii=False
            )
        )
    return train, valid, info


def construct_mechanism_features(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy(deep=True)
    operation_missing = {col: int(out[col].isna().sum()) for col in OPERATION_COLUMNS if out[col].isna().any()}
    if operation_missing:
        raise ValueError("运行工况存在缺失：" + json.dumps(operation_missing, ensure_ascii=False))
    aerobic_fraction = 1.0 - out["Anaerobic_Volume_Fraction"] - out["Anoxic_Volume_Fraction"]
    if (aerobic_fraction <= 0).any():
        bad = out.loc[aerobic_fraction <= 0, "样本序号"].tolist()
        raise ValueError(f"好氧区容积分数不大于0，样本序号：{bad[:20]}")
    out["Anaerobic_HRT"] = out["HRT"] * out["Anaerobic_Volume_Fraction"]
    out["Anoxic_HRT"] = out["HRT"] * out["Anoxic_Volume_Fraction"]
    out["Aerobic_HRT"] = out["HRT"] * aerobic_fraction
    out["C_N_ratio"] = out["Inf.COD"] / (out["Inf.TN"] + EPS)
    out["N_P_ratio"] = out["Inf.TN"] / (out["Inf.TP"] + EPS)
    out["COD_volumetric_load"] = out["Inf.COD"] / (out["HRT"] + EPS)
    out["TN_volumetric_load"] = out["Inf.TN"] / (out["HRT"] + EPS)
    out["NH4_N_volumetric_load"] = out["Inf.NH4+-N"] / (out["HRT"] + EPS)
    out["TP_volumetric_load"] = out["Inf.TP"] / (out["HRT"] + EPS)
    return out.replace([np.inf, -np.inf], np.nan)


def prepare_split(
    train_raw: pd.DataFrame,
    valid_raw: pd.DataFrame,
    cfg: Mapping[str, Any],
) -> Dict[str, Any]:
    features = feature_list_for_target(cfg["history"])
    train_imp, valid_imp, impute_info = fit_train_imputer_and_apply(train_raw, valid_raw)

    if FEATURE_GROUP in {"F1", "F3"}:
        train_feat = construct_mechanism_features(train_imp)
        valid_feat = construct_mechanism_features(valid_imp)
    else:
        train_feat = train_imp.copy()
        valid_feat = valid_imp.copy()

    missing_info: Dict[str, Dict[str, int]] = {}
    for col in features:
        if col not in train_feat.columns or col not in valid_feat.columns:
            missing_info[col] = {"train": -1, "valid": -1}
            continue
        train_missing = int(train_feat[col].isna().sum())
        valid_missing = int(valid_feat[col].isna().sum())
        if train_missing or valid_missing:
            missing_info[col] = {"train": train_missing, "valid": valid_missing}
    if missing_info:
        raise ValueError(
            f"{FEATURE_GROUP} 特征构造后仍有缺失或字段不存在："
            + json.dumps(missing_info, ensure_ascii=False)
        )

    X_train_raw = train_feat[features].astype(float)
    X_valid_raw = valid_feat[features].astype(float)
    scaler = StandardScaler()
    X_train = scaler.fit_transform(X_train_raw)
    X_valid = scaler.transform(X_valid_raw)

    return {
        "X_train": X_train,
        "X_valid": X_valid,
        "y_train": train_feat[cfg["target_col"]].astype(float).to_numpy(),
        "y_valid": valid_feat[cfg["target_col"]].astype(float).to_numpy(),
        "train_feature_df": train_feat,
        "valid_feature_df": valid_feat,
        "impute_info": impute_info,
        "scaler": scaler,
        "feature_names": features,
    }


def prepare_full_features_for_diagnostics(df: pd.DataFrame) -> pd.DataFrame:
    """仅用于工况分组阈值诊断；不用于模型训练或泛化评价。"""
    ordered = df.sort_values("_row_id", kind="mergesort").copy()
    for col in IMPUTE_COLUMNS:
        original = convert_to_numeric_series(ordered[col]).astype(float)
        median_value = original.median(skipna=True)
        if pd.isna(median_value):
            raise ValueError(f"全量诊断数据变量 {col} 全部缺失。")
        interpolated = interpolate_short_internal_gaps(
            original, ordered["样本序号"], MAX_INTERPOLATION_GAP
        )
        ordered[col] = interpolated.fillna(float(median_value))
    return construct_mechanism_features(ordered)


# =============================================================================
# 4. 六类模型与参数网格
# =============================================================================


def build_parameter_grid() -> List[Dict[str, Any]]:
    candidates: List[Dict[str, Any]] = []

    for alpha, l1_ratio in product(
        [0.0001, 0.0003, 0.001, 0.003, 0.01, 0.03, 0.1, 0.3, 1.0],
        [0.1, 0.3, 0.5, 0.7, 0.9],
    ):
        candidates.append({"模型": "ElasticNet", "参数": {"alpha": alpha, "l1_ratio": l1_ratio}})

    for C, epsilon, gamma in product(
        [0.3, 1.0, 3.0, 10.0, 30.0],
        [0.01, 0.05, 0.10],
        ["scale", 0.01, 0.03],
    ):
        candidates.append({
            "模型": "SVR",
            "参数": {"kernel": "rbf", "C": C, "epsilon": epsilon, "gamma": gamma},
        })

    for kernel_type, length_scale, alpha in product(
        ["RBF", "Matern"], [1.0, 2.0, 5.0], [0.001, 0.01]
    ):
        candidates.append({
            "模型": "GPR",
            "参数": {"kernel_type": kernel_type, "length_scale": length_scale, "alpha": alpha},
        })

    for n_estimators, max_depth, min_samples_leaf, min_samples_split, max_features in product(
        [300, 600], [None, 6, 12], [1, 3, 5], [2, 5], [0.7, 1.0]
    ):
        candidates.append({
            "模型": "RandomForest",
            "参数": {
                "n_estimators": n_estimators,
                "max_depth": max_depth,
                "min_samples_leaf": min_samples_leaf,
                "min_samples_split": min_samples_split,
                "max_features": max_features,
            },
        })

    for n_estimators, max_depth, learning_rate, subsample, colsample_bytree in product(
        [200, 500], [2, 3], [0.03, 0.05, 0.10], [0.8, 1.0], [0.8, 1.0]
    ):
        candidates.append({
            "模型": "XGBoost",
            "参数": {
                "n_estimators": n_estimators,
                "max_depth": max_depth,
                "learning_rate": learning_rate,
                "subsample": subsample,
                "colsample_bytree": colsample_bytree,
            },
        })

    for hidden_layer_sizes, activation, alpha, learning_rate_init in product(
        [(8,), (16,), (32,), (16, 8), (32, 16)],
        ["relu", "tanh"], [0.0001, 0.001, 0.01], [0.0003, 0.001]
    ):
        candidates.append({
            "模型": "MLP",
            "参数": {
                "hidden_layer_sizes": hidden_layer_sizes,
                "activation": activation,
                "alpha": alpha,
                "learning_rate_init": learning_rate_init,
                "batch_size": 32,
                "early_stopping": True,
                "max_iter": 1000,
                "n_iter_no_change": 30,
            },
        })

    model_counter: Dict[str, int] = {}
    output: List[Dict[str, Any]] = []
    for item in candidates:
        model_name = item["模型"]
        model_counter[model_name] = model_counter.get(model_name, 0) + 1
        item = deepcopy(item)
        item["候选参数ID"] = f"{model_name}_{model_counter[model_name]:03d}"
        item["参数序号"] = model_counter[model_name]
        item["参数JSON"] = params_to_json(item["参数"])
        output.append(item)

    if CANDIDATE_LIMIT_PER_MODEL > 0:
        limited: List[Dict[str, Any]] = []
        by_model: Dict[str, List[Dict[str, Any]]] = {m: [] for m in MODEL_ORDER}
        for item in output:
            by_model[item["模型"]].append(item)
        for model_name in MODEL_ORDER:
            limited.extend(by_model[model_name][:CANDIDATE_LIMIT_PER_MODEL])
        output = limited

    for global_index, item in enumerate(output, start=1):
        item["全局参数序号"] = global_index
    return output


def build_model(model_name: str, params: Mapping[str, Any], random_state: int):
    p = deepcopy(dict(params))
    if model_name == "ElasticNet":
        return ElasticNet(
            alpha=p["alpha"], l1_ratio=p["l1_ratio"], max_iter=50000,
            random_state=random_state,
        )
    if model_name == "SVR":
        return SVR(kernel=p["kernel"], C=p["C"], epsilon=p["epsilon"], gamma=p["gamma"])
    if model_name == "GPR":
        if p["kernel_type"] == "RBF":
            kernel = ConstantKernel(1.0, constant_value_bounds="fixed") * RBF(
                length_scale=p["length_scale"], length_scale_bounds="fixed"
            )
        else:
            kernel = ConstantKernel(1.0, constant_value_bounds="fixed") * Matern(
                length_scale=p["length_scale"], length_scale_bounds="fixed", nu=1.5
            )
        return GaussianProcessRegressor(
            kernel=kernel, alpha=p["alpha"], optimizer=None,
            normalize_y=True, random_state=random_state,
        )
    if model_name == "RandomForest":
        if SMOKE_TEST:
            p["n_estimators"] = min(int(p["n_estimators"]), 20)
        return RandomForestRegressor(
            n_estimators=p["n_estimators"], max_depth=p["max_depth"],
            min_samples_leaf=p["min_samples_leaf"], min_samples_split=p["min_samples_split"],
            max_features=p["max_features"], random_state=random_state, n_jobs=N_JOBS_TREE,
        )
    if model_name == "XGBoost":
        if SMOKE_TEST:
            p["n_estimators"] = min(int(p["n_estimators"]), 20)
        if not XGBOOST_AVAILABLE:
            raise ImportError("当前 Python 环境未安装 xgboost，请执行：pip install xgboost")
        return XGBRegressor(
            objective="reg:squarederror",
            n_estimators=p["n_estimators"], max_depth=p["max_depth"],
            learning_rate=p["learning_rate"], subsample=p["subsample"],
            colsample_bytree=p["colsample_bytree"], random_state=random_state,
            n_jobs=N_JOBS_TREE, reg_lambda=1.0, verbosity=0,
        )
    if model_name == "MLP":
        if SMOKE_TEST:
            p["max_iter"] = min(int(p["max_iter"]), 80)
            p["n_iter_no_change"] = min(int(p["n_iter_no_change"]), 8)
        return MLPRegressor(
            hidden_layer_sizes=p["hidden_layer_sizes"], activation=p["activation"],
            alpha=p["alpha"], learning_rate_init=p["learning_rate_init"],
            batch_size=p["batch_size"], early_stopping=p["early_stopping"],
            validation_fraction=0.15, max_iter=p["max_iter"],
            n_iter_no_change=p["n_iter_no_change"], random_state=random_state,
        )
    raise ValueError(f"未知模型：{model_name}")


def get_model_random_state(
    target_name: str,
    model_name: str,
    candidate_id: str,
    stage_name: str,
    repeat_id: int,
) -> int:
    # 特征组故意不进入种子文本，使 F0-F3 的同一目标/算法/参数/阶段使用相同随机种子。
    base = RANDOM_SEED + stable_seed_from_text(
        f"{target_name}-{model_name}-{candidate_id}-{stage_name}"
    )
    if model_name == "MLP":
        base += 1_000_003
    return int((base + repeat_id * 1009) % (2**31 - 1))


def get_stability_random_state(target_name: str, model_name: str, candidate_id: str, repeat_id: int) -> int:
    base = RANDOM_SEED + stable_seed_from_text(
        f"{target_name}-{model_name}-{candidate_id}-seed-stability"
    )
    return int((base + repeat_id * 7919) % (2**31 - 1))


def fit_model_safely(model: Any, X_train: np.ndarray, y_train: np.ndarray) -> str:
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        warnings.filterwarnings("always", category=ConvergenceWarning)
        model.fit(X_train, y_train)
    if not caught:
        return ""
    return " | ".join(str(item.message)[:300] for item in caught[:5])


# =============================================================================
# 5. 指标、Bootstrap、时间顺序与扩展窗口验证
# =============================================================================


def calc_metrics(y_true: Iterable[float], y_pred: Iterable[float]) -> Dict[str, float]:
    true = np.asarray(list(y_true), dtype=float)
    pred = np.asarray(list(y_pred), dtype=float)
    mask = np.isfinite(true) & np.isfinite(pred)
    true = true[mask]
    pred = pred[mask]
    n = len(true)
    if n == 0:
        return {"n": 0, "R2": np.nan, "MSE": np.nan, "RMSE": np.nan, "MAE": np.nan, "sMAPE": np.nan, "MAPE": np.nan}
    mse = float(mean_squared_error(true, pred))
    mae = float(mean_absolute_error(true, pred))
    rmse = float(math.sqrt(mse))
    smape = float(np.mean(2.0 * np.abs(pred - true) / (np.abs(true) + np.abs(pred) + EPS)) * 100.0)
    mape = float(np.mean(np.abs(pred - true) / (np.abs(true) + EPS)) * 100.0)
    r2 = float(r2_score(true, pred)) if n >= 2 and np.nanstd(true) > 0 else np.nan
    return {"n": int(n), "R2": r2, "MSE": mse, "RMSE": rmse, "MAE": mae, "sMAPE": smape, "MAPE": mape}


def calc_632plus_error(app_err: float, oob_err: float, noinfo_err: float) -> Tuple[float, float, float]:
    if not np.isfinite(app_err) or not np.isfinite(oob_err):
        return np.nan, np.nan, np.nan
    if not np.isfinite(noinfo_err) or noinfo_err <= app_err:
        weight = 0.632
        return float((1.0 - weight) * app_err + weight * oob_err), weight, np.nan
    denominator = noinfo_err - app_err
    relative_overfit = 0.0 if abs(denominator) < EPS else (oob_err - app_err) / denominator
    relative_overfit = float(np.clip(relative_overfit, 0.0, 1.0))
    weight = 0.632 / (1.0 - 0.368 * relative_overfit)
    result = (1.0 - weight) * app_err + weight * oob_err
    return float(result), float(weight), relative_overfit


def make_bootstrap_indices(n_samples: int, n_bootstrap: int, seed: int) -> List[Tuple[np.ndarray, np.ndarray]]:
    rng = np.random.default_rng(seed)
    all_indices = np.arange(n_samples)
    output: List[Tuple[np.ndarray, np.ndarray]] = []
    for _ in range(n_bootstrap):
        for _attempt in range(100):
            inbag = rng.choice(all_indices, size=n_samples, replace=True)
            oob = np.setdiff1d(all_indices, np.unique(inbag), assume_unique=False)
            if len(oob) >= 5:
                output.append((inbag, oob))
                break
        else:
            raise RuntimeError("无法生成至少含5个OOB样本的Bootstrap重复。")
    return output


def evaluate_bootstrap_stage(
    target_name: str,
    cfg: Mapping[str, Any],
    df: pd.DataFrame,
    candidates: Sequence[Mapping[str, Any]],
    bootstrap_indices: Sequence[Tuple[np.ndarray, np.ndarray]],
    stage_name: str,
) -> pd.DataFrame:
    rows: List[Dict[str, Any]] = []
    total_boot = len(bootstrap_indices)
    features = feature_list_for_target(cfg["history"])

    for boot_id, (inbag_idx, oob_idx) in enumerate(bootstrap_indices, start=1):
        if boot_id == 1 or boot_id % 10 == 0 or boot_id == total_boot:
            log(f"{target_name} | {FEATURE_GROUP} | {stage_name} | Bootstrap {boot_id}/{total_boot}")
        try:
            prepared = prepare_split(df.iloc[inbag_idx].copy(), df.iloc[oob_idx].copy(), cfg)
        except Exception as exc:
            for candidate in candidates:
                rows.append({
                    "阶段": stage_name, "目标指标": target_name, "特征组": FEATURE_GROUP,
                    "特征数": len(features), "模型": candidate["模型"],
                    "候选参数ID": candidate["候选参数ID"], "参数JSON": candidate["参数JSON"],
                    "Bootstrap编号": boot_id, "是否成功": "否",
                    "错误信息": "预处理失败：" + repr(exc), "警告信息": "",
                })
            continue

        X_train, X_oob = prepared["X_train"], prepared["X_valid"]
        y_train, y_oob = prepared["y_train"], prepared["y_valid"]
        noinfo = calc_metrics(y_oob, np.full_like(y_oob, float(np.mean(y_train)), dtype=float))

        for candidate in candidates:
            start = time.time()
            model_name = candidate["模型"]
            base = {
                "阶段": stage_name, "目标指标": target_name, "特征组": FEATURE_GROUP,
                "特征数": len(features), "模型": model_name,
                "候选参数ID": candidate["候选参数ID"], "参数JSON": candidate["参数JSON"],
                "Bootstrap编号": boot_id, "是否成功": "是", "错误信息": "", "警告信息": "",
            }
            try:
                random_state = get_model_random_state(
                    target_name, model_name, candidate["候选参数ID"], stage_name, boot_id
                )
                model = build_model(model_name, candidate["参数"], random_state)
                warning_text = fit_model_safely(model, X_train, y_train)
                pred_app = model.predict(X_train)
                pred_oob = model.predict(X_oob)
                app = calc_metrics(y_train, pred_app)
                oob = calc_metrics(y_oob, pred_oob)
                mse632, weight, relative_overfit = calc_632plus_error(app["MSE"], oob["MSE"], noinfo["MSE"])
                mae632, _, _ = calc_632plus_error(app["MAE"], oob["MAE"], noinfo["MAE"])
                smape632, _, _ = calc_632plus_error(app["sMAPE"], oob["sMAPE"], noinfo["sMAPE"])
                mape632, _, _ = calc_632plus_error(app["MAPE"], oob["MAPE"], noinfo["MAPE"])
                r2_632 = 1.0 - mse632 / noinfo["MSE"] if np.isfinite(mse632) and noinfo["MSE"] > 0 else np.nan
                rows.append({
                    **base,
                    "App_n": app["n"], "OOB_n": oob["n"],
                    "App_R2": app["R2"], "App_RMSE": app["RMSE"], "App_MAE": app["MAE"],
                    "OOB_R2": oob["R2"], "OOB_RMSE": oob["RMSE"], "OOB_MAE": oob["MAE"],
                    "NoInfo_MSE": noinfo["MSE"], "R2_632plus": r2_632,
                    "RMSE_632plus": math.sqrt(mse632) if np.isfinite(mse632) else np.nan,
                    "MAE_632plus": mae632, "sMAPE_632plus": smape632, "MAPE_632plus": mape632,
                    "632plus权重_MSE": weight, "632plus相对过拟合率_MSE": relative_overfit,
                    "训练耗时秒": time.time() - start, "警告信息": warning_text,
                })
            except Exception as exc:
                rows.append({
                    **base, "是否成功": "否",
                    "错误信息": repr(exc) + "\n" + traceback.format_exc(limit=3),
                    "训练耗时秒": time.time() - start,
                })
    return pd.DataFrame(rows)


def summarize_bootstrap(detail: pd.DataFrame) -> pd.DataFrame:
    success = detail.loc[detail["是否成功"] == "是"].copy()
    if success.empty:
        return pd.DataFrame()
    keys = ["阶段", "目标指标", "特征组", "特征数", "模型", "候选参数ID", "参数JSON"]
    rows: List[Dict[str, Any]] = []
    metrics = [
        "App_R2", "OOB_R2", "R2_632plus", "RMSE_632plus", "MAE_632plus",
        "sMAPE_632plus", "MAPE_632plus", "App_RMSE", "OOB_RMSE", "App_MAE", "OOB_MAE",
    ]
    for key_values, group in success.groupby(keys, dropna=False, sort=False):
        row = dict(zip(keys, key_values if isinstance(key_values, tuple) else (key_values,)))
        row["有效Bootstrap次数"] = int(group["Bootstrap编号"].nunique())
        row["平均训练耗时秒"] = float(group["训练耗时秒"].mean())
        for metric in metrics:
            values = pd.to_numeric(group[metric], errors="coerce").dropna()
            row[f"{metric}均值"] = values.mean() if len(values) else np.nan
            row[f"{metric}标准差"] = values.std(ddof=1) if len(values) > 1 else 0.0
            if metric in {"R2_632plus", "RMSE_632plus", "MAE_632plus", "sMAPE_632plus"}:
                row[f"{metric}最小值"] = values.min() if len(values) else np.nan
                row[f"{metric}2.5分位"] = values.quantile(0.025) if len(values) else np.nan
                row[f"{metric}中位数"] = values.median() if len(values) else np.nan
                row[f"{metric}97.5分位"] = values.quantile(0.975) if len(values) else np.nan
                row[f"{metric}最大值"] = values.max() if len(values) else np.nan
        row["App_OOB_R2差距"] = abs(safe_float(row.get("App_R2均值")) - safe_float(row.get("OOB_R2均值")))
        row["App_OOB_RMSE差距"] = abs(safe_float(row.get("OOB_RMSE均值")) - safe_float(row.get("App_RMSE均值")))
        rows.append(row)
    return pd.DataFrame(rows)


def build_prediction_frame(
    target_name: str,
    candidate: Mapping[str, Any],
    validation_type: str,
    valid_features: pd.DataFrame,
    y_true: np.ndarray,
    y_pred: np.ndarray,
    window_id: Optional[int] = None,
) -> pd.DataFrame:
    output = pd.DataFrame({
        "验证类型": validation_type, "目标指标": target_name, "特征组": FEATURE_GROUP,
        "模型": candidate["模型"], "候选参数ID": candidate["候选参数ID"],
        "样本序号": valid_features["样本序号"].to_numpy(),
        "Date": valid_features["Date"].to_numpy(),
        "源位置": valid_features["_row_id"].to_numpy(),
        "实测值": y_true, "预测值": y_pred,
        "残差": y_pred - y_true, "绝对误差": np.abs(y_pred - y_true),
    })
    if window_id is not None:
        output["窗口编号"] = int(window_id)
    return output


def evaluate_time_stage(
    target_name: str,
    cfg: Mapping[str, Any],
    df: pd.DataFrame,
    candidates: Sequence[Mapping[str, Any]],
    stage_name: str,
    return_predictions: bool,
) -> Tuple[pd.DataFrame, pd.DataFrame]:
    split_idx = int(math.floor(len(df) * TIME_TRAIN_RATIO))
    if split_idx < 10 or len(df) - split_idx < 5:
        raise ValueError("时间顺序验证训练集或验证集样本不足。")
    train_raw = df.iloc[:split_idx].copy()
    valid_raw = df.iloc[split_idx:].copy()
    prepared = prepare_split(train_raw, valid_raw, cfg)
    X_train, X_valid = prepared["X_train"], prepared["X_valid"]
    y_train, y_valid = prepared["y_train"], prepared["y_valid"]
    valid_feat = prepared["valid_feature_df"]

    rows: List[Dict[str, Any]] = []
    pred_parts: List[pd.DataFrame] = []
    for index, candidate in enumerate(candidates, start=1):
        model_name = candidate["模型"]
        start = time.time()
        base = {
            "阶段": stage_name, "目标指标": target_name, "特征组": FEATURE_GROUP,
            "特征数": X_train.shape[1], "模型": model_name,
            "候选参数ID": candidate["候选参数ID"], "参数JSON": candidate["参数JSON"],
            "时间训练样本数": len(train_raw), "时间验证样本数": len(valid_raw),
            "时间验证起始样本序号": valid_raw["样本序号"].min(),
            "时间验证结束样本序号": valid_raw["样本序号"].max(),
            "是否成功": "是", "错误信息": "", "警告信息": "",
        }
        try:
            random_state = get_model_random_state(target_name, model_name, candidate["候选参数ID"], stage_name, 0)
            model = build_model(model_name, candidate["参数"], random_state)
            warning_text = fit_model_safely(model, X_train, y_train)
            pred_train = model.predict(X_train)
            pred_valid = model.predict(X_valid)
            train_metrics = calc_metrics(y_train, pred_train)
            valid_metrics = calc_metrics(y_valid, pred_valid)
            rows.append({
                **base,
                "Train_R2": train_metrics["R2"], "Train_RMSE": train_metrics["RMSE"],
                "Train_MAE": train_metrics["MAE"], "Train_sMAPE": train_metrics["sMAPE"],
                "Time_R2": valid_metrics["R2"], "Time_RMSE": valid_metrics["RMSE"],
                "Time_MAE": valid_metrics["MAE"], "Time_sMAPE": valid_metrics["sMAPE"],
                "Time_MAPE": valid_metrics["MAPE"],
                "Train_Time_R2差距": abs(train_metrics["R2"] - valid_metrics["R2"])
                if np.isfinite(train_metrics["R2"]) and np.isfinite(valid_metrics["R2"]) else np.nan,
                "训练耗时秒": time.time() - start, "警告信息": warning_text,
            })
            if return_predictions:
                pred_parts.append(build_prediction_frame(
                    target_name, candidate, stage_name, valid_feat, y_valid, pred_valid
                ))
        except Exception as exc:
            rows.append({
                **base, "是否成功": "否",
                "错误信息": repr(exc) + "\n" + traceback.format_exc(limit=3),
                "训练耗时秒": time.time() - start,
            })
    return pd.DataFrame(rows), pd.concat(pred_parts, ignore_index=True) if pred_parts else pd.DataFrame()


def prepare_expanding_windows(df: pd.DataFrame, cfg: Mapping[str, Any]) -> List[Dict[str, Any]]:
    initial_train = int(math.floor(len(df) * EXPANDING_INITIAL_TRAIN_RATIO))
    windows: List[Dict[str, Any]] = []
    for window_id, start_idx in enumerate(range(initial_train, len(df), EXPANDING_STEP_SIZE), start=1):
        end_idx = min(start_idx + EXPANDING_VALID_WINDOW_SIZE, len(df))
        if end_idx - start_idx < 2:
            continue
        train_raw = df.iloc[:start_idx].copy()
        valid_raw = df.iloc[start_idx:end_idx].copy()
        try:
            prepared = prepare_split(train_raw, valid_raw, cfg)
            prepared.update({
                "窗口编号": window_id,
                "训练样本数": len(train_raw), "验证样本数": len(valid_raw),
                "窗口起始样本序号": valid_raw["样本序号"].min(),
                "窗口结束样本序号": valid_raw["样本序号"].max(),
                "窗口起始日期": valid_raw["Date"].min(), "窗口结束日期": valid_raw["Date"].max(),
            })
            windows.append(prepared)
        except Exception as exc:
            windows.append({
                "窗口编号": window_id, "错误": repr(exc),
                "窗口起始样本序号": valid_raw["样本序号"].min(),
                "窗口结束样本序号": valid_raw["样本序号"].max(),
            })
    return windows


def evaluate_expanding_stage(
    target_name: str,
    cfg: Mapping[str, Any],
    df: pd.DataFrame,
    candidates: Sequence[Mapping[str, Any]],
    return_predictions: bool,
) -> Tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    windows = prepare_expanding_windows(df, cfg)
    accum: Dict[str, Dict[str, List[float]]] = {
        c["候选参数ID"]: {"true": [], "pred": []} for c in candidates
    }
    window_rows: List[Dict[str, Any]] = []
    pred_parts: List[pd.DataFrame] = []
    candidate_map = {c["候选参数ID"]: c for c in candidates}

    for window in windows:
        window_id = int(window["窗口编号"])
        if "错误" in window:
            for candidate in candidates:
                window_rows.append({
                    "阶段": "扩展窗口验证", "目标指标": target_name, "特征组": FEATURE_GROUP,
                    "模型": candidate["模型"], "候选参数ID": candidate["候选参数ID"],
                    "窗口编号": window_id, "是否成功": "否", "错误信息": "预处理失败：" + window["错误"],
                    "警告信息": "",
                })
            continue
        for candidate in candidates:
            model_name = candidate["模型"]
            start = time.time()
            base = {
                "阶段": "扩展窗口验证", "目标指标": target_name, "特征组": FEATURE_GROUP,
                "模型": model_name, "候选参数ID": candidate["候选参数ID"],
                "参数JSON": candidate["参数JSON"], "窗口编号": window_id,
                "训练样本数": window["训练样本数"], "验证样本数": window["验证样本数"],
                "窗口起始样本序号": window["窗口起始样本序号"],
                "窗口结束样本序号": window["窗口结束样本序号"],
                "窗口起始日期": window["窗口起始日期"], "窗口结束日期": window["窗口结束日期"],
                "是否成功": "是", "错误信息": "", "警告信息": "",
            }
            try:
                random_state = get_model_random_state(
                    target_name, model_name, candidate["候选参数ID"], "扩展窗口验证", window_id
                )
                model = build_model(model_name, candidate["参数"], random_state)
                warning_text = fit_model_safely(model, window["X_train"], window["y_train"])
                pred = model.predict(window["X_valid"])
                metrics = calc_metrics(window["y_valid"], pred)
                window_rows.append({
                    **base, **metrics, "训练耗时秒": time.time() - start, "警告信息": warning_text,
                })
                accum[candidate["候选参数ID"]]["true"].extend(window["y_valid"].tolist())
                accum[candidate["候选参数ID"]]["pred"].extend(pred.tolist())
                if return_predictions:
                    pred_parts.append(build_prediction_frame(
                        target_name, candidate, "扩展窗口验证", window["valid_feature_df"],
                        window["y_valid"], pred, window_id=window_id,
                    ))
            except Exception as exc:
                window_rows.append({
                    **base, "是否成功": "否",
                    "错误信息": repr(exc) + "\n" + traceback.format_exc(limit=3),
                    "训练耗时秒": time.time() - start,
                })

    window_df = pd.DataFrame(window_rows)
    summary_rows: List[Dict[str, Any]] = []
    for candidate_id, values in accum.items():
        if not values["true"]:
            continue
        candidate = candidate_map[candidate_id]
        metrics = calc_metrics(values["true"], values["pred"])
        successful_windows = window_df.loc[
            (window_df["候选参数ID"] == candidate_id) & (window_df["是否成功"] == "是")
        ].copy()
        row = {
            "阶段": "扩展窗口验证", "目标指标": target_name, "特征组": FEATURE_GROUP,
            "特征数": len(feature_list_for_target(cfg["history"])), "模型": candidate["模型"],
            "候选参数ID": candidate_id, "参数JSON": candidate["参数JSON"],
            "扩展窗口数": int(successful_windows["窗口编号"].nunique()),
            "Expanding_n": metrics["n"], "Expanding_R2": metrics["R2"],
            "Expanding_RMSE": metrics["RMSE"], "Expanding_MAE": metrics["MAE"],
            "Expanding_sMAPE": metrics["sMAPE"], "Expanding_MAPE": metrics["MAPE"],
            "是否成功": "是", "错误信息": "",
        }
        for metric_name in ["RMSE", "MAE", "R2"]:
            values_series = pd.to_numeric(successful_windows.get(metric_name, pd.Series(dtype=float)), errors="coerce").dropna()
            row[f"Window_{metric_name}均值"] = values_series.mean() if len(values_series) else np.nan
            row[f"Window_{metric_name}标准差"] = values_series.std(ddof=1) if len(values_series) > 1 else 0.0
            row[f"Window_{metric_name}中位数"] = values_series.median() if len(values_series) else np.nan
            row[f"Window_{metric_name}2.5分位"] = values_series.quantile(0.025) if len(values_series) else np.nan
            row[f"Window_{metric_name}97.5分位"] = values_series.quantile(0.975) if len(values_series) else np.nan
        summary_rows.append(row)
    return (
        pd.DataFrame(summary_rows),
        window_df,
        pd.concat(pred_parts, ignore_index=True) if pred_parts else pd.DataFrame(),
    )


# =============================================================================
# 6. 参数筛选、正式复核与六算法代表配置
# =============================================================================


def rank_bootstrap_screen(summary: pd.DataFrame) -> pd.DataFrame:
    if summary.empty:
        return pd.DataFrame()
    parts: List[pd.DataFrame] = []
    specs = [
        ("R2_632plus均值", False), ("RMSE_632plus均值", True),
        ("MAE_632plus均值", True), ("sMAPE_632plus均值", True),
        ("App_OOB_R2差距", True), ("App_OOB_RMSE差距", True),
    ]
    for model_name, group in summary.groupby("模型", sort=False):
        part = group.copy()
        rank_cols: List[str] = []
        for col, ascending in specs:
            rank_col = f"{col}_Bootstrap排名"
            part[rank_col] = part[col].rank(method="average", ascending=ascending, na_option="bottom")
            rank_cols.append(rank_col)
        part["Bootstrap综合排序分数"] = part[rank_cols].sum(axis=1)
        part["Bootstrap模型内排名"] = part["Bootstrap综合排序分数"].rank(
            method="min", ascending=True, na_option="bottom"
        ).astype(int)
        parts.append(part.sort_values(
            ["Bootstrap综合排序分数", "RMSE_632plus均值", "R2_632plus均值"],
            ascending=[True, True, False],
        ))
    return pd.concat(parts, ignore_index=True)


def select_formal_candidates(
    screen_ranked: pd.DataFrame,
    quick_time: pd.DataFrame,
    candidates: Sequence[Mapping[str, Any]],
) -> Tuple[List[Dict[str, Any]], pd.DataFrame]:
    candidate_map = {c["候选参数ID"]: dict(c) for c in candidates}
    selected_ids: set[str] = set()
    source_rows: List[Dict[str, Any]] = []

    for model_name in MODEL_ORDER:
        boot_group = screen_ranked.loc[screen_ranked["模型"] == model_name].copy()
        top_boot = boot_group.sort_values(
            ["Bootstrap综合排序分数", "RMSE_632plus均值"], ascending=[True, True]
        ).head(KEEP_TOP_BOOTSTRAP_PER_MODEL)
        for _, row in top_boot.iterrows():
            cid = row["候选参数ID"]
            selected_ids.add(cid)
            source_rows.append({"模型": model_name, "候选参数ID": cid, "入选来源": "Bootstrap综合排名前3", "来源排序值": row["Bootstrap综合排序分数"]})

        time_group = quick_time.loc[
            (quick_time["模型"] == model_name) & (quick_time["是否成功"] == "是")
        ].copy()
        top_r2 = time_group.sort_values(["Time_R2", "Time_RMSE"], ascending=[False, True]).head(KEEP_TOP_TIME_R2_PER_MODEL)
        for _, row in top_r2.iterrows():
            cid = row["候选参数ID"]
            selected_ids.add(cid)
            source_rows.append({"模型": model_name, "候选参数ID": cid, "入选来源": "Time_R2前3", "来源排序值": row["Time_R2"]})
        top_rmse = time_group.sort_values(["Time_RMSE", "Time_R2"], ascending=[True, False]).head(KEEP_TOP_TIME_RMSE_PER_MODEL)
        for _, row in top_rmse.iterrows():
            cid = row["候选参数ID"]
            selected_ids.add(cid)
            source_rows.append({"模型": model_name, "候选参数ID": cid, "入选来源": "Time_RMSE前3", "来源排序值": row["Time_RMSE"]})

        if not any(candidate_map[cid]["模型"] == model_name for cid in selected_ids if cid in candidate_map):
            raise RuntimeError(f"{model_name} 未选出正式复核候选。")

    selected = [c for c in candidates if c["候选参数ID"] in selected_ids]
    source_df = pd.DataFrame(source_rows)
    if not source_df.empty:
        source_df = source_df.drop_duplicates().sort_values(["模型", "候选参数ID", "入选来源"]).reset_index(drop=True)
    return [dict(c) for c in selected], source_df


def merge_formal_results(
    bootstrap_summary: pd.DataFrame,
    time_result: pd.DataFrame,
    expanding_result: pd.DataFrame,
) -> pd.DataFrame:
    keys = ["目标指标", "特征组", "特征数", "模型", "候选参数ID", "参数JSON"]
    boot = bootstrap_summary.drop(columns=["阶段"], errors="ignore")
    time_df = time_result.drop(columns=["阶段"], errors="ignore")
    exp = expanding_result.drop(columns=["阶段"], errors="ignore")
    merged = boot.merge(time_df, on=keys, how="left", suffixes=("", "_time"))
    merged = merged.merge(exp, on=keys, how="left", suffixes=("", "_exp"))
    return merged


def add_reference_diagnostics(formal_merged: pd.DataFrame) -> pd.DataFrame:
    output = formal_merged.copy()
    output["旧版R2阈值参考"] = np.where(
        (output["Time_R2"] >= LEGACY_TIME_R2_REFERENCE)
        & (output["Expanding_R2"] >= LEGACY_EXPANDING_R2_REFERENCE),
        "达到", "未达到",
    )
    rank_specs = [
        ("Time_R2", False, 1.2), ("Time_RMSE", True, 1.2),
        ("Expanding_R2", False, 1.2), ("Expanding_RMSE", True, 1.2),
        ("R2_632plus均值", False, 0.6), ("RMSE_632plus均值", True, 0.6),
        ("Train_Time_R2差距", True, 0.4),
    ]
    parts: List[pd.DataFrame] = []
    for model_name, group in output.groupby("模型", sort=False):
        part = group.copy()
        score_cols: List[str] = []
        for col, ascending, weight in rank_specs:
            rc = f"旧版参考_{col}_加权排名"
            part[rc] = part[col].rank(method="average", ascending=ascending, na_option="bottom") * weight
            score_cols.append(rc)
        part["旧版加权参考分"] = part[score_cols].sum(axis=1)
        parts.append(part)
    return pd.concat(parts, ignore_index=True) if parts else output


def choose_algorithm_representatives(formal_scored: pd.DataFrame) -> pd.DataFrame:
    rows: List[pd.DataFrame] = []
    for model_name in MODEL_ORDER:
        group = formal_scored.loc[formal_scored["模型"] == model_name].copy()
        if group.empty:
            raise RuntimeError(f"{model_name} 正式复核结果为空。")
        valid = group.loc[
            pd.to_numeric(group["Expanding_RMSE"], errors="coerce").notna()
            & pd.to_numeric(group["Time_RMSE"], errors="coerce").notna()
            & pd.to_numeric(group["RMSE_632plus均值"], errors="coerce").notna()
        ].copy()
        if valid.empty:
            valid = group.copy()
            rule = "扩展窗口/时间/Bootstrap存在缺失，按可用的时间RMSE、Bootstrap RMSE兜底选择算法代表参数"
        else:
            rule = (
                "算法内部超参数选择：扩展窗口RMSE优先，时间顺序RMSE次之，Bootstrap .632+ RMSE再次；"
                "再以扩展窗口MAE、时间MAE及训练-时间R2差距作为并列时的稳定性参考。旧版加权分和R2阈值仅保留为诊断。"
            )
        sort_cols = ["Expanding_RMSE", "Time_RMSE", "RMSE_632plus均值", "Expanding_MAE", "Time_MAE", "Train_Time_R2差距"]
        for col in sort_cols:
            if col not in valid.columns:
                valid[col] = np.nan
        best = valid.sort_values(sort_cols, ascending=[True] * len(sort_cols), na_position="last").head(1).copy()
        best["代表配置选择规则"] = rule
        best["Bootstrap_R2"] = best["R2_632plus均值"]
        best["Bootstrap_RMSE"] = best["RMSE_632plus均值"]
        best["Bootstrap_MAE"] = best["MAE_632plus均值"]
        best["Bootstrap_sMAPE"] = best["sMAPE_632plus均值"]
        rows.append(best)
    representatives = pd.concat(rows, ignore_index=True)
    representatives["模型顺序"] = representatives["模型"].map({m: i for i, m in enumerate(MODEL_ORDER)})
    return representatives.sort_values("模型顺序").drop(columns=["模型顺序"]).reset_index(drop=True)


def filter_representative_rows(table: pd.DataFrame, representatives: pd.DataFrame) -> pd.DataFrame:
    if table.empty or representatives.empty:
        return pd.DataFrame()
    ids = set(representatives["候选参数ID"].astype(str))
    return table.loc[table["候选参数ID"].astype(str).isin(ids)].copy().reset_index(drop=True)


# =============================================================================
# 7. 朴素动态基线、随机种子稳定性与工况分组误差
# =============================================================================


def build_naive_baselines(
    target_name: str,
    cfg: Mapping[str, Any],
    df: pd.DataFrame,
) -> Tuple[pd.DataFrame, pd.DataFrame]:
    lag1, _, roll3 = cfg["history"]
    rows: List[Dict[str, Any]] = []
    window_rows: List[Dict[str, Any]] = []

    split_idx = int(math.floor(len(df) * TIME_TRAIN_RATIO))
    time_valid = df.iloc[split_idx:].copy()
    for baseline_name, pred_col in [("lag1持续性基线", lag1), ("roll3滚动均值基线", roll3)]:
        metrics = calc_metrics(time_valid[cfg["target_col"]], time_valid[pred_col])
        rows.append({"目标指标": target_name, "验证方式": "时间顺序验证", "基线": baseline_name, "预测字段": pred_col, **metrics})

    all_true: Dict[str, List[float]] = {"lag1持续性基线": [], "roll3滚动均值基线": []}
    all_pred: Dict[str, List[float]] = {"lag1持续性基线": [], "roll3滚动均值基线": []}
    initial_train = int(math.floor(len(df) * EXPANDING_INITIAL_TRAIN_RATIO))
    for window_id, start_idx in enumerate(range(initial_train, len(df), EXPANDING_STEP_SIZE), start=1):
        end_idx = min(start_idx + EXPANDING_VALID_WINDOW_SIZE, len(df))
        valid = df.iloc[start_idx:end_idx].copy()
        if len(valid) < 2:
            continue
        for baseline_name, pred_col in [("lag1持续性基线", lag1), ("roll3滚动均值基线", roll3)]:
            metrics = calc_metrics(valid[cfg["target_col"]], valid[pred_col])
            all_true[baseline_name].extend(valid[cfg["target_col"]].astype(float).tolist())
            all_pred[baseline_name].extend(valid[pred_col].astype(float).tolist())
            window_rows.append({
                "目标指标": target_name, "基线": baseline_name, "窗口编号": window_id,
                "窗口起始样本序号": valid["样本序号"].min(), "窗口结束样本序号": valid["样本序号"].max(),
                **metrics,
            })
    for baseline_name in ["lag1持续性基线", "roll3滚动均值基线"]:
        metrics = calc_metrics(all_true[baseline_name], all_pred[baseline_name])
        rows.append({
            "目标指标": target_name, "验证方式": "扩展窗口验证", "基线": baseline_name,
            "预测字段": lag1 if baseline_name.startswith("lag1") else roll3, **metrics,
        })
    return pd.DataFrame(rows), pd.DataFrame(window_rows)


def build_baseline_comparison(representatives: pd.DataFrame, baseline: pd.DataFrame) -> pd.DataFrame:
    rows: List[Dict[str, Any]] = []
    for _, rep in representatives.iterrows():
        for validation_type in ["时间顺序验证", "扩展窗口验证"]:
            if validation_type == "时间顺序验证":
                model_metrics = {"R2": rep["Time_R2"], "RMSE": rep["Time_RMSE"], "MAE": rep["Time_MAE"]}
            else:
                model_metrics = {"R2": rep["Expanding_R2"], "RMSE": rep["Expanding_RMSE"], "MAE": rep["Expanding_MAE"]}
            for _, base in baseline.loc[baseline["验证方式"] == validation_type].iterrows():
                base_rmse = safe_float(base["RMSE"])
                base_mae = safe_float(base["MAE"])
                rows.append({
                    "特征组": FEATURE_GROUP, "模型": rep["模型"], "候选参数ID": rep["候选参数ID"],
                    "验证方式": validation_type, "基线": base["基线"],
                    "模型R2": model_metrics["R2"], "基线R2": base["R2"],
                    "R2提高量": model_metrics["R2"] - base["R2"],
                    "模型RMSE": model_metrics["RMSE"], "基线RMSE": base_rmse,
                    "RMSE降低率百分比": (base_rmse - model_metrics["RMSE"]) / base_rmse * 100.0 if abs(base_rmse) > EPS else np.nan,
                    "模型MAE": model_metrics["MAE"], "基线MAE": base_mae,
                    "MAE降低率百分比": (base_mae - model_metrics["MAE"]) / base_mae * 100.0 if abs(base_mae) > EPS else np.nan,
                    "模型是否优于基线_RMSE": "是" if model_metrics["RMSE"] < base_rmse else "否",
                })
    return pd.DataFrame(rows)


def build_window_baseline_comparison(
    representative_windows: pd.DataFrame,
    baseline_windows: pd.DataFrame,
) -> pd.DataFrame:
    rows: List[Dict[str, Any]] = []
    if representative_windows.empty or baseline_windows.empty:
        return pd.DataFrame()
    for (model_name, candidate_id), model_group in representative_windows.loc[
        representative_windows["是否成功"] == "是"
    ].groupby(["模型", "候选参数ID"], sort=False):
        for baseline_name, base_group in baseline_windows.groupby("基线", sort=False):
            merged = model_group[["窗口编号", "RMSE", "MAE"]].merge(
                base_group[["窗口编号", "RMSE", "MAE"]], on="窗口编号", suffixes=("_模型", "_基线")
            )
            if merged.empty:
                continue
            rows.append({
                "特征组": FEATURE_GROUP, "模型": model_name, "候选参数ID": candidate_id,
                "基线": baseline_name, "可比较窗口数": len(merged),
                "RMSE优于基线窗口比例百分比": float((merged["RMSE_模型"] < merged["RMSE_基线"]).mean() * 100.0),
                "MAE优于基线窗口比例百分比": float((merged["MAE_模型"] < merged["MAE_基线"]).mean() * 100.0),
                "窗口RMSE差值均值_模型减基线": float((merged["RMSE_模型"] - merged["RMSE_基线"]).mean()),
            })
    return pd.DataFrame(rows)


def evaluate_seed_stability(
    target_name: str,
    cfg: Mapping[str, Any],
    df: pd.DataFrame,
    representatives: pd.DataFrame,
) -> pd.DataFrame:
    stochastic = representatives.loc[representatives["模型"].isin(STOCHASTIC_MODELS)].copy()
    if stochastic.empty or SEED_STABILITY_REPEATS <= 0:
        return pd.DataFrame()

    split_idx = int(math.floor(len(df) * TIME_TRAIN_RATIO))
    time_prepared = prepare_split(df.iloc[:split_idx].copy(), df.iloc[split_idx:].copy(), cfg)
    expanding_windows = prepare_expanding_windows(df, cfg)
    rows: List[Dict[str, Any]] = []

    for _, rep in stochastic.iterrows():
        params = json.loads(rep["参数JSON"])
        if rep["模型"] == "MLP" and isinstance(params.get("hidden_layer_sizes"), list):
            params["hidden_layer_sizes"] = tuple(params["hidden_layer_sizes"])
        for repeat_id in range(1, SEED_STABILITY_REPEATS + 1):
            seed = get_stability_random_state(target_name, rep["模型"], rep["候选参数ID"], repeat_id)
            model = build_model(rep["模型"], params, seed)
            warning_text = fit_model_safely(model, time_prepared["X_train"], time_prepared["y_train"])
            time_pred = model.predict(time_prepared["X_valid"])
            time_metrics = calc_metrics(time_prepared["y_valid"], time_pred)

            all_true: List[float] = []
            all_pred: List[float] = []
            for window in expanding_windows:
                if "错误" in window:
                    continue
                model_w = build_model(rep["模型"], params, seed)
                fit_model_safely(model_w, window["X_train"], window["y_train"])
                pred = model_w.predict(window["X_valid"])
                all_true.extend(window["y_valid"].tolist())
                all_pred.extend(pred.tolist())
            exp_metrics = calc_metrics(all_true, all_pred)
            rows.append({
                "目标指标": target_name, "特征组": FEATURE_GROUP, "模型": rep["模型"],
                "候选参数ID": rep["候选参数ID"], "稳定性重复": repeat_id, "随机种子": seed,
                "Time_R2": time_metrics["R2"], "Time_RMSE": time_metrics["RMSE"], "Time_MAE": time_metrics["MAE"],
                "Expanding_R2": exp_metrics["R2"], "Expanding_RMSE": exp_metrics["RMSE"], "Expanding_MAE": exp_metrics["MAE"],
                "警告信息": warning_text,
            })

    detail = pd.DataFrame(rows)
    if detail.empty:
        return detail
    summary_rows: List[Dict[str, Any]] = []
    for (model_name, candidate_id), group in detail.groupby(["模型", "候选参数ID"], sort=False):
        row: Dict[str, Any] = {
            "目标指标": target_name, "特征组": FEATURE_GROUP, "模型": model_name,
            "候选参数ID": candidate_id, "重复次数": int(group["稳定性重复"].nunique()),
            "随机种子列表": ",".join(str(int(x)) for x in group["随机种子"]),
        }
        for metric in ["Time_R2", "Time_RMSE", "Time_MAE", "Expanding_R2", "Expanding_RMSE", "Expanding_MAE"]:
            values = pd.to_numeric(group[metric], errors="coerce").dropna()
            row[f"{metric}均值"] = values.mean() if len(values) else np.nan
            row[f"{metric}标准差"] = values.std(ddof=1) if len(values) > 1 else 0.0
            row[f"{metric}最小值"] = values.min() if len(values) else np.nan
            row[f"{metric}最大值"] = values.max() if len(values) else np.nan
        summary_rows.append(row)
    summary = pd.DataFrame(summary_rows)
    summary.attrs["detail"] = detail
    return summary


def build_global_diagnostic_annotations(
    target_name: str,
    cfg: Mapping[str, Any],
    df: pd.DataFrame,
    full_features: pd.DataFrame,
) -> Tuple[pd.DataFrame, pd.DataFrame]:
    annotations = df[["_row_id", "样本序号", "Date"]].copy()
    n = len(annotations)
    first_cut = int(math.floor(n / 3))
    second_cut = int(math.floor(2 * n / 3))
    annotations["运行时间阶段"] = "后期"
    annotations.loc[annotations.index < second_cut, "运行时间阶段"] = "中期"
    annotations.loc[annotations.index < first_cut, "运行时间阶段"] = "前期"

    operation_values = df[OPERATION_COLUMNS].astype(float)
    change_flags = operation_values.ne(operation_values.shift(1)).any(axis=1)
    change_flags.iloc[0] = False
    transition_labels: List[str] = []
    for index in range(n):
        recent_start = max(0, index - 2)
        transition_labels.append("工况调整后1–3个样本" if bool(change_flags.iloc[recent_start:index + 1].any()) else "相对稳定期")
    annotations["工况过渡状态"] = transition_labels
    if "Stage_ID" in df.columns:
        annotations["Stage_ID"] = df["Stage_ID"].values

    threshold_rows: List[Dict[str, Any]] = []

    # 运行变量：水平数较少时保留实际设定水平；连续取值较多时分为低/中/高三档。
    for column in OPERATION_COLUMNS:
        values = df[column].astype(float)
        group_column = f"{column}_分组"
        unique_count = int(values.nunique(dropna=True))
        if unique_count <= 10:
            annotations[group_column] = values.map(lambda x: f"{float(x):.6g}" if np.isfinite(x) else "缺失")
            threshold_rows.append({
                "分组字段": column, "分组列": group_column, "分组方式": "实际记录水平",
                "低中分界值": np.nan, "中高分界值": np.nan,
                "分组规则": f"该变量共有{unique_count}个实际记录水平，直接按实际水平分组；仅用于样本外误差诊断。",
            })
        else:
            lower = float(values.quantile(1 / 3))
            upper = float(values.quantile(2 / 3))
            if lower < upper:
                annotations[group_column] = pd.cut(
                    values, bins=[-np.inf, lower, upper, np.inf],
                    labels=["低", "中", "高"], include_lowest=True,
                ).astype("string")
            else:
                annotations[group_column] = "无法区分"
            threshold_rows.append({
                "分组字段": column, "分组列": group_column, "分组方式": "三分位水平",
                "低中分界值": lower, "中高分界值": upper,
                "分组规则": "连续取值较多，按全体动态可用样本三分位分为低/中/高；仅用于样本外误差诊断。",
            })

    # 进水与负荷变量统一按三分位分组，避免逐个连续数值形成过细小样本组。
    target_specific = {
        "TP": ["Inf.TP", "TP_volumetric_load"],
        "TN": ["Inf.TN", "TN_volumetric_load"],
        "NH4+-N": ["Inf.NH4+-N", "NH4_N_volumetric_load"],
        "COD": ["Inf.COD", "COD_volumetric_load"],
    }[target_name]
    threshold_columns = list(dict.fromkeys(["Inf.COD", "Inf.TN"] + target_specific))
    for column in threshold_columns:
        values = full_features[column].astype(float)
        lower = float(values.quantile(1 / 3))
        upper = float(values.quantile(2 / 3))
        group_column = f"{column}_分组"
        if lower < upper:
            annotations[group_column] = pd.cut(
                values, bins=[-np.inf, lower, upper, np.inf],
                labels=["低", "中", "高"], include_lowest=True,
            ).astype("string")
        else:
            annotations[group_column] = "无法区分"
        threshold_rows.append({
            "分组字段": column, "分组列": group_column, "分组方式": "三分位水平",
            "低中分界值": lower, "中高分界值": upper,
            "分组规则": "基于全体动态可用样本三分位阈值，仅用于样本外误差诊断，不参与训练或模型选择。",
        })
    return annotations, pd.DataFrame(threshold_rows)


def metrics_with_reliability(y_true: Sequence[float], y_pred: Sequence[float]) -> Tuple[Dict[str, float], str]:
    metrics = calc_metrics(y_true, y_pred)
    n = int(metrics["n"])
    if n < 5:
        metrics["R2"] = np.nan
        return metrics, "极小样本（n<5），仅供核查"
    if n < 10:
        metrics["R2"] = np.nan
        return metrics, "小样本（5≤n<10），不解释R2"
    return metrics, "n≥10，可完整报告指标"


def group_prediction_error(predictions: pd.DataFrame, group_column: str, group_dimension: str) -> pd.DataFrame:
    if predictions.empty or group_column not in predictions.columns:
        return pd.DataFrame()
    rows: List[Dict[str, Any]] = []
    available = predictions.loc[predictions[group_column].notna()].copy()
    group_keys = ["模型", "候选参数ID", "验证类型", group_column]
    for key_values, group in available.groupby(group_keys, dropna=False):
        model_name, candidate_id, validation_type, group_value = key_values
        metrics, reliability = metrics_with_reliability(group["实测值"], group["预测值"])
        rows.append({
            "特征组": FEATURE_GROUP, "模型": model_name, "候选参数ID": candidate_id,
            "验证类型": validation_type, "分组维度": group_dimension, "分组字段": group_column,
            "分组值": str(group_value), "样本数": metrics["n"], "R2": metrics["R2"],
            "RMSE": metrics["RMSE"], "MAE": metrics["MAE"], "sMAPE": metrics["sMAPE"],
            "平均实测值": group["实测值"].mean(), "平均预测值": group["预测值"].mean(),
            "平均残差": group["残差"].mean(), "平均绝对误差": group["绝对误差"].mean(),
            "可靠性说明": reliability,
        })
    return pd.DataFrame(rows)


def build_group_error_analysis(
    representative_predictions: pd.DataFrame,
    annotations: pd.DataFrame,
    thresholds: pd.DataFrame,
) -> pd.DataFrame:
    if representative_predictions.empty:
        return pd.DataFrame()
    enriched = representative_predictions.merge(
        annotations, left_on="源位置", right_on="_row_id", how="left", suffixes=("", "_注释")
    )
    tables: List[pd.DataFrame] = []

    if not thresholds.empty:
        for _, spec in thresholds.drop_duplicates(subset=["分组列"]).iterrows():
            group_column = spec["分组列"]
            dimension = f"{spec['分组字段']}（{spec['分组方式']}）"
            table = group_prediction_error(enriched, group_column, dimension)
            if not table.empty:
                tables.append(table)

    for column, dimension in [("运行时间阶段", "运行时间阶段"), ("工况过渡状态", "工况调整过渡期"), ("Stage_ID", "实验阶段标识")]:
        table = group_prediction_error(enriched, column, dimension)
        if not table.empty:
            tables.append(table)
    return pd.concat(tables, ignore_index=True) if tables else pd.DataFrame()


# =============================================================================
# 8. 输出信息、错误/警告整理
# =============================================================================


def collect_error_warning_tables(*tables: pd.DataFrame) -> Tuple[pd.DataFrame, pd.DataFrame]:
    error_parts: List[pd.DataFrame] = []
    warning_parts: List[pd.DataFrame] = []
    for table in tables:
        if table is None or table.empty:
            continue
        if "是否成功" in table.columns:
            errors = table.loc[table["是否成功"] == "否"].copy()
            if not errors.empty:
                error_parts.append(errors)
        if "警告信息" in table.columns:
            warns = table.loc[table["警告信息"].fillna("").astype(str).str.len() > 0].copy()
            if not warns.empty:
                keep = [c for c in ["阶段", "目标指标", "特征组", "模型", "候选参数ID", "警告信息"] if c in warns.columns]
                warning_parts.append(warns[keep])
    error_df = pd.concat(error_parts, ignore_index=True) if error_parts else pd.DataFrame()
    warning_df = pd.concat(warning_parts, ignore_index=True) if warning_parts else pd.DataFrame()
    if not warning_df.empty:
        group_cols = [c for c in warning_df.columns if c != "警告信息"] + ["警告信息"]
        warning_df = warning_df.groupby(group_cols, dropna=False).size().reset_index(name="出现次数")
    return error_df, warning_df


def build_info_tables(
    target_name: str,
    cfg: Mapping[str, Any],
    df: pd.DataFrame,
    candidates: Sequence[Mapping[str, Any]],
    formal_candidates: Sequence[Mapping[str, Any]],
) -> Tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    features = feature_list_for_target(cfg["history"])
    info_rows: List[Dict[str, Any]] = [
        {"项目": "目标指标", "内容": target_name},
        {"项目": "目标变量", "内容": cfg["target_col"]},
        {"项目": "特征组", "内容": FEATURE_GROUP},
        {"项目": "输入文件", "内容": str(INPUT_FILE)},
        {"项目": "输入工作表", "内容": cfg["sheet"]},
        {"项目": "输出目录", "内容": str(OUTPUT_DIR)},
        {"项目": "动态可用样本数", "内容": len(df)},
        {"项目": "样本序号范围", "内容": f"{int(df['样本序号'].min())}–{int(df['样本序号'].max())}"},
        {"项目": "日期范围", "内容": f"{df['Date'].min().date()}–{df['Date'].max().date()}"},
        {"项目": "特征数", "内容": len(features)},
        {"项目": "特征顺序", "内容": "、".join(features)},
        {"项目": "候选算法", "内容": "、".join(MODEL_ORDER)},
        {"项目": "主随机种子", "内容": RANDOM_SEED},
        {"项目": "初筛Bootstrap次数", "内容": BOOTSTRAP_SCREEN_N},
        {"项目": "正式Bootstrap次数", "内容": BOOTSTRAP_FINAL_N},
        {"项目": "时间顺序验证", "内容": f"前{TIME_TRAIN_RATIO:.0%}训练、后{1-TIME_TRAIN_RATIO:.0%}验证"},
        {"项目": "扩展窗口验证", "内容": f"初始训练{EXPANDING_INITIAL_TRAIN_RATIO:.0%}、窗口{EXPANDING_VALID_WINDOW_SIZE}、步长{EXPANDING_STEP_SIZE}"},
        {"项目": "正式候选保留规则", "内容": "每种算法 Bootstrap综合排名前3 ∪ Time_R2前3 ∪ Time_RMSE前3"},
        {"项目": "训练内插补", "内容": f"原始样本跨度≤{MAX_INTERPOLATION_GAP}的内部短缺口线性插值 + 训练中位数兜底"},
        {"项目": "验证/OOB插补", "内容": "仅使用训练中位数；不利用验证序列或未来观测进行双向插值"},
        {"项目": "标准化", "内容": "每次训练/验证划分内由训练集拟合StandardScaler，并应用于验证集；保持原项目设置"},
        {"项目": "Q", "内容": "不进入模型"},
        {"项目": "本阶段模型文件", "内容": "不保存joblib；仅输出六算法代表配置及完整验证证据"},
        {"项目": "全部候选参数组合数", "内容": len(candidates)},
        {"项目": "正式复核候选数", "内容": len(formal_candidates)},
    ]
    for col in IMPUTE_COLUMNS:
        info_rows.append({"项目": f"{col}真实缺失数", "内容": int(df[col].isna().sum())})

    feature_rows = []
    for idx, feature in enumerate(features, start=1):
        feature_type = "原始输入" if feature in RAW_FEATURES else ("机理衍生" if feature in MECHANISM_FEATURES else "历史状态")
        feature_rows.append({"特征组": FEATURE_GROUP, "特征数": len(features), "组内顺序": idx, "特征类型": feature_type, "特征名称": feature})

    formal_ids = set(c["候选参数ID"] for c in formal_candidates)
    parameter_rows = [{
        "全局参数序号": c["全局参数序号"], "模型": c["模型"], "参数序号": c["参数序号"],
        "候选参数ID": c["候选参数ID"], "参数JSON": c["参数JSON"],
        "是否进入正式复核": "是" if c["候选参数ID"] in formal_ids else "否",
    } for c in candidates]
    return pd.DataFrame(info_rows), pd.DataFrame(feature_rows), pd.DataFrame(parameter_rows)


def best_group_error_summary(group_error: pd.DataFrame) -> str:
    if group_error.empty:
        return "未生成工况分组误差。"
    eligible = group_error.loc[group_error["样本数"] >= 10].copy()
    if eligible.empty:
        return "无n≥10的稳定分组。"
    worst = eligible.sort_values("RMSE", ascending=False).head(3)
    parts = []
    for _, row in worst.iterrows():
        parts.append(f"{row['模型']}/{row['验证类型']}/{row['分组维度']}/{row['分组值']} RMSE={row['RMSE']:.4g}")
    return "；".join(parts)


# =============================================================================
# 9. 单目标完整流程
# =============================================================================


def process_target(target_name: str) -> Dict[str, Any]:
    cfg = TARGET_CONFIGS[target_name]
    output_excel = OUTPUT_DIR / f"{cfg['file_prefix']}_{FEATURE_GROUP}_六模型比较结果.xlsx"
    start_time = time.time()
    print_section(f"{target_name} | {FEATURE_GROUP} | 六模型比较开始")

    if SKIP_COMPLETED_TARGETS and output_excel.exists() and not SMOKE_TEST:
        log(f"发现既有输出并启用跳过：{output_excel}")
        reps = pd.read_excel(output_excel, sheet_name="六模型代表结果")
        baseline = pd.read_excel(output_excel, sheet_name="朴素动态基线")
        return {"target": target_name, "output_excel": output_excel, "representatives": reps, "baseline": baseline, "elapsed": 0.0, "skipped": True}

    df = read_target_table(target_name, cfg)
    candidates = build_parameter_grid()
    fingerprint = config_fingerprint(target_name, cfg, candidates)
    log(f"样本数={len(df)}；特征组={FEATURE_GROUP}；特征数={len(feature_list_for_target(cfg['history']))}；参数组合={len(candidates)}")

    screen_indices = make_bootstrap_indices(len(df), BOOTSTRAP_SCREEN_N, RANDOM_SEED + 100)
    screen_detail = load_or_compute_checkpoint(
        checkpoint_path(target_name, "01_screen_bootstrap"), fingerprint,
        lambda: evaluate_bootstrap_stage(target_name, cfg, df, candidates, screen_indices, "初筛Bootstrap"),
    )
    screen_summary = summarize_bootstrap(screen_detail)
    screen_ranked = rank_bootstrap_screen(screen_summary)

    quick_time, _ = load_or_compute_checkpoint(
        checkpoint_path(target_name, "02_quick_time"), fingerprint,
        lambda: evaluate_time_stage(target_name, cfg, df, candidates, "快速时间顺序验证", False),
    )

    formal_candidates, formal_sources = select_formal_candidates(screen_ranked, quick_time, candidates)
    log("正式候选数量：" + ", ".join(
        f"{m}={sum(c['模型'] == m for c in formal_candidates)}" for m in MODEL_ORDER
    ))

    final_indices = make_bootstrap_indices(len(df), BOOTSTRAP_FINAL_N, RANDOM_SEED + 1000)
    formal_boot_detail = load_or_compute_checkpoint(
        checkpoint_path(target_name, "03_formal_bootstrap"), fingerprint,
        lambda: evaluate_bootstrap_stage(target_name, cfg, df, formal_candidates, final_indices, "正式Bootstrap"),
    )
    formal_boot_summary = summarize_bootstrap(formal_boot_detail)

    formal_time, formal_time_pred = load_or_compute_checkpoint(
        checkpoint_path(target_name, "04_formal_time"), fingerprint,
        lambda: evaluate_time_stage(target_name, cfg, df, formal_candidates, "正式时间顺序验证", True),
    )

    formal_expanding, expanding_windows, formal_expanding_pred = load_or_compute_checkpoint(
        checkpoint_path(target_name, "05_formal_expanding"), fingerprint,
        lambda: evaluate_expanding_stage(target_name, cfg, df, formal_candidates, True),
    )

    log(f"{target_name} | 合并正式验证结果")
    formal_merged = merge_formal_results(formal_boot_summary, formal_time, formal_expanding)
    formal_scored = add_reference_diagnostics(formal_merged)
    representatives = choose_algorithm_representatives(formal_scored)
    log(f"{target_name} | 已确定六算法代表配置")
    representative_time_pred = filter_representative_rows(formal_time_pred, representatives)
    representative_exp_pred = filter_representative_rows(formal_expanding_pred, representatives)
    representative_predictions = pd.concat([representative_time_pred, representative_exp_pred], ignore_index=True, sort=False)
    representative_windows = filter_representative_rows(expanding_windows, representatives)

    log(f"{target_name} | 计算朴素动态基线")
    baseline, baseline_windows = build_naive_baselines(target_name, cfg, df)
    baseline_compare = build_baseline_comparison(representatives, baseline)
    window_baseline_compare = build_window_baseline_comparison(representative_windows, baseline_windows)

    log(f"{target_name} | 随机种子稳定性复核")
    seed_stability = evaluate_seed_stability(target_name, cfg, df, representatives)
    seed_stability_detail = seed_stability.attrs.get("detail", pd.DataFrame()) if not seed_stability.empty else pd.DataFrame()

    log(f"{target_name} | 工况分组误差诊断")
    full_features = prepare_full_features_for_diagnostics(df)
    annotations, group_thresholds = build_global_diagnostic_annotations(target_name, cfg, df, full_features)
    group_error = build_group_error_analysis(representative_predictions, annotations, group_thresholds)

    errors, warning_df = collect_error_warning_tables(
        screen_detail, quick_time, formal_boot_detail, formal_time, expanding_windows, seed_stability_detail
    )
    info_df, feature_df, parameter_df = build_info_tables(target_name, cfg, df, candidates, formal_candidates)

    tables: Dict[str, pd.DataFrame] = {
        "样本与特征信息": info_df,
        "特征定义": feature_df,
        "参数网格": parameter_df,
        "初筛Bootstrap汇总": screen_ranked,
        "初筛Bootstrap明细": screen_detail,
        "快速时间验证": quick_time,
        "正式候选来源": formal_sources,
        "正式Bootstrap汇总": formal_boot_summary,
        "正式Bootstrap明细": formal_boot_detail,
        "正式时间验证": formal_time,
        "正式扩展窗口": formal_expanding,
        "扩展窗口逐窗": expanding_windows,
        "正式候选综合信息": formal_scored,
        "六模型代表结果": representatives,
        "代表模型时间预测": representative_time_pred,
        "代表模型扩展预测": representative_exp_pred,
        "朴素动态基线": baseline,
        "六模型基线对比": baseline_compare,
        "扩展窗口基线胜率": window_baseline_compare,
        "随机种子稳定性": seed_stability,
        "随机种子稳定性明细": seed_stability_detail,
        "工况分组误差": group_error,
        "工况分组阈值": group_thresholds,
        "错误记录": errors,
        "警告汇总": warning_df,
    }
    log(f"{target_name} | 写出Excel")
    actual_excel = safe_write_excel(output_excel, tables)
    elapsed = time.time() - start_time
    remove_target_checkpoints(target_name)

    log(f"{target_name} 完成：{actual_excel}；耗时 {elapsed / 60:.2f} 分钟")
    return {
        "target": target_name, "output_excel": actual_excel,
        "representatives": representatives, "baseline": baseline,
        "baseline_compare": baseline_compare, "window_baseline_compare": window_baseline_compare,
        "seed_stability": seed_stability, "group_error": group_error,
        "errors": errors, "warnings": warning_df,
        "sample_n": len(df), "elapsed": elapsed, "skipped": False,
    }


# =============================================================================
# 10. 四指标综合 TXT 与入口
# =============================================================================


def build_combined_text(results: Mapping[str, Mapping[str, Any]], failures: Sequence[Mapping[str, Any]], total_elapsed: float) -> str:
    feature_desc = {
        "F0": "9个原始进水/运行/结构特征",
        "F1": "F0 + 9个机理衍生特征，共18个特征",
        "F2": "F0 + 当前目标lag1、lag2、roll3，共12个特征",
        "F3": "F0 + 9个机理衍生特征 + 当前目标3个历史状态特征，共21个特征",
    }[FEATURE_GROUP]
    lines: List[str] = []
    add = lines.append
    add(f"A²/O {FEATURE_GROUP} 四指标六模型比较结果说明")
    add("=" * 100)
    add("")
    add("一、运行与方法设置")
    add(f"1. 完成时间：{now_str()}")
    add(f"2. 总运行耗时：{total_elapsed / 60:.2f} 分钟")
    add(f"3. 输入文件：{INPUT_FILE}")
    add(f"4. 输出目录：{OUTPUT_DIR}")
    add(f"5. 特征组：{FEATURE_GROUP}，{feature_desc}。")
    add(f"6. 主随机种子：{RANDOM_SEED}。F0-F3脚本故意使用相同Bootstrap索引规则和不含特征组名称的模型子随机种子，便于后续成对比较。")
    add(f"7. 初筛Bootstrap={BOOTSTRAP_SCREEN_N}次；正式Bootstrap .632+={BOOTSTRAP_FINAL_N}次。")
    add(f"8. 时间顺序验证：前{TIME_TRAIN_RATIO:.0%}训练、后{1-TIME_TRAIN_RATIO:.0%}验证。")
    add(f"9. 扩展窗口：初始训练{EXPANDING_INITIAL_TRAIN_RATIO:.0%}，验证窗口{EXPANDING_VALID_WINDOW_SIZE}，步长{EXPANDING_STEP_SIZE}。")
    add("10. 每种算法的正式候选保留规则：Bootstrap综合排名前3 ∪ Time_R2前3 ∪ Time_RMSE前3。")
    add("11. 六类算法：ElasticNet、SVR、GPR、RandomForest、XGBoost、MLP。")
    add("")

    add("二、缺失值处理与信息泄漏控制")
    add(f"1. 四个进水变量在每一次训练/验证划分内部重新处理；训练唯一原始样本中，原始样本跨度不超过{MAX_INTERPOLATION_GAP}的内部短缺口允许线性插值。")
    add("2. 训练中仍无法插值的位置使用训练中位数兜底。")
    add("3. OOB、时间验证和扩展窗口验证样本不参与训练插补规则拟合，真实缺失仅使用训练中位数。")
    add("4. 验证样本不使用自身之后的未来观测进行双向线性插值。")
    add("5. 标准化器也只在训练数据上拟合，再应用于验证数据；Q不进入模型。")
    add("")

    add("三、算法代表配置的含义")
    add("1. 本版不选择全局最终机器学习模型，也不保存joblib模型文件。")
    add("2. 每个目标在当前特征组下，为六种算法分别选出1个代表超参数配置，因此每个目标得到6行代表结果。")
    add("3. 算法内部代表参数以扩展窗口RMSE为首要标准，时间顺序RMSE次之，Bootstrap .632+ RMSE再次；MAE和训练-时间R2差距用于并列时稳定性参考。")
    add("4. 旧版Time_R2≥0.20、Expanding_R2≥0及旧加权分仅保留为诊断列，不作为本阶段强制淘汰规则。")
    add("5. 等F0-F3全部完成后，每个目标再把四版各6个代表配置合并为24个候选，按时间性能、稳定性和简约性规则确定最佳机器学习模型。")
    add("")

    section_no = 4
    for target_name in TARGET_ORDER:
        if target_name not in results:
            continue
        result = results[target_name]
        add(f"{section_no}、{target_name}结果")
        section_no += 1
        add(f"1. 输出Excel：{result['output_excel']}")
        if result.get("sample_n") is not None:
            add(f"2. 动态可用样本数：{result['sample_n']}。")
        reps = result.get("representatives", pd.DataFrame())
        if reps is None or reps.empty:
            add("3. 未生成六算法代表结果。")
        else:
            add("3. 六算法代表配置：")
            for _, row in reps.iterrows():
                add(
                    f"- {row['模型']} | {row['候选参数ID']} | "
                    f"Bootstrap RMSE={safe_float(row.get('Bootstrap_RMSE')):.6g}，"
                    f"Time RMSE={safe_float(row.get('Time_RMSE')):.6g}，"
                    f"Expanding RMSE={safe_float(row.get('Expanding_RMSE')):.6g}，"
                    f"Time R2={safe_float(row.get('Time_R2')):.6g}，"
                    f"Expanding R2={safe_float(row.get('Expanding_R2')):.6g}。"
                )
        baseline = result.get("baseline", pd.DataFrame())
        if baseline is not None and not baseline.empty:
            add("4. 朴素动态基线：")
            for _, row in baseline.iterrows():
                add(f"- {row['验证方式']} / {row['基线']}：RMSE={safe_float(row['RMSE']):.6g}，R2={safe_float(row['R2']):.6g}。")
        seed_stability = result.get("seed_stability", pd.DataFrame())
        if seed_stability is not None and not seed_stability.empty:
            add(f"5. 随机模型稳定性复核：共{len(seed_stability)}个随机算法代表配置，详见Excel。")
        group_error = result.get("group_error", pd.DataFrame())
        add(f"6. 工况分组误差概览：{best_group_error_summary(group_error)}")
        errors = result.get("errors", pd.DataFrame())
        warnings_df = result.get("warnings", pd.DataFrame())
        add(f"7. 错误记录数={0 if errors is None else len(errors)}；去重后警告记录数={0 if warnings_df is None else len(warnings_df)}。")
        add("")

    if failures:
        add(f"{section_no}、运行失败记录")
        section_no += 1
        for item in failures:
            add(f"- {item['目标指标']}：{item['错误']}")
        add("")

    add(f"{section_no}、后续使用边界")
    add("1. 当前F版本的6个算法代表配置只是24候选中的一部分，不能单独称为该指标最终预测模型。")
    add("2. lag1和roll3为独立朴素动态基线，不参与当前机器学习参数内部筛选。")
    add("3. F0-F3四版全部跑完后，应在完全一致的样本、Bootstrap重复和时间窗口上进行24配置综合比较。")
    add("4. 若最终最佳机器学习模型稳定优于两个历史基线，可定义为最终预测模型；否则同时保留最佳机器学习模型与表现更好的朴素动态基线。")
    add("5. SHAP和过程响应分析应在四版综合选择完成后再针对最终保留的最佳机器学习模型开展。")
    return "\n".join(lines)


def main() -> None:
    total_start = time.time()
    print_section(f"A²/O {FEATURE_GROUP} 四指标六模型比较总程序开始")
    if not XGBOOST_AVAILABLE:
        raise ImportError("当前环境未检测到xgboost，请执行：pip install xgboost")
    if BOOTSTRAP_SCREEN_N < 2 or BOOTSTRAP_FINAL_N < 2:
        raise ValueError("Bootstrap次数至少为2。")
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    log(f"输入文件：{INPUT_FILE}")
    log(f"输出目录：{OUTPUT_DIR}")
    log(f"特征组：{FEATURE_GROUP}")
    log(f"SMOKE_TEST={SMOKE_TEST}")

    results: Dict[str, Dict[str, Any]] = {}
    failures: List[Dict[str, Any]] = []
    for target_name in TARGET_ORDER:
        try:
            results[target_name] = process_target(target_name)
        except Exception as exc:
            failures.append({
                "目标指标": target_name,
                "错误": repr(exc),
                "错误堆栈": traceback.format_exc(),
            })
            log(f"{target_name} 运行失败：{exc}")
            traceback.print_exc()

    total_elapsed = time.time() - total_start
    combined_text = build_combined_text(results, failures, total_elapsed)
    text_path = OUTPUT_DIR / f"{FEATURE_GROUP}_四指标模型比较说明.txt"
    actual_text = safe_write_text(text_path, combined_text)

    if failures:
        failure_df = pd.DataFrame(failures)
        safe_write_excel(OUTPUT_DIR / f"{FEATURE_GROUP}_运行失败记录.xlsx", {"失败记录": failure_df})

    print_section(f"A²/O {FEATURE_GROUP} 四指标六模型比较结束")
    log(f"综合说明：{actual_text}")
    log(f"成功目标：{list(results.keys())}")
    if failures:
        log(f"失败目标：{[x['目标指标'] for x in failures]}")
    log(f"总耗时：{total_elapsed / 60:.2f} 分钟")


if __name__ == "__main__":
    try:
        main()
    except Exception:
        print("\n程序运行失败，错误堆栈如下：")
        traceback.print_exc()
        raise