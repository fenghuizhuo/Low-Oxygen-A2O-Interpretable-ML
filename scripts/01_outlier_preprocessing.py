# -*- coding: utf-8 -*-
"""
A²/O 污水处理数据预处理：固定物理边界清洗 + 3×IQR 稳健异常值筛查

默认输入文件：
${PROJECT_ROOT}/data/raw

默认输出目录：
${PROJECT_ROOT}/outputs/preprocessing/outlier

输出文件：
1. 异常值处理结果.xlsx
2. 异常值处理说明.txt

运行前安装依赖：
    pip install pandas numpy openpyxl

直接运行：
    python 01_异常值处理.py

也可以临时覆盖路径或指定工作表：
    python 01_异常值处理.py --input "[input/output path]" --output-dir "[input/output path]"
    python 01_异常值处理.py --sheet "数据总表"

处理原则：
- 不删除整行样本。
- 固定物理边界超界值置为空值。
- 3×IQR 仅用于 Inf.TP、Inf.COD、Inf.NH4+-N 和 Inf.TN 四个进水水质变量。
- 出水目标变量不进行 3×IQR 筛查。
- 本脚本只完成异常值处理，不进行缺失值插补。
- HRT、分段进水比例和反应区容积分数只做一致性核查，不使用统计方法自动修正。
- Q 不作为本阶段输入或筛查变量。
"""

from __future__ import annotations

import argparse
import math
import sys
import traceback
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


# =============================================================================
# 1. 用户配置区
# =============================================================================

PROJECT_ROOT = Path(__file__).resolve().parents[1]

DEFAULT_INPUT_FILE = Path(
    str(PROJECT_ROOT / "data/raw/数据总表.xlsx")
)
DEFAULT_OUTPUT_DIR = Path(
    str(PROJECT_ROOT / "outputs/preprocessing/outlier")
)

OUTPUT_EXCEL_NAME = "异常值处理结果.xlsx"
OUTPUT_TEXT_NAME = "异常值处理说明.txt"

# None 表示自动寻找包含全部必要字段的工作表；也可直接填写中文工作表名。
DEFAULT_SOURCE_SHEET: str | int | None = None

# 是否在控制台打印更详细的信息。
VERBOSE = True


# =============================================================================
# 2. 字段、规则与常量
# =============================================================================

ID_COLUMN = "样本序号"
DATE_COLUMN = "Date"

INFLUENT_COLUMNS = ["Inf.TP", "Inf.COD", "Inf.NH4+-N", "Inf.TN"]
EFFLUENT_COLUMNS = ["eff.TP", "eff.COD", "eff.NH4+-N", "eff.TN"]

OPERATING_COLUMNS = [
    "Aerobic_DO",
    "Anaerobic_Feed_Fraction",
    "HRT",
    "Anaerobic_Volume_Fraction",
    "Anoxic_Volume_Fraction",
]

REQUIRED_COLUMNS = [
    ID_COLUMN,
    DATE_COLUMN,
    "Aerobic_DO",
    "Anaerobic_Feed_Fraction",
    "Inf.TP",
    "eff.TP",
    "Inf.COD",
    "eff.COD",
    "Inf.NH4+-N",
    "eff.NH4+-N",
    "Inf.TN",
    "eff.TN",
    "HRT",
    "Anaerobic_Volume_Fraction",
    "Anoxic_Volume_Fraction",
]

NUMERIC_COLUMNS = [column for column in REQUIRED_COLUMNS if column != DATE_COLUMN]

# 3×IQR 只筛查这四个进水变量。
IQR_COLUMNS = INFLUENT_COLUMNS.copy()
IQR_MULTIPLIER = 3.0


@dataclass(frozen=True)
class BoundaryRule:
    lower: float
    upper: float
    lower_inclusive: bool
    upper_inclusive: bool
    variable_type: str
    unit: str
    explanation: str

    @property
    def rule_text(self) -> str:
        left = "≤" if self.lower_inclusive else "<"
        right = "≤" if self.upper_inclusive else "<"
        return f"{self.lower:g} {left} x {right} {self.upper:g}"


# 固定物理边界与修订流程一致。
FIXED_BOUNDARIES: dict[str, BoundaryRule] = {
    "Aerobic_DO": BoundaryRule(
        lower=0,
        upper=5,
        lower_inclusive=True,
        upper_inclusive=True,
        variable_type="运行工况",
        unit="mg/L",
        explanation="低氧运行条件下允许 0；小于 0 或大于 5 视为异常。",
    ),
    "Inf.TP": BoundaryRule(
        lower=0,
        upper=10,
        lower_inclusive=False,
        upper_inclusive=True,
        variable_type="进水水质",
        unit="mg/L",
        explanation="浓度应大于 0；大于 10 超出实验合理范围。",
    ),
    "eff.TP": BoundaryRule(
        lower=0,
        upper=5,
        lower_inclusive=False,
        upper_inclusive=True,
        variable_type="出水目标",
        unit="mg/L",
        explanation="浓度应大于 0；大于 5 超出实验合理范围。",
    ),
    "Inf.COD": BoundaryRule(
        lower=0,
        upper=300,
        lower_inclusive=False,
        upper_inclusive=True,
        variable_type="进水水质",
        unit="mg/L",
        explanation="浓度应大于 0；大于 300 超出人工配水实验合理范围。",
    ),
    "eff.COD": BoundaryRule(
        lower=0,
        upper=60,
        lower_inclusive=False,
        upper_inclusive=True,
        variable_type="出水目标",
        unit="mg/L",
        explanation="大于 0 且不高于 60；低于 30 但大于 0 的读数保留。",
    ),
    "Inf.NH4+-N": BoundaryRule(
        lower=0,
        upper=100,
        lower_inclusive=False,
        upper_inclusive=True,
        variable_type="进水水质",
        unit="mg/L",
        explanation="浓度应大于 0；大于 100 超出实验合理范围。",
    ),
    "eff.NH4+-N": BoundaryRule(
        lower=0,
        upper=10,
        lower_inclusive=False,
        upper_inclusive=True,
        variable_type="出水目标",
        unit="mg/L",
        explanation="浓度应大于 0；大于 10 超出实验合理范围。",
    ),
    "Inf.TN": BoundaryRule(
        lower=0,
        upper=100,
        lower_inclusive=False,
        upper_inclusive=True,
        variable_type="进水水质",
        unit="mg/L",
        explanation="浓度应大于 0；大于 100 超出实验合理范围。",
    ),
    "eff.TN": BoundaryRule(
        lower=0,
        upper=20,
        lower_inclusive=False,
        upper_inclusive=True,
        variable_type="出水目标",
        unit="mg/L",
        explanation="浓度应大于 0；大于 20 超出实验合理范围。",
    ),
}


# =============================================================================
# 3. 通用辅助函数
# =============================================================================


def parse_arguments() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="完成固定物理边界清洗和 3×IQR 稳健异常值筛查。"
    )
    parser.add_argument(
        "--input",
        type=str,
        default=str(DEFAULT_INPUT_FILE),
        help="原始 Excel 文件路径。",
    )
    parser.add_argument(
        "--output-dir",
        type=str,
        default=str(DEFAULT_OUTPUT_DIR),
        help="结果输出目录。",
    )
    parser.add_argument(
        "--sheet",
        type=str,
        default=None,
        help="输入工作表名；不填写时自动寻找包含必要字段的工作表。",
    )
    return parser.parse_args()


def normalize_column_name(value: Any) -> str:
    """清理字段名前后的普通空格、全角空格、不换行空格和换行。"""
    text = str(value)
    text = text.replace("\u3000", " ").replace("\xa0", " ")
    text = text.replace("\r", " ").replace("\n", " ")
    return " ".join(text.strip().split())


def blank_mask(series: pd.Series) -> pd.Series:
    """识别真正空值和只有空白字符的单元格。"""
    as_text = series.astype("string")
    return series.isna() | as_text.str.strip().isin(["", "<NA>", "nan", "None", "NaT"])


def safe_display_value(value: Any) -> Any:
    if pd.isna(value):
        return ""
    if isinstance(value, pd.Timestamp):
        return value.strftime("%Y-%m-%d")
    return value


def mask_outside_boundary(series: pd.Series, rule: BoundaryRule) -> pd.Series:
    valid = series.notna()
    if rule.lower_inclusive:
        lower_bad = series < rule.lower
    else:
        lower_bad = series <= rule.lower

    if rule.upper_inclusive:
        upper_bad = series > rule.upper
    else:
        upper_bad = series >= rule.upper

    return valid & (lower_bad | upper_bad)


def format_number(value: Any, digits: int = 6) -> str:
    if value is None or pd.isna(value):
        return "无"
    value = float(value)
    if math.isfinite(value):
        text = f"{value:.{digits}f}".rstrip("0").rstrip(".")
        return text if text else "0"
    return str(value)


def sample_list_text(values: Iterable[Any], limit: int = 40) -> str:
    cleaned = [str(safe_display_value(value)) for value in values if not pd.isna(value)]
    if not cleaned:
        return "无"
    if len(cleaned) <= limit:
        return "、".join(cleaned)
    return "、".join(cleaned[:limit]) + f"……（共 {len(cleaned)} 个）"


def ensure_unique_columns(columns: list[str]) -> None:
    duplicated = pd.Series(columns)[pd.Series(columns).duplicated()].tolist()
    if duplicated:
        raise ValueError(
            "字段名清理后出现重复列：" + "、".join(sorted(set(duplicated)))
        )


def choose_source_sheet(
    input_path: Path,
    requested_sheet: str | int | None,
) -> tuple[str, pd.DataFrame]:
    """自动选择包含全部必要字段的工作表。"""
    excel_file = pd.ExcelFile(input_path, engine="openpyxl")

    if requested_sheet is not None:
        if isinstance(requested_sheet, str) and requested_sheet.isdigit():
            sheet_index = int(requested_sheet)
            if not 0 <= sheet_index < len(excel_file.sheet_names):
                raise ValueError(
                    f"工作表序号 {sheet_index} 超出范围；当前文件共 "
                    f"{len(excel_file.sheet_names)} 个工作表。"
                )
            sheet_name = excel_file.sheet_names[sheet_index]
        else:
            sheet_name = str(requested_sheet)
            if sheet_name not in excel_file.sheet_names:
                raise ValueError(
                    f"未找到工作表“{sheet_name}”。可用工作表："
                    + "、".join(excel_file.sheet_names)
                )

        data = pd.read_excel(input_path, sheet_name=sheet_name, engine="openpyxl")
        data.columns = [normalize_column_name(column) for column in data.columns]
        ensure_unique_columns(data.columns.tolist())
        return sheet_name, data

    audit_messages: list[str] = []
    for sheet_name in excel_file.sheet_names:
        preview = pd.read_excel(
            input_path,
            sheet_name=sheet_name,
            nrows=5,
            engine="openpyxl",
        )
        preview.columns = [normalize_column_name(column) for column in preview.columns]
        preview_columns = preview.columns.tolist()
        missing = [column for column in REQUIRED_COLUMNS if column not in preview_columns]
        if not missing:
            data = pd.read_excel(
                input_path,
                sheet_name=sheet_name,
                engine="openpyxl",
            )
            data.columns = [normalize_column_name(column) for column in data.columns]
            ensure_unique_columns(data.columns.tolist())
            return sheet_name, data
        audit_messages.append(f"{sheet_name}：缺少 {', '.join(missing)}")

    raise ValueError(
        "没有找到同时包含全部必要字段的工作表。逐表检查结果：\n"
        + "\n".join(audit_messages)
    )


def add_anomaly_records(
    records: list[dict[str, Any]],
    raw_data: pd.DataFrame,
    working_data: pd.DataFrame,
    mask: pd.Series,
    stage: str,
    anomaly_type: str,
    variable: str,
    lower: Any,
    upper: Any,
    rule_text: str,
    action: str,
) -> None:
    indices = working_data.index[mask]
    for index in indices:
        records.append(
            {
                "处理阶段": stage,
                "异常类型": anomaly_type,
                "变量": variable,
                "样本序号": safe_display_value(working_data.at[index, ID_COLUMN]),
                "日期": safe_display_value(working_data.at[index, DATE_COLUMN]),
                "原始Excel行号": int(working_data.at[index, "__原始Excel行号__"]),
                "原始值": safe_display_value(raw_data.at[index, variable]),
                "判定下限": safe_display_value(lower),
                "判定上限": safe_display_value(upper),
                "判定规则": rule_text,
                "处理方式": action,
            }
        )


# =============================================================================
# 4. 数据处理主函数
# =============================================================================


def run_processing(
    input_path: Path,
    output_dir: Path,
    requested_sheet: str | int | None,
) -> dict[str, Any]:
    start_time = datetime.now()
    output_dir.mkdir(parents=True, exist_ok=True)

    output_excel = output_dir / OUTPUT_EXCEL_NAME
    output_text = output_dir / OUTPUT_TEXT_NAME

    if not input_path.exists():
        raise FileNotFoundError(
            f"未找到输入文件：{input_path}\n"
            "请确认路径、文件名和扩展名是否正确。"
        )

    source_sheet, raw_data = choose_source_sheet(input_path, requested_sheet)
    original_columns = raw_data.columns.tolist()

    missing_required = [
        column for column in REQUIRED_COLUMNS if column not in raw_data.columns
    ]
    if missing_required:
        raise ValueError(
            "输入表缺少必要字段：" + "、".join(missing_required)
        )

    if raw_data.empty:
        raise ValueError("输入工作表没有数据行。")

    data = raw_data.copy(deep=True)
    data["__原始Excel行号__"] = np.arange(2, len(data) + 2)

    anomaly_records: list[dict[str, Any]] = []
    consistency_records: list[dict[str, Any]] = []

    # 保存每个变量在不同阶段新增为空的掩码，用于缺失值统计。
    original_missing_masks: dict[str, pd.Series] = {}
    format_error_masks: dict[str, pd.Series] = {}
    fixed_boundary_masks: dict[str, pd.Series] = {}
    iqr_masks: dict[str, pd.Series] = {}

    for column in original_columns:
        original_missing_masks[column] = blank_mask(raw_data[column])
        format_error_masks[column] = pd.Series(False, index=data.index)
        fixed_boundary_masks[column] = pd.Series(False, index=data.index)
        iqr_masks[column] = pd.Series(False, index=data.index)

    # -------------------------------------------------------------------------
    # 4.1 字段格式核查与类型转换
    # -------------------------------------------------------------------------
    raw_date = raw_data[DATE_COLUMN].copy()
    date_blank = blank_mask(raw_date)
    parsed_date = pd.to_datetime(raw_date, errors="coerce")
    date_format_error = (~date_blank) & parsed_date.isna()
    format_error_masks[DATE_COLUMN] = date_format_error
    data[DATE_COLUMN] = parsed_date

    if date_format_error.any():
        add_anomaly_records(
            anomaly_records,
            raw_data,
            data,
            date_format_error,
            stage="数据格式核查",
            anomaly_type="日期格式异常",
            variable=DATE_COLUMN,
            lower="",
            upper="",
            rule_text="应能转换为有效日期",
            action="置为空值，需回查原始记录",
        )

    for column in NUMERIC_COLUMNS:
        raw_series = raw_data[column].copy()
        is_blank = blank_mask(raw_series)
        converted = pd.to_numeric(raw_series, errors="coerce")
        format_error = (~is_blank) & converted.isna()
        format_error_masks[column] = format_error
        data[column] = converted

        if format_error.any():
            add_anomaly_records(
                anomaly_records,
                raw_data,
                data,
                format_error,
                stage="数据格式核查",
                anomaly_type="非数值格式",
                variable=column,
                lower="",
                upper="",
                rule_text="应为可转换的数值",
                action="置为空值，需回查原始记录",
            )

    # 样本序号优先使用可空整数格式，避免有效整数显示为 1.0。
    if data[ID_COLUMN].dropna().apply(lambda value: float(value).is_integer()).all():
        data[ID_COLUMN] = data[ID_COLUMN].astype("Int64")

    # -------------------------------------------------------------------------
    # 4.2 固定物理边界清洗
    # -------------------------------------------------------------------------
    boundary_summary_rows: list[dict[str, Any]] = []

    for variable, rule in FIXED_BOUNDARIES.items():
        before_nonmissing = int(data[variable].notna().sum())
        boundary_mask = mask_outside_boundary(data[variable], rule)
        fixed_boundary_masks[variable] = boundary_mask.copy()

        if boundary_mask.any():
            add_anomaly_records(
                anomaly_records,
                raw_data,
                data,
                boundary_mask,
                stage="固定物理边界清洗",
                anomaly_type="固定边界超界",
                variable=variable,
                lower=rule.lower,
                upper=rule.upper,
                rule_text=rule.rule_text,
                action="置为空值，不删除整行样本",
            )
            data.loc[boundary_mask, variable] = np.nan

        boundary_summary_rows.append(
            {
                "变量": variable,
                "变量类别": rule.variable_type,
                "单位": rule.unit,
                "固定下限": rule.lower,
                "固定上限": rule.upper,
                "下限是否包含": "是" if rule.lower_inclusive else "否",
                "上限是否包含": "是" if rule.upper_inclusive else "否",
                "判定规则": rule.rule_text,
                "边界清洗前非缺失数": before_nonmissing,
                "固定边界异常数": int(boundary_mask.sum()),
                "边界清洗后非缺失数": int(data[variable].notna().sum()),
                "边界清洗后缺失数": int(data[variable].isna().sum()),
                "规则说明": rule.explanation,
            }
        )

    boundary_summary = pd.DataFrame(boundary_summary_rows)

    # -------------------------------------------------------------------------
    # 4.3 3×IQR 稳健异常值筛查
    # -------------------------------------------------------------------------
    iqr_summary_rows: list[dict[str, Any]] = []

    for variable in IQR_COLUMNS:
        valid_series = data[variable].dropna()
        valid_count = int(valid_series.shape[0])

        if valid_count < 4:
            q1 = q3 = iqr = lower_bound = upper_bound = np.nan
            lower_mask = pd.Series(False, index=data.index)
            upper_mask = pd.Series(False, index=data.index)
            iqr_mask = pd.Series(False, index=data.index)
            note = "有效观测少于 4 个，未执行 3×IQR 筛查。"
        else:
            q1 = float(valid_series.quantile(0.25, interpolation="linear"))
            q3 = float(valid_series.quantile(0.75, interpolation="linear"))
            iqr = q3 - q1
            lower_bound = q1 - IQR_MULTIPLIER * iqr
            upper_bound = q3 + IQR_MULTIPLIER * iqr

            lower_mask = data[variable].notna() & (data[variable] < lower_bound)
            upper_mask = data[variable].notna() & (data[variable] > upper_bound)
            iqr_mask = lower_mask | upper_mask

            if iqr == 0:
                note = (
                    "IQR 为 0；按 Q1±3×IQR 规则筛查。"
                    "建议结合变量实际设定水平人工复核。"
                )
            else:
                note = "先完成固定物理边界清洗，再基于剩余有效值计算。"

        iqr_masks[variable] = iqr_mask.copy()

        if iqr_mask.any():
            add_anomaly_records(
                anomaly_records,
                raw_data,
                data,
                iqr_mask,
                stage="3×IQR稳健筛查",
                anomaly_type="统计极端值",
                variable=variable,
                lower=lower_bound,
                upper=upper_bound,
                rule_text="x < Q1-3×IQR 或 x > Q3+3×IQR",
                action="置为空值，不删除整行样本",
            )
            data.loc[iqr_mask, variable] = np.nan

        iqr_summary_rows.append(
            {
                "变量": variable,
                "有效观测数": valid_count,
                "Q1": q1,
                "Q3": q3,
                "IQR": iqr,
                "倍数": IQR_MULTIPLIER,
                "下边界": lower_bound,
                "上边界": upper_bound,
                "下侧异常数": int(lower_mask.sum()),
                "上侧异常数": int(upper_mask.sum()),
                "异常总数": int(iqr_mask.sum()),
                "处理方式": "异常值置为空值",
                "说明": note,
            }
        )

    iqr_summary = pd.DataFrame(iqr_summary_rows)

    # -------------------------------------------------------------------------
    # 4.4 基本一致性核查：只记录，不做统计插补或自动修正
    # -------------------------------------------------------------------------
    def record_consistency(mask: pd.Series, item: str, variable: str, rule: str) -> None:
        for index in data.index[mask]:
            consistency_records.append(
                {
                    "核查项目": item,
                    "变量": variable,
                    "样本序号": safe_display_value(data.at[index, ID_COLUMN]),
                    "日期": safe_display_value(data.at[index, DATE_COLUMN]),
                    "原始Excel行号": int(data.at[index, "__原始Excel行号__"]),
                    "当前值": safe_display_value(data.at[index, variable])
                    if variable in data.columns
                    else "",
                    "核查规则": rule,
                    "建议处理": "回查实验记录或运行日志，不自动插补",
                }
            )

    # 样本序号缺失或重复。
    record_consistency(
        data[ID_COLUMN].isna(),
        "样本序号缺失或格式异常",
        ID_COLUMN,
        "样本序号应存在且可转换为数值",
    )
    duplicate_sample_mask = data[ID_COLUMN].notna() & data[ID_COLUMN].duplicated(keep=False)
    record_consistency(
        duplicate_sample_mask,
        "样本序号重复",
        ID_COLUMN,
        "样本序号原则上应唯一",
    )

    # 日期缺失/格式异常和重复。
    record_consistency(
        data[DATE_COLUMN].isna(),
        "日期缺失或格式异常",
        DATE_COLUMN,
        "Date 应存在且可转换为有效日期",
    )
    duplicate_date_mask = data[DATE_COLUMN].notna() & data[DATE_COLUMN].duplicated(keep=False)
    record_consistency(
        duplicate_date_mask,
        "日期重复",
        DATE_COLUMN,
        "同一连续运行记录中应核实重复日期是否合理",
    )

    # 运行工况与结构变量。
    feed = data["Anaerobic_Feed_Fraction"]
    record_consistency(
        feed.isna(),
        "分段进水比例缺失",
        "Anaerobic_Feed_Fraction",
        "应由实验运行记录核实，不采用统计插补",
    )
    record_consistency(
        feed.notna() & ((feed < 0) | (feed > 1)),
        "分段进水比例越界",
        "Anaerobic_Feed_Fraction",
        "应满足 0≤Anaerobic_Feed_Fraction≤1",
    )

    hrt = data["HRT"]
    record_consistency(
        hrt.isna(),
        "HRT缺失",
        "HRT",
        "HRT 为可信运行设定，应由实验记录核实",
    )
    record_consistency(
        hrt.notna() & (hrt <= 0),
        "HRT非正值",
        "HRT",
        "HRT 应大于 0",
    )

    for variable in ["Anaerobic_Volume_Fraction", "Anoxic_Volume_Fraction"]:
        series = data[variable]
        record_consistency(
            series.isna(),
            "反应区容积分数缺失",
            variable,
            "应由反应器结构记录核实",
        )
        record_consistency(
            series.notna() & ((series < 0) | (series > 1)),
            "反应区容积分数越界",
            variable,
            "容积分数应满足 0≤x≤1",
        )

    volume_sum = (
        data["Anaerobic_Volume_Fraction"]
        + data["Anoxic_Volume_Fraction"]
    )
    bad_volume_sum = volume_sum.notna() & (volume_sum >= 1)
    for index in data.index[bad_volume_sum]:
        consistency_records.append(
            {
                "核查项目": "反应区容积分数之和不合理",
                "变量": "Anaerobic_Volume_Fraction + Anoxic_Volume_Fraction",
                "样本序号": safe_display_value(data.at[index, ID_COLUMN]),
                "日期": safe_display_value(data.at[index, DATE_COLUMN]),
                "原始Excel行号": int(data.at[index, "__原始Excel行号__"]),
                "当前值": safe_display_value(volume_sum.at[index]),
                "核查规则": "厌氧区容积分数 + 缺氧区容积分数 < 1",
                "建议处理": "回查反应器结构记录，不自动修正",
            }
        )

    # 日期排序核查：仅给出汇总，不逐行制造大量记录。
    nonmissing_dates = data[DATE_COLUMN].dropna()
    dates_monotonic = bool(nonmissing_dates.is_monotonic_increasing)
    if not dates_monotonic:
        consistency_records.append(
            {
                "核查项目": "日期顺序非递增",
                "变量": DATE_COLUMN,
                "样本序号": "汇总",
                "日期": "",
                "原始Excel行号": "",
                "当前值": "",
                "核查规则": "连续运行数据应按样本序号或日期排序",
                "建议处理": "建模前按样本序号或 Date 重新排序并核对",
            }
        )

    consistency_df = pd.DataFrame(consistency_records)
    if consistency_df.empty:
        consistency_df = pd.DataFrame(
            [
                {
                    "核查项目": "未发现需人工复核的基本一致性问题",
                    "变量": "",
                    "样本序号": "",
                    "日期": "",
                    "原始Excel行号": "",
                    "当前值": "",
                    "核查规则": "",
                    "建议处理": "",
                }
            ]
        )

    # -------------------------------------------------------------------------
    # 4.5 汇总表
    # -------------------------------------------------------------------------
    anomaly_detail = pd.DataFrame(anomaly_records)
    if anomaly_detail.empty:
        anomaly_detail = pd.DataFrame(
            columns=[
                "处理阶段",
                "异常类型",
                "变量",
                "样本序号",
                "日期",
                "原始Excel行号",
                "原始值",
                "判定下限",
                "判定上限",
                "判定规则",
                "处理方式",
            ]
        )
    else:
        anomaly_detail = anomaly_detail.sort_values(
            by=["原始Excel行号", "处理阶段", "变量"],
            kind="stable",
        ).reset_index(drop=True)

    missing_rows: list[dict[str, Any]] = []
    for column in original_columns:
        original_missing_count = int(original_missing_masks[column].sum())
        format_count = int(format_error_masks[column].sum())
        boundary_count = int(fixed_boundary_masks[column].sum())
        iqr_count = int(iqr_masks[column].sum())
        final_missing_count = int(data[column].isna().sum())

        missing_rows.append(
            {
                "变量": column,
                "原始缺失数": original_missing_count,
                "格式异常新增缺失数": format_count,
                "固定边界新增缺失数": boundary_count,
                "3×IQR新增缺失数": iqr_count,
                "异常处理新增缺失总数": format_count + boundary_count + iqr_count,
                "最终缺失数": final_missing_count,
                "最终缺失率(%)": final_missing_count / len(data) * 100,
                "最终非缺失数": int(data[column].notna().sum()),
            }
        )
    missing_summary = pd.DataFrame(missing_rows)

    # 发生新增清洗的行数。
    any_new_missing = pd.Series(False, index=data.index)
    for column in original_columns:
        any_new_missing |= format_error_masks[column]
        any_new_missing |= fixed_boundary_masks[column]
        any_new_missing |= iqr_masks[column]

    overview_rows = [
        {
            "项目": "输入文件",
            "结果": str(input_path),
            "说明": "原始文件只读取，不覆盖。",
        },
        {
            "项目": "输入工作表",
            "结果": source_sheet,
            "说明": "自动选择或由命令行指定。",
        },
        {
            "项目": "数据行数",
            "结果": len(data),
            "说明": "异常处理不删除整行。",
        },
        {
            "项目": "原始字段数",
            "结果": len(original_columns),
            "说明": "额外字段会随清洗后数据一并保留。",
        },
        {
            "项目": "格式异常单元格数",
            "结果": sum(int(mask.sum()) for mask in format_error_masks.values()),
            "说明": "非空但无法转换为预期日期或数值的单元格。",
        },
        {
            "项目": "固定物理边界异常单元格数",
            "结果": sum(int(mask.sum()) for mask in fixed_boundary_masks.values()),
            "说明": "超出固定实验合理边界后置为空值。",
        },
        {
            "项目": "3×IQR异常单元格数",
            "结果": sum(int(mask.sum()) for mask in iqr_masks.values()),
            "说明": "仅针对四个进水水质变量。",
        },
        {
            "项目": "存在新增清洗值的样本行数",
            "结果": int(any_new_missing.sum()),
            "说明": "同一行可能有多个异常单元格。",
        },
        {
            "项目": "一致性核查记录数",
            "结果": 0 if consistency_records == [] else len(consistency_records),
            "说明": "此类记录不自动修改，需人工核实。",
        },
        {
            "项目": "最终清洗原则",
            "结果": "异常值置为空值，保留样本行",
            "说明": "本阶段不做缺失值插补。",
        },
    ]
    overview = pd.DataFrame(overview_rows)

    # 清洗后数据只输出原始字段，不输出内部 Excel 行号辅助列。
    cleaned_data = data[original_columns].copy()

    # -------------------------------------------------------------------------
    # 4.6 写入 Excel
    # -------------------------------------------------------------------------
    try:
        with pd.ExcelWriter(output_excel, engine="openpyxl") as writer:
            cleaned_data.to_excel(writer, sheet_name="清洗后数据", index=False)
            overview.to_excel(writer, sheet_name="处理概览", index=False)
            boundary_summary.to_excel(writer, sheet_name="固定边界汇总", index=False)
            iqr_summary.to_excel(writer, sheet_name="三倍IQR阈值", index=False)
            anomaly_detail.to_excel(writer, sheet_name="异常值明细", index=False)
            missing_summary.to_excel(writer, sheet_name="缺失值统计", index=False)
            consistency_df.to_excel(writer, sheet_name="一致性核查", index=False)
    except PermissionError as exc:
        raise PermissionError(
            f"无法写入 {output_excel}。请确认该 Excel 文件没有被其他程序打开。"
        ) from exc

    style_workbook(output_excel)

    # -------------------------------------------------------------------------
    # 4.7 写入详细 TXT 说明
    # -------------------------------------------------------------------------
    end_time = datetime.now()
    duration_seconds = (end_time - start_time).total_seconds()

    text_lines: list[str] = []
    text_lines.append("A²/O 数据异常值处理结果说明")
    text_lines.append("=" * 72)
    text_lines.append("")
    text_lines.append("一、任务与文件信息")
    text_lines.append(f"1. 执行时间：{end_time.strftime('%Y-%m-%d %H:%M:%S')}")
    text_lines.append(f"2. 运行耗时：{duration_seconds:.2f} 秒")
    text_lines.append(f"3. 输入文件：{input_path}")
    text_lines.append(f"4. 输入工作表：{source_sheet}")
    text_lines.append(f"5. 输出 Excel：{output_excel}")
    text_lines.append(f"6. 输出说明：{output_text}")
    text_lines.append("")

    text_lines.append("二、数据结构核查")
    text_lines.append(f"1. 原始数据行数：{len(raw_data)}")
    text_lines.append(f"2. 原始字段数：{len(original_columns)}")
    text_lines.append("3. 原始字段：" + "、".join(original_columns))
    extra_columns = [column for column in original_columns if column not in REQUIRED_COLUMNS]
    text_lines.append(
        "4. 额外保留字段：" + ("、".join(extra_columns) if extra_columns else "无")
    )
    text_lines.append(
        "5. 必要字段核查：全部必要字段存在。"
    )
    text_lines.append("")

    text_lines.append("三、处理顺序和基本原则")
    text_lines.append("1. 先进行字段格式核查和数值/日期类型转换。")
    text_lines.append("2. 再进行固定物理边界清洗。")
    text_lines.append("3. 最后对四个进水变量执行 3×IQR 稳健异常值筛查。")
    text_lines.append("4. 异常值仅置为空值，不删除整行样本。")
    text_lines.append("5. 出水目标变量只执行固定物理边界清洗，不执行 3×IQR。")
    text_lines.append("6. 本脚本不进行缺失值插补，不生成提前全量插补数据。")
    text_lines.append("7. HRT 和结构变量只做一致性核查，不用均值或 IQR 自动修正。")
    text_lines.append("8. Q 不作为本阶段输入或异常筛查变量。")
    text_lines.append("")

    text_lines.append("四、字段格式核查结果")
    total_format_errors = sum(int(mask.sum()) for mask in format_error_masks.values())
    text_lines.append(f"1. 格式异常单元格总数：{total_format_errors}")
    if total_format_errors == 0:
        text_lines.append("2. 未发现非空但无法转换为日期或数值的单元格。")
    else:
        for column in original_columns:
            count = int(format_error_masks[column].sum())
            if count:
                samples = data.loc[format_error_masks[column], ID_COLUMN].tolist()
                text_lines.append(
                    f"2. {column}：{count} 个格式异常；样本序号："
                    f"{sample_list_text(samples)}"
                )
    text_lines.append("")

    text_lines.append("五、固定物理边界清洗结果")
    total_fixed = sum(int(mask.sum()) for mask in fixed_boundary_masks.values())
    text_lines.append(f"固定物理边界异常单元格总数：{total_fixed}")
    for _, row in boundary_summary.iterrows():
        variable = row["变量"]
        count = int(row["固定边界异常数"])
        samples = data.loc[fixed_boundary_masks[variable], ID_COLUMN].tolist()
        text_lines.append(
            f"- {variable}：规则 {row['判定规则']} {row['单位']}；"
            f"异常 {count} 个；样本序号：{sample_list_text(samples)}。"
        )
        text_lines.append(f"  说明：{row['规则说明']}")
    text_lines.append("")

    text_lines.append("六、3×IQR 稳健异常值筛查结果")
    text_lines.append(
        "四分位距定义：IQR = Q3 - Q1；下边界 = Q1 - 3×IQR；"
        "上边界 = Q3 + 3×IQR。"
    )
    text_lines.append(
        "Q1 和 Q3 在这里表示第一、第三四分位数，与进水流量 Q 无关。"
    )
    total_iqr = sum(int(mask.sum()) for mask in iqr_masks.values())
    text_lines.append(f"3×IQR 异常单元格总数：{total_iqr}")
    for _, row in iqr_summary.iterrows():
        variable = row["变量"]
        samples = data.loc[iqr_masks[variable], ID_COLUMN].tolist()
        text_lines.append(
            f"- {variable}：有效观测数 {int(row['有效观测数'])}；"
            f"Q1={format_number(row['Q1'])}；Q3={format_number(row['Q3'])}；"
            f"IQR={format_number(row['IQR'])}；"
            f"下边界={format_number(row['下边界'])}；"
            f"上边界={format_number(row['上边界'])}；"
            f"下侧异常 {int(row['下侧异常数'])} 个；"
            f"上侧异常 {int(row['上侧异常数'])} 个；"
            f"异常总数 {int(row['异常总数'])} 个；"
            f"样本序号：{sample_list_text(samples)}。"
        )
    text_lines.append("")

    text_lines.append("七、缺失值变化")
    original_missing_total = int(
        sum(mask.sum() for mask in original_missing_masks.values())
    )
    final_missing_total = int(cleaned_data.isna().sum().sum())
    new_missing_total = total_format_errors + total_fixed + total_iqr
    text_lines.append(f"1. 原始缺失单元格总数：{original_missing_total}")
    text_lines.append(f"2. 异常处理新增缺失单元格总数：{new_missing_total}")
    text_lines.append(f"3. 最终缺失单元格总数：{final_missing_total}")
    text_lines.append(
        f"4. 存在新增清洗值的样本行数：{int(any_new_missing.sum())}"
    )
    text_lines.append("5. 各变量详细缺失变化见 Excel 的“缺失值统计”副表。")
    text_lines.append("")

    text_lines.append("八、基本一致性核查")
    if not consistency_records:
        text_lines.append("未发现需人工复核的样本序号、日期、HRT、分段进水比例或容积分数问题。")
    else:
        consistency_counter = (
            pd.DataFrame(consistency_records)["核查项目"].value_counts().to_dict()
        )
        text_lines.append(f"需人工复核的记录总数：{len(consistency_records)}")
        for item, count in consistency_counter.items():
            text_lines.append(f"- {item}：{count} 条。")
        text_lines.append("具体记录见 Excel 的“一致性核查”副表。")
    text_lines.append("")

    text_lines.append("九、Excel 工作表说明")
    text_lines.append("1. 清洗后数据：完成固定边界和 3×IQR 后的最终数据，异常值已置空。")
    text_lines.append("2. 处理概览：文件、行列数以及各处理阶段异常数量的总览。")
    text_lines.append("3. 固定边界汇总：每个固定边界变量的规则、异常数量和处理后情况。")
    text_lines.append("4. 三倍IQR阈值：四个进水变量的 Q1、Q3、IQR、上下边界和异常数量。")
    text_lines.append("5. 异常值明细：逐单元格列出异常来源、原始值、规则和处理方式。")
    text_lines.append("6. 缺失值统计：原始缺失和各处理阶段新增缺失的变量级汇总。")
    text_lines.append("7. 一致性核查：不自动修正、需要回查实验记录的问题。")
    text_lines.append("")

    text_lines.append("十、重要解释边界与下一步建议")
    text_lines.append("1. 固定物理边界是实验合理边界，不是排放达标筛选标准。")
    text_lines.append("2. 出水超标或真实运行波动不应仅因排放限值而被删除。")
    text_lines.append("3. 3×IQR 只用于进水水质统计极端值，不用于出水目标变量。")
    text_lines.append("4. 本阶段没有执行缺失值插补，清洗后的空值需在后续插补实验中处理。")
    text_lines.append("5. 后续建模时，插补规则必须在训练/验证划分内部重新拟合。")
    text_lines.append("6. 建议人工查看“异常值明细”和“一致性核查”，确认是否存在录入或单位问题。")
    text_lines.append("7. 清洗后数据可作为后续四个单目标动态可用基础表构建的起点。")
    text_lines.append("")
    text_lines.append("处理完成。")

    output_text.write_text("\n".join(text_lines), encoding="utf-8-sig")

    if VERBOSE:
        print("=" * 72)
        print("异常值处理完成")
        print(f"输入工作表：{source_sheet}")
        print(f"数据行数：{len(data)}")
        print(f"格式异常：{total_format_errors}")
        print(f"固定边界异常：{total_fixed}")
        print(f"3×IQR异常：{total_iqr}")
        print(f"Excel 输出：{output_excel}")
        print(f"TXT 输出：{output_text}")
        print("=" * 72)

    return {
        "input_file": input_path,
        "source_sheet": source_sheet,
        "output_excel": output_excel,
        "output_text": output_text,
        "rows": len(data),
        "columns": len(original_columns),
        "format_errors": total_format_errors,
        "fixed_boundary_errors": total_fixed,
        "iqr_errors": total_iqr,
        "consistency_records": len(consistency_records),
    }


# =============================================================================
# 5. Excel 格式美化
# =============================================================================


def style_workbook(excel_path: Path) -> None:
    workbook = load_workbook(excel_path)

    header_fill = PatternFill("solid", fgColor="4472C4")
    header_font = Font(color="FFFFFF", bold=True, name="微软雅黑", size=10)
    body_font = Font(name="微软雅黑", size=10)
    thin_gray = Side(style="thin", color="D9E1F2")
    border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)

    preferred_widths: dict[str, dict[str, float]] = {
        "清洗后数据": {
            "样本序号": 10,
            "Date": 13,
        },
        "异常值明细": {
            "处理阶段": 18,
            "异常类型": 16,
            "变量": 25,
            "判定规则": 38,
            "处理方式": 28,
        },
        "一致性核查": {
            "核查项目": 26,
            "变量": 38,
            "核查规则": 42,
            "建议处理": 34,
        },
        "处理概览": {
            "项目": 28,
            "结果": 55,
            "说明": 45,
        },
    }

    for sheet in workbook.worksheets:
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "A2"

        max_row = sheet.max_row
        max_col = sheet.max_column
        if max_row >= 1 and max_col >= 1:
            sheet.auto_filter.ref = f"A1:{get_column_letter(max_col)}{max_row}"

        # 表头格式。
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
        sheet.row_dimensions[1].height = 28

        # 正文格式。
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = body_font
                cell.border = border
                cell.alignment = Alignment(vertical="center", wrap_text=False)

        header_map = {
            sheet.cell(row=1, column=column_index).value: column_index
            for column_index in range(1, max_col + 1)
        }

        # 日期格式。
        if "Date" in header_map:
            column_index = header_map["Date"]
            for row_index in range(2, max_row + 1):
                sheet.cell(row=row_index, column=column_index).number_format = "yyyy-mm-dd"
        if "日期" in header_map:
            column_index = header_map["日期"]
            for row_index in range(2, max_row + 1):
                sheet.cell(row=row_index, column=column_index).number_format = "yyyy-mm-dd"

        # 数值格式。
        numeric_headers = {
            "Aerobic_DO",
            "Anaerobic_Feed_Fraction",
            "Inf.TP",
            "eff.TP",
            "Inf.COD",
            "eff.COD",
            "Inf.NH4+-N",
            "eff.NH4+-N",
            "Inf.TN",
            "eff.TN",
            "HRT",
            "Anaerobic_Volume_Fraction",
            "Anoxic_Volume_Fraction",
            "固定下限",
            "固定上限",
            "Q1",
            "Q3",
            "IQR",
            "下边界",
            "上边界",
            "判定下限",
            "判定上限",
            "原始值",
            "当前值",
        }
        for header, column_index in header_map.items():
            if header in numeric_headers:
                for row_index in range(2, max_row + 1):
                    sheet.cell(row=row_index, column=column_index).number_format = "0.0000"
            if header == "最终缺失率(%)":
                for row_index in range(2, max_row + 1):
                    sheet.cell(row=row_index, column=column_index).number_format = "0.00"

        # 自适应列宽，并限制最大宽度，避免极端长文本把表格拉得过宽。
        for column_index in range(1, max_col + 1):
            letter = get_column_letter(column_index)
            header = sheet.cell(row=1, column=column_index).value

            preferred = preferred_widths.get(sheet.title, {}).get(str(header))
            if preferred is not None:
                sheet.column_dimensions[letter].width = preferred
                continue

            max_length = 0
            inspect_rows = min(max_row, 300)
            for row_index in range(1, inspect_rows + 1):
                value = sheet.cell(row=row_index, column=column_index).value
                if value is None:
                    continue
                text = str(value)
                max_length = max(max_length, len(text))

            width = max(10, min(max_length * 1.15 + 2, 32))
            sheet.column_dimensions[letter].width = width

        # 文本较多的副表允许自动换行。
        if sheet.title in {"处理概览", "固定边界汇总", "三倍IQR阈值", "异常值明细", "一致性核查"}:
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)

    # 默认打开“处理概览”。
    if "处理概览" in workbook.sheetnames:
        workbook.active = workbook.sheetnames.index("处理概览")

    try:
        workbook.save(excel_path)
    except PermissionError as exc:
        raise PermissionError(
            f"无法保存格式化后的文件 {excel_path}。请确认该文件未被打开。"
        ) from exc


# =============================================================================
# 6. 程序入口
# =============================================================================


def main() -> int:
    arguments = parse_arguments()
    input_path = Path(arguments.input).expanduser()
    output_dir = Path(arguments.output_dir).expanduser()
    requested_sheet: str | int | None = arguments.sheet

    try:
        run_processing(
            input_path=input_path,
            output_dir=output_dir,
            requested_sheet=requested_sheet,
        )
        return 0
    except Exception as exc:  # noqa: BLE001 - 需要向用户提供完整错误上下文
        print("\n异常值处理失败：", file=sys.stderr)
        print(str(exc), file=sys.stderr)
        print("\n完整错误信息：", file=sys.stderr)
        traceback.print_exc()
        print(
            "\n请检查：输入路径、工作表字段名、文件是否被占用，以及 pandas、numpy、openpyxl 是否已安装。",
            file=sys.stderr,
        )
        return 1


if __name__ == "__main__":
    raise SystemExit(main())