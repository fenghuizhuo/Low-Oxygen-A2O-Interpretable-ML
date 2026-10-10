# -*- coding: utf-8 -*-
r"""
A²/O 四指标最终模型全量训练与固化
==================================

目的
----
本脚本只做“最终模型固化”，不重新搜索超参数、不重新选择模型、
不重新执行 Bootstrap / Time / Expanding 验证。

它读取已经完成的统一终选结果，严格锁定四个最终机器学习配置：
- TP       ：F2 + SVR         | SVR_032
- TN       ：F2 + SVR         | SVR_018
- NH4+-N   ：F2 + ElasticNet  | ElasticNet_025
- COD      ：F2 + ElasticNet  | ElasticNet_040

然后对每个目标：
1. 读取对应动态可用数据；
2. 按正式建模一致的“全历史训练”规则处理四个进水变量缺失：
   短内部缺口线性插值（最大原样本跨度 7）+ 全量训练中位数兜底；
3. 构造最终 F2 的 12 个输入特征（9 个原始当前输入 + 目标自身 lag1/lag2/roll3）；
4. 用全量历史样本拟合 StandardScaler；
5. 用已经锁定的最终参数拟合最终模型；
6. 保存完整 joblib 模型包；
7. 统一输出一个 Excel 汇总文件和一个详细 TXT 说明文件。

重要边界
--------
- 全量拟合 R²/RMSE/MAE 只用于固化核查，不是泛化性能证据。
- 泛化性能仍以统一终选文件中的 Bootstrap .632+、Time、Expanding 为准。
- 后续 SHAP、Permutation Importance、PDP/ICE、正向预测和 NSGA-II 应直接加载
  本脚本保存的 joblib 模型包，不应再次自行重训另一套最终模型。
- 对“未来新样本”的缺失进水变量，只能使用模型包中保存的训练中位数等
  当时可获得的信息，不能使用未来观测做双向插值。

默认输入
--------
最终模型终选结果：
${PROJECT_ROOT}/outputs/final_selection/F0-F3_统一终选与随机稳定性复核结果.xlsx

动态可用数据：
${PROJECT_ROOT}/outputs/preprocessing/dynamic/动态可用数据与插补比较结果.xlsx

默认输出目录（按要求）
----------------------
${PROJECT_ROOT}/outputs/final_selection/moxingbaocun

主要输出
--------
1. TP_最终模型.joblib
2. TN_最终模型.joblib
3. NH4_N_最终模型.joblib
4. COD_最终模型.joblib
5. 四指标_最终模型固化汇总.xlsx
6. 四指标_最终模型固化说明.txt

依赖
----
pip install pandas numpy openpyxl scikit-learn joblib

可选环境变量
------------
A2O_FINAL_SELECTION_FILE  最终终选 Excel 路径
A2O_DYNAMIC_INPUT_FILE    动态可用数据 Excel 路径
A2O_MODEL_SAVE_DIR        输出目录
A2O_MODEL_SAVE_STRICT=1   是否严格核对已经锁定的四个模型配置；正式运行建议保持 1
"""

from __future__ import annotations

import hashlib
import json
import math
import os
import re
import tempfile
import warnings
from copy import deepcopy
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Mapping, Sequence, Tuple

import joblib
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sklearn.exceptions import ConvergenceWarning
from sklearn.linear_model import ElasticNet
from sklearn.metrics import mean_absolute_error, mean_squared_error, r2_score
from sklearn.preprocessing import StandardScaler
from sklearn.svm import SVR


# =============================================================================
# 1. 路径、常量与最终配置锁定
# =============================================================================

PROJECT_ROOT = Path(__file__).resolve().parents[1]

FINAL_SELECTION_FILE = Path(
    os.environ.get(
        "A2O_FINAL_SELECTION_FILE",
        str(PROJECT_ROOT / "outputs/final_selection/F0-F3_统一终选与随机稳定性复核结果.xlsx"),
    )
)

DYNAMIC_INPUT_FILE = Path(
    os.environ.get(
        "A2O_DYNAMIC_INPUT_FILE",
        str(PROJECT_ROOT / "outputs/preprocessing/dynamic/动态可用数据与插补比较结果.xlsx"),
    )
)

OUTPUT_DIR = Path(
    os.environ.get(
        "A2O_MODEL_SAVE_DIR",
        str(PROJECT_ROOT / "outputs/final_selection/moxingbaocun"),
    )
)

OUTPUT_EXCEL = OUTPUT_DIR / "四指标_最终模型固化汇总.xlsx"
OUTPUT_TEXT = OUTPUT_DIR / "四指标_最终模型固化说明.txt"

RANDOM_SEED = 20260620
MAX_INTERPOLATION_GAP = 7
STRICT_MODE = os.environ.get("A2O_MODEL_SAVE_STRICT", "1").strip() != "0"
ROUNDTRIP_PRED_TOL = 1e-12

TARGET_ORDER = ["TP", "TN", "NH4+-N", "COD"]

RAW_FEATURES = [
    "Aerobic_DO",
    "Anaerobic_Feed_Fraction",
    "Inf.TP",
    "Inf.COD",
    "Inf.NH4+-N",
    "Inf.TN",
    "HRT",
    "Anaerobic_Volume_Fraction",
    "Anoxic_Volume_Fraction",
]

IMPUTE_COLUMNS = ["Inf.TP", "Inf.COD", "Inf.NH4+-N", "Inf.TN"]

OPERATION_COLUMNS = [
    "Aerobic_DO",
    "Anaerobic_Feed_Fraction",
    "HRT",
    "Anaerobic_Volume_Fraction",
    "Anoxic_Volume_Fraction",
]

TEXTUAL_MISSING_VALUES = {
    "", " ", "NA", "N/A", "na", "n/a", "NaN", "nan", "None", "none",
    "NULL", "null", "-", "—", "–", "－", "/", "\\", "未测", "缺失",
}

TARGET_CONFIG: Dict[str, Dict[str, Any]] = {
    "TP": {
        "sheet": "TP动态可用表",
        "target_col": "eff.TP",
        "history": ["eff.TP_lag1", "eff.TP_lag2", "eff.TP_roll3"],
        "joblib_name": "TP_最终模型.joblib",
    },
    "TN": {
        "sheet": "TN动态可用表",
        "target_col": "eff.TN",
        "history": ["eff.TN_lag1", "eff.TN_lag2", "eff.TN_roll3"],
        "joblib_name": "TN_最终模型.joblib",
    },
    "NH4+-N": {
        "sheet": "氨氮动态可用表",
        "target_col": "eff.NH4+-N",
        "history": ["eff.NH4+-N_lag1", "eff.NH4+-N_lag2", "eff.NH4+-N_roll3"],
        "joblib_name": "NH4_N_最终模型.joblib",
    },
    "COD": {
        "sheet": "COD动态可用表",
        "target_col": "eff.COD",
        "history": ["eff.COD_lag1", "eff.COD_lag2", "eff.COD_roll3"],
        "joblib_name": "COD_最终模型.joblib",
    },
}

# 这四项来自已经完成并确认的统一终选结果。
# STRICT_MODE=1 时，输入终选表必须与下列结果完全一致，否则停止，避免保存错模型。
EXPECTED_FINAL_CONFIG: Dict[str, Dict[str, Any]] = {
    "TP": {
        "feature_group": "F2",
        "feature_count": 12,
        "algorithm": "SVR",
        "candidate_id": "SVR_032",
        "parameters": {"C": 10.0, "epsilon": 0.05, "gamma": 0.01, "kernel": "rbf"},
    },
    "TN": {
        "feature_group": "F2",
        "feature_count": 12,
        "algorithm": "SVR",
        "candidate_id": "SVR_018",
        "parameters": {"C": 1.0, "epsilon": 0.1, "gamma": 0.03, "kernel": "rbf"},
    },
    "NH4+-N": {
        "feature_group": "F2",
        "feature_count": 12,
        "algorithm": "ElasticNet",
        "candidate_id": "ElasticNet_025",
        "parameters": {"alpha": 0.01, "l1_ratio": 0.9},
    },
    "COD": {
        "feature_group": "F2",
        "feature_count": 12,
        "algorithm": "ElasticNet",
        "candidate_id": "ElasticNet_040",
        "parameters": {"alpha": 0.3, "l1_ratio": 0.9},
    },
}


# =============================================================================
# 2. 通用工具
# =============================================================================

def now_str() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def log(message: str) -> None:
    print(f"[{now_str()}] {message}", flush=True)


def clean_column_name(value: Any) -> str:
    text = str(value)
    text = text.replace("\ufeff", "").replace("\xa0", " ").replace("\u3000", " ")
    text = text.replace("\n", "").replace("\r", "").strip()
    return re.sub(r"\s+", "", text)


def convert_to_numeric_series(series: pd.Series) -> pd.Series:
    output = series.copy()
    if output.dtype == "object" or str(output.dtype).startswith("string"):
        output = (
            output.astype("string")
            .str.strip()
            .str.replace(",", "", regex=False)
            .str.replace("，", "", regex=False)
        )
        output = output.mask(output.isin(TEXTUAL_MISSING_VALUES), np.nan)
    return pd.to_numeric(output, errors="coerce")


def safe_float(value: Any) -> float:
    try:
        x = float(value)
    except Exception:
        return np.nan
    return x if np.isfinite(x) else np.nan


def sha256_file(path: Path, block_size: int = 1024 * 1024) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while True:
            block = handle.read(block_size)
            if not block:
                break
            digest.update(block)
    return digest.hexdigest()


def json_normalize(value: Any) -> Any:
    if isinstance(value, Mapping):
        return {str(k): json_normalize(v) for k, v in sorted(value.items(), key=lambda x: str(x[0]))}
    if isinstance(value, (list, tuple)):
        return [json_normalize(v) for v in value]
    if isinstance(value, (np.integer,)):
        return int(value)
    if isinstance(value, (np.floating,)):
        return float(value)
    return value


def params_equal(left: Mapping[str, Any], right: Mapping[str, Any]) -> bool:
    if set(left.keys()) != set(right.keys()):
        return False
    for key in left:
        a, b = left[key], right[key]
        if isinstance(a, (int, float, np.integer, np.floating)) and isinstance(
            b, (int, float, np.integer, np.floating)
        ):
            if not np.isclose(float(a), float(b), rtol=1e-12, atol=1e-12):
                return False
        else:
            if str(a) != str(b):
                return False
    return True


def feature_category(feature: str, history_features: Sequence[str]) -> str:
    if feature in history_features:
        return "历史状态"
    if feature in {"Inf.TP", "Inf.COD", "Inf.NH4+-N", "Inf.TN"}:
        return "进水水质"
    return "运行与结构变量"


def calc_metrics(y_true: Iterable[float], y_pred: Iterable[float]) -> Dict[str, float]:
    true = np.asarray(list(y_true), dtype=float)
    pred = np.asarray(list(y_pred), dtype=float)
    mask = np.isfinite(true) & np.isfinite(pred)
    true = true[mask]
    pred = pred[mask]
    if len(true) == 0:
        return {"n": 0, "R2": np.nan, "RMSE": np.nan, "MAE": np.nan}
    mse = float(mean_squared_error(true, pred))
    return {
        "n": int(len(true)),
        "R2": float(r2_score(true, pred)) if len(true) >= 2 and np.std(true) > 0 else np.nan,
        "RMSE": float(math.sqrt(mse)),
        "MAE": float(mean_absolute_error(true, pred)),
    }


def atomic_joblib_dump(obj: Any, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary = output_path.with_name(output_path.name + ".part")
    if temporary.exists():
        temporary.unlink()
    joblib.dump(obj, temporary, compress=3)
    os.replace(temporary, output_path)


def atomic_write_text(output_path: Path, content: str) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary = output_path.with_name(output_path.name + ".part")
    temporary.write_text(content, encoding="utf-8-sig")
    os.replace(temporary, output_path)


# =============================================================================
# 3. 读取与严格核对最终终选结果
# =============================================================================

def load_final_selection() -> Tuple[pd.DataFrame, pd.DataFrame, str]:
    if not FINAL_SELECTION_FILE.exists():
        raise FileNotFoundError(f"未找到最终终选结果：{FINAL_SELECTION_FILE}")

    excel = pd.ExcelFile(FINAL_SELECTION_FILE, engine="openpyxl")
    required_sheets = {"四指标最终ML", "朴素动态基线"}
    missing_sheets = sorted(required_sheets - set(excel.sheet_names))
    if missing_sheets:
        raise ValueError(
            f"最终终选Excel缺少必要工作表：{missing_sheets}。现有工作表：{excel.sheet_names}"
        )

    final_ml = pd.read_excel(
        FINAL_SELECTION_FILE, sheet_name="四指标最终ML", engine="openpyxl"
    )
    baselines = pd.read_excel(
        FINAL_SELECTION_FILE, sheet_name="朴素动态基线", engine="openpyxl"
    )

    required_columns = [
        "目标指标", "特征组", "特征数", "模型", "候选参数ID", "参数JSON",
        "Bootstrap_R2", "Bootstrap_RMSE", "Bootstrap_MAE",
        "Eval_Time_R2", "Eval_Time_RMSE", "Eval_Time_MAE",
        "Eval_Expanding_R2", "Eval_Expanding_RMSE", "Eval_Expanding_MAE",
        "ML是否在两种时间验证中均优于最佳历史基线", "最终预测结论",
    ]
    missing_columns = [c for c in required_columns if c not in final_ml.columns]
    if missing_columns:
        raise ValueError(f"“四指标最终ML”缺少字段：{missing_columns}")

    final_ml = final_ml.loc[final_ml["目标指标"].isin(TARGET_ORDER)].copy()
    if len(final_ml) != 4 or final_ml["目标指标"].nunique() != 4:
        raise RuntimeError(
            f"“四指标最终ML”应恰好包含TP、TN、NH4+-N、COD各1行，当前行数={len(final_ml)}。"
        )

    for target in TARGET_ORDER:
        row = final_ml.loc[final_ml["目标指标"] == target].iloc[0]
        expected = EXPECTED_FINAL_CONFIG[target]
        try:
            params = json.loads(str(row["参数JSON"]))
        except Exception as exc:
            raise ValueError(f"{target} 参数JSON无法解析：{row['参数JSON']}") from exc

        actual_signature = (
            str(row["特征组"]),
            int(row["特征数"]),
            str(row["模型"]),
            str(row["候选参数ID"]),
        )
        expected_signature = (
            expected["feature_group"],
            expected["feature_count"],
            expected["algorithm"],
            expected["candidate_id"],
        )

        if STRICT_MODE and actual_signature != expected_signature:
            raise RuntimeError(
                f"{target} 最终配置与已确认结果不一致。\n"
                f"实际：{actual_signature}\n预期：{expected_signature}\n"
                "为避免固化错误模型，脚本已停止。"
            )
        if STRICT_MODE and not params_equal(params, expected["parameters"]):
            raise RuntimeError(
                f"{target} 最终参数与已确认结果不一致。\n"
                f"实际：{json.dumps(params, ensure_ascii=False, sort_keys=True)}\n"
                f"预期：{json.dumps(expected['parameters'], ensure_ascii=False, sort_keys=True)}"
            )

        if str(row["特征组"]) != "F2" or int(row["特征数"]) != 12:
            raise RuntimeError(f"{target} 当前最终模型不是F2的12特征配置，停止固化。")

    source_hash = sha256_file(FINAL_SELECTION_FILE)
    return final_ml, baselines, source_hash


def best_baseline_for_target(
    baselines: pd.DataFrame,
    target_name: str,
    validation_keyword: str,
) -> Dict[str, Any]:
    group = baselines.loc[
        (baselines["目标指标"].astype(str) == target_name)
        & (baselines["验证方式"].astype(str).str.contains(validation_keyword, na=False))
    ].copy()
    if group.empty:
        raise RuntimeError(f"{target_name} 未找到包含“{validation_keyword}”的朴素动态基线。")
    group["RMSE"] = pd.to_numeric(group["RMSE"], errors="coerce")
    group = group.loc[group["RMSE"].notna()].sort_values("RMSE")
    if group.empty:
        raise RuntimeError(f"{target_name}-{validation_keyword} 朴素基线RMSE无有效值。")
    row = group.iloc[0]
    return {
        "name": str(row["基线"]),
        "prediction_field": str(row.get("预测字段", "")),
        "n": int(row["n"]) if pd.notna(row.get("n", np.nan)) else None,
        "R2": safe_float(row.get("R2", np.nan)),
        "RMSE": safe_float(row.get("RMSE", np.nan)),
        "MAE": safe_float(row.get("MAE", np.nan)),
    }


# =============================================================================
# 4. 动态数据与全量训练预处理
# =============================================================================

def read_dynamic_table(target_name: str) -> pd.DataFrame:
    if not DYNAMIC_INPUT_FILE.exists():
        raise FileNotFoundError(f"未找到动态可用数据：{DYNAMIC_INPUT_FILE}")

    cfg = TARGET_CONFIG[target_name]
    excel = pd.ExcelFile(DYNAMIC_INPUT_FILE, engine="openpyxl")
    if cfg["sheet"] not in excel.sheet_names:
        raise ValueError(
            f"动态数据缺少 {target_name} 工作表“{cfg['sheet']}”。现有：{excel.sheet_names}"
        )

    df = pd.read_excel(
        DYNAMIC_INPUT_FILE,
        sheet_name=cfg["sheet"],
        engine="openpyxl",
        na_values=list(TEXTUAL_MISSING_VALUES),
        keep_default_na=True,
    )
    df.columns = [clean_column_name(c) for c in df.columns]

    required = (
        ["样本序号", "Date"]
        + RAW_FEATURES
        + [cfg["target_col"]]
        + list(cfg["history"])
    )
    missing = [c for c in required if c not in df.columns]
    if missing:
        raise ValueError(f"{target_name} 动态可用表缺少必要字段：{missing}")

    df["Date"] = pd.to_datetime(df["Date"], errors="coerce")
    for col in [c for c in required if c != "Date"]:
        df[col] = convert_to_numeric_series(df[col])

    df = df.sort_values(["样本序号", "Date"], kind="mergesort").reset_index(drop=True)

    if df.empty:
        raise ValueError(f"{target_name} 动态可用表为空。")
    if df["样本序号"].isna().any() or df["Date"].isna().any():
        raise ValueError(f"{target_name} 存在无效样本序号或日期。")
    if df["样本序号"].duplicated().any():
        raise ValueError(f"{target_name} 存在重复样本序号。")
    if not df["样本序号"].is_monotonic_increasing:
        raise ValueError(f"{target_name} 样本序号未按递增排序。")

    must_complete = [cfg["target_col"]] + list(cfg["history"]) + OPERATION_COLUMNS
    missing_counts = {c: int(df[c].isna().sum()) for c in must_complete if df[c].isna().any()}
    if missing_counts:
        raise ValueError(
            f"{target_name} 的目标、历史状态或运行/结构变量存在缺失："
            + json.dumps(missing_counts, ensure_ascii=False)
        )
    if (df["HRT"] <= 0).any():
        raise ValueError(f"{target_name} 存在HRT<=0。")

    return df


def interpolate_short_internal_gaps(
    series: pd.Series,
    sequence_coordinates: pd.Series,
    max_gap: int,
) -> pd.Series:
    """
    与统一终选/正式建模脚本保持一致：
    只填两侧均有观测、且按原始“样本序号”计算的内部缺口跨度不超过max_gap的缺失。
    """
    output = convert_to_numeric_series(series).astype(float).copy()
    coordinates = convert_to_numeric_series(sequence_coordinates).astype(float).to_numpy()
    values = output.to_numpy(copy=True)
    n = len(values)
    position = 0

    while position < n:
        if np.isfinite(values[position]):
            position += 1
            continue
        start = position
        while position < n and not np.isfinite(values[position]):
            position += 1
        end = position - 1
        left = start - 1
        right = end + 1

        if (
            left >= 0
            and right < n
            and np.isfinite(values[left])
            and np.isfinite(values[right])
            and np.isfinite(coordinates[left])
            and np.isfinite(coordinates[right])
        ):
            original_gap_span = coordinates[right] - coordinates[left] - 1.0
            coordinate_span = coordinates[right] - coordinates[left]
            if 0 <= original_gap_span <= max_gap and coordinate_span > 0:
                left_value = float(values[left])
                right_value = float(values[right])
                for index in range(start, end + 1):
                    fraction = (coordinates[index] - coordinates[left]) / coordinate_span
                    values[index] = left_value + fraction * (right_value - left_value)

    return pd.Series(values, index=output.index, name=output.name)


def prepare_full_training_data(
    raw_data: pd.DataFrame,
    target_name: str,
) -> Tuple[pd.DataFrame, List[str], Dict[str, float], pd.DataFrame]:
    data = raw_data.copy(deep=True)
    cfg = TARGET_CONFIG[target_name]
    feature_names = list(RAW_FEATURES) + list(cfg["history"])

    if len(feature_names) != 12 or len(set(feature_names)) != 12:
        raise RuntimeError(f"{target_name} F2特征清单不是12个唯一特征：{feature_names}")
    if "Q" in feature_names:
        raise RuntimeError("Q 不允许进入最终模型。")

    medians: Dict[str, float] = {}
    imputation_rows: List[Dict[str, Any]] = []

    for col in IMPUTE_COLUMNS:
        original = convert_to_numeric_series(data[col]).astype(float)
        median_value = float(original.median(skipna=True))
        if not np.isfinite(median_value):
            raise ValueError(f"{target_name} 全量训练变量 {col} 全部缺失，无法固化模型。")

        interpolated = interpolate_short_internal_gaps(
            original,
            data["样本序号"],
            MAX_INTERPOLATION_GAP,
        )
        linear_filled = int(original.isna().sum() - interpolated.isna().sum())
        median_filled = int(interpolated.isna().sum())
        data[col] = interpolated.fillna(median_value).astype(float)
        medians[col] = median_value

        imputation_rows.append({
            "目标指标": target_name,
            "变量": col,
            "原缺失数": int(original.isna().sum()),
            "短内部线性插值数": linear_filled,
            "训练中位数兜底数": median_filled,
            "保存的全量训练中位数": median_value,
            "处理后缺失数": int(data[col].isna().sum()),
            "训练阶段规则": f"内部短缺口线性插值（原样本跨度≤{MAX_INTERPOLATION_GAP}）+训练中位数",
            "未来新样本缺失规则": "仅使用已保存训练信息（默认训练中位数）；不得使用未来观测双向插值",
        })

    remaining = {f: int(data[f].isna().sum()) for f in feature_names if data[f].isna().any()}
    if remaining:
        raise ValueError(
            f"{target_name} 全量训练特征处理后仍存在缺失："
            + json.dumps(remaining, ensure_ascii=False)
        )

    return data, feature_names, medians, pd.DataFrame(imputation_rows)


# =============================================================================
# 5. 最终模型构建、训练与模型包
# =============================================================================

def build_final_model(algorithm: str, params: Mapping[str, Any]):
    p = deepcopy(dict(params))
    if algorithm == "SVR":
        return SVR(
            kernel=str(p["kernel"]),
            C=float(p["C"]),
            epsilon=float(p["epsilon"]),
            gamma=p["gamma"],
        )
    if algorithm == "ElasticNet":
        return ElasticNet(
            alpha=float(p["alpha"]),
            l1_ratio=float(p["l1_ratio"]),
            max_iter=50000,
            random_state=RANDOM_SEED,
        )
    raise ValueError(f"当前最终固化脚本只允许SVR/ElasticNet，实际算法={algorithm}")


def fit_model_safely(model: Any, x_train: np.ndarray, y_train: np.ndarray) -> str:
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        warnings.filterwarnings("always", category=ConvergenceWarning)
        model.fit(x_train, y_train)
    return " | ".join(str(item.message)[:500] for item in caught[:10])


def build_source_validation_metrics(
    selected_row: pd.Series,
    baselines: pd.DataFrame,
    target_name: str,
) -> Dict[str, Any]:
    time_baseline = best_baseline_for_target(baselines, target_name, "时间")
    expanding_baseline = best_baseline_for_target(baselines, target_name, "扩展")

    return {
        "bootstrap_632plus": {
            "R2": safe_float(selected_row["Bootstrap_R2"]),
            "RMSE": safe_float(selected_row["Bootstrap_RMSE"]),
            "MAE": safe_float(selected_row["Bootstrap_MAE"]),
        },
        "time": {
            "R2": safe_float(selected_row["Eval_Time_R2"]),
            "RMSE": safe_float(selected_row["Eval_Time_RMSE"]),
            "MAE": safe_float(selected_row["Eval_Time_MAE"]),
        },
        "expanding": {
            "R2": safe_float(selected_row["Eval_Expanding_R2"]),
            "RMSE": safe_float(selected_row["Eval_Expanding_RMSE"]),
            "MAE": safe_float(selected_row["Eval_Expanding_MAE"]),
        },
        "best_historical_baseline": {
            "time": time_baseline,
            "expanding": expanding_baseline,
        },
        "beats_best_historical_baseline_in_both_time_validations": (
            str(selected_row["ML是否在两种时间验证中均优于最佳历史基线"]).strip() == "是"
        ),
        "final_selection_conclusion": str(selected_row["最终预测结论"]),
        "source_file": str(FINAL_SELECTION_FILE),
        "source_sheet": "四指标最终ML",
    }


def interpretation_note_for_target(target_name: str, source_metrics: Mapping[str, Any]) -> str:
    common = (
        "该joblib为统一终选后固定参数在全部历史动态可用样本上的最终固化模型。"
        "全量拟合指标仅用于模型固化核查，不替代Bootstrap .632+、Time和Expanding泛化结果。"
        "后续SHAP/PDP/Permutation Importance等解释属于模型关联解释，不应直接表述为因果效应。"
    )
    if target_name == "NH4+-N":
        return common + (
            "NH4+-N的最佳机器学习模型未在Time与Expanding两种验证中同时超过最佳roll3历史基线；"
            "因此该模型可用于最佳ML内部解释和运行响应探索，但论文需明确这一预测性能边界。"
        )
    if source_metrics["beats_best_historical_baseline_in_both_time_validations"]:
        return common + "该机器学习模型在Time和Expanding两种验证中均优于对应最佳历史基线。"
    return common


def make_package(
    target_name: str,
    selected_row: pd.Series,
    processed_data: pd.DataFrame,
    feature_names: Sequence[str],
    medians: Mapping[str, float],
    scaler: StandardScaler,
    model: Any,
    source_metrics: Mapping[str, Any],
    full_fit_metrics: Mapping[str, Any],
    dynamic_hash: str,
    selection_hash: str,
) -> Dict[str, Any]:
    cfg = TARGET_CONFIG[target_name]
    params = json.loads(str(selected_row["参数JSON"]))

    return {
        "package_type": "A2O_final_regression_model",
        "package_version": "2.0",
        "created_at": now_str(),
        "random_seed": RANDOM_SEED,
        "target_name": target_name,
        "target_column": cfg["target_col"],
        "input_sheet": cfg["sheet"],
        "selected_feature_group": str(selected_row["特征组"]),
        "feature_count": int(len(feature_names)),
        "feature_names": list(feature_names),
        "raw_features_in_model": list(RAW_FEATURES),
        "history_features_in_model": list(cfg["history"]),
        "mechanism_features_in_model": [],
        "algorithm": str(selected_row["模型"]),
        "candidate_id": str(selected_row["候选参数ID"]),
        "parameters": json_normalize(params),
        "preprocessing": {
            "version": "full_history_fit_v1",
            "impute_columns": list(IMPUTE_COLUMNS),
            "max_interpolation_gap": MAX_INTERPOLATION_GAP,
            "training_medians": {str(k): float(v) for k, v in medians.items()},
            "training_imputation_rule": (
                "全量历史训练样本：两侧均有训练观测且原样本跨度不超过7的内部缺口线性插值，"
                "剩余缺失用全量训练中位数兜底。"
            ),
            "future_prediction_missing_rule": (
                "新/未来记录不得使用未来观测双向插值；缺失进水变量只能使用当时已知历史信息，"
                "本模型包的默认兜底值为保存的全量训练中位数。"
            ),
            "required_raw_columns_for_prediction": list(feature_names),
            "scaler_class": "StandardScaler",
            "scaler_mean": [float(x) for x in np.asarray(scaler.mean_, dtype=float)],
            "scaler_scale": [float(x) for x in np.asarray(scaler.scale_, dtype=float)],
        },
        "scaler": scaler,
        "model": model,
        "source_validation_metrics": json_normalize(source_metrics),
        "full_fit_metrics": json_normalize(full_fit_metrics),
        "training_data": {
            "sample_count": int(len(processed_data)),
            "sample_ids": processed_data["样本序号"].astype(int).tolist(),
            "start_date": processed_data["Date"].min().strftime("%Y-%m-%d"),
            "end_date": processed_data["Date"].max().strftime("%Y-%m-%d"),
            "dynamic_input_file": str(DYNAMIC_INPUT_FILE),
            "dynamic_input_file_sha256": dynamic_hash,
            "final_selection_file": str(FINAL_SELECTION_FILE),
            "final_selection_file_sha256": selection_hash,
        },
        "interpretation_note": interpretation_note_for_target(target_name, source_metrics),
    }


def verify_saved_package(
    output_file: Path,
    target_name: str,
    feature_names: Sequence[str],
    x_scaled: np.ndarray,
    original_predictions: np.ndarray,
) -> Tuple[str, float]:
    loaded = joblib.load(output_file)
    required_keys = {
        "package_type", "package_version", "target_name", "target_column", "input_sheet",
        "selected_feature_group", "feature_count", "feature_names", "algorithm", "candidate_id",
        "parameters", "preprocessing", "scaler", "model", "source_validation_metrics",
        "full_fit_metrics", "training_data", "interpretation_note",
    }
    missing = sorted(required_keys - set(loaded.keys()))
    if missing:
        raise RuntimeError(f"{target_name} 保存后的模型包缺少字段：{missing}")
    if loaded["target_name"] != target_name:
        raise RuntimeError(f"{target_name} 模型包目标名称回读不一致。")
    if list(loaded["feature_names"]) != list(feature_names):
        raise RuntimeError(f"{target_name} 模型包特征顺序回读不一致。")

    reloaded_predictions = np.asarray(loaded["model"].predict(x_scaled), dtype=float)
    max_diff = float(np.max(np.abs(reloaded_predictions - original_predictions)))
    if max_diff > ROUNDTRIP_PRED_TOL:
        raise RuntimeError(
            f"{target_name} joblib保存/回读后预测发生变化，最大绝对差={max_diff:.3e}。"
        )
    return sha256_file(output_file), max_diff


# =============================================================================
# 6. Excel与TXT统一输出
# =============================================================================

def style_workbook(path: Path) -> None:
    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        if ws.max_row >= 1 and ws.max_column >= 1:
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        sample_end = min(ws.max_row, 1000)
        for col_idx in range(1, ws.max_column + 1):
            max_len = 0
            for row_idx in range(1, sample_end + 1):
                value = ws.cell(row=row_idx, column=col_idx).value
                if value is not None:
                    max_len = max(max_len, len(str(value)))
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 10), 42)

        # 常用数值统一显示6位小数；日期列保持日期格式。
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = "0.000000"
                elif isinstance(cell.value, datetime):
                    cell.number_format = "yyyy-mm-dd"
    wb.save(path)


def write_excel_atomic(
    overview: pd.DataFrame,
    feature_table: pd.DataFrame,
    imputation_table: pd.DataFrame,
    fit_detail: pd.DataFrame,
) -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    temporary = OUTPUT_EXCEL.with_name(OUTPUT_EXCEL.stem + ".part.xlsx")
    if temporary.exists():
        temporary.unlink()

    with pd.ExcelWriter(temporary, engine="openpyxl") as writer:
        overview.to_excel(writer, sheet_name="模型总览", index=False)
        feature_table.to_excel(writer, sheet_name="特征与Scaler", index=False)
        imputation_table.to_excel(writer, sheet_name="预处理记录", index=False)
        fit_detail.to_excel(writer, sheet_name="全量拟合明细", index=False)

    style_workbook(temporary)
    os.replace(temporary, OUTPUT_EXCEL)


def build_text_report(
    overview: pd.DataFrame,
    feature_table: pd.DataFrame,
    imputation_table: pd.DataFrame,
    selection_hash: str,
    dynamic_hash: str,
) -> str:
    lines: List[str] = []
    lines.append("A²/O 四指标最终模型全量训练与固化说明")
    lines.append("=" * 106)
    lines.append("")

    lines.append("一、任务与文件")
    lines.append(f"1. 完成时间：{now_str()}")
    lines.append(f"2. 最终终选结果：{FINAL_SELECTION_FILE}")
    lines.append(f"   SHA256：{selection_hash}")
    lines.append(f"3. 动态可用数据：{DYNAMIC_INPUT_FILE}")
    lines.append(f"   SHA256：{dynamic_hash}")
    lines.append(f"4. 输出目录：{OUTPUT_DIR}")
    lines.append(f"5. 汇总Excel：{OUTPUT_EXCEL}")
    lines.append(f"6. 说明TXT：{OUTPUT_TEXT}")
    lines.append("7. 本脚本没有重新进行模型筛选、超参数搜索或Bootstrap/Time/Expanding复核。")
    lines.append("")

    lines.append("二、统一固化原则")
    lines.append("1. 四个目标均使用已经锁定的最终F2机器学习模型，特征数均为12。")
    lines.append("2. F2统一由9个当前原始输入 + 各目标自身lag1、lag2、roll3组成；Q不进入模型。")
    lines.append(
        f"3. 全量历史训练缺失处理：四个进水变量先做内部短缺口线性插值（原样本跨度≤{MAX_INTERPOLATION_GAP}），"
        "剩余缺失用该目标全量训练中位数兜底。"
    )
    lines.append("4. StandardScaler只在该目标全部历史训练样本上拟合，并与模型一起保存到joblib。")
    lines.append("5. 全量拟合R²/RMSE/MAE只用于固化核查；正式泛化证据仍为原终选表中的Bootstrap .632+、Time、Expanding。")
    lines.append("6. 新/未来样本如有进水缺失，不得使用未来观测做双向插值；应使用当时已知历史信息，默认兜底为模型包保存的训练中位数。")
    lines.append("")

    lines.append("三、四个最终模型固化结果")
    for target in TARGET_ORDER:
        row = overview.loc[overview["目标指标"] == target].iloc[0]
        lines.append("")
        lines.append(f"【{target}】")
        lines.append(
            f"1. 最终模型：{row['特征组']}-{row['算法']} | {row['候选参数ID']}；参数={row['参数JSON']}"
        )
        lines.append(
            f"2. 全量训练样本：n={int(row['样本数'])}；日期={row['开始日期']}–{row['结束日期']}。"
        )
        lines.append("3. 原终选泛化指标：")
        lines.append(
            f"   Bootstrap .632+：R²={row['Bootstrap_R2']:.6f}，RMSE={row['Bootstrap_RMSE']:.6f}，MAE={row['Bootstrap_MAE']:.6f}。"
        )
        lines.append(
            f"   Time：R²={row['Time_R2']:.6f}，RMSE={row['Time_RMSE']:.6f}，MAE={row['Time_MAE']:.6f}。"
        )
        lines.append(
            f"   Expanding：R²={row['Expanding_R2']:.6f}，RMSE={row['Expanding_RMSE']:.6f}，MAE={row['Expanding_MAE']:.6f}。"
        )
        lines.append(
            f"4. 最佳历史基线：Time={row['Time最佳历史基线']}（RMSE={row['Time基线_RMSE']:.6f}）；"
            f"Expanding={row['Expanding最佳历史基线']}（RMSE={row['Expanding基线_RMSE']:.6f}）。"
        )
        lines.append(f"   ML是否在两种时间验证中均优于最佳历史基线：{row['两种时间验证均优于历史基线']}。")
        lines.append("5. 全量拟合核查指标（不作为泛化证据）：")
        lines.append(
            f"   R²={row['全量拟合_R2']:.6f}，RMSE={row['全量拟合_RMSE']:.6f}，MAE={row['全量拟合_MAE']:.6f}。"
        )
        lines.append(f"6. 训练警告：{row['训练警告'] if str(row['训练警告']).strip() else '无'}。")
        lines.append(f"7. 模型包：{row['模型包文件']}")
        lines.append(f"   SHA256：{row['模型包SHA256']}")
        lines.append(f"   保存/回读预测最大绝对差：{row['回读预测最大绝对差']:.3e}。")
        lines.append(f"8. 终选结论：{row['最终预测结论']}")
        if target == "NH4+-N":
            lines.append(
                "9. NH4+-N特别说明：该F2-ElasticNet是最佳机器学习模型，但roll3历史基线在Time和Expanding中仍更强；"
                "后续可用于最佳ML的SHAP/PDP/运行响应解释，但论文必须保留这一性能边界。"
            )

    lines.append("")
    lines.append("四、12个F2输入和Scaler")
    for target in TARGET_ORDER:
        sub = feature_table.loc[feature_table["目标指标"] == target].sort_values("特征顺序")
        lines.append(f"- {target}：" + "、".join(sub["特征"].astype(str).tolist()) + "。")
    lines.append("- 每个特征对应的Scaler mean/scale已写入Excel“特征与Scaler”，并同时保存在joblib模型包中。")
    lines.append("")

    lines.append("五、缺失处理固化记录")
    for target in TARGET_ORDER:
        sub = imputation_table.loc[imputation_table["目标指标"] == target]
        pieces = []
        for _, r in sub.iterrows():
            pieces.append(
                f"{r['变量']}：原缺失{int(r['原缺失数'])}，线性插值{int(r['短内部线性插值数'])}，"
                f"中位数兜底{int(r['训练中位数兜底数'])}，训练中位数={float(r['保存的全量训练中位数']):.6f}"
            )
        lines.append(f"- {target}：" + "；".join(pieces) + "。")
    lines.append("")

    lines.append("六、后续使用规则")
    lines.append("1. SHAP + Permutation Importance：直接加载这里保存的四个joblib，不再重新拟合一套所谓“最终模型”。")
    lines.append("2. PDP/ICE、具体状态正向预测和NSGA-II同样读取同一joblib，确保解释、预测和优化使用完全同一模型。")
    lines.append("3. joblib中已经保存：模型对象、StandardScaler、特征顺序、四个训练中位数、最终参数、来源泛化指标、训练样本ID/日期和输入文件哈希。")
    lines.append("4. 使用模型时必须严格保持feature_names中的输入顺序；不能把F1/F3机理衍生变量加入F2模型。")
    lines.append("5. 不应把本次全量拟合R²写成未来时间泛化R²；论文性能结果继续引用统一终选的Bootstrap .632+、Time、Expanding。")
    lines.append("6. SHAP/PDP等结果描述模型关联关系，不直接作为因果效应或长期稳态机理结论。")
    lines.append("")

    lines.append("七、输出完整性")
    lines.append(f"1. 统一Excel包含4个工作表：模型总览、特征与Scaler、预处理记录、全量拟合明细。")
    lines.append("2. 四个joblib均已执行保存后回读预测一致性检查，并在模型总览中保存SHA256。")
    lines.append("3. 如果本TXT已正常生成，说明四个目标均完成模型训练、模型包保存、回读核查和统一汇总。")

    return "\n".join(lines)


# =============================================================================
# 7. 主流程
# =============================================================================

def main() -> None:
    start_time = now_str()
    log("A²/O 四指标最终模型全量训练与固化开始。")
    log(f"最终终选结果：{FINAL_SELECTION_FILE}")
    log(f"动态可用数据：{DYNAMIC_INPUT_FILE}")
    log(f"输出目录：{OUTPUT_DIR}")

    if not DYNAMIC_INPUT_FILE.exists():
        raise FileNotFoundError(f"动态可用数据不存在：{DYNAMIC_INPUT_FILE}")

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    final_ml, baselines, selection_hash = load_final_selection()
    dynamic_hash = sha256_file(DYNAMIC_INPUT_FILE)

    overview_rows: List[Dict[str, Any]] = []
    feature_rows: List[Dict[str, Any]] = []
    imputation_frames: List[pd.DataFrame] = []
    fit_detail_frames: List[pd.DataFrame] = []

    for target_name in TARGET_ORDER:
        log(f"开始固化 {target_name}。")
        selected_row = final_ml.loc[final_ml["目标指标"] == target_name].iloc[0]
        params = json.loads(str(selected_row["参数JSON"]))

        raw_data = read_dynamic_table(target_name)
        processed, feature_names, medians, imputation_table = prepare_full_training_data(
            raw_data, target_name
        )

        x_raw = processed[feature_names].astype(float).to_numpy()
        y = processed[TARGET_CONFIG[target_name]["target_col"]].astype(float).to_numpy()

        scaler = StandardScaler()
        x_scaled = scaler.fit_transform(x_raw)

        model = build_final_model(str(selected_row["模型"]), params)
        warning_text = fit_model_safely(model, x_scaled, y)
        predictions = np.asarray(model.predict(x_scaled), dtype=float)
        full_metrics = calc_metrics(y, predictions)

        source_metrics = build_source_validation_metrics(
            selected_row, baselines, target_name
        )

        package = make_package(
            target_name=target_name,
            selected_row=selected_row,
            processed_data=processed,
            feature_names=feature_names,
            medians=medians,
            scaler=scaler,
            model=model,
            source_metrics=source_metrics,
            full_fit_metrics=full_metrics,
            dynamic_hash=dynamic_hash,
            selection_hash=selection_hash,
        )

        model_file = OUTPUT_DIR / TARGET_CONFIG[target_name]["joblib_name"]
        atomic_joblib_dump(package, model_file)
        model_hash, roundtrip_diff = verify_saved_package(
            model_file,
            target_name,
            feature_names,
            x_scaled,
            predictions,
        )

        time_base = source_metrics["best_historical_baseline"]["time"]
        exp_base = source_metrics["best_historical_baseline"]["expanding"]

        overview_rows.append({
            "目标指标": target_name,
            "特征组": str(selected_row["特征组"]),
            "特征数": int(selected_row["特征数"]),
            "算法": str(selected_row["模型"]),
            "候选参数ID": str(selected_row["候选参数ID"]),
            "参数JSON": json.dumps(json_normalize(params), ensure_ascii=False, sort_keys=True),
            "样本数": int(len(processed)),
            "开始日期": processed["Date"].min().strftime("%Y-%m-%d"),
            "结束日期": processed["Date"].max().strftime("%Y-%m-%d"),
            "Bootstrap_R2": source_metrics["bootstrap_632plus"]["R2"],
            "Bootstrap_RMSE": source_metrics["bootstrap_632plus"]["RMSE"],
            "Bootstrap_MAE": source_metrics["bootstrap_632plus"]["MAE"],
            "Time_R2": source_metrics["time"]["R2"],
            "Time_RMSE": source_metrics["time"]["RMSE"],
            "Time_MAE": source_metrics["time"]["MAE"],
            "Expanding_R2": source_metrics["expanding"]["R2"],
            "Expanding_RMSE": source_metrics["expanding"]["RMSE"],
            "Expanding_MAE": source_metrics["expanding"]["MAE"],
            "Time最佳历史基线": time_base["name"],
            "Time基线_R2": time_base["R2"],
            "Time基线_RMSE": time_base["RMSE"],
            "Expanding最佳历史基线": exp_base["name"],
            "Expanding基线_R2": exp_base["R2"],
            "Expanding基线_RMSE": exp_base["RMSE"],
            "两种时间验证均优于历史基线": (
                "是" if source_metrics["beats_best_historical_baseline_in_both_time_validations"] else "否"
            ),
            "全量拟合_R2": full_metrics["R2"],
            "全量拟合_RMSE": full_metrics["RMSE"],
            "全量拟合_MAE": full_metrics["MAE"],
            "训练警告": warning_text,
            "模型包文件": str(model_file),
            "模型包SHA256": model_hash,
            "模型包大小字节": model_file.stat().st_size,
            "回读预测最大绝对差": roundtrip_diff,
            "最终预测结论": source_metrics["final_selection_conclusion"],
            "解释边界": package["interpretation_note"],
        })

        for idx, feature in enumerate(feature_names):
            feature_rows.append({
                "目标指标": target_name,
                "特征组": "F2",
                "特征顺序": idx + 1,
                "特征": feature,
                "特征类别": feature_category(feature, TARGET_CONFIG[target_name]["history"]),
                "Scaler_mean": float(scaler.mean_[idx]),
                "Scaler_scale": float(scaler.scale_[idx]),
                "是否进水缺失处理变量": "是" if feature in IMPUTE_COLUMNS else "否",
                "训练中位数": medians.get(feature, np.nan),
            })

        imputation_frames.append(imputation_table)

        detail = pd.DataFrame({
            "目标指标": target_name,
            "样本序号": processed["样本序号"].astype(int).values,
            "Date": processed["Date"].values,
            "实测值": y,
            "全量拟合预测值": predictions,
            "残差_实测减预测": y - predictions,
            "绝对误差": np.abs(y - predictions),
        })
        fit_detail_frames.append(detail)

        log(
            f"{target_name} 完成：{selected_row['特征组']}-{selected_row['模型']} "
            f"{selected_row['候选参数ID']}；全量拟合R²={full_metrics['R2']:.6f}；"
            f"joblib={model_file.name}"
        )

    overview = pd.DataFrame(overview_rows)
    feature_table = pd.DataFrame(feature_rows)
    imputation_all = pd.concat(imputation_frames, ignore_index=True)
    fit_detail_all = pd.concat(fit_detail_frames, ignore_index=True)

    # 最终完整性检查：四个模型均存在且SHA256可读取。
    if set(overview["目标指标"]) != set(TARGET_ORDER) or len(overview) != 4:
        raise RuntimeError("最终模型总览未完整包含四个目标。")
    for target_name in TARGET_ORDER:
        model_path = OUTPUT_DIR / TARGET_CONFIG[target_name]["joblib_name"]
        if not model_path.exists() or model_path.stat().st_size <= 0:
            raise RuntimeError(f"{target_name} joblib未成功生成：{model_path}")

    write_excel_atomic(overview, feature_table, imputation_all, fit_detail_all)
    text = build_text_report(
        overview,
        feature_table,
        imputation_all,
        selection_hash,
        dynamic_hash,
    )
    atomic_write_text(OUTPUT_TEXT, text)

    log("四指标最终模型固化全部完成。")
    print("\n输出结果：")
    for target_name in TARGET_ORDER:
        print(" -", OUTPUT_DIR / TARGET_CONFIG[target_name]["joblib_name"])
    print(" -", OUTPUT_EXCEL)
    print(" -", OUTPUT_TEXT)
    print(f"开始时间：{start_time}")
    print(f"完成时间：{now_str()}")


if __name__ == "__main__":
    main()