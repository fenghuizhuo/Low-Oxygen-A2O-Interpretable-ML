# -*- coding: utf-8 -*-
r"""
A²/O 四指标最终模型 SHAP + Permutation Importance 综合分析
=============================================================

目的
----
本脚本针对已经完成“最终模型统一终选 + 全量训练固化”的四个模型开展解释分析。
不重新选择模型、不重新调参，也不生成任何图片。

正式解释对象（严格核对）
------------------------
- TP       ：F2 + SVR         | SVR_032
- TN       ：F2 + SVR         | SVR_018
- NH4+-N   ：F2 + ElasticNet  | ElasticNet_025
- COD      ：F2 + ElasticNet  | ElasticNet_040

分析框架
--------
A. 最终固化模型 SHAP（解释“最终冻结模型如何使用输入信息”）
   - TP/TN（SVR）：model-agnostic Permutation SHAP，使用分层真实背景样本；
     使用多个固定随机种子重复计算后取样本级SHAP均值，并记录近似稳定性。
   - NH4+-N/COD（ElasticNet）：Independent-background Linear SHAP；
     同时用解析公式 (x - background_mean) * coef_ 做一致性核查。
   - 四个目标都输出解释样本逐特征SHAP长表（SVR为分层代表样本、ElasticNet为全部样本）、全局重要性、方向、三类特征组贡献、
     时间阶段汇总、高相关特征及相关簇。

B. Permutation Importance（解释“破坏某项信息后，样本外预测性能损失多少”）
   - 不在全量训练集上做PI作为主证据；而是复现原模型评价的：
     1) Time：前80%训练、后20%验证；
     2) Expanding：初始60%训练，验证窗10，步长10。
   - 每个时间切分仅使用已经锁定的同一算法和同一超参数重新拟合，绝不重新调参。
   - 训练内缺失处理、Scaler均严格只用训练段拟合；验证段不使用未来信息。
   - 对验证集做重复置换，报告 ΔRMSE、ΔMAE、R²下降及95%区间。
   - 正式PI先检查“可置换性”：Time验证段无变化则标记不可评估；Expanding逐窗口检查，
     只有存在组内变化的窗口才计算PI，窗口过少时标记有限可评估且不进入正式Top-k。
   - 零/负PI不进入Top-k；95%区间下限仍>0才标记为“稳健正重要性”。
   - 同时做：单特征PI、三大组PI、|Spearman rho|>=0.80相关簇PI。
   - 分组/相关簇置换时，组内特征使用同一行排列，尽量保留组内联合结构，
     只破坏该组与目标/其他特征的对应关系。

C. SHAP × PI 综合核查
   - 不建立人为加权总分。
   - SHAP零贡献不参与排名；PI只有在可评估且ΔRMSE>数值容差时才参与正式排名。
   - 并列报告 SHAP排名、Time PI排名、Expanding PI排名及一致性标签。
   - 高相关特征可能发生SHAP贡献分摊、单变量PI替代，必须结合“相关簇PI”解释。

重要解释边界
------------
1. SHAP解释模型关联贡献，不代表因果效应。
2. Permutation Importance是模型依赖性检验，不等同于变量的真实因果效应。
3. 最终SHAP基于全量固化模型；Time/Expanding PI基于固定超参数的样本外时间重拟合。
   这些时间重拟合仅用于解释稳定性/样本外变量依赖，不改变已固化最终joblib。
4. NH4+-N最佳ML仍未在Time和Expanding中同时超过roll3历史基线，解释时必须保留该边界。
5. 四个最终模型全部为F2，只解释12个F2输入；F1/F3机理衍生特征不进入本阶段最终模型解释。

默认输入
--------
模型目录：
${PROJECT_ROOT}/outputs/final_selection/moxingbaocun

动态数据：
${PROJECT_ROOT}/outputs/preprocessing/dynamic/动态可用数据与插补比较结果.xlsx

默认输出目录（按用户要求）
--------------------------
${PROJECT_ROOT}/outputs/interpretation/shap_pi

正式输出（仅两个文件，不生成图）
--------------------------------
1. 四指标_SHAP与PermutationImportance综合结果.xlsx
2. 四指标_SHAP与PermutationImportance综合说明.txt

依赖
----
pip install pandas numpy openpyxl xlsxwriter scikit-learn joblib shap

可选环境变量
------------
A2O_EXPLAIN_MODEL_DIR                 最终joblib模型所在目录
A2O_EXPLAIN_DYNAMIC_INPUT             动态数据Excel
A2O_EXPLAIN_OUTPUT_DIR                输出目录
A2O_EXPLAIN_PI_REPEATS=100            每个PI项目重复置换次数
A2O_EXPLAIN_SHAP_BACKGROUND_N=30      SHAP分层背景真实样本数
A2O_EXPLAIN_SVR_SHAP_REPEATS=2        SVR Permutation SHAP固定种子重复次数（近似稳定性核查）
A2O_EXPLAIN_SVR_SHAP_CYCLES=1         每次SVR SHAP的完整permutation cycle数；正式默认1
A2O_EXPLAIN_SVR_SHAP_EXPLAIN_N=300     SVR正式SHAP分层解释样本数；ElasticNet仍解释全部样本
A2O_EXPLAIN_HIGH_CORR=0.80            高相关阈值
A2O_EXPLAIN_SHAP_ZERO_TOL=1e-12       SHAP零贡献数值容差
A2O_EXPLAIN_PI_ZERO_TOL=1e-12         PI零变化数值容差
A2O_EXPLAIN_EXP_MIN_WINDOWS=3          Expanding至少3个可置换窗口才进入正式PI排名
A2O_EXPLAIN_RESUME=1                  是否使用检查点续跑
A2O_EXPLAIN_DELETE_CHECKPOINT=1       全部成功后是否删除检查点
A2O_EXPLAIN_STRICT=1                  严格核对四个锁定模型、哈希和时间验证
A2O_EXPLAIN_SMOKE_TEST=1              仅用于程序测试：减少SHAP样本和PI重复；正式运行不要开启
"""

from __future__ import annotations

import hashlib
import json
import math
import os
import pickle
import re
import shutil
import time
import traceback
import warnings
from copy import deepcopy
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Mapping, Optional, Sequence, Tuple

import joblib
import numpy as np
import pandas as pd
import shap
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sklearn.exceptions import ConvergenceWarning
from sklearn.linear_model import ElasticNet
from sklearn.metrics import mean_absolute_error, mean_squared_error, r2_score
from sklearn.preprocessing import StandardScaler
from sklearn.svm import SVR


# =============================================================================
# 1. 路径与统一参数
# =============================================================================

PROJECT_ROOT = Path(__file__).resolve().parents[1]

MODEL_DIR = Path(
    os.environ.get(
        "A2O_EXPLAIN_MODEL_DIR",
        str(PROJECT_ROOT / "outputs/final_selection/moxingbaocun"),
    )
)

MODEL_SUMMARY_EXCEL = MODEL_DIR / "四指标_最终模型固化汇总.xlsx"

DYNAMIC_INPUT_FILE = Path(
    os.environ.get(
        "A2O_EXPLAIN_DYNAMIC_INPUT",
        str(PROJECT_ROOT / "outputs/preprocessing/dynamic/动态可用数据与插补比较结果.xlsx"),
    )
)

OUTPUT_DIR = Path(
    os.environ.get(
        "A2O_EXPLAIN_OUTPUT_DIR",
        str(PROJECT_ROOT / "outputs/interpretation/shap_pi"),
    )
)
OUTPUT_EXCEL = OUTPUT_DIR / "四指标_SHAP与PermutationImportance综合结果.xlsx"
OUTPUT_TEXT = OUTPUT_DIR / "四指标_SHAP与PermutationImportance综合说明.txt"
CHECKPOINT_DIR = OUTPUT_DIR / "_checkpoint_SHAP_PI"

RANDOM_SEED = 20260620
EPS = 1e-12
TIME_TRAIN_RATIO = 0.80
EXPANDING_INITIAL_TRAIN_RATIO = 0.60
EXPANDING_VALID_WINDOW_SIZE = 10
EXPANDING_STEP_SIZE = 10
VALIDATION_METRIC_TOL = 1e-8
SHAP_ADDITIVITY_TOL = 1e-6
LINEAR_SHAP_TOL = 1e-8

PI_REPEATS = int(os.environ.get("A2O_EXPLAIN_PI_REPEATS", "100"))
SHAP_BACKGROUND_N = int(os.environ.get("A2O_EXPLAIN_SHAP_BACKGROUND_N", "30"))
SVR_SHAP_REPEATS = int(os.environ.get("A2O_EXPLAIN_SVR_SHAP_REPEATS", "2"))
SVR_SHAP_CYCLES = int(os.environ.get("A2O_EXPLAIN_SVR_SHAP_CYCLES", "1"))
SVR_SHAP_EXPLAIN_N = int(os.environ.get("A2O_EXPLAIN_SVR_SHAP_EXPLAIN_N", "300"))
HIGH_CORRELATION_THRESHOLD = float(os.environ.get("A2O_EXPLAIN_HIGH_CORR", "0.80"))
SHAP_ZERO_TOL = float(os.environ.get("A2O_EXPLAIN_SHAP_ZERO_TOL", "1e-12"))
PI_ZERO_TOL = float(os.environ.get("A2O_EXPLAIN_PI_ZERO_TOL", "1e-12"))
EXPANDING_MIN_EVALUABLE_WINDOWS = int(os.environ.get("A2O_EXPLAIN_EXP_MIN_WINDOWS", "3"))
STRICT_MODE = os.environ.get("A2O_EXPLAIN_STRICT", "1").strip() != "0"
RESUME_CHECKPOINT = os.environ.get("A2O_EXPLAIN_RESUME", "1").strip() != "0"
DELETE_CHECKPOINT_AFTER_SUCCESS = (
    os.environ.get("A2O_EXPLAIN_DELETE_CHECKPOINT", "1").strip() != "0"
)
SMOKE_TEST = os.environ.get("A2O_EXPLAIN_SMOKE_TEST", "0").strip() == "1"

if PI_REPEATS < 2:
    raise ValueError("A2O_EXPLAIN_PI_REPEATS至少应为2。")
if SHAP_BACKGROUND_N < 5:
    raise ValueError("A2O_EXPLAIN_SHAP_BACKGROUND_N至少应为5。")
if SVR_SHAP_REPEATS < 1:
    raise ValueError("A2O_EXPLAIN_SVR_SHAP_REPEATS至少应为1。")
if SVR_SHAP_CYCLES < 1:
    raise ValueError("A2O_EXPLAIN_SVR_SHAP_CYCLES至少应为1。")
if SVR_SHAP_EXPLAIN_N < 60:
    raise ValueError("A2O_EXPLAIN_SVR_SHAP_EXPLAIN_N至少应为60。")
if not (0.0 < HIGH_CORRELATION_THRESHOLD < 1.0):
    raise ValueError("A2O_EXPLAIN_HIGH_CORR必须位于(0,1)。")
if SHAP_ZERO_TOL < 0 or PI_ZERO_TOL < 0:
    raise ValueError("SHAP/PI零值容差不能为负数。")
if EXPANDING_MIN_EVALUABLE_WINDOWS < 1:
    raise ValueError("A2O_EXPLAIN_EXP_MIN_WINDOWS至少为1。")

# Smoke test只用于脚本结构测试，不得作为正式解释结果。
if SMOKE_TEST:
    PI_REPEATS = min(PI_REPEATS, 4)
    SHAP_BACKGROUND_N = min(SHAP_BACKGROUND_N, 10)
    SVR_SHAP_REPEATS = 1
    SVR_SHAP_CYCLES = 1

TARGET_ORDER = ["TP", "TN", "NH4+-N", "COD"]

RAW_FEATURES = [
    "Aerobic_DO",
    "Anaerobic_Feed_Fraction",
    "Inf.TP",
    "Inf.COD",
    "Inf.NH4+-N",
    "Inf.TN",
    "HRT",
    "Anaerobic_Volume_Fraction",
    "Anoxic_Volume_Fraction",
]

INFLUENT_FEATURES = ["Inf.TP", "Inf.COD", "Inf.NH4+-N", "Inf.TN"]
OPERATION_FEATURES = [
    "Aerobic_DO",
    "Anaerobic_Feed_Fraction",
    "HRT",
    "Anaerobic_Volume_Fraction",
    "Anoxic_Volume_Fraction",
]
IMPUTE_COLUMNS = list(INFLUENT_FEATURES)

TEXTUAL_MISSING_VALUES = {
    "", " ", "NA", "N/A", "na", "n/a", "NaN", "nan", "None", "none",
    "NULL", "null", "-", "—", "–", "－", "/", "\\", "未测", "缺失",
}

TARGET_CONFIG: Dict[str, Dict[str, Any]] = {
    "TP": {
        "sheet": "TP动态可用表",
        "target_col": "eff.TP",
        "history": ["eff.TP_lag1", "eff.TP_lag2", "eff.TP_roll3"],
        "joblib": "TP_最终模型.joblib",
    },
    "TN": {
        "sheet": "TN动态可用表",
        "target_col": "eff.TN",
        "history": ["eff.TN_lag1", "eff.TN_lag2", "eff.TN_roll3"],
        "joblib": "TN_最终模型.joblib",
    },
    "NH4+-N": {
        "sheet": "氨氮动态可用表",
        "target_col": "eff.NH4+-N",
        "history": ["eff.NH4+-N_lag1", "eff.NH4+-N_lag2", "eff.NH4+-N_roll3"],
        "joblib": "NH4_N_最终模型.joblib",
    },
    "COD": {
        "sheet": "COD动态可用表",
        "target_col": "eff.COD",
        "history": ["eff.COD_lag1", "eff.COD_lag2", "eff.COD_roll3"],
        "joblib": "COD_最终模型.joblib",
    },
}

# 与已经锁定的最终模型保持一致。STRICT_MODE=1时任何不一致都直接停止。
EXPECTED_FINAL_CONFIG: Dict[str, Dict[str, Any]] = {
    "TP": {
        "feature_group": "F2", "feature_count": 12,
        "algorithm": "SVR", "candidate_id": "SVR_032",
        "parameters": {"C": 10.0, "epsilon": 0.05, "gamma": 0.01, "kernel": "rbf"},
    },
    "TN": {
        "feature_group": "F2", "feature_count": 12,
        "algorithm": "SVR", "candidate_id": "SVR_018",
        "parameters": {"C": 1.0, "epsilon": 0.1, "gamma": 0.03, "kernel": "rbf"},
    },
    "NH4+-N": {
        "feature_group": "F2", "feature_count": 12,
        "algorithm": "ElasticNet", "candidate_id": "ElasticNet_025",
        "parameters": {"alpha": 0.01, "l1_ratio": 0.9},
    },
    "COD": {
        "feature_group": "F2", "feature_count": 12,
        "algorithm": "ElasticNet", "candidate_id": "ElasticNet_040",
        "parameters": {"alpha": 0.3, "l1_ratio": 0.9},
    },
}

GROUP_ORDER = ["运行与结构变量", "进水水质", "历史状态"]


# =============================================================================
# 2. 通用工具
# =============================================================================

def now_str() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def log(message: str) -> None:
    print(f"[{now_str()}] {message}", flush=True)


def clean_column_name(value: Any) -> str:
    text = str(value)
    text = text.replace("\ufeff", "").replace("\xa0", " ").replace("\u3000", " ")
    text = text.replace("\n", "").replace("\r", "").strip()
    return re.sub(r"\s+", "", text)


def convert_to_numeric_series(series: pd.Series) -> pd.Series:
    output = series.copy()
    if output.dtype == "object" or str(output.dtype).startswith("string"):
        output = (
            output.astype("string")
            .str.strip()
            .str.replace(",", "", regex=False)
            .str.replace("，", "", regex=False)
        )
        output = output.mask(output.isin(TEXTUAL_MISSING_VALUES), np.nan)
    return pd.to_numeric(output, errors="coerce")


def safe_float(value: Any) -> float:
    try:
        x = float(value)
    except Exception:
        return np.nan
    return x if np.isfinite(x) else np.nan

def zero_small(values: Any, tol: float) -> np.ndarray:
    """把绝对值不超过tol的浮点噪声严格归零，避免1e-17被误判成正重要性。"""
    arr = np.asarray(values, dtype=float).copy()
    arr[np.isfinite(arr) & (np.abs(arr) <= tol)] = 0.0
    return arr


def positive_importance_ranks(values: pd.Series, tol: float, eligible: Optional[pd.Series] = None) -> pd.Series:
    """只给真正大于零容差且满足可评估条件的项目排名；零/负/不可评估均保持NaN。"""
    numeric = pd.to_numeric(values, errors="coerce")
    mask = numeric.notna() & np.isfinite(numeric) & (numeric > tol)
    if eligible is not None:
        mask &= eligible.fillna(False).astype(bool)
    ranks = pd.Series(np.nan, index=values.index, dtype=float)
    if mask.any():
        ranks.loc[mask] = numeric.loc[mask].rank(method="min", ascending=False).astype(float)
    return ranks


def joint_unique_pattern_count(frame: pd.DataFrame, members: Sequence[str]) -> int:
    """返回验证段中一组变量的联合唯一模式数；单变量时就是唯一值数。"""
    cols = [str(c) for c in members]
    if not cols:
        return 0
    return int(frame[cols].drop_duplicates().shape[0])


def rank_text(value: Any, unavailable_text: str = "无正式排名") -> str:
    x = safe_float(value)
    return f"#{int(x)}" if np.isfinite(x) else unavailable_text



def sha256_file(path: Path, block_size: int = 1024 * 1024) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while True:
            block = handle.read(block_size)
            if not block:
                break
            digest.update(block)
    return digest.hexdigest()


def stable_seed(text: str) -> int:
    digest = hashlib.md5(str(text).encode("utf-8")).hexdigest()
    return int(digest[:8], 16) % (2**31 - 1)


def parameters_equal(actual: Mapping[str, Any], expected: Mapping[str, Any]) -> bool:
    if set(actual.keys()) != set(expected.keys()):
        return False
    for key in expected:
        a, e = actual[key], expected[key]
        if isinstance(e, (int, float)) and not isinstance(e, bool):
            try:
                if not np.isclose(float(a), float(e), rtol=1e-12, atol=1e-12):
                    return False
            except Exception:
                return False
        elif str(a) != str(e):
            return False
    return True


def calc_metrics(y_true: Sequence[float], y_pred: Sequence[float]) -> Dict[str, float]:
    y = np.asarray(y_true, dtype=float)
    p = np.asarray(y_pred, dtype=float)
    mask = np.isfinite(y) & np.isfinite(p)
    y = y[mask]
    p = p[mask]
    if len(y) == 0:
        return {"n": 0, "R2": np.nan, "RMSE": np.nan, "MAE": np.nan, "SSE": np.nan}
    errors = y - p
    sse = float(np.sum(errors**2))
    rmse = float(np.sqrt(np.mean(errors**2)))
    mae = float(np.mean(np.abs(errors)))
    r2 = float(r2_score(y, p)) if len(y) >= 2 and np.std(y) > 0 else np.nan
    return {"n": int(len(y)), "R2": r2, "RMSE": rmse, "MAE": mae, "SSE": sse}


def quantile_ci(values: Sequence[float], low: float = 0.025, high: float = 0.975) -> Tuple[float, float]:
    arr = np.asarray(values, dtype=float)
    arr = arr[np.isfinite(arr)]
    if len(arr) == 0:
        return np.nan, np.nan
    return float(np.quantile(arr, low)), float(np.quantile(arr, high))


def sanitize_excel_value(value: Any) -> Any:
    if isinstance(value, str):
        return re.sub(r"[\x00-\x08\x0B\x0C\x0E-\x1F]", "", value)
    return value


def sanitize_dataframe(frame: pd.DataFrame) -> pd.DataFrame:
    out = frame.copy()
    for col in out.select_dtypes(include=["object", "string"]).columns:
        out[col] = out[col].map(sanitize_excel_value)
    return out


def atomic_write_text(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + ".tmp")
    tmp.write_text(text, encoding="utf-8-sig")
    os.replace(tmp, path)


def atomic_pickle_dump(obj: Any, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + ".tmp")
    with tmp.open("wb") as handle:
        pickle.dump(obj, handle, protocol=pickle.HIGHEST_PROTOCOL)
    os.replace(tmp, path)


def write_excel(path: Path, tables: Mapping[str, pd.DataFrame]) -> None:
    """
    优先使用XlsxWriter写大型长表，避免openpyxl在PI逐次明细/SHAP长表上耗时过长。
    若本机未安装xlsxwriter，则退回openpyxl（功能相同但写入可能较慢）。
    """
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(path.stem + ".part.xlsx")
    if tmp.exists():
        tmp.unlink()

    try:
        import xlsxwriter  # noqa: F401
        engine = "xlsxwriter"
    except Exception:
        engine = "openpyxl"

    if engine == "xlsxwriter":
        with pd.ExcelWriter(
            tmp,
            engine="xlsxwriter",
            date_format="yyyy-mm-dd",
            datetime_format="yyyy-mm-dd",
            engine_kwargs={"options": {"constant_memory": False}},
        ) as writer:
            workbook = writer.book
            header_format = workbook.add_format({
                "bold": True, "bg_color": "#D9EAF7", "border": 1,
                "align": "center", "valign": "vcenter", "text_wrap": True,
            })
            number_format = workbook.add_format({"num_format": "0.000000"})
            integer_format = workbook.add_format({"num_format": "0"})
            date_format = workbook.add_format({"num_format": "yyyy-mm-dd"})

            used: set[str] = set()
            for original_name, frame in tables.items():
                if frame is None or frame.empty:
                    continue
                name = original_name[:31]
                counter = 1
                while name in used:
                    suffix = f"_{counter}"
                    name = original_name[: 31 - len(suffix)] + suffix
                    counter += 1
                used.add(name)

                clean = sanitize_dataframe(frame)
                clean.to_excel(writer, sheet_name=name, index=False)
                ws = writer.sheets[name]
                ws.freeze_panes(1, 0)
                ws.autofilter(0, 0, len(clean), max(len(clean.columns) - 1, 0))
                ws.set_row(0, 30, header_format)

                sample = clean.head(800)
                for col_idx, col in enumerate(clean.columns):
                    header = str(col)
                    width = len(header)
                    if len(sample):
                        try:
                            width = max(width, int(sample[col].astype(str).str.len().max()))
                        except Exception:
                            pass
                    width = min(max(width + 2, 10), 42)
                    series = clean[col]
                    if pd.api.types.is_datetime64_any_dtype(series):
                        fmt = date_format
                    elif pd.api.types.is_integer_dtype(series):
                        fmt = integer_format
                    elif pd.api.types.is_numeric_dtype(series):
                        fmt = number_format
                    else:
                        fmt = None
                    ws.set_column(col_idx, col_idx, width, fmt)
    else:
        with pd.ExcelWriter(tmp, engine="openpyxl") as writer:
            used: set[str] = set()
            for original_name, frame in tables.items():
                if frame is None or frame.empty:
                    continue
                name = original_name[:31]
                counter = 1
                while name in used:
                    suffix = f"_{counter}"
                    name = original_name[: 31 - len(suffix)] + suffix
                    counter += 1
                used.add(name)
                sanitize_dataframe(frame).to_excel(writer, sheet_name=name, index=False)

        wb = load_workbook(tmp)
        header_fill = PatternFill("solid", fgColor="D9EAF7")
        for ws in wb.worksheets:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            sample_end = min(ws.max_row, 500)
            for idx in range(1, ws.max_column + 1):
                width = 10
                for r in range(1, sample_end + 1):
                    value = ws.cell(r, idx).value
                    if value is not None:
                        width = max(width, min(len(str(value)) + 2, 42))
                ws.column_dimensions[get_column_letter(idx)].width = min(width, 42)
        wb.save(tmp)

    os.replace(tmp, path)


def concat_frames(frames: Iterable[pd.DataFrame]) -> pd.DataFrame:
    valid = [f for f in frames if f is not None and not f.empty]
    return pd.concat(valid, ignore_index=True) if valid else pd.DataFrame()


# =============================================================================
# 3. 模型包与数据核查
# =============================================================================

def model_path(target_name: str) -> Path:
    return MODEL_DIR / TARGET_CONFIG[target_name]["joblib"]


def load_model_summary() -> pd.DataFrame:
    if not MODEL_SUMMARY_EXCEL.exists():
        raise FileNotFoundError(f"未找到最终模型固化汇总：{MODEL_SUMMARY_EXCEL}")
    excel = pd.ExcelFile(MODEL_SUMMARY_EXCEL, engine="openpyxl")
    if "模型总览" not in excel.sheet_names:
        raise ValueError(f"{MODEL_SUMMARY_EXCEL} 缺少“模型总览”工作表。")
    summary = pd.read_excel(MODEL_SUMMARY_EXCEL, sheet_name="模型总览", engine="openpyxl")
    required = ["目标指标", "特征组", "特征数", "算法", "候选参数ID", "模型包SHA256"]
    missing = [c for c in required if c not in summary.columns]
    if missing:
        raise ValueError(f"模型总览缺少字段：{missing}")
    return summary


def load_package(target_name: str, summary: pd.DataFrame) -> Tuple[Path, Dict[str, Any], pd.Series]:
    path = model_path(target_name)
    if not path.exists():
        raise FileNotFoundError(f"未找到 {target_name} 最终模型包：{path}")
    package = joblib.load(path)
    if not isinstance(package, dict):
        raise TypeError(f"{target_name} joblib不是字典模型包。")

    required = {
        "package_type", "package_version", "target_name", "target_column", "input_sheet",
        "selected_feature_group", "feature_count", "feature_names", "algorithm", "candidate_id",
        "parameters", "preprocessing", "scaler", "model", "source_validation_metrics",
        "full_fit_metrics", "training_data", "interpretation_note",
    }
    missing = sorted(required - set(package.keys()))
    if missing:
        raise ValueError(f"{target_name} 模型包缺少字段：{missing}")

    rows = summary.loc[summary["目标指标"].astype(str) == target_name]
    if len(rows) != 1:
        raise ValueError(f"模型总览中 {target_name} 应恰好1行，实际{len(rows)}行。")
    srow = rows.iloc[0]

    actual_hash = sha256_file(path)
    recorded_hash = str(srow["模型包SHA256"]).strip().lower()
    if actual_hash.lower() != recorded_hash:
        raise ValueError(
            f"{target_name} joblib SHA256与固化汇总不一致：\n"
            f"实际={actual_hash}\n记录={recorded_hash}"
        )

    cfg = EXPECTED_FINAL_CONFIG[target_name]
    actual_config = (
        str(package["selected_feature_group"]),
        int(package["feature_count"]),
        str(package["algorithm"]),
        str(package["candidate_id"]),
    )
    expected_config = (
        cfg["feature_group"], cfg["feature_count"], cfg["algorithm"], cfg["candidate_id"]
    )
    if STRICT_MODE and actual_config != expected_config:
        raise ValueError(f"{target_name} 模型配置不一致：实际={actual_config}，预期={expected_config}")
    if STRICT_MODE and not parameters_equal(package["parameters"], cfg["parameters"]):
        raise ValueError(
            f"{target_name} 最终参数不一致：实际={package['parameters']}，预期={cfg['parameters']}"
        )

    feature_names = list(package["feature_names"])
    expected_features = RAW_FEATURES + TARGET_CONFIG[target_name]["history"]
    if feature_names != expected_features:
        raise ValueError(
            f"{target_name} F2特征顺序不一致。\n实际={feature_names}\n预期={expected_features}"
        )
    if len(feature_names) != 12 or "Q" in feature_names:
        raise ValueError(f"{target_name} 最终F2特征数量/Q检查失败。")
    return path, package, srow


def read_dynamic_table(target_name: str, package: Mapping[str, Any]) -> pd.DataFrame:
    if not DYNAMIC_INPUT_FILE.exists():
        raise FileNotFoundError(f"未找到动态数据：{DYNAMIC_INPUT_FILE}")

    actual_hash = sha256_file(DYNAMIC_INPUT_FILE)
    expected_hash = str(package["training_data"].get("dynamic_input_file_sha256", "")).strip()
    if expected_hash and actual_hash.lower() != expected_hash.lower():
        raise ValueError(
            f"{target_name} 当前动态数据SHA256与模型训练记录不一致。\n"
            f"当前={actual_hash}\n模型记录={expected_hash}"
        )

    sheet = str(package["input_sheet"])
    excel = pd.ExcelFile(DYNAMIC_INPUT_FILE, engine="openpyxl")
    if sheet not in excel.sheet_names:
        raise ValueError(f"动态数据缺少工作表：{sheet}")
    data = pd.read_excel(
        DYNAMIC_INPUT_FILE,
        sheet_name=sheet,
        engine="openpyxl",
        na_values=list(TEXTUAL_MISSING_VALUES),
        keep_default_na=True,
    )
    data.columns = [clean_column_name(c) for c in data.columns]

    required = ["样本序号", "Date", str(package["target_column"])] + list(package["feature_names"])
    required = list(dict.fromkeys(required))
    missing = [c for c in required if c not in data.columns]
    if missing:
        raise ValueError(f"{target_name} 动态表缺少字段：{missing}")

    data["Date"] = pd.to_datetime(data["Date"], errors="coerce")
    for col in [c for c in required if c != "Date"]:
        data[col] = convert_to_numeric_series(data[col])
    data = data.sort_values(["样本序号", "Date"], kind="mergesort").reset_index(drop=True)
    data["_row_id"] = np.arange(len(data), dtype=int)

    if data["样本序号"].isna().any() or data["Date"].isna().any():
        raise ValueError(f"{target_name} 存在无效样本序号或日期。")
    if data["样本序号"].duplicated().any():
        raise ValueError(f"{target_name} 存在重复样本序号。")

    expected_ids = [int(x) for x in package["training_data"]["sample_ids"]]
    actual_ids = data["样本序号"].astype(int).tolist()
    if actual_ids != expected_ids:
        raise ValueError(f"{target_name} 当前动态表样本序号与模型包训练样本不完全一致。")
    if len(data) != int(package["training_data"]["sample_count"]):
        raise ValueError(f"{target_name} 当前样本数与模型包不一致。")
    return data


# =============================================================================
# 4. 与正式建模一致的预处理
# =============================================================================

def interpolate_short_internal_gaps(
    series: pd.Series,
    sequence_coordinates: pd.Series,
    max_gap: int,
) -> pd.Series:
    output = convert_to_numeric_series(series).astype(float).copy()
    coordinates = convert_to_numeric_series(sequence_coordinates).astype(float).to_numpy()
    values = output.to_numpy(copy=True)
    n = len(values)
    position = 0
    while position < n:
        if np.isfinite(values[position]):
            position += 1
            continue
        start = position
        while position < n and not np.isfinite(values[position]):
            position += 1
        end = position - 1
        left = start - 1
        right = end + 1
        if (
            left >= 0 and right < n
            and np.isfinite(values[left]) and np.isfinite(values[right])
            and np.isfinite(coordinates[left]) and np.isfinite(coordinates[right])
        ):
            original_gap_span = coordinates[right] - coordinates[left] - 1.0
            coordinate_span = coordinates[right] - coordinates[left]
            if 0 <= original_gap_span <= max_gap and coordinate_span > 0:
                for idx in range(start, end + 1):
                    frac = (coordinates[idx] - coordinates[left]) / coordinate_span
                    values[idx] = values[left] + frac * (values[right] - values[left])
    return pd.Series(values, index=output.index, name=output.name)


def rebuild_full_feature_data(
    target_name: str,
    raw_data: pd.DataFrame,
    package: Mapping[str, Any],
) -> Tuple[pd.DataFrame, np.ndarray, np.ndarray, np.ndarray, pd.DataFrame]:
    data = raw_data.copy(deep=True)
    saved_medians = {
        str(k): float(v) for k, v in package["preprocessing"]["training_medians"].items()
    }
    max_gap = int(package["preprocessing"].get("max_interpolation_gap", 7))
    records: List[Dict[str, Any]] = []

    for col in package["preprocessing"]["impute_columns"]:
        original = convert_to_numeric_series(data[col]).astype(float)
        current_median = float(original.median(skipna=True))
        saved = float(saved_medians[col])
        if not np.isclose(current_median, saved, rtol=1e-10, atol=1e-10):
            raise ValueError(
                f"{target_name} {col} 当前中位数与模型包保存值不一致：当前={current_median}，保存={saved}"
            )
        interpolated = interpolate_short_internal_gaps(original, data["样本序号"], max_gap)
        linear_filled = int(original.isna().sum() - interpolated.isna().sum())
        median_filled = int(interpolated.isna().sum())
        data[col] = interpolated.fillna(saved).astype(float)
        records.append({
            "目标指标": target_name,
            "变量": col,
            "原缺失数": int(original.isna().sum()),
            "短内部线性插值数": linear_filled,
            "中位数兜底数": median_filled,
            "模型包训练中位数": saved,
            "处理后缺失数": int(data[col].isna().sum()),
        })

    features = list(package["feature_names"])
    remaining = {f: int(data[f].isna().sum()) for f in features if data[f].isna().any()}
    if remaining:
        raise ValueError(f"{target_name} 全量解释特征处理后仍有缺失：{remaining}")

    x_raw = data[features].astype(float).to_numpy()
    x_scaled = package["scaler"].transform(x_raw)
    predictions = np.asarray(package["model"].predict(x_scaled), dtype=float)
    y = data[str(package["target_column"])].astype(float).to_numpy()

    metrics = calc_metrics(y, predictions)
    stored = package["full_fit_metrics"]
    for key in ["R2", "RMSE", "MAE"]:
        if not np.isclose(metrics[key], float(stored[key]), rtol=1e-9, atol=1e-9):
            raise RuntimeError(
                f"{target_name} 重建后的全量拟合{key}与模型包不一致：当前={metrics[key]}，保存={stored[key]}"
            )
    return data, x_raw, x_scaled, predictions, pd.DataFrame(records)


def prepare_temporal_split(
    train_raw: pd.DataFrame,
    valid_raw: pd.DataFrame,
    package: Mapping[str, Any],
) -> Dict[str, Any]:
    train = train_raw.copy(deep=True)
    valid = valid_raw.copy(deep=True)
    max_gap = int(package["preprocessing"].get("max_interpolation_gap", 7))

    medians: Dict[str, float] = {}
    for col in package["preprocessing"]["impute_columns"]:
        original = convert_to_numeric_series(train[col]).astype(float)
        median = float(original.median(skipna=True))
        if not np.isfinite(median):
            raise ValueError(f"时间训练段 {col} 全部缺失。")
        train[col] = interpolate_short_internal_gaps(
            original, train["样本序号"], max_gap
        ).fillna(median)
        valid[col] = convert_to_numeric_series(valid[col]).astype(float).fillna(median)
        medians[col] = median

    features = list(package["feature_names"])
    for feature in features:
        if train[feature].isna().any() or valid[feature].isna().any():
            raise ValueError(f"时间切分后特征 {feature} 仍有缺失。")

    scaler = StandardScaler()
    x_train = scaler.fit_transform(train[features].astype(float).to_numpy())
    x_valid = scaler.transform(valid[features].astype(float).to_numpy())
    y_train = train[str(package["target_column"])].astype(float).to_numpy()
    y_valid = valid[str(package["target_column"])].astype(float).to_numpy()
    return {
        "X_train": x_train,
        "X_valid": x_valid,
        "y_train": y_train,
        "y_valid": y_valid,
        "valid_ids": valid["样本序号"].astype(int).to_numpy(),
        "valid_dates": valid["Date"].to_numpy(),
        "valid_feature_frame": valid[features].astype(float).reset_index(drop=True),
        "feature_names": features,
        "scaler": scaler,
        "medians": medians,
    }


def build_locked_model(package: Mapping[str, Any]):
    params = deepcopy(dict(package["parameters"]))
    algorithm = str(package["algorithm"])
    if algorithm == "SVR":
        return SVR(
            kernel=str(params["kernel"]),
            C=float(params["C"]),
            epsilon=float(params["epsilon"]),
            gamma=params["gamma"],
        )
    if algorithm == "ElasticNet":
        return ElasticNet(
            alpha=float(params["alpha"]),
            l1_ratio=float(params["l1_ratio"]),
            max_iter=50000,
            random_state=RANDOM_SEED,
        )
    raise ValueError(f"当前解释脚本只支持锁定的SVR/ElasticNet，实际={algorithm}")


def fit_locked_model(model: Any, x: np.ndarray, y: np.ndarray) -> str:
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        warnings.filterwarnings("always", category=ConvergenceWarning)
        model.fit(x, y)
    return " | ".join(str(w.message)[:500] for w in caught[:10])


# =============================================================================
# 5. 特征语义与背景抽样
# =============================================================================

def feature_group(feature: str, package: Mapping[str, Any]) -> str:
    if feature in package.get("history_features_in_model", []):
        return "历史状态"
    if feature in INFLUENT_FEATURES:
        return "进水水质"
    return "运行与结构变量"


def semantic_feature(target_name: str, feature: str) -> str:
    if feature.endswith("_lag1"):
        return "目标自身_lag1"
    if feature.endswith("_lag2"):
        return "目标自身_lag2"
    if feature.endswith("_roll3"):
        return "目标自身_roll3"
    return feature


def assign_time_stage(data: pd.DataFrame) -> pd.Series:
    labels = np.empty(len(data), dtype=object)
    for name, idx in zip(["前期", "中期", "后期"], np.array_split(np.arange(len(data)), 3)):
        labels[idx] = name
    return pd.Series(labels, index=data.index, name="时间阶段")


def stratified_sample_indices(
    data: pd.DataFrame,
    target_column: str,
    n_samples: int,
    seed: int,
) -> np.ndarray:
    if n_samples >= len(data):
        return np.arange(len(data), dtype=int)
    rng = np.random.default_rng(seed)
    frame = pd.DataFrame(index=data.index)
    frame["time"] = assign_time_stage(data).astype(str)
    try:
        frame["level"] = pd.qcut(
            data[target_column].rank(method="first"), q=3, labels=["低", "中", "高"]
        ).astype(str)
    except Exception:
        frame["level"] = "全部"
    frame["stratum"] = frame["time"] + "_" + frame["level"]
    groups = [g.index.to_numpy(dtype=int) for _, g in frame.groupby("stratum", sort=True)]

    selected: List[int] = []
    cursor = 0
    while len(selected) < n_samples and groups:
        group = groups[cursor % len(groups)]
        available = np.setdiff1d(group, np.asarray(selected, dtype=int), assume_unique=False)
        if available.size:
            selected.append(int(rng.choice(available)))
        cursor += 1
        if cursor > n_samples * max(len(groups), 1) * 20:
            break
    if len(selected) < n_samples:
        remaining = np.setdiff1d(np.arange(len(data)), np.asarray(selected, dtype=int))
        extra = rng.choice(remaining, size=n_samples - len(selected), replace=False)
        selected.extend(int(i) for i in extra)
    return np.asarray(sorted(selected), dtype=int)


# =============================================================================
# 6. SHAP计算
# =============================================================================

def extract_explanation_values(explanation: Any, n_rows: int, n_features: int) -> Tuple[np.ndarray, np.ndarray]:
    values = np.asarray(explanation.values, dtype=float)
    if values.ndim == 3 and values.shape[-1] == 1:
        values = values[:, :, 0]
    if values.shape != (n_rows, n_features):
        raise ValueError(f"SHAP矩阵形状异常：{values.shape}，预期={(n_rows, n_features)}")
    base = np.asarray(explanation.base_values, dtype=float)
    if base.ndim > 1:
        base = base.reshape(n_rows, -1)[:, 0]
    else:
        base = base.reshape(-1)
    if base.size == 1:
        base = np.repeat(base[0], n_rows)
    if base.size != n_rows:
        raise ValueError(f"SHAP base_values长度异常：{base.size}，预期={n_rows}")
    return values, base


def compute_svr_permutation_shap(
    target_name: str,
    package: Mapping[str, Any],
    x_explain: np.ndarray,
    background: np.ndarray,
) -> Tuple[np.ndarray, np.ndarray, pd.DataFrame, str]:
    background = np.asarray(background, dtype=float)
    x_explain = np.asarray(x_explain, dtype=float)
    n_rows, n_features = x_explain.shape
    max_evals = SVR_SHAP_CYCLES * (2 * n_features + 1)
    values_per_repeat: List[np.ndarray] = []
    base_per_repeat: List[np.ndarray] = []
    stability_rows: List[Dict[str, Any]] = []

    # 预热numba/SHAP路径，减少第一次正式重复的编译时间影响。
    warm_n = min(1, n_rows)
    if warm_n:
        try:
            warm_masker = shap.maskers.Independent(background[: min(5, len(background))], max_samples=min(5, len(background)))
            warm = shap.PermutationExplainer(package["model"].predict, warm_masker, seed=RANDOM_SEED)
            warm(x_explain[:warm_n], max_evals=2 * n_features + 1, silent=True)
        except Exception:
            pass

    for rep in range(1, SVR_SHAP_REPEATS + 1):
        seed = int((RANDOM_SEED + stable_seed(f"{target_name}-SVRSHAP-{rep}")) % (2**31 - 1))
        start = time.time()
        masker = shap.maskers.Independent(background, max_samples=len(background))
        try:
            explainer = shap.PermutationExplainer(
                package["model"].predict,
                masker,
                seed=seed,
            )
        except TypeError:
            np.random.seed(seed)
            explainer = shap.PermutationExplainer(package["model"].predict, masker)

        explanation = explainer(
            x_explain,
            max_evals=max_evals,
            batch_size=256,
            silent=True,
        )
        phi, base = extract_explanation_values(explanation, n_rows, n_features)
        prediction = np.asarray(package["model"].predict(x_explain), dtype=float)
        add_err = np.max(np.abs(base + phi.sum(axis=1) - prediction))
        if add_err > SHAP_ADDITIVITY_TOL:
            raise RuntimeError(
                f"{target_name} SVR Permutation SHAP第{rep}次加和误差超限：{add_err:.3e}"
            )
        values_per_repeat.append(phi)
        base_per_repeat.append(base)
        mean_abs = zero_small(np.mean(np.abs(phi), axis=0), SHAP_ZERO_TOL)
        rank_series = positive_importance_ranks(pd.Series(mean_abs), SHAP_ZERO_TOL)
        ranks = rank_series.to_numpy()
        for j, feature in enumerate(package["feature_names"]):
            stability_rows.append({
                "目标指标": target_name,
                "SHAP重复": rep,
                "随机种子": seed,
                "特征": feature,
                "平均绝对SHAP": float(mean_abs[j]),
                "是否有效SHAP贡献": "是" if mean_abs[j] > SHAP_ZERO_TOL else "否",
                "重要性排名": int(ranks[j]) if np.isfinite(ranks[j]) else np.nan,
                "SHAP加和最大绝对误差": float(add_err),
                "运行耗时秒": float(time.time() - start),
            })
        log(f"{target_name} SVR Permutation SHAP：{rep}/{SVR_SHAP_REPEATS} 完成。")

    phi_mean = np.mean(np.stack(values_per_repeat, axis=0), axis=0)
    base_mean = np.mean(np.stack(base_per_repeat, axis=0), axis=0)
    method = (
        f"PermutationSHAP_IndependentBackground_n{len(background)}_"
        f"repeats{SVR_SHAP_REPEATS}_cycles{SVR_SHAP_CYCLES}"
    )
    return phi_mean, base_mean, pd.DataFrame(stability_rows), method


def compute_linear_shap(
    target_name: str,
    package: Mapping[str, Any],
    x_explain: np.ndarray,
    background: np.ndarray,
) -> Tuple[np.ndarray, np.ndarray, pd.DataFrame, str, float]:
    background = np.asarray(background, dtype=float)
    x_explain = np.asarray(x_explain, dtype=float)
    model = package["model"]
    coefficients = np.asarray(model.coef_, dtype=float)
    background_mean = np.mean(background, axis=0)
    analytic = (x_explain - background_mean) * coefficients
    analytic_base = float(model.intercept_ + np.dot(background_mean, coefficients))
    base = np.repeat(analytic_base, len(x_explain))

    method = "LinearSHAP_IndependentBackground_analytic"
    max_difference = 0.0
    try:
        masker = shap.maskers.Independent(background, max_samples=len(background))
        explainer = shap.LinearExplainer(model, masker=masker)
        explanation = explainer(x_explain)
        phi_lib, base_lib = extract_explanation_values(
            explanation, len(x_explain), x_explain.shape[1]
        )
        max_difference = float(max(
            np.max(np.abs(phi_lib - analytic)),
            np.max(np.abs(base_lib - base)),
        ))
        if max_difference > LINEAR_SHAP_TOL:
            raise RuntimeError(
                f"{target_name} LinearSHAP与解析公式不一致，最大差={max_difference:.3e}"
            )
        phi = phi_lib
        base = base_lib
        method = "LinearSHAP_IndependentBackground_library+analytic_check"
    except Exception as exc:
        # 解析式本身就是Independent背景下线性模型的精确SHAP；
        # 仅当shap版本接口问题时退回解析式，不掩盖数值不一致错误。
        if isinstance(exc, RuntimeError):
            raise
        phi = analytic
        method = "LinearSHAP_IndependentBackground_analytic_fallback"

    prediction = np.asarray(model.predict(x_explain), dtype=float)
    add_err = float(np.max(np.abs(base + phi.sum(axis=1) - prediction)))
    if add_err > SHAP_ADDITIVITY_TOL:
        raise RuntimeError(f"{target_name} Linear SHAP加和误差超限：{add_err:.3e}")

    mean_abs = zero_small(np.mean(np.abs(phi), axis=0), SHAP_ZERO_TOL)
    ranks = positive_importance_ranks(pd.Series(mean_abs), SHAP_ZERO_TOL).to_numpy()
    stability = pd.DataFrame([
        {
            "目标指标": target_name,
            "SHAP重复": 1,
            "随机种子": np.nan,
            "特征": feature,
            "平均绝对SHAP": float(mean_abs[j]),
            "是否有效SHAP贡献": "是" if mean_abs[j] > SHAP_ZERO_TOL else "否",
            "重要性排名": int(ranks[j]) if np.isfinite(ranks[j]) else np.nan,
            "SHAP加和最大绝对误差": add_err,
            "运行耗时秒": 0.0,
        }
        for j, feature in enumerate(package["feature_names"])
    ])
    return phi, base, stability, method, max_difference


def direction_metrics(raw_values: np.ndarray, shap_values: np.ndarray) -> Dict[str, Any]:
    x = np.asarray(raw_values, dtype=float)
    phi = np.asarray(shap_values, dtype=float)
    valid = np.isfinite(x) & np.isfinite(phi)
    x = x[valid]
    phi = phi[valid]
    if len(x) < 5 or np.std(x) <= EPS:
        return {
            "特征值_SHAP_Spearman": np.nan,
            "低四分位SHAP中位数": np.nan,
            "高四分位SHAP中位数": np.nan,
            "高低四分位SHAP差": np.nan,
            "方向结论": "特征变化不足或样本过少",
        }
    rho = float(pd.Series(x).corr(pd.Series(phi), method="spearman")) if np.std(phi) > EPS else np.nan
    q25, q75 = np.quantile(x, [0.25, 0.75])
    low = phi[x <= q25]
    high = phi[x >= q75]
    low_med = float(np.median(low)) if len(low) else np.nan
    high_med = float(np.median(high)) if len(high) else np.nan
    diff = high_med - low_med
    weak = max(EPS, float(np.mean(np.abs(phi))) * 0.05)
    if abs(diff) <= weak:
        conclusion = "方向弱或不明确"
    elif np.isfinite(rho) and rho >= 0.15 and diff > 0:
        conclusion = "高值倾向提高模型预测值"
    elif np.isfinite(rho) and rho <= -0.15 and diff < 0:
        conclusion = "高值倾向降低模型预测值"
    else:
        conclusion = "可能非单调，或受交互/相关特征影响"
    return {
        "特征值_SHAP_Spearman": rho,
        "低四分位SHAP中位数": low_med,
        "高四分位SHAP中位数": high_med,
        "高低四分位SHAP差": diff,
        "方向结论": conclusion,
    }


def summarize_shap(
    target_name: str,
    package: Mapping[str, Any],
    feature_data: pd.DataFrame,
    x_scaled: np.ndarray,
    predictions: np.ndarray,
    phi: np.ndarray,
    base_values: np.ndarray,
    stability_detail: pd.DataFrame,
) -> Tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    features = list(package["feature_names"])
    target_col = str(package["target_column"])
    mean_abs = zero_small(np.mean(np.abs(phi), axis=0), SHAP_ZERO_TOL)
    total = float(mean_abs.sum())
    ranks = positive_importance_ranks(pd.Series(mean_abs, index=features), SHAP_ZERO_TOL)
    reconstructed = base_values + phi.sum(axis=1)
    add_error = reconstructed - predictions
    time_stage = assign_time_stage(feature_data)

    global_rows: List[Dict[str, Any]] = []
    long_parts: List[pd.DataFrame] = []
    for j, feature in enumerate(features):
        values = phi[:, j]
        raw = feature_data[feature].astype(float).to_numpy()
        direction = direction_metrics(raw, values)
        rep = stability_detail.loc[stability_detail["特征"] == feature].copy()
        rep_ranks = pd.to_numeric(rep["重要性排名"], errors="coerce").dropna()
        rep_importance = pd.to_numeric(rep["平均绝对SHAP"], errors="coerce").dropna()
        global_rows.append({
            "目标指标": target_name,
            "最终特征组": package["selected_feature_group"],
            "算法": package["algorithm"],
            "候选参数ID": package["candidate_id"],
            "特征": feature,
            "统一语义特征": semantic_feature(target_name, feature),
            "特征类别": feature_group(feature, package),
            "平均绝对SHAP": float(mean_abs[j]),
            "中位绝对SHAP": float(np.median(np.abs(values))),
            "SHAP标准差_样本间": float(np.std(values, ddof=1)),
            "平均有符号SHAP": float(np.mean(values)),
            "正SHAP样本比例百分比": float(np.mean(values > 0) * 100.0),
            "负SHAP样本比例百分比": float(np.mean(values < 0) * 100.0),
            "重要性占比百分比": float(mean_abs[j] / total * 100.0) if total > 0 else np.nan,
            "是否有效SHAP贡献": "是" if mean_abs[j] > SHAP_ZERO_TOL else "否",
            "SHAP重要性排名": int(ranks.loc[feature]) if np.isfinite(ranks.loc[feature]) else np.nan,
            "SHAP重复次数": int(rep["SHAP重复"].nunique()) if not rep.empty else 0,
            "重复平均绝对SHAP标准差": float(rep_importance.std(ddof=1)) if len(rep_importance) > 1 else 0.0,
            "重复排名中位数": float(rep_ranks.median()) if len(rep_ranks) else np.nan,
            "重复排名IQR": float(rep_ranks.quantile(.75) - rep_ranks.quantile(.25)) if len(rep_ranks) > 1 else 0.0,
            "Top5重复出现频率百分比": float(np.mean(rep_ranks <= 5) * 100.0) if len(rep_ranks) else np.nan,
            **direction,
        })
        long_parts.append(pd.DataFrame({
            "目标指标": target_name,
            "样本序号": feature_data["样本序号"].astype(int).to_numpy(),
            "Date": feature_data["Date"].to_numpy(),
            "时间阶段": time_stage.to_numpy(),
            "特征": feature,
            "统一语义特征": semantic_feature(target_name, feature),
            "特征类别": feature_group(feature, package),
            "原始特征值": raw,
            "标准化特征值": x_scaled[:, j],
            "SHAP值": values,
            "SHAP绝对值": np.abs(values),
            "实测值": feature_data[target_col].astype(float).to_numpy(),
            "模型预测值": predictions,
            "SHAP基准值": base_values,
            "SHAP重构预测值": reconstructed,
            "SHAP加和误差": add_error,
        }))

    global_df = pd.DataFrame(global_rows).sort_values(
        ["目标指标", "SHAP重要性排名"]
    ).reset_index(drop=True)
    long_df = pd.concat(long_parts, ignore_index=True)

    group_rows: List[Dict[str, Any]] = []
    for group_name in GROUP_ORDER:
        idx = [j for j, f in enumerate(features) if feature_group(f, package) == group_name]
        if not idx:
            continue
        abs_sum = np.abs(phi[:, idx]).sum(axis=1)
        signed_sum = phi[:, idx].sum(axis=1)
        importance = float(np.mean(abs_sum))
        group_rows.append({
            "目标指标": target_name,
            "特征组": group_name,
            "特征数": len(idx),
            "组内特征": "、".join(features[j] for j in idx),
            "组内绝对SHAP总和均值": importance,
            "组内绝对SHAP总和中位数": float(np.median(abs_sum)),
            "组净SHAP均值": float(np.mean(signed_sum)),
            "组净SHAP标准差": float(np.std(signed_sum, ddof=1)),
            "绝对贡献占比百分比": float(importance / total * 100.0) if total > 0 else np.nan,
        })
    group_df = pd.DataFrame(group_rows)

    stage_rows: List[Dict[str, Any]] = []
    stage_values = time_stage.to_numpy()
    for stage in ["前期", "中期", "后期"]:
        mask = stage_values == stage
        for j, feature in enumerate(features):
            vals = phi[mask, j]
            stage_rows.append({
                "目标指标": target_name,
                "时间阶段": stage,
                "样本数": int(mask.sum()),
                "特征": feature,
                "特征类别": feature_group(feature, package),
                "平均绝对SHAP": float(np.mean(np.abs(vals))),
                "平均有符号SHAP": float(np.mean(vals)),
            })
    stage_df = pd.DataFrame(stage_rows)
    if not stage_df.empty:
        stage_df["平均绝对SHAP"] = zero_small(stage_df["平均绝对SHAP"].to_numpy(), SHAP_ZERO_TOL)
        stage_df["阶段内重要性排名"] = np.nan
        for _, idx in stage_df.groupby(["目标指标", "时间阶段"], sort=False).groups.items():
            sub = stage_df.loc[idx]
            stage_df.loc[idx, "阶段内重要性排名"] = positive_importance_ranks(
                sub["平均绝对SHAP"], SHAP_ZERO_TOL
            )

    additivity_df = pd.DataFrame([{
        "目标指标": target_name,
        "样本数": len(feature_data),
        "SHAP基准值均值": float(np.mean(base_values)),
        "SHAP基准值标准差": float(np.std(base_values, ddof=1)) if len(base_values) > 1 else 0.0,
        "加和误差平均绝对值": float(np.mean(np.abs(add_error))),
        "加和误差95%分位数": float(np.quantile(np.abs(add_error), .95)),
        "加和误差最大绝对值": float(np.max(np.abs(add_error))),
        "允许最大绝对误差": SHAP_ADDITIVITY_TOL,
        "是否通过": "是" if np.max(np.abs(add_error)) <= SHAP_ADDITIVITY_TOL else "否",
    }])
    return global_df, long_df, group_df, stage_df, additivity_df


# =============================================================================
# 7. 相关性与相关簇
# =============================================================================

def build_correlation_tables(
    target_name: str,
    package: Mapping[str, Any],
    feature_data: pd.DataFrame,
) -> Tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, Dict[str, int]]:
    features = list(package["feature_names"])
    corr = feature_data[features].astype(float).corr(method="spearman")
    long_rows: List[Dict[str, Any]] = []
    high_rows: List[Dict[str, Any]] = []

    adjacency: Dict[str, set[str]] = {f: set() for f in features}
    for i in range(len(features)):
        for j in range(i + 1, len(features)):
            a, b = features[i], features[j]
            rho = safe_float(corr.loc[a, b])
            row = {
                "目标指标": target_name,
                "特征1": a,
                "特征2": b,
                "Spearman_rho": rho,
                "绝对rho": abs(rho) if np.isfinite(rho) else np.nan,
                "是否高相关": "是" if np.isfinite(rho) and abs(rho) >= HIGH_CORRELATION_THRESHOLD else "否",
            }
            long_rows.append(row)
            if row["是否高相关"] == "是":
                high_rows.append(row.copy())
                adjacency[a].add(b)
                adjacency[b].add(a)

    # 基于高相关边的连通分量构造相关簇；单特征也保留簇号，PI仅对>=2成员的簇计算。
    visited: set[str] = set()
    mapping: Dict[str, int] = {}
    cluster_rows: List[Dict[str, Any]] = []
    cluster_id = 0
    for feature in features:
        if feature in visited:
            continue
        cluster_id += 1
        stack = [feature]
        members: List[str] = []
        while stack:
            current = stack.pop()
            if current in visited:
                continue
            visited.add(current)
            members.append(current)
            stack.extend(sorted(adjacency[current] - visited))
        members = [f for f in features if f in members]
        for member in members:
            mapping[member] = cluster_id
        cluster_rows.append({
            "目标指标": target_name,
            "相关簇编号": cluster_id,
            "簇内特征数": len(members),
            "簇内特征": "、".join(members),
            "是否多特征高相关簇": "是" if len(members) >= 2 else "否",
        })

    return (
        pd.DataFrame(long_rows),
        pd.DataFrame(high_rows),
        pd.DataFrame(cluster_rows),
        mapping,
    )


# =============================================================================
# 8. Time / Expanding 时间模型复现与基线核查
# =============================================================================

def prepare_temporal_models(
    target_name: str,
    raw_data: pd.DataFrame,
    package: Mapping[str, Any],
) -> Tuple[Dict[str, Any], List[Dict[str, Any]], pd.DataFrame]:
    warnings_rows: List[Dict[str, Any]] = []

    split_idx = int(math.floor(len(raw_data) * TIME_TRAIN_RATIO))
    time_split = prepare_temporal_split(
        raw_data.iloc[:split_idx].copy(), raw_data.iloc[split_idx:].copy(), package
    )
    time_model = build_locked_model(package)
    time_warning = fit_locked_model(time_model, time_split["X_train"], time_split["y_train"])
    time_pred = np.asarray(time_model.predict(time_split["X_valid"]), dtype=float)
    time_metrics = calc_metrics(time_split["y_valid"], time_pred)
    time_obj = {
        **time_split,
        "model": time_model,
        "baseline_pred": time_pred,
        "baseline_metrics": time_metrics,
        "warning": time_warning,
    }
    if time_warning:
        warnings_rows.append({"目标指标": target_name, "阶段": "Time重拟合", "警告或错误": time_warning})

    windows: List[Dict[str, Any]] = []
    initial = int(math.floor(len(raw_data) * EXPANDING_INITIAL_TRAIN_RATIO))
    all_true: List[float] = []
    all_pred: List[float] = []
    for window_id, start in enumerate(
        range(initial, len(raw_data), EXPANDING_STEP_SIZE), start=1
    ):
        end = min(start + EXPANDING_VALID_WINDOW_SIZE, len(raw_data))
        if end - start < 2:
            continue
        prepared = prepare_temporal_split(
            raw_data.iloc[:start].copy(), raw_data.iloc[start:end].copy(), package
        )
        model = build_locked_model(package)
        warning_text = fit_locked_model(model, prepared["X_train"], prepared["y_train"])
        pred = np.asarray(model.predict(prepared["X_valid"]), dtype=float)
        metrics = calc_metrics(prepared["y_valid"], pred)
        obj = {
            **prepared,
            "model": model,
            "baseline_pred": pred,
            "baseline_metrics": metrics,
            "warning": warning_text,
            "window_id": window_id,
            "train_n": start,
            "valid_n": end - start,
            "valid_start_date": pd.to_datetime(raw_data.iloc[start]["Date"]),
            "valid_end_date": pd.to_datetime(raw_data.iloc[end - 1]["Date"]),
        }
        windows.append(obj)
        all_true.extend(prepared["y_valid"].tolist())
        all_pred.extend(pred.tolist())
        if warning_text:
            warnings_rows.append({
                "目标指标": target_name,
                "阶段": f"Expanding窗口{window_id}重拟合",
                "警告或错误": warning_text,
            })

    expanding_metrics = calc_metrics(all_true, all_pred)

    # 与终选阶段来源指标做严格复现核查。
    source_time = package["source_validation_metrics"]["time"]
    source_exp = package["source_validation_metrics"]["expanding"]
    check_rows = []
    for validation, current, source in [
        ("Time", time_metrics, source_time),
        ("Expanding", expanding_metrics, source_exp),
    ]:
        row = {
            "目标指标": target_name,
            "验证方式": validation,
            "当前复现_R2": current["R2"],
            "来源_R2": safe_float(source["R2"]),
            "R2绝对差": abs(current["R2"] - safe_float(source["R2"])),
            "当前复现_RMSE": current["RMSE"],
            "来源_RMSE": safe_float(source["RMSE"]),
            "RMSE绝对差": abs(current["RMSE"] - safe_float(source["RMSE"])),
            "当前复现_MAE": current["MAE"],
            "来源_MAE": safe_float(source["MAE"]),
            "MAE绝对差": abs(current["MAE"] - safe_float(source["MAE"])),
        }
        max_diff = max(row["R2绝对差"], row["RMSE绝对差"], row["MAE绝对差"])
        row["允许绝对差"] = VALIDATION_METRIC_TOL
        row["是否通过"] = "是" if max_diff <= VALIDATION_METRIC_TOL else "否"
        check_rows.append(row)
        if STRICT_MODE and row["是否通过"] != "是":
            raise RuntimeError(
                f"{target_name} {validation}时间验证无法复现原终选指标，最大差={max_diff:.3e}"
            )

    return time_obj, windows, pd.DataFrame(check_rows)


# =============================================================================
# 9. Permutation Importance
# =============================================================================

def build_permutation_items(
    package: Mapping[str, Any],
    cluster_mapping: Mapping[str, int],
) -> List[Dict[str, Any]]:
    features = list(package["feature_names"])
    items: List[Dict[str, Any]] = []

    # 单特征
    for j, feature in enumerate(features):
        items.append({
            "项目类型": "单特征",
            "项目名称": feature,
            "项目语义": semantic_feature(str(package["target_name"]), feature),
            "成员特征": feature,
            "成员数": 1,
            "member_list": [feature],
            "indices": [j],
        })

    # 三类主分组
    for group_name in GROUP_ORDER:
        members = [f for f in features if feature_group(f, package) == group_name]
        if members:
            items.append({
                "项目类型": "主分组",
                "项目名称": group_name,
                "项目语义": group_name,
                "成员特征": "、".join(members),
                "成员数": len(members),
                "member_list": list(members),
                "indices": [features.index(f) for f in members],
            })

    # 高相关多特征簇
    clusters: Dict[int, List[str]] = {}
    for f, cid in cluster_mapping.items():
        clusters.setdefault(int(cid), []).append(f)
    for cid in sorted(clusters):
        members = [f for f in features if f in clusters[cid]]
        if len(members) >= 2:
            items.append({
                "项目类型": "相关簇",
                "项目名称": f"相关簇{cid}",
                "项目语义": f"相关簇{cid}",
                "成员特征": "、".join(members),
                "成员数": len(members),
                "member_list": list(members),
                "indices": [features.index(f) for f in members],
            })
    return items


def generate_permutations(n: int, repeats: int, seed: int) -> np.ndarray:
    rng = np.random.default_rng(seed)
    return np.vstack([rng.permutation(n) for _ in range(repeats)])


def vectorized_permuted_predictions(
    model: Any,
    x_valid: np.ndarray,
    indices: Sequence[int],
    permutations: np.ndarray,
) -> np.ndarray:
    repeats, n = permutations.shape
    p = x_valid.shape[1]
    stacked = np.tile(x_valid, (repeats, 1))
    for r in range(repeats):
        rows = slice(r * n, (r + 1) * n)
        perm = permutations[r]
        # 同一排列应用于组内全部特征，保持组内联合模式。
        stacked[rows, list(indices)] = x_valid[perm][:, list(indices)]
    pred = np.asarray(model.predict(stacked), dtype=float)
    return pred.reshape(repeats, n)


def repeated_metrics_from_predictions(
    y: np.ndarray,
    predictions: np.ndarray,
) -> Tuple[np.ndarray, np.ndarray, np.ndarray]:
    # predictions shape=(repeats, n)
    errors = y[None, :] - predictions
    mse = np.mean(errors**2, axis=1)
    rmse = np.sqrt(mse)
    mae = np.mean(np.abs(errors), axis=1)
    sst = float(np.sum((y - np.mean(y))**2))
    sse = np.sum(errors**2, axis=1)
    r2 = 1.0 - sse / sst if sst > 0 else np.full(len(rmse), np.nan)
    return rmse, mae, r2


def summarize_pi_repeats(
    target_name: str,
    validation: str,
    item: Mapping[str, Any],
    baseline: Mapping[str, float],
    rmse: np.ndarray,
    mae: np.ndarray,
    r2: np.ndarray,
) -> Tuple[pd.DataFrame, Dict[str, Any]]:
    # 浮点噪声先按统一容差归零，避免1e-17之类数值被误判成“正重要性”。
    delta_rmse = zero_small(rmse - float(baseline["RMSE"]), PI_ZERO_TOL)
    delta_mae = zero_small(mae - float(baseline["MAE"]), PI_ZERO_TOL)
    r2_drop = zero_small(float(baseline["R2"]) - r2, PI_ZERO_TOL)
    detail = pd.DataFrame({
        "目标指标": target_name,
        "验证方式": validation,
        "项目类型": item["项目类型"],
        "项目名称": item["项目名称"],
        "成员特征": item["成员特征"],
        "成员数": item["成员数"],
        "置换重复": np.arange(1, len(rmse) + 1),
        "基线_RMSE": baseline["RMSE"],
        "置换后_RMSE": rmse,
        "RMSE增加": delta_rmse,
        "基线_MAE": baseline["MAE"],
        "置换后_MAE": mae,
        "MAE增加": delta_mae,
        "基线_R2": baseline["R2"],
        "置换后_R2": r2,
        "R2下降": r2_drop,
    })
    rmse_lo, rmse_hi = quantile_ci(delta_rmse)
    mae_lo, mae_hi = quantile_ci(delta_mae)
    r2_lo, r2_hi = quantile_ci(r2_drop)
    summary = {
        "目标指标": target_name,
        "验证方式": validation,
        "项目类型": item["项目类型"],
        "项目名称": item["项目名称"],
        "项目语义": item["项目语义"],
        "成员特征": item["成员特征"],
        "成员数": item["成员数"],
        "置换重复次数": len(rmse),
        "基线_R2": baseline["R2"],
        "基线_RMSE": baseline["RMSE"],
        "基线_MAE": baseline["MAE"],
        "RMSE增加均值": float(np.mean(delta_rmse)),
        "RMSE增加标准差": float(np.std(delta_rmse, ddof=1)) if len(delta_rmse) > 1 else 0.0,
        "RMSE增加2.5%分位": rmse_lo,
        "RMSE增加97.5%分位": rmse_hi,
        "RMSE增加为正比例百分比": float(np.mean(delta_rmse > PI_ZERO_TOL) * 100.0),
        "MAE增加均值": float(np.mean(delta_mae)),
        "MAE增加2.5%分位": mae_lo,
        "MAE增加97.5%分位": mae_hi,
        "R2下降均值": float(np.mean(r2_drop)),
        "R2下降2.5%分位": r2_lo,
        "R2下降97.5%分位": r2_hi,
    }
    return detail, summary


def unevaluable_pi_summary(
    target_name: str,
    validation: str,
    item: Mapping[str, Any],
    baseline: Mapping[str, float],
    evaluability: str,
    unique_patterns: float = np.nan,
    evaluable_windows: float = np.nan,
    total_windows: float = np.nan,
) -> Dict[str, Any]:
    return {
        "目标指标": target_name,
        "验证方式": validation,
        "项目类型": item["项目类型"],
        "项目名称": item["项目名称"],
        "项目语义": item["项目语义"],
        "成员特征": item["成员特征"],
        "成员数": item["成员数"],
        "置换重复次数": 0,
        "基线_R2": baseline.get("R2", np.nan),
        "基线_RMSE": baseline.get("RMSE", np.nan),
        "基线_MAE": baseline.get("MAE", np.nan),
        "RMSE增加均值": np.nan,
        "RMSE增加标准差": np.nan,
        "RMSE增加2.5%分位": np.nan,
        "RMSE增加97.5%分位": np.nan,
        "RMSE增加为正比例百分比": np.nan,
        "MAE增加均值": np.nan,
        "MAE增加2.5%分位": np.nan,
        "MAE增加97.5%分位": np.nan,
        "R2下降均值": np.nan,
        "R2下降2.5%分位": np.nan,
        "R2下降97.5%分位": np.nan,
        "验证段联合唯一模式数": unique_patterns,
        "可置换窗口数": evaluable_windows,
        "总窗口数": total_windows,
        "可置换窗口比例百分比": (
            float(evaluable_windows / total_windows * 100.0)
            if np.isfinite(evaluable_windows) and np.isfinite(total_windows) and total_windows > 0
            else np.nan
        ),
        "PI可评估性": evaluability,
    }


def run_time_permutation_importance(
    target_name: str,
    time_obj: Mapping[str, Any],
    items: Sequence[Mapping[str, Any]],
) -> Tuple[pd.DataFrame, pd.DataFrame]:
    detail_parts: List[pd.DataFrame] = []
    summaries: List[Dict[str, Any]] = []
    x_valid = np.asarray(time_obj["X_valid"], dtype=float)
    y = np.asarray(time_obj["y_valid"], dtype=float)
    baseline = time_obj["baseline_metrics"]
    model = time_obj["model"]
    valid_frame = time_obj["valid_feature_frame"]

    for item in items:
        unique_patterns = joint_unique_pattern_count(valid_frame, item["member_list"])
        if unique_patterns < 2:
            summaries.append(unevaluable_pi_summary(
                target_name, "Time", item, baseline,
                evaluability="不可评估：Time验证段无可置换变化",
                unique_patterns=float(unique_patterns),
            ))
            continue

        seed = int((RANDOM_SEED + stable_seed(
            f"{target_name}-TimePI-{item['项目类型']}-{item['项目名称']}"
        )) % (2**31 - 1))
        perms = generate_permutations(len(y), PI_REPEATS, seed)
        pred = vectorized_permuted_predictions(model, x_valid, item["indices"], perms)
        rmse, mae, r2 = repeated_metrics_from_predictions(y, pred)
        detail, summary = summarize_pi_repeats(
            target_name, "Time", item, baseline, rmse, mae, r2
        )
        detail["随机种子"] = seed
        detail["验证段联合唯一模式数"] = unique_patterns
        detail["PI可评估性"] = "可评估"
        summary.update({
            "验证段联合唯一模式数": unique_patterns,
            "可置换窗口数": np.nan,
            "总窗口数": np.nan,
            "可置换窗口比例百分比": np.nan,
            "PI可评估性": "可评估",
        })
        detail_parts.append(detail)
        summaries.append(summary)
    return concat_frames(detail_parts), pd.DataFrame(summaries)


def run_expanding_permutation_importance(
    target_name: str,
    windows: Sequence[Mapping[str, Any]],
    items: Sequence[Mapping[str, Any]],
) -> Tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    if not windows:
        raise RuntimeError(f"{target_name} 没有Expanding窗口。")

    # 所有项目共享同一个完整Expanding样本外基线，保证不同特征/特征组的ΔRMSE可横向比较。
    all_y = np.concatenate([np.asarray(w["y_valid"], dtype=float) for w in windows])
    all_baseline_pred = np.concatenate([np.asarray(w["baseline_pred"], dtype=float) for w in windows])
    pooled_baseline = calc_metrics(all_y, all_baseline_pred)
    n_total = len(all_y)
    mean_y = float(np.mean(all_y))
    sst_total = float(np.sum((all_y - mean_y) ** 2))

    detail_parts: List[pd.DataFrame] = []
    summaries: List[Dict[str, Any]] = []
    window_summary_rows: List[Dict[str, Any]] = []
    total_windows = len(windows)

    for item in items:
        pattern_counts: Dict[int, int] = {
            int(window["window_id"]): joint_unique_pattern_count(
                window["valid_feature_frame"], item["member_list"]
            )
            for window in windows
        }
        n_eval = int(sum(count >= 2 for count in pattern_counts.values()))
        eval_ratio = float(n_eval / total_windows * 100.0)

        if n_eval == 0:
            # 全部窗口恒定时不能置换，因此不能把“0”当作无重要性证据。
            for window in windows:
                window_summary_rows.append({
                    "目标指标": target_name,
                    "窗口编号": window["window_id"],
                    "训练样本数": window["train_n"],
                    "验证样本数": window["valid_n"],
                    "验证开始日期": window["valid_start_date"],
                    "验证结束日期": window["valid_end_date"],
                    "项目类型": item["项目类型"],
                    "项目名称": item["项目名称"],
                    "成员特征": item["成员特征"],
                    "验证段联合唯一模式数": pattern_counts[int(window["window_id"])],
                    "PI可评估性": "不可评估：该窗口无可置换变化",
                    "窗口基线_RMSE": window["baseline_metrics"]["RMSE"],
                    "RMSE增加均值": np.nan,
                    "RMSE增加2.5%分位": np.nan,
                    "RMSE增加97.5%分位": np.nan,
                    "RMSE增加为正比例百分比": np.nan,
                    "MAE增加均值": np.nan,
                    "R2下降均值": np.nan,
                })
            summaries.append(unevaluable_pi_summary(
                target_name, "Expanding", item, pooled_baseline,
                evaluability="不可评估：全部Expanding窗口无可置换变化",
                evaluable_windows=0.0,
                total_windows=float(total_windows),
            ))
            continue

        # 完整Expanding池：有变化的窗口做置换；无变化窗口保持原基线预测。
        # 这样不会把“不可评估窗口”当作变量无效，同时不同项目仍共享同一总体基线和样本权重。
        pooled_sse = np.zeros(PI_REPEATS, dtype=float)
        pooled_sae = np.zeros(PI_REPEATS, dtype=float)

        for window in windows:
            x = np.asarray(window["X_valid"], dtype=float)
            y = np.asarray(window["y_valid"], dtype=float)
            base = window["baseline_metrics"]
            pattern_count = pattern_counts[int(window["window_id"])]

            if pattern_count < 2:
                base_pred = np.asarray(window["baseline_pred"], dtype=float)
                base_errors = y - base_pred
                pooled_sse += float(np.sum(base_errors ** 2))
                pooled_sae += float(np.sum(np.abs(base_errors)))
                window_summary_rows.append({
                    "目标指标": target_name,
                    "窗口编号": window["window_id"],
                    "训练样本数": window["train_n"],
                    "验证样本数": window["valid_n"],
                    "验证开始日期": window["valid_start_date"],
                    "验证结束日期": window["valid_end_date"],
                    "项目类型": item["项目类型"],
                    "项目名称": item["项目名称"],
                    "成员特征": item["成员特征"],
                    "验证段联合唯一模式数": pattern_count,
                    "PI可评估性": "不可评估：该窗口无可置换变化",
                    "窗口基线_RMSE": base["RMSE"],
                    "RMSE增加均值": np.nan,
                    "RMSE增加2.5%分位": np.nan,
                    "RMSE增加97.5%分位": np.nan,
                    "RMSE增加为正比例百分比": np.nan,
                    "MAE增加均值": np.nan,
                    "R2下降均值": np.nan,
                })
                continue

            seed = int((
                RANDOM_SEED
                + stable_seed(
                    f"{target_name}-ExpPI-W{window['window_id']}-{item['项目类型']}-{item['项目名称']}"
                )
            ) % (2**31 - 1))
            perms = generate_permutations(len(y), PI_REPEATS, seed)
            pred = vectorized_permuted_predictions(window["model"], x, item["indices"], perms)
            errors = y[None, :] - pred
            sse = np.sum(errors ** 2, axis=1)
            sae = np.sum(np.abs(errors), axis=1)
            pooled_sse += sse
            pooled_sae += sae

            window_rmse = np.sqrt(sse / len(y))
            window_mae = sae / len(y)
            window_sst = float(np.sum((y - np.mean(y)) ** 2))
            window_r2 = 1.0 - sse / window_sst if window_sst > 0 else np.full(PI_REPEATS, np.nan)
            delta = zero_small(window_rmse - float(base["RMSE"]), PI_ZERO_TOL)
            lo, hi = quantile_ci(delta)
            window_summary_rows.append({
                "目标指标": target_name,
                "窗口编号": window["window_id"],
                "训练样本数": window["train_n"],
                "验证样本数": window["valid_n"],
                "验证开始日期": window["valid_start_date"],
                "验证结束日期": window["valid_end_date"],
                "项目类型": item["项目类型"],
                "项目名称": item["项目名称"],
                "成员特征": item["成员特征"],
                "验证段联合唯一模式数": pattern_count,
                "PI可评估性": "可评估",
                "窗口基线_RMSE": base["RMSE"],
                "RMSE增加均值": float(np.mean(delta)),
                "RMSE增加2.5%分位": lo,
                "RMSE增加97.5%分位": hi,
                "RMSE增加为正比例百分比": float(np.mean(delta > PI_ZERO_TOL) * 100.0),
                "MAE增加均值": float(np.mean(zero_small(window_mae - float(base["MAE"]), PI_ZERO_TOL))),
                "R2下降均值": float(np.nanmean(zero_small(float(base["R2"]) - window_r2, PI_ZERO_TOL))) if np.any(np.isfinite(window_r2)) else np.nan,
            })

        pooled_rmse = np.sqrt(pooled_sse / n_total)
        pooled_mae = pooled_sae / n_total
        pooled_r2 = 1.0 - pooled_sse / sst_total if sst_total > 0 else np.full(PI_REPEATS, np.nan)
        detail, summary = summarize_pi_repeats(
            target_name, "Expanding", item, pooled_baseline,
            pooled_rmse, pooled_mae, pooled_r2
        )
        evaluability = (
            "可评估"
            if n_eval >= EXPANDING_MIN_EVALUABLE_WINDOWS
            else f"有限可评估：仅{n_eval}/{total_windows}个窗口有变化"
        )
        detail["可置换窗口数"] = n_eval
        detail["总窗口数"] = total_windows
        detail["可置换窗口比例百分比"] = eval_ratio
        detail["PI可评估性"] = evaluability
        summary.update({
            "验证段联合唯一模式数": np.nan,
            "可置换窗口数": n_eval,
            "总窗口数": total_windows,
            "可置换窗口比例百分比": eval_ratio,
            "PI可评估性": evaluability,
        })
        detail_parts.append(detail)
        summaries.append(summary)

    return concat_frames(detail_parts), pd.DataFrame(summaries), pd.DataFrame(window_summary_rows)


def add_pi_ranks(summary: pd.DataFrame) -> pd.DataFrame:
    out = summary.copy()
    if out.empty:
        return out

    # 只有“可评估”且平均ΔRMSE真正大于零容差的项目进入正式重要性排名。
    # “有限可评估”保留数值供参考，但不进入正式Top-k，避免仅1个阶段窗口支撑的变量被误称为稳定Top5。
    out["是否正PI信号"] = np.where(
        pd.to_numeric(out["RMSE增加均值"], errors="coerce") > PI_ZERO_TOL,
        "是", "否"
    )
    out.loc[out["RMSE增加均值"].isna(), "是否正PI信号"] = "不可判断"
    out["是否稳健正PI"] = np.where(
        (out["PI可评估性"] == "可评估")
        & (pd.to_numeric(out["RMSE增加2.5%分位"], errors="coerce") > PI_ZERO_TOL),
        "是", "否"
    )
    out.loc[out["RMSE增加均值"].isna(), "是否稳健正PI"] = "不可判断"

    conclusions: List[str] = []
    for _, row in out.iterrows():
        status = str(row.get("PI可评估性", ""))
        mean_delta = safe_float(row.get("RMSE增加均值"))
        lower = safe_float(row.get("RMSE增加2.5%分位"))
        if status.startswith("不可评估"):
            conclusions.append(status)
        elif status.startswith("有限可评估"):
            if np.isfinite(mean_delta) and mean_delta > PI_ZERO_TOL:
                conclusions.append(status + "；存在正PI信号但不作正式排名")
            else:
                conclusions.append(status + "；未见明确正PI信号")
        elif np.isfinite(lower) and lower > PI_ZERO_TOL:
            conclusions.append("稳健正重要性：95%置换区间下限仍>0")
        elif np.isfinite(mean_delta) and mean_delta > PI_ZERO_TOL:
            conclusions.append("正重要性信号，但95%区间跨0")
        else:
            conclusions.append("无明确正重要性信号")
    out["PI证据结论"] = conclusions

    out["RMSE增加重要性排名"] = np.nan
    for _, idx in out.groupby(["目标指标", "验证方式", "项目类型"], sort=False).groups.items():
        sub = out.loc[idx]
        eligible = sub["PI可评估性"].eq("可评估")
        ranks = positive_importance_ranks(sub["RMSE增加均值"], PI_ZERO_TOL, eligible)
        out.loc[idx, "RMSE增加重要性排名"] = ranks

    return out.sort_values(
        ["目标指标", "验证方式", "项目类型", "RMSE增加重要性排名", "RMSE增加均值"],
        ascending=[True, True, True, True, False],
        na_position="last",
    ).reset_index(drop=True)


# =============================================================================
# 10. SHAP × PI 对照与跨指标汇总
# =============================================================================

def build_shap_pi_feature_comparison(
    target_name: str,
    shap_global: pd.DataFrame,
    pi_summary: pd.DataFrame,
    high_corr: pd.DataFrame,
) -> pd.DataFrame:
    single = pi_summary.loc[pi_summary["项目类型"] == "单特征"].copy()
    time_pi = single.loc[single["验证方式"] == "Time"].set_index("项目名称")
    exp_pi = single.loc[single["验证方式"] == "Expanding"].set_index("项目名称")

    correlated_features: set[str] = set()
    if not high_corr.empty:
        correlated_features.update(high_corr["特征1"].astype(str).tolist())
        correlated_features.update(high_corr["特征2"].astype(str).tolist())

    rows: List[Dict[str, Any]] = []
    for _, shap_row in shap_global.iterrows():
        feature = str(shap_row["特征"])
        t = time_pi.loc[feature] if feature in time_pi.index else None
        e = exp_pi.loc[feature] if feature in exp_pi.index else None
        srank = safe_float(shap_row.get("SHAP重要性排名"))
        trank = safe_float(t.get("RMSE增加重要性排名")) if t is not None else np.nan
        erank = safe_float(e.get("RMSE增加重要性排名")) if e is not None else np.nan

        top_shap = np.isfinite(srank) and srank <= 5
        top_time = np.isfinite(trank) and trank <= 5
        top_exp = np.isfinite(erank) and erank <= 5
        n_top = int(top_shap) + int(top_time) + int(top_exp)
        if top_shap and top_time and top_exp:
            agreement = "SHAP与Time/Expanding PI共同Top5"
        elif n_top >= 2:
            agreement = "三种证据中至少两种进入Top5"
        elif top_shap and not top_time and not top_exp:
            agreement = "SHAP较高，但Time/Expanding正式PI未同时支持"
        elif (not top_shap) and top_time and top_exp:
            agreement = "两种PI较高但SHAP未进Top5"
        elif not np.isfinite(srank):
            agreement = "SHAP无有效贡献"
        else:
            agreement = "一般或方法间差异"
        if feature in correlated_features and "SHAP较高" in agreement:
            agreement += "；存在高相关替代可能"

        rows.append({
            "目标指标": target_name,
            "特征": feature,
            "统一语义特征": shap_row["统一语义特征"],
            "特征类别": shap_row["特征类别"],
            "SHAP是否有效贡献": shap_row.get("是否有效SHAP贡献", ""),
            "SHAP排名": srank,
            "SHAP重要性占比百分比": shap_row["重要性占比百分比"],
            "SHAP方向结论": shap_row["方向结论"],
            "Time_PI可评估性": str(t.get("PI可评估性")) if t is not None else "缺少结果",
            "Time_PI证据结论": str(t.get("PI证据结论")) if t is not None else "缺少结果",
            "Time_PI排名": trank,
            "Time_RMSE增加均值": safe_float(t.get("RMSE增加均值")) if t is not None else np.nan,
            "Time_RMSE增加95%下限": safe_float(t.get("RMSE增加2.5%分位")) if t is not None else np.nan,
            "Time_RMSE增加95%上限": safe_float(t.get("RMSE增加97.5%分位")) if t is not None else np.nan,
            "Time_RMSE增加为正比例百分比": safe_float(t.get("RMSE增加为正比例百分比")) if t is not None else np.nan,
            "Time_验证段联合唯一模式数": safe_float(t.get("验证段联合唯一模式数")) if t is not None else np.nan,
            "Expanding_PI可评估性": str(e.get("PI可评估性")) if e is not None else "缺少结果",
            "Expanding_PI证据结论": str(e.get("PI证据结论")) if e is not None else "缺少结果",
            "Expanding_PI排名": erank,
            "Expanding_RMSE增加均值": safe_float(e.get("RMSE增加均值")) if e is not None else np.nan,
            "Expanding_RMSE增加95%下限": safe_float(e.get("RMSE增加2.5%分位")) if e is not None else np.nan,
            "Expanding_RMSE增加95%上限": safe_float(e.get("RMSE增加97.5%分位")) if e is not None else np.nan,
            "Expanding_RMSE增加为正比例百分比": safe_float(e.get("RMSE增加为正比例百分比")) if e is not None else np.nan,
            "Expanding_可置换窗口数": safe_float(e.get("可置换窗口数")) if e is not None else np.nan,
            "Expanding_总窗口数": safe_float(e.get("总窗口数")) if e is not None else np.nan,
            "Expanding_可置换窗口比例百分比": safe_float(e.get("可置换窗口比例百分比")) if e is not None else np.nan,
            "是否涉及高相关特征": "是" if feature in correlated_features else "否",
            "一致性解释": agreement,
        })
    return pd.DataFrame(rows).sort_values(
        ["目标指标", "SHAP排名", "SHAP重要性占比百分比"],
        ascending=[True, True, False],
        na_position="last",
    ).reset_index(drop=True)


def build_method_agreement(target_name: str, comparison: pd.DataFrame) -> pd.DataFrame:
    if comparison.empty:
        return pd.DataFrame()

    def corr_where(a_col: str, b_col: str, mask: pd.Series) -> float:
        a = pd.to_numeric(comparison.loc[mask, a_col], errors="coerce")
        b = pd.to_numeric(comparison.loc[mask, b_col], errors="coerce")
        valid = a.notna() & b.notna() & np.isfinite(a) & np.isfinite(b)
        if int(valid.sum()) < 3:
            return np.nan
        return safe_float(a.loc[valid].corr(b.loc[valid], method="spearman"))

    time_ok = comparison["Time_PI可评估性"].eq("可评估")
    exp_ok = comparison["Expanding_PI可评估性"].eq("可评估")
    shap_ok = comparison["SHAP是否有效贡献"].eq("是")
    return pd.DataFrame([{
        "目标指标": target_name,
        "特征数": len(comparison),
        "有效SHAP特征数": int(shap_ok.sum()),
        "Time_PI正式可评估特征数": int(time_ok.sum()),
        "Expanding_PI正式可评估特征数": int(exp_ok.sum()),
        "SHAP与Time_PI_Spearman": corr_where(
            "SHAP重要性占比百分比", "Time_RMSE增加均值", shap_ok & time_ok
        ),
        "SHAP与Expanding_PI_Spearman": corr_where(
            "SHAP重要性占比百分比", "Expanding_RMSE增加均值", shap_ok & exp_ok
        ),
        "Time_PI与Expanding_PI_Spearman": corr_where(
            "Time_RMSE增加均值", "Expanding_RMSE增加均值", time_ok & exp_ok
        ),
        "三者共同Top5特征数": int((
            comparison["SHAP排名"].notna()
            & comparison["Time_PI排名"].notna()
            & comparison["Expanding_PI排名"].notna()
            & (comparison["SHAP排名"] <= 5)
            & (comparison["Time_PI排名"] <= 5)
            & (comparison["Expanding_PI排名"] <= 5)
        ).sum()),
        "说明": "Spearman仅在对应方法正式可评估且SHAP有效的特征子集上计算；不作为加权综合评分。",
    }])


def build_cross_target_table(comparison_all: pd.DataFrame) -> pd.DataFrame:
    rows: List[Dict[str, Any]] = []
    for semantic, group in comparison_all.groupby("统一语义特征", sort=False):
        row: Dict[str, Any] = {
            "统一语义特征": semantic,
            "出现指标数": int(group["目标指标"].nunique()),
            "SHAP_Top5指标数": int((group["SHAP排名"].notna() & (group["SHAP排名"] <= 5)).sum()),
            "Time_PI_Top5指标数": int((group["Time_PI排名"].notna() & (group["Time_PI排名"] <= 5)).sum()),
            "Expanding_PI_Top5指标数": int((group["Expanding_PI排名"].notna() & (group["Expanding_PI排名"] <= 5)).sum()),
        }
        for target in TARGET_ORDER:
            sub = group.loc[group["目标指标"] == target]
            if sub.empty:
                row[f"{target}_SHAP排名"] = np.nan
                row[f"{target}_TimePI排名"] = np.nan
                row[f"{target}_ExpPI排名"] = np.nan
            else:
                item = sub.iloc[0]
                row[f"{target}_SHAP排名"] = item["SHAP排名"]
                row[f"{target}_TimePI排名"] = item["Time_PI排名"]
                row[f"{target}_ExpPI排名"] = item["Expanding_PI排名"]
        rows.append(row)
    return pd.DataFrame(rows).sort_values(
        ["SHAP_Top5指标数", "Time_PI_Top5指标数", "Expanding_PI_Top5指标数", "出现指标数"],
        ascending=[False, False, False, False],
    ).reset_index(drop=True)


# =============================================================================
# 11. 单目标分析
# =============================================================================

def process_target(
    target_name: str,
    summary: pd.DataFrame,
    fingerprint: str,
) -> Dict[str, pd.DataFrame]:
    checkpoint = CHECKPOINT_DIR / f"{target_name.replace('+','plus').replace('-','_')}.pkl"
    if RESUME_CHECKPOINT and checkpoint.exists():
        try:
            with checkpoint.open("rb") as handle:
                saved = pickle.load(handle)
            if saved.get("fingerprint") == fingerprint:
                log(f"{target_name} 使用已有检查点，跳过重复计算。")
                return saved["result"]
        except Exception:
            log(f"{target_name} 检查点读取失败，重新计算。")

    start = time.time()
    path, package, summary_row = load_package(target_name, summary)
    raw_data = read_dynamic_table(target_name, package)
    feature_data, x_raw, x_scaled, predictions, imputation = rebuild_full_feature_data(
        target_name, raw_data, package
    )

    # SVR的model-agnostic SHAP计算成本明显高于线性模型。
    # 正式分析对SVR使用按“时间阶段×目标水平”分层抽取的代表样本；
    # ElasticNet解析SHAP成本低，仍解释全部样本。后续beeswarm可直接使用该代表样本长表。
    shap_feature_data = feature_data
    shap_x_scaled = x_scaled
    shap_predictions = predictions
    if package["algorithm"] == "SVR" and len(feature_data) > SVR_SHAP_EXPLAIN_N:
        n_explain = 60 if SMOKE_TEST else SVR_SHAP_EXPLAIN_N
        idx = stratified_sample_indices(
            feature_data, str(package["target_column"]), min(n_explain, len(feature_data)),
            RANDOM_SEED + stable_seed(f"{target_name}-svr-shap-explain")
        )
        shap_feature_data = feature_data.iloc[idx].reset_index(drop=True)
        shap_x_scaled = x_scaled[idx]
        shap_predictions = predictions[idx]
    elif SMOKE_TEST and len(feature_data) > 60:
        idx = stratified_sample_indices(
            feature_data, str(package["target_column"]), 60,
            RANDOM_SEED + stable_seed(f"{target_name}-smoke-shap")
        )
        shap_feature_data = feature_data.iloc[idx].reset_index(drop=True)
        shap_x_scaled = x_scaled[idx]
        shap_predictions = predictions[idx]

    background_indices_full = stratified_sample_indices(
        feature_data,
        str(package["target_column"]),
        min(SHAP_BACKGROUND_N, len(feature_data)),
        RANDOM_SEED + stable_seed(f"{target_name}-SHAP-background"),
    )
    # 若smoke抽样，背景仍从完整历史中取真实样本，模型和背景不改变。
    background_x = x_scaled[background_indices_full]
    background_table = feature_data.iloc[background_indices_full][[
        "样本序号", "Date", str(package["target_column"])
    ]].copy()
    background_table.insert(0, "目标指标", target_name)
    background_table = background_table.rename(columns={str(package["target_column"]): "目标实测值"})
    background_table["背景样本顺序"] = np.arange(1, len(background_table) + 1)

    log(f"{target_name} 开始最终固化模型SHAP。")
    if package["algorithm"] == "SVR":
        phi, base_values, shap_stability, shap_method = compute_svr_permutation_shap(
            target_name, package, shap_x_scaled, background_x
        )
        linear_consistency = np.nan
    else:
        phi, base_values, shap_stability, shap_method, linear_consistency = compute_linear_shap(
            target_name, package, shap_x_scaled, background_x
        )

    shap_global, shap_long, shap_group, shap_stage, additivity = summarize_shap(
        target_name, package, shap_feature_data, shap_x_scaled, shap_predictions,
        phi, base_values, shap_stability
    )

    corr_long, high_corr, clusters, cluster_mapping = build_correlation_tables(
        target_name, package, feature_data
    )

    log(f"{target_name} 复现Time/Expanding固定超参数时间模型。")
    time_obj, windows, validation_check = prepare_temporal_models(
        target_name, raw_data, package
    )
    items = build_permutation_items(package, cluster_mapping)

    log(f"{target_name} 计算Time Permutation Importance（{PI_REPEATS}次置换）。")
    time_detail, time_summary = run_time_permutation_importance(
        target_name, time_obj, items
    )
    log(f"{target_name} 计算Expanding Permutation Importance（{PI_REPEATS}次置换）。")
    exp_detail, exp_summary, window_pi = run_expanding_permutation_importance(
        target_name, windows, items
    )
    pi_detail = concat_frames([time_detail, exp_detail])
    pi_summary = add_pi_ranks(concat_frames([time_summary, exp_summary]))

    comparison = build_shap_pi_feature_comparison(
        target_name, shap_global, pi_summary, high_corr
    )
    method_agreement = build_method_agreement(target_name, comparison)

    coef_rows: List[Dict[str, Any]] = []
    if package["algorithm"] == "ElasticNet":
        coefficients = zero_small(np.asarray(package["model"].coef_, dtype=float), SHAP_ZERO_TOL)
        abs_coef = np.abs(coefficients)
        ranks = positive_importance_ranks(pd.Series(abs_coef), SHAP_ZERO_TOL)
        for j, feature in enumerate(package["feature_names"]):
            coef_rows.append({
                "目标指标": target_name,
                "特征": feature,
                "特征类别": feature_group(feature, package),
                "标准化ElasticNet系数": float(coefficients[j]),
                "系数绝对值": float(abs_coef[j]),
                "是否非零系数": "是" if abs_coef[j] > SHAP_ZERO_TOL else "否",
                "系数绝对值排名": int(ranks.iloc[j]) if np.isfinite(ranks.iloc[j]) else np.nan,
                "系数方向": "正" if coefficients[j] > SHAP_ZERO_TOL else ("负" if coefficients[j] < -SHAP_ZERO_TOL else "0"),
            })
    coef_df = pd.DataFrame(coef_rows)

    package_metrics = package["source_validation_metrics"]
    overview = pd.DataFrame([{
        "目标指标": target_name,
        "最终特征组": package["selected_feature_group"],
        "算法": package["algorithm"],
        "候选参数ID": package["candidate_id"],
        "参数JSON": json.dumps(package["parameters"], ensure_ascii=False, sort_keys=True),
        "特征数": package["feature_count"],
        "全量历史样本数": len(feature_data),
        "SHAP解释样本数": len(shap_feature_data),
        "SHAP背景样本数": len(background_indices_full),
        "SHAP方法": shap_method,
        "SVR_SHAP重复次数": SVR_SHAP_REPEATS if package["algorithm"] == "SVR" else 0,
        "SVR_SHAP_cycle数": SVR_SHAP_CYCLES if package["algorithm"] == "SVR" else 0,
        "SVR正式分层解释目标样本数": SVR_SHAP_EXPLAIN_N if package["algorithm"] == "SVR" else len(feature_data),
        "LinearSHAP解析一致性最大差": linear_consistency,
        "PI重复次数": PI_REPEATS,
        "SHAP零贡献容差": SHAP_ZERO_TOL,
        "PI零变化容差": PI_ZERO_TOL,
        "Expanding正式排名最少可置换窗口数": EXPANDING_MIN_EVALUABLE_WINDOWS,
        "Time验证样本数": len(time_obj["y_valid"]),
        "Expanding窗口数": len(windows),
        "Expanding累计验证样本数": int(sum(len(w["y_valid"]) for w in windows)),
        "Bootstrap_R2": package_metrics["bootstrap_632plus"]["R2"],
        "Bootstrap_RMSE": package_metrics["bootstrap_632plus"]["RMSE"],
        "Time_R2": package_metrics["time"]["R2"],
        "Time_RMSE": package_metrics["time"]["RMSE"],
        "Expanding_R2": package_metrics["expanding"]["R2"],
        "Expanding_RMSE": package_metrics["expanding"]["RMSE"],
        "两种时间验证均优于最佳历史基线": (
            "是" if package_metrics["beats_best_historical_baseline_in_both_time_validations"] else "否"
        ),
        "模型包文件": str(path),
        "模型包SHA256": sha256_file(path),
        "动态数据SHA256": sha256_file(DYNAMIC_INPUT_FILE),
        "解释边界": package.get("interpretation_note", ""),
        "运行耗时分钟": (time.time() - start) / 60.0,
        "是否SmokeTest": "是" if SMOKE_TEST else "否",
    }])

    warnings_df = pd.DataFrame()
    warning_rows: List[Dict[str, Any]] = []
    if time_obj.get("warning"):
        warning_rows.append({"目标指标": target_name, "阶段": "Time重拟合", "警告或错误": time_obj["warning"]})
    for w in windows:
        if w.get("warning"):
            warning_rows.append({
                "目标指标": target_name,
                "阶段": f"Expanding窗口{w['window_id']}重拟合",
                "警告或错误": w["warning"],
            })
    warnings_df = pd.DataFrame(warning_rows)

    result = {
        "分析与模型核查": overview,
        "插补重建核查": imputation,
        "SHAP背景样本": background_table,
        "SHAP全局汇总": shap_global,
        "SHAP样本长表": shap_long,
        "SHAP分组贡献": shap_group,
        "SHAP时间阶段": shap_stage,
        "SHAP近似稳定性": shap_stability,
        "SHAP加和核查": additivity,
        "ElasticNet系数": coef_df,
        "Spearman相关性": corr_long,
        "高相关特征对": high_corr,
        "相关特征簇": clusters,
        "时间验证复现核查": validation_check,
        "PI逐次明细": pi_detail,
        "PI汇总": pi_summary,
        "Expanding逐窗口PI": window_pi,
        "SHAP_PI单特征对照": comparison,
        "方法一致性": method_agreement,
        "错误与警告": warnings_df,
    }

    atomic_pickle_dump({"fingerprint": fingerprint, "result": result}, checkpoint)
    log(f"{target_name} SHAP + PI完成，耗时{(time.time()-start)/60:.2f}分钟。")
    return result


# =============================================================================
# 12. 综合文字说明
# =============================================================================

def top_lines(frame: pd.DataFrame, value_col: str, rank_col: str, n: int = 5) -> List[pd.Series]:
    if frame.empty:
        return []
    return [row for _, row in frame.sort_values(rank_col).head(n).iterrows()]


def build_text_report(tables: Mapping[str, pd.DataFrame], elapsed_minutes: float) -> str:
    overview = tables["分析与模型核查"]
    shap_global = tables["SHAP全局汇总"]
    shap_group = tables["SHAP分组贡献"]
    pi = tables["PI汇总"]
    comparison = tables["SHAP_PI单特征对照"]
    high_corr = tables.get("高相关特征对", pd.DataFrame())
    agreement = tables["方法一致性"]
    validation = tables["时间验证复现核查"]
    warnings_df = tables.get("错误与警告", pd.DataFrame())

    def fmt_num(value: Any, digits: int = 6) -> str:
        x = safe_float(value)
        return f"{x:.{digits}g}" if np.isfinite(x) else "NA"

    def fmt_corr(value: Any) -> str:
        x = safe_float(value)
        return f"{x:.3f}" if np.isfinite(x) else "NA（有效可比特征不足）"

    lines: List[str] = []
    lines.append("A²/O 四指标最终模型 SHAP + Permutation Importance 综合说明（最终修正版）")
    lines.append("=" * 118)
    lines.append("")
    lines.append("一、任务与运行信息")
    lines.append(f"1. 完成时间：{now_str()}")
    lines.append(f"2. 总运行耗时：{elapsed_minutes:.2f} 分钟")
    lines.append(f"3. 最终模型目录：{MODEL_DIR}")
    lines.append(f"4. 动态数据：{DYNAMIC_INPUT_FILE}")
    lines.append(f"5. 输出Excel：{OUTPUT_EXCEL}")
    lines.append(f"6. 输出TXT：{OUTPUT_TEXT}")
    lines.append("7. 本阶段不生成任何PNG/PDF/SVG图件。")
    lines.append(f"8. 正式PI重复次数：{PI_REPEATS}；高相关阈值：|rho|>={HIGH_CORRELATION_THRESHOLD:.2f}。")
    lines.append(
        f"9. SVR SHAP：背景样本{SHAP_BACKGROUND_N}，固定种子重复{SVR_SHAP_REPEATS}次，"
        f"每次Permutation cycle={SVR_SHAP_CYCLES}；SVR正式分层解释样本目标数={SVR_SHAP_EXPLAIN_N}。"
    )
    lines.append(
        f"10. 数值零容差：SHAP={SHAP_ZERO_TOL:g}，PI Δ指标={PI_ZERO_TOL:g}；"
        f"Expanding至少{EXPANDING_MIN_EVALUABLE_WINDOWS}个窗口有组内变化才进入正式PI排名。"
    )
    lines.append(f"11. Smoke test状态：{'开启（仅程序测试，不可作为正式结果）' if SMOKE_TEST else '关闭（正式分析）'}。")
    lines.append("")

    lines.append("二、解释方法、可评估性和证据边界")
    lines.append("1. 最终SHAP直接解释已经冻结的全量历史joblib模型；不会在SHAP阶段重新选择模型或参数。")
    lines.append(
        f"   TP/TN的SVR按时间阶段×目标水平分层抽取最多{SVR_SHAP_EXPLAIN_N}条代表历史样本；"
        "NH4+-N/COD的ElasticNet解释全部历史样本。"
    )
    lines.append("2. TP/TN采用model-agnostic Permutation SHAP；NH4+-N/COD采用Independent-background Linear SHAP并做解析一致性核查。")
    lines.append("3. SHAP平均绝对贡献≤数值容差的特征记为‘无有效SHAP贡献’，不参与Top5；不会再用零值并列填满Top5。")
    lines.append("4. Permutation Importance主证据来自样本外时间验证，不在全量训练拟合上宣称泛化重要性。")
    lines.append("5. Time PI：若验证段某特征/特征组只有1个联合模式，则置换无法改变数据，明确标记‘不可评估’，不把PI=0解释为无重要性。")
    lines.append(
        f"6. Expanding PI：逐窗口先检查是否存在至少2个联合模式；有变化窗口执行置换，无变化窗口保持原基线预测，所有项目共享同一完整Expanding样本外基线。"
        f"0个窗口为不可评估，1–{EXPANDING_MIN_EVALUABLE_WINDOWS-1}个窗口为有限可评估且不进入正式排名，"
        f"≥{EXPANDING_MIN_EVALUABLE_WINDOWS}个窗口才进入正式PI排名。"
    )
    lines.append("7. PI平均ΔRMSE>0仅称为‘正重要性信号’；只有95%置换区间下限仍>0且可评估性充分时，才标记‘稳健正重要性’。")
    lines.append("8. 单变量PI容易受到高相关变量替代；因此同时报告三类主分组PI和高相关特征簇PI。")
    lines.append("9. SHAP与PI不合成加权总分；只比较排名、方向、时间稳定性和方法一致性。")
    lines.append("10. 所有结果均是模型依赖/关联解释，不等于真实因果效应；后续PDP、正向预测和优化仍需结合历史支持域与物理可行性。")
    lines.append("")

    lines.append("三、模型包与时间验证复现核查")
    for target in TARGET_ORDER:
        row = overview.loc[overview["目标指标"] == target].iloc[0]
        v = validation.loc[validation["目标指标"] == target]
        status = "、".join(f"{r['验证方式']}={r['是否通过']}" for _, r in v.iterrows())
        lines.append(
            f"- {target}：{row['最终特征组']}-{row['算法']}（{row['候选参数ID']}），"
            f"12个F2特征；Time/Expanding复现核查：{status}；SHAP方法={row['SHAP方法']}。"
        )
    lines.append("")

    section_titles = {"TP": "四", "TN": "五", "NH4+-N": "六", "COD": "七"}
    for target in TARGET_ORDER:
        lines.append(f"{section_titles[target]}、{target} 解释结果")
        ov = overview.loc[overview["目标指标"] == target].iloc[0]
        lines.append(
            f"1. 最终ML：{ov['最终特征组']}-{ov['算法']}（{ov['候选参数ID']}）；"
            f"Bootstrap R²={ov['Bootstrap_R2']:.6f}, RMSE={ov['Bootstrap_RMSE']:.6f}；"
            f"Time R²={ov['Time_R2']:.6f}, RMSE={ov['Time_RMSE']:.6f}；"
            f"Expanding R²={ov['Expanding_R2']:.6f}, RMSE={ov['Expanding_RMSE']:.6f}。"
        )

        sg_all = shap_global.loc[shap_global["目标指标"] == target].copy()
        sg = sg_all.loc[
            sg_all["是否有效SHAP贡献"].eq("是") & sg_all["SHAP重要性排名"].notna()
        ].sort_values("SHAP重要性排名")
        n_eff = len(sg)
        lines.append(f"2. SHAP有效贡献特征Top5（有效特征共{n_eff}个；不足5个时不以零贡献特征补位）：")
        if sg.empty:
            lines.append("   - 没有超过数值零容差的有效SHAP特征。")
        else:
            for _, r in sg.head(5).iterrows():
                lines.append(
                    f"   #{int(r['SHAP重要性排名'])} {r['特征']}：mean(|SHAP|)={r['平均绝对SHAP']:.6g}，"
                    f"占比={r['重要性占比百分比']:.2f}%；{r['方向结论']}。"
                )

        gg = shap_group.loc[shap_group["目标指标"] == target].sort_values(
            "绝对贡献占比百分比", ascending=False
        )
        lines.append("3. 三类信息的SHAP绝对贡献占比：")
        for _, r in gg.iterrows():
            lines.append(
                f"   - {r['特征组']}：{r['绝对贡献占比百分比']:.2f}%（{int(r['特征数'])}个特征）。"
            )

        single = pi.loc[(pi["目标指标"] == target) & (pi["项目类型"] == "单特征")].copy()
        for step_no, validation_name in [(4, "Time"), (5, "Expanding")]:
            sub = single.loc[single["验证方式"] == validation_name].copy()
            ranked = sub.loc[sub["RMSE增加重要性排名"].notna()].sort_values("RMSE增加重要性排名").head(5)
            lines.append(f"{step_no}. {validation_name}样本外Permutation Importance正式正PI Top5：")
            if ranked.empty:
                lines.append("   - 没有同时满足‘正式可评估 + 平均ΔRMSE>0’的单特征。")
            else:
                for _, r in ranked.iterrows():
                    robust = "稳健正PI" if r["是否稳健正PI"] == "是" else "正PI但95%区间跨0"
                    eval_extra = ""
                    if validation_name == "Time":
                        eval_extra = f"，验证段联合唯一模式={int(r['验证段联合唯一模式数'])}"
                    else:
                        eval_extra = (
                            f"，可置换窗口={int(r['可置换窗口数'])}/{int(r['总窗口数'])}"
                            f" ({r['可置换窗口比例百分比']:.1f}%)"
                        )
                    lines.append(
                        f"   #{int(r['RMSE增加重要性排名'])} {r['项目名称']}：ΔRMSE={r['RMSE增加均值']:.6g} "
                        f"[{r['RMSE增加2.5%分位']:.6g}, {r['RMSE增加97.5%分位']:.6g}]，"
                        f"变差比例={r['RMSE增加为正比例百分比']:.1f}%，{robust}{eval_extra}。"
                    )
            unavailable = sub.loc[sub["PI可评估性"].astype(str).str.startswith("不可评估")]
            limited = sub.loc[sub["PI可评估性"].astype(str).str.startswith("有限可评估")]
            if not unavailable.empty:
                lines.append("   - 不可评估单特征：" + "、".join(unavailable["项目名称"].tolist()) + "。")
            if not limited.empty:
                detail = "；".join(
                    f"{r['项目名称']}({int(r['可置换窗口数'])}/{int(r['总窗口数'])}窗口)"
                    for _, r in limited.iterrows()
                )
                lines.append("   - 有限可评估、因此不进入正式排名：" + detail + "。")

        group_pi = pi.loc[(pi["目标指标"] == target) & (pi["项目类型"] == "主分组")]
        lines.append("6. 三类信息的分组Permutation Importance：")
        for group_name in GROUP_ORDER:
            trow = group_pi.loc[(group_pi["验证方式"] == "Time") & (group_pi["项目名称"] == group_name)]
            erow = group_pi.loc[(group_pi["验证方式"] == "Expanding") & (group_pi["项目名称"] == group_name)]
            if not trow.empty and not erow.empty:
                t = trow.iloc[0]; e = erow.iloc[0]
                lines.append(
                    f"   - {group_name}：Time ΔRMSE={fmt_num(t['RMSE增加均值'])}（{t['PI证据结论']}）；"
                    f"Expanding ΔRMSE={fmt_num(e['RMSE增加均值'])}（{e['PI证据结论']}，"
                    f"可置换窗口={int(e['可置换窗口数']) if np.isfinite(safe_float(e['可置换窗口数'])) else 'NA'}/"
                    f"{int(e['总窗口数']) if np.isfinite(safe_float(e['总窗口数'])) else 'NA'}）。"
                )

        comp = comparison.loc[comparison["目标指标"] == target].copy()
        common = comp.loc[
            comp["SHAP排名"].notna() & comp["Time_PI排名"].notna() & comp["Expanding_PI排名"].notna()
            & (comp["SHAP排名"] <= 5) & (comp["Time_PI排名"] <= 5) & (comp["Expanding_PI排名"] <= 5)
        ]
        lines.append(
            "7. SHAP与两种正式PI共同Top5："
            + ("、".join(common["特征"].tolist()) if not common.empty else "无完全共同Top5特征")
            + "。"
        )

        corr = high_corr.loc[high_corr["目标指标"] == target] if not high_corr.empty else pd.DataFrame()
        if corr.empty:
            lines.append(f"8. |Spearman rho|>={HIGH_CORRELATION_THRESHOLD:.2f}的高相关特征对：无。")
        else:
            top_corr = corr.sort_values("绝对rho", ascending=False).head(8)
            desc = "；".join(
                f"{r['特征1']}–{r['特征2']} (rho={r['Spearman_rho']:.3f})" for _, r in top_corr.iterrows()
            )
            lines.append(f"8. 主要高相关特征对：{desc}。单变量SHAP/PI应结合相关簇PI解读。")

        operation = comp.loc[comp["特征"].isin(OPERATION_FEATURES)].copy()
        operation["候选证据数"] = (
            operation["SHAP排名"].notna().astype(int)
            + operation["Time_PI排名"].notna().astype(int)
            + operation["Expanding_PI排名"].notna().astype(int)
        )
        operation = operation.sort_values(
            ["候选证据数", "SHAP排名", "Time_PI排名", "Expanding_PI排名"],
            ascending=[False, True, True, True], na_position="last"
        )
        supported = operation.loc[operation["候选证据数"] > 0]
        lines.append("9. 后续PDP/ICE运行变量证据：")
        if supported.empty:
            lines.append("   - 当前最终模型中没有运行变量获得有效SHAP贡献或正式可评估的正PI排名，不宜强行把运行变量PDP作为主解释。")
        else:
            for _, r in supported.iterrows():
                shap_desc = rank_text(r["SHAP排名"], "无有效SHAP贡献")
                time_status = str(r["Time_PI可评估性"])
                exp_status = str(r["Expanding_PI可评估性"])
                time_fallback = "可评估但无正PI排名" if time_status == "可评估" else time_status
                exp_fallback = "可评估但无正PI排名" if exp_status == "可评估" else exp_status
                time_desc = rank_text(r["Time_PI排名"], time_fallback)
                exp_desc = rank_text(r["Expanding_PI排名"], exp_fallback)
                lines.append(
                    f"   - {r['特征']}：SHAP={shap_desc}；Time PI={time_desc}；Expanding PI={exp_desc}。"
                )
            lines.append("   - 这里的PDP候选只表示模型响应值得进一步查看；PI不可评估/有限可评估不等同于变量真实无效。")

        if target == "NH4+-N":
            lines.append(
                "10. NH4+-N边界：该F2-ElasticNet是最佳机器学习模型，但Time/Expanding预测性能仍未同时超过roll3；"
                "若运行变量SHAP为0且PI无正式正排名，应明确解释为该最终ML未学习到可用于运行响应分析的稳定工况信号，"
                "不能强行从该模型推导工况优化效应。"
            )
        else:
            lines.append("10. 解释边界：以上均为模型依赖/关联解释，不应直接写成运行变量的因果效应。")
        lines.append("")

    lines.append("八、SHAP与Permutation Importance方法一致性")
    for _, r in agreement.iterrows():
        lines.append(
            f"- {r['目标指标']}：有效SHAP={int(r['有效SHAP特征数'])}个，"
            f"Time正式可评估PI={int(r['Time_PI正式可评估特征数'])}个，"
            f"Expanding正式可评估PI={int(r['Expanding_PI正式可评估特征数'])}个；"
            f"SHAP-Time PI Spearman={fmt_corr(r['SHAP与Time_PI_Spearman'])}；"
            f"SHAP-Expanding PI={fmt_corr(r['SHAP与Expanding_PI_Spearman'])}；"
            f"Time PI-Expanding PI={fmt_corr(r['Time_PI与Expanding_PI_Spearman'])}；"
            f"三者共同Top5={int(r['三者共同Top5特征数'])}个。"
        )
    lines.append("- 相关系数只在对应方法具备正式可评估性的特征子集上计算，仅描述解释方法一致程度，不建立加权总分。")
    lines.append("")

    lines.append("九、后续论文制图所需数据位置")
    lines.append("1. SHAP重要性条形图：Excel“SHAP全局汇总”；只使用‘是否有效SHAP贡献=是’的特征。")
    lines.append("2. SHAP beeswarm：Excel“SHAP样本长表”（SVR为分层代表样本，ElasticNet为全部样本）。")
    lines.append("3. 三类信息贡献图：Excel“SHAP分组贡献”。")
    lines.append("4. 时间阶段SHAP：Excel“SHAP时间阶段”；零贡献特征不进入阶段内排名。")
    lines.append("5. 单变量/分组/相关簇PI：Excel“PI汇总”；同时读取‘PI可评估性’、‘PI证据结论’和95%区间。")
    lines.append("6. Expanding时间稳定性/工况可置换性：Excel“Expanding逐窗口PI”。")
    lines.append("7. SHAP与PI排名对照：Excel“SHAP_PI单特征对照”。")
    lines.append("8. 高相关性/相关簇：Excel“高相关特征对”和“相关特征簇”。")
    lines.append("")

    lines.append("十、结果完整性和警告")
    fail_checks = validation.loc[validation["是否通过"] != "是"] if not validation.empty else pd.DataFrame()
    lines.append(f"1. Time/Expanding复现未通过项：{len(fail_checks)}。")
    lines.append(f"2. 训练或运行警告记录：{len(warnings_df)}。")
    lines.append("3. 正式运行时如任一模型哈希、F2特征顺序、样本ID、全量拟合核查或时间验证复现失败，脚本会停止，不生成正式Excel/TXT。")
    lines.append("4. 零SHAP/零PI不再参与Top5；阶段性恒定工况不再被错误解释为PI=0的‘无重要性’。")
    lines.append("5. 本阶段没有重新搜索超参数，也没有用SHAP/PI结果反向修改最终模型。")
    return "\n".join(lines)


# =============================================================================
# 13. 主流程
# =============================================================================

def make_fingerprint(summary: pd.DataFrame) -> str:
    payload = {
        "script_version": "A2O_SHAP_PI_v3.1_final_rank_evaluability_support_fix",
        "dynamic_hash": sha256_file(DYNAMIC_INPUT_FILE) if DYNAMIC_INPUT_FILE.exists() else "",
        "model_hashes": {
            target: sha256_file(model_path(target)) if model_path(target).exists() else ""
            for target in TARGET_ORDER
        },
        "pi_repeats": PI_REPEATS,
        "shap_background_n": SHAP_BACKGROUND_N,
        "svr_shap_repeats": SVR_SHAP_REPEATS,
        "svr_shap_cycles": SVR_SHAP_CYCLES,
        "svr_shap_explain_n": SVR_SHAP_EXPLAIN_N,
        "shap_zero_tol": SHAP_ZERO_TOL,
        "pi_zero_tol": PI_ZERO_TOL,
        "expanding_min_evaluable_windows": EXPANDING_MIN_EVALUABLE_WINDOWS,
        "high_corr": HIGH_CORRELATION_THRESHOLD,
        "smoke": SMOKE_TEST,
        "shap_version": getattr(shap, "__version__", "unknown"),
    }
    raw = json.dumps(payload, ensure_ascii=False, sort_keys=True).encode("utf-8")
    return hashlib.sha256(raw).hexdigest()


def validate_environment() -> pd.DataFrame:
    if not MODEL_DIR.exists():
        raise FileNotFoundError(f"最终模型目录不存在：{MODEL_DIR}")
    if not DYNAMIC_INPUT_FILE.exists():
        raise FileNotFoundError(f"动态数据不存在：{DYNAMIC_INPUT_FILE}")
    summary = load_model_summary()
    rows = []
    for target in TARGET_ORDER:
        path, package, srow = load_package(target, summary)
        rows.append({
            "目标指标": target,
            "模型文件": str(path),
            "SHA256核查": "通过",
            "最终特征组": package["selected_feature_group"],
            "算法": package["algorithm"],
            "候选参数ID": package["candidate_id"],
            "特征数": package["feature_count"],
        })
    return pd.DataFrame(rows)


def main() -> None:
    overall_start = time.time()
    print("=" * 112)
    print("A²/O 四指标最终模型 SHAP + Permutation Importance 综合分析开始")
    print("=" * 112)
    log(f"模型目录：{MODEL_DIR}")
    log(f"动态数据：{DYNAMIC_INPUT_FILE}")
    log(f"输出目录：{OUTPUT_DIR}")
    log(f"SHAP版本：{getattr(shap, '__version__', 'unknown')}")
    log(f"PI重复次数：{PI_REPEATS}")
    log(f"SVR SHAP：背景={SHAP_BACKGROUND_N}，重复={SVR_SHAP_REPEATS}，cycles={SVR_SHAP_CYCLES}，解释样本={SVR_SHAP_EXPLAIN_N}")
    log(f"零值容差：SHAP={SHAP_ZERO_TOL:g}，PI={PI_ZERO_TOL:g}；Expanding正式排名最少可置换窗口={EXPANDING_MIN_EVALUABLE_WINDOWS}")
    log(f"Smoke test：{SMOKE_TEST}")

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    CHECKPOINT_DIR.mkdir(parents=True, exist_ok=True)

    environment = validate_environment()
    summary = load_model_summary()
    fingerprint = make_fingerprint(summary)

    result_by_target: Dict[str, Dict[str, pd.DataFrame]] = {}
    fatal_errors: List[Dict[str, Any]] = []
    for target in TARGET_ORDER:
        try:
            result_by_target[target] = process_target(target, summary, fingerprint)
        except Exception as exc:
            fatal_errors.append({
                "目标指标": target,
                "阶段": "SHAP+PI综合分析",
                "错误": repr(exc),
                "Traceback": traceback.format_exc(),
            })
            traceback.print_exc()
            if STRICT_MODE:
                break

    if fatal_errors:
        error_file = OUTPUT_DIR / "SHAP_PI分析失败记录.txt"
        text = "\n\n".join(
            f"目标：{row['目标指标']}\n错误：{row['错误']}\n{row['Traceback']}"
            for row in fatal_errors
        )
        atomic_write_text(error_file, text)
        raise RuntimeError(f"SHAP+PI分析失败，正式结果未生成。错误记录：{error_file}")

    table_names = sorted({
        name for result in result_by_target.values() for name in result.keys()
    })
    combined: Dict[str, pd.DataFrame] = {}
    for name in table_names:
        combined[name] = concat_frames(
            result_by_target[target].get(name, pd.DataFrame()) for target in TARGET_ORDER
        )

    combined["环境与文件核查"] = environment
    combined["跨指标特征比较"] = build_cross_target_table(combined["SHAP_PI单特征对照"])
    if "PI汇总" in combined and not combined["PI汇总"].empty:
        eval_cols = [
            "目标指标", "验证方式", "项目类型", "项目名称", "成员特征",
            "验证段联合唯一模式数", "可置换窗口数", "总窗口数", "可置换窗口比例百分比",
            "PI可评估性", "是否正PI信号", "是否稳健正PI", "PI证据结论",
            "RMSE增加均值", "RMSE增加2.5%分位", "RMSE增加97.5%分位", "RMSE增加重要性排名",
        ]
        existing = [c for c in eval_cols if c in combined["PI汇总"].columns]
        combined["PI可评估性"] = combined["PI汇总"][existing].copy()

    preferred_order = [
        "环境与文件核查",
        "分析与模型核查",
        "插补重建核查",
        "时间验证复现核查",
        "SHAP背景样本",
        "SHAP全局汇总",
        "SHAP分组贡献",
        "SHAP时间阶段",
        "SHAP近似稳定性",
        "SHAP加和核查",
        "ElasticNet系数",
        "SHAP样本长表",
        "Spearman相关性",
        "高相关特征对",
        "相关特征簇",
        "PI汇总",
        "PI可评估性",
        "PI逐次明细",
        "Expanding逐窗口PI",
        "SHAP_PI单特征对照",
        "方法一致性",
        "跨指标特征比较",
        "错误与警告",
    ]
    output_tables = {
        name: combined[name]
        for name in preferred_order
        if name in combined and combined[name] is not None and not combined[name].empty
    }

    elapsed = (time.time() - overall_start) / 60.0
    report = build_text_report(combined, elapsed)

    if SMOKE_TEST:
        smoke_excel = OUTPUT_DIR / "SMOKE_TEST_四指标_SHAP与PI结果.xlsx"
        smoke_text = OUTPUT_DIR / "SMOKE_TEST_四指标_SHAP与PI说明.txt"
        write_excel(smoke_excel, output_tables)
        atomic_write_text(smoke_text, report)
        log("Smoke test完成；未覆盖正式结果文件。")
        print(f"\nSmoke Excel：{smoke_excel}\nSmoke TXT：{smoke_text}")
    else:
        write_excel(OUTPUT_EXCEL, output_tables)
        atomic_write_text(OUTPUT_TEXT, report)
        if DELETE_CHECKPOINT_AFTER_SUCCESS and CHECKPOINT_DIR.exists():
            shutil.rmtree(CHECKPOINT_DIR, ignore_errors=True)
        log("四指标SHAP + Permutation Importance正式分析全部完成。")
        print(f"\nExcel：{OUTPUT_EXCEL}\nTXT：{OUTPUT_TEXT}")


if __name__ == "__main__":
    main()