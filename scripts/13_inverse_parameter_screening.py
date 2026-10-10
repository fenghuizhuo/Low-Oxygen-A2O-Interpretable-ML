# -*- coding: utf-8 -*-
r"""
A²/O 四指标一级 A 参考目标反向筛选与运行参数推荐
=====================================================

定位
----
本脚本直接读取正向预测阶段生成的 11 个结果 Excel：
    1. 四指标_DO-HRT-进水分配联合正向预测_第1部分.xlsx
    2. 四指标_DO-HRT-进水分配联合正向预测_第2部分.xlsx
    3. 四指标_DO-HRT-进水分配联合正向预测_第3部分.xlsx
    4. 四指标_DO-HRT-进水分配联合正向预测_第4部分.xlsx
    5. 四指标_DO-HRT-进水分配联合正向预测_第5部分.xlsx
    6. 四指标_DO-HRT-进水分配联合正向预测_第6部分.xlsx
    7. 四指标_DO-HRT-进水分配联合正向预测_第7部分.xlsx
    8. 四指标_DO-HRT-进水分配联合正向预测_第8部分.xlsx
    9. 四指标_DO-HRT-进水分配联合正向预测_第9部分.xlsx
    10. 四指标_DO-HRT-进水分配联合正向预测_第10部分.xlsx
    11. 四指标_DO-HRT-进水分配联合正向预测_第11部分.xlsx

不重新训练模型，不重新调用模型预测，不使用 NSGA-II，不做 Pareto 搜索。

正式逻辑
--------
1. 只读取十一个正向 Excel 中所有“明细_”副表。
2. 正式反向分析只使用“网格来源 == 第一轮粗扫描”的记录。
   “当前实际工况补充”只保留数量用于核对，不参与参数范围、组合统计和推荐。
3. 采用四项一级 A 工程参考目标，四项必须同时满足：
       TP_预测      <= 0.5 mg/L
       TN_预测      <= 15.0 mg/L
       NH4_预测     <= 5.0 mg/L
       COD_预测     <= 50.0 mg/L
4. 正向阶段标记为预测物理/数值检查不通过的记录，不进入正式反向筛选。
5. 反向程序根据四项原始预测值重新计算一级 A 判定，并与正向文件中保存的
   5 个一级 A 判定字段逐条核对；正向标签只用于核查，不直接替代反向计算。
6. 不使用历史距离、历史接近程度、是否历史真实工况等字段设置额外门槛，
   也不使用预测误差去改写一级 A 参考目标。
7. 正式粗网格应为：631 个背景 × 3780 个 DO×HRT×Feed 组合 = 2,385,180 条。
8. 分别统计四项在正式有效记录中的满足数量，并统计四项同时满足数量和比例。
9. 用自然方式说明每一项条件对最终结果的影响：
   暂时不要求某一项时，先看其他三项都满足的记录有多少；再加入该项，
   记录减少多少。这里不使用“约束贡献”等容易造成误解的术语。
10. 对全部 3780 个三参数组合，分别统计其在 631 个历史背景中的：
    物理/数值检查通过数、四项各自满足数、四项同时满足背景数及比例，
    以及四项预测的最小值、中位数和最大值。
11. 推荐顺序严格按既定流程：
    A. 先找“四项同时满足背景数最多”的组合；并列时依次选较低 DO、较短 HRT、较小 Feed。
    B. 再保留“四项同时满足背景数达到最佳组合 90% 以上”的组合，作为接近最佳表现的候选组合。
    C. 在这些接近最佳表现的组合中选“最低 DO 方案”：DO升序 -> 满足背景数降序 -> HRT升序 -> Feed升序。
    D. 在这些接近最佳表现的组合中选“最短 HRT 方案”：HRT升序 -> 满足背景数降序 -> DO升序 -> Feed升序。
    最终报告三类代表组合：适用背景最多、接近最佳表现的最低DO、接近最佳表现的最短HRT。
12. 如果没有任何正式粗网格记录能够四项同时满足，则如实报告无可推荐组合，
    不虚构参数范围或运行方案。
13. 所有结果均为冻结模型的一步条件预测筛选，不是长期稳态模拟，也不等同于法规意义上的日均达标判定。
14. 在上述总体结果基础上，再按 631 个历史背景的进水 TP 实际分布进行三分位分组：低、中、高三个进水 TP 水平。
    每个进水 TP 组内重新统计 3780 组 DO×HRT×Feed 参数在该组历史背景中的四项同时达标次数和比例。
15. 每个进水 TP 组内，以该组表现最佳参数组合的达标背景数为参照，保留达到最佳表现 90% 以上的组合，
    并汇总这些组合的 DO、HRT 和厌氧区进水分配比例范围。该范围用于说明高达标组合主要分布在哪里，
    不表示三个参数区间内任意自由组合都能达到要求。

默认输入目录
------------
${PROJECT_ROOT}/outputs/screening/forward

默认输出目录
------------
${PROJECT_ROOT}/outputs/screening/inverse

主要输出
--------
1. 四指标_目标出水反向优化结果.xlsx
   - 01_说明与目标
   - 02_筛选汇总
   - 03_各项条件影响
   - 04_参数范围
   - 05_参数组合统计
   - 06_接近最佳组合
   - 07_三个代表方案
   - 08_背景筛选统计
   - 09_方法检查
   - 10_进水TP分组汇总
   - 11_各组参数范围
   - 12_各组参数组合
2. 四指标_目标出水反向优化说明.txt

说明：
- 输出路径和文件名沿用原反向优化目录，不改变。
- 不输出百万级满足记录明细，因此一个 Excel 足够。
- 不再生成旧版“结果检查包.xlsx”。

依赖
----
pip install pandas numpy openpyxl xlsxwriter

可选环境变量
------------
A2O_INVERSE_FORWARD_DIR
A2O_INVERSE_FORWARD_PART1
A2O_INVERSE_FORWARD_PART2
A2O_INVERSE_FORWARD_PART3
A2O_INVERSE_FORWARD_PART4
A2O_INVERSE_FORWARD_PART5
A2O_INVERSE_FORWARD_PART6
A2O_INVERSE_FORWARD_PART7
A2O_INVERSE_FORWARD_PART8
A2O_INVERSE_FORWARD_PART9
A2O_INVERSE_FORWARD_PART10
A2O_INVERSE_FORWARD_PART11
A2O_INVERSE_OUTPUT_DIR
A2O_INVERSE_STRICT=1
A2O_INVERSE_EXPECTED_BACKGROUND_N=631
A2O_INVERSE_EXPECTED_COMBINATION_N=3780
A2O_INVERSE_PROGRESS_EVERY=50

四项一级 A 工程参考目标在本脚本中固定，不通过环境变量修改，避免误运行。
"""

from __future__ import annotations

import hashlib
import math
import os
import re
import time
import traceback
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterator, List, Mapping, Optional, Sequence, Tuple

import numpy as np
import pandas as pd
from openpyxl import load_workbook


# =============================================================================
# 1. 路径、固定目标与严格配置
# =============================================================================

PROJECT_ROOT = Path(__file__).resolve().parents[1]

FORWARD_DIR = Path(
    os.environ.get(
        "A2O_INVERSE_FORWARD_DIR",
        str(PROJECT_ROOT / "outputs/screening/forward"),
    )
)

FORWARD_FILES: List[Path] = [
    Path(os.environ.get("A2O_INVERSE_FORWARD_PART1", str(FORWARD_DIR / "四指标_DO-HRT-进水分配联合正向预测_第1部分.xlsx"))),
    Path(os.environ.get("A2O_INVERSE_FORWARD_PART2", str(FORWARD_DIR / "四指标_DO-HRT-进水分配联合正向预测_第2部分.xlsx"))),
    Path(os.environ.get("A2O_INVERSE_FORWARD_PART3", str(FORWARD_DIR / "四指标_DO-HRT-进水分配联合正向预测_第3部分.xlsx"))),
    Path(os.environ.get("A2O_INVERSE_FORWARD_PART4", str(FORWARD_DIR / "四指标_DO-HRT-进水分配联合正向预测_第4部分.xlsx"))),
    Path(os.environ.get("A2O_INVERSE_FORWARD_PART5", str(FORWARD_DIR / "四指标_DO-HRT-进水分配联合正向预测_第5部分.xlsx"))),
    Path(os.environ.get("A2O_INVERSE_FORWARD_PART6", str(FORWARD_DIR / "四指标_DO-HRT-进水分配联合正向预测_第6部分.xlsx"))),
    Path(os.environ.get("A2O_INVERSE_FORWARD_PART7", str(FORWARD_DIR / "四指标_DO-HRT-进水分配联合正向预测_第7部分.xlsx"))),
    Path(os.environ.get("A2O_INVERSE_FORWARD_PART8", str(FORWARD_DIR / "四指标_DO-HRT-进水分配联合正向预测_第8部分.xlsx"))),
    Path(os.environ.get("A2O_INVERSE_FORWARD_PART9", str(FORWARD_DIR / "四指标_DO-HRT-进水分配联合正向预测_第9部分.xlsx"))),
    Path(os.environ.get("A2O_INVERSE_FORWARD_PART10", str(FORWARD_DIR / "四指标_DO-HRT-进水分配联合正向预测_第10部分.xlsx"))),
    Path(os.environ.get("A2O_INVERSE_FORWARD_PART11", str(FORWARD_DIR / "四指标_DO-HRT-进水分配联合正向预测_第11部分.xlsx"))),
]

OUTPUT_DIR = Path(
    os.environ.get(
        "A2O_INVERSE_OUTPUT_DIR",
        str(PROJECT_ROOT / "outputs/screening/inverse"),
    )
)

OUTPUT_EXCEL = OUTPUT_DIR / "四指标_目标出水反向优化结果.xlsx"
OUTPUT_TEXT = OUTPUT_DIR / "四指标_目标出水反向优化说明.txt"
LEGACY_CHECK_EXCEL = OUTPUT_DIR / "四指标_目标出水反向优化_结果检查包.xlsx"

SCRIPT_VERSION = "A2O_INVERSE_A1_FOUR_TARGETS_v3.4_eleven_forward_files_tp_tertiles"
BUILD_FINGERPRINT = "A1_TP05_TN15_NH45_COD50_NEAR_BEST90_TP_TERTILES_20260922"

STRICT_MODE = os.environ.get("A2O_INVERSE_STRICT", "1").strip() != "0"
EXPECTED_BACKGROUND_N = int(os.environ.get("A2O_INVERSE_EXPECTED_BACKGROUND_N", "631"))
EXPECTED_COMBINATION_N = int(os.environ.get("A2O_INVERSE_EXPECTED_COMBINATION_N", "3780"))
PROGRESS_EVERY = int(os.environ.get("A2O_INVERSE_PROGRESS_EVERY", "50"))

if EXPECTED_BACKGROUND_N <= 0:
    raise ValueError("A2O_INVERSE_EXPECTED_BACKGROUND_N 必须为正整数。")
if EXPECTED_COMBINATION_N <= 0:
    raise ValueError("A2O_INVERSE_EXPECTED_COMBINATION_N 必须为正整数。")
if PROGRESS_EVERY <= 0:
    raise ValueError("A2O_INVERSE_PROGRESS_EVERY 必须为正整数。")

TP_TARGET = 0.5
TN_TARGET = 15.0
NH4_TARGET = 5.0
COD_TARGET = 50.0
NEAR_BEST_RATIO = 0.90

SATISFY_TOL = 1e-12
COMBO_ROUND_DIGITS = 10
Q_CHECK_TOL = 1e-9
REACTOR_VOLUME_L = 57.6
COARSE_SOURCE = "第一轮粗扫描"
CURRENT_SUPPLEMENT_SOURCE = "当前实际工况补充"

FORWARD_FLAG_COLUMNS = [
    "TP_一级A参考满足",
    "TN_一级A参考满足",
    "NH4_一级A参考满足",
    "COD_一级A参考满足",
    "四项一级A参考同时满足",
]

REQUIRED_DETAIL_COLS = [
    "历史状态ID", "样本序号", "Date", "网格来源",
    "候选_DO", "候选_HRT", "候选_Feed", "候选_Q_L_h", "候选_Q_L_d",
    "预测物理检查", "负预测指标",
    "TP_预测", "TN_预测", "NH4_预测", "COD_预测",
] + FORWARD_FLAG_COLUMNS

NUMERIC_DETAIL_COLS = [
    "样本序号", "候选_DO", "候选_HRT", "候选_Feed", "候选_Q_L_h", "候选_Q_L_d",
    "TP_预测", "TN_预测", "NH4_预测", "COD_预测",
]

def now_str() -> str:
    return time.strftime("%Y-%m-%d %H:%M:%S")


def log(message: str) -> None:
    print(f"[{now_str()}] {message}", flush=True)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def safe_float(value: Any, default: float = np.nan) -> float:
    try:
        result = float(value)
        return result if np.isfinite(result) else default
    except Exception:
        return default


def fmt_number(value: Any, digits: int = 6) -> str:
    try:
        number = float(value)
        if not np.isfinite(number):
            return ""
        return f"{number:.{digits}g}"
    except Exception:
        return str(value)


def atomic_write_text(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(path.stem + ".part.txt")
    tmp.write_text(text, encoding="utf-8-sig")
    os.replace(tmp, path)


def cleanup_previous_outputs() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    candidates = [
        OUTPUT_EXCEL,
        OUTPUT_EXCEL.with_name(OUTPUT_EXCEL.stem + ".part.xlsx"),
        OUTPUT_TEXT,
        OUTPUT_TEXT.with_name(OUTPUT_TEXT.stem + ".part.txt"),
        LEGACY_CHECK_EXCEL,
        LEGACY_CHECK_EXCEL.with_name(LEGACY_CHECK_EXCEL.stem + ".part.xlsx"),
    ]
    for path in candidates:
        try:
            if path.exists():
                path.unlink()
        except FileNotFoundError:
            pass


def excel_date_to_timestamp(value: Any) -> pd.Timestamp:
    """兼容 Python 日期、Excel 1900 序列号和普通日期字符串。"""
    if value is None:
        return pd.NaT
    try:
        if pd.isna(value):
            return pd.NaT
    except Exception:
        pass

    if isinstance(value, (pd.Timestamp, datetime, np.datetime64)):
        try:
            ts = pd.Timestamp(value)
            return ts.normalize() if not pd.isna(ts) else pd.NaT
        except Exception:
            return pd.NaT

    numeric_value: Optional[float] = None
    if isinstance(value, (np.integer, int, np.floating, float)) and not isinstance(value, bool):
        try:
            numeric_value = float(value)
        except Exception:
            numeric_value = None
    elif isinstance(value, str):
        text = value.strip()
        if not text:
            return pd.NaT
        try:
            numeric_value = float(text)
        except Exception:
            numeric_value = None

    if numeric_value is not None and np.isfinite(numeric_value) and 1.0 <= numeric_value <= 100000.0:
        try:
            ts = pd.to_datetime(numeric_value, unit="D", origin="1899-12-30", errors="coerce")
            return pd.Timestamp(ts).normalize() if not pd.isna(ts) else pd.NaT
        except Exception:
            return pd.NaT

    try:
        ts = pd.to_datetime(value, errors="coerce")
        return pd.Timestamp(ts).normalize() if not pd.isna(ts) else pd.NaT
    except Exception:
        return pd.NaT


def combo_key(do_value: float, hrt_value: float, feed_value: float) -> Tuple[float, float, float]:
    return (
        round(float(do_value), COMBO_ROUND_DIGITS),
        round(float(hrt_value), COMBO_ROUND_DIGITS),
        round(float(feed_value), COMBO_ROUND_DIGITS),
    )


def recommendation_combo_id(do_value: float, hrt_value: float, feed_value: float) -> str:
    return f"DO{do_value:.6g}_HRT{hrt_value:.6g}_Feed{feed_value:.6g}"


# =============================================================================
# 3. 十一个正向结果文件的流式读取
# =============================================================================

def check_input_files() -> pd.DataFrame:
    rows: List[Dict[str, Any]] = []
    missing: List[str] = []
    for i, path in enumerate(FORWARD_FILES, start=1):
        exists = path.exists()
        if not exists:
            missing.append(str(path))
        rows.append(
            {
                "结果部分": f"第{i}部分",
                "输入文件": str(path),
                "文件是否存在": "是" if exists else "否",
                "文件大小_MB": path.stat().st_size / 1024 / 1024 if exists else np.nan,
                "SHA256": sha256_file(path) if exists else "",
            }
        )
    if missing:
        raise FileNotFoundError("以下正向预测输入文件不存在：\n" + "\n".join(missing))
    return pd.DataFrame(rows)



def read_background_influent_tp(path: Path) -> pd.DataFrame:
    """
    从第1部分正向结果的“03_历史状态”读取每个历史背景的进水TP。

    该表保存的是历史状态原始值。若 Inf.TP 存在短内部缺失，则按照正向阶段
    相同的样本序号线性插值思路（最大缺口 7 个样本序号）补齐；本研究当前
    数据中该规则可补齐全部缺失。补齐后的值只用于进水TP分组，不改写任何
    已经完成的正向预测结果。
    """
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet_names = [name for name in wb.sheetnames if str(name).strip() == "03_历史状态"]
        if not sheet_names:
            raise ValueError(f"第1部分正向结果未找到‘03_历史状态’：{path}")
        ws = wb[sheet_names[0]]
        rows = ws.iter_rows(values_only=True)
        try:
            headers = ["" if value is None else str(value).strip() for value in next(rows)]
        except StopIteration:
            raise ValueError("‘03_历史状态’为空。")

        required = ["历史状态ID", "样本序号", "Inf.TP"]
        missing = [column for column in required if column not in headers]
        if missing:
            raise ValueError(f"‘03_历史状态’缺少进水TP分组所需字段：{missing}")

        idx_bg = headers.index("历史状态ID")
        idx_sample = headers.index("样本序号")
        idx_tp = headers.index("Inf.TP")
        records: List[Dict[str, Any]] = []
        for row in rows:
            bg_id = "" if row[idx_bg] is None else str(row[idx_bg]).strip()
            if not bg_id:
                continue
            records.append({
                "历史状态ID": bg_id,
                "样本序号": safe_float(row[idx_sample]),
                "进水TP原始值_mg_L": safe_float(row[idx_tp]),
            })
    finally:
        wb.close()

    out = pd.DataFrame(records)
    if out.empty:
        raise ValueError("‘03_历史状态’没有可用于进水TP分组的历史背景。")
    if out["历史状态ID"].duplicated().any():
        dup = out.loc[out["历史状态ID"].duplicated(), "历史状态ID"].astype(str).tolist()[:10]
        raise ValueError(f"‘03_历史状态’存在重复历史状态ID：{dup}")
    if out["样本序号"].isna().any():
        raise ValueError("‘03_历史状态’存在无效样本序号，无法按正向规则补齐进水TP。")

    out = out.sort_values(["样本序号", "历史状态ID"], kind="mergesort").reset_index(drop=True)
    values = out["进水TP原始值_mg_L"].to_numpy(dtype=float).copy()
    coordinates = out["样本序号"].to_numpy(dtype=float)
    original_missing = int(np.sum(~np.isfinite(values)))

    # 与正向预处理一致：只补内部短缺口，按样本序号坐标做线性插值。
    i = 0
    n = len(values)
    while i < n:
        if np.isfinite(values[i]):
            i += 1
            continue
        start = i
        while i + 1 < n and not np.isfinite(values[i + 1]):
            i += 1
        end = i
        left = start - 1
        right = end + 1
        if (
            left >= 0 and right < n
            and np.isfinite(values[left]) and np.isfinite(values[right])
        ):
            original_gap_span = coordinates[right] - coordinates[left] - 1.0
            coordinate_span = coordinates[right] - coordinates[left]
            if 0 <= original_gap_span <= 7 and coordinate_span > 0:
                left_value = float(values[left])
                right_value = float(values[right])
                for pos in range(start, end + 1):
                    fraction = (coordinates[pos] - coordinates[left]) / coordinate_span
                    values[pos] = left_value + fraction * (right_value - left_value)
        i += 1

    remaining_missing = int(np.sum(~np.isfinite(values)))
    if remaining_missing > 0:
        if STRICT_MODE:
            raise ValueError(
                f"进水TP按正向短缺口插值后仍有 {remaining_missing} 个缺失值；"
                "严格模式下不使用新的临时填补规则，请先核对正向‘03_历史状态’。"
            )
        fallback = float(np.nanmedian(out["进水TP原始值_mg_L"].to_numpy(dtype=float)))
        values = np.where(np.isfinite(values), values, fallback)

    out["进水TP分组使用值_mg_L"] = values
    out["进水TP原始缺失是否已补齐"] = np.where(
        out["进水TP原始值_mg_L"].isna(), "是", "否"
    )
    out.attrs["original_missing_n"] = original_missing
    out.attrs["remaining_missing_n"] = remaining_missing
    return out


def get_detail_sheet_names_from_workbook(wb: Any, path: Path) -> List[str]:
    names = [name for name in wb.sheetnames if str(name).startswith("明细_")]
    if not names:
        raise ValueError(f"正向预测文件未找到‘明细_’副表：{path}")
    return names


def iter_background_candidate_frames(path: Path) -> Iterator[Tuple[str, pd.DataFrame, List[str]]]:
    """
    只读、流式读取一个正向结果 Excel 的全部明细副表，并按历史状态连续分组。

    正向脚本按历史状态连续写出；即使某个背景恰好跨明细副表边界，
    current_rows 也会跨副表继续累积，直到历史状态ID改变才 yield。
    """
    wb = load_workbook(path, read_only=True, data_only=True)
    detail_names = get_detail_sheet_names_from_workbook(wb, path)

    current_id: Optional[str] = None
    current_rows: List[List[Any]] = []
    canonical_headers = list(REQUIRED_DETAIL_COLS)

    try:
        for sheet_name in detail_names:
            ws = wb[sheet_name]
            row_iter = ws.iter_rows(values_only=True)
            try:
                headers = ["" if value is None else str(value).strip() for value in next(row_iter)]
            except StopIteration:
                continue

            missing = [column for column in REQUIRED_DETAIL_COLS if column not in headers]
            if missing:
                raise ValueError(f"{path.name}/{sheet_name} 缺少新版反向分析必需字段：{missing}")

            selected_indices = [headers.index(column) for column in canonical_headers]

            for raw_row in row_iter:
                values = [raw_row[idx] for idx in selected_indices]
                bg_id = "" if values[0] is None else str(values[0]).strip()
                if not bg_id:
                    continue

                if current_id is None:
                    current_id = bg_id
                if bg_id != current_id:
                    frame = pd.DataFrame(current_rows, columns=canonical_headers)
                    yield current_id, frame, detail_names
                    current_id = bg_id
                    current_rows = []
                current_rows.append(values)

        if current_id is not None and current_rows:
            frame = pd.DataFrame(current_rows, columns=canonical_headers)
            yield current_id, frame, detail_names
    finally:
        wb.close()


def prepare_background_frame(frame: pd.DataFrame) -> pd.DataFrame:
    df = frame.copy()
    for column in NUMERIC_DETAIL_COLS:
        df[column] = pd.to_numeric(df[column], errors="coerce")
    return df


# =============================================================================
# 4. Excel 输出
# =============================================================================

def _column_width(column_name: str, series: pd.Series) -> float:
    name = str(column_name)
    if "说明" in name or "规则" in name or "路径" in name or "备注" in name or "核查" in name:
        return 34
    if "SHA256" in name:
        return 68
    if "Date" in name:
        return 13
    if "历史状态ID" in name or "组合ID" in name or "推荐标签" in name:
        return 20
    if "比例" in name or "预测" in name or "目标" in name:
        return 16
    if pd.api.types.is_numeric_dtype(series):
        return 14
    try:
        max_len = max([len(name)] + [len(str(x)) for x in series.dropna().astype(str).head(200)])
        return min(max(max_len + 2, 12), 28)
    except Exception:
        return 14




# =============================================================================
# 4. Excel 输出
# =============================================================================

def _column_width(column_name: str, series: pd.Series) -> float:
    name = str(column_name)
    if any(key in name for key in ["说明", "规则", "路径", "备注", "核查", "含义"]):
        return 36
    if "SHA256" in name:
        return 68
    if "Date" in name:
        return 13
    if "历史状态ID" in name or "组合ID" in name or "推荐标签" in name:
        return 22
    if "比例" in name or "预测" in name or "目标" in name or "记录数" in name or "背景数" in name:
        return 18
    if pd.api.types.is_numeric_dtype(series):
        return 14
    try:
        max_len = max([len(name)] + [len(str(x)) for x in series.dropna().astype(str).head(200)])
        return min(max(max_len + 2, 12), 32)
    except Exception:
        return 14


def write_result_excel(
    settings_df: pd.DataFrame,
    summary_df: pd.DataFrame,
    condition_df: pd.DataFrame,
    range_df: pd.DataFrame,
    combo_df: pd.DataFrame,
    near_best_df: pd.DataFrame,
    recommendation_df: pd.DataFrame,
    background_df: pd.DataFrame,
    checks_df: pd.DataFrame,
    tp_group_summary_df: pd.DataFrame,
    tp_group_range_df: pd.DataFrame,
    tp_group_combo_df: pd.DataFrame,
) -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    tmp = OUTPUT_EXCEL.with_name(OUTPUT_EXCEL.stem + ".part.xlsx")
    if tmp.exists():
        tmp.unlink()

    with pd.ExcelWriter(
        tmp,
        engine="xlsxwriter",
        datetime_format="yyyy-mm-dd",
        engine_kwargs={"options": {"strings_to_urls": False, "nan_inf_to_errors": True}},
    ) as writer:
        tables = [
            ("01_说明与目标", settings_df),
            ("02_筛选汇总", summary_df),
            ("03_各项条件影响", condition_df),
            ("04_参数范围", range_df),
            ("05_参数组合统计", combo_df),
            ("06_接近最佳组合", near_best_df),
            ("07_三个代表方案", recommendation_df),
            ("08_背景筛选统计", background_df),
            ("09_方法检查", checks_df),
            ("10_进水TP分组汇总", tp_group_summary_df),
            ("11_各组参数范围", tp_group_range_df),
            ("12_各组参数组合", tp_group_combo_df),
        ]

        workbook = writer.book
        header_fmt = workbook.add_format({
            "bold": True, "bg_color": "#D9EAF7", "border": 1,
            "align": "center", "valign": "vcenter", "text_wrap": True,
        })
        percent_fmt = workbook.add_format({"num_format": "0.0000"})
        decimal_fmt = workbook.add_format({"num_format": "0.000000"})
        integer_fmt = workbook.add_format({"num_format": "0"})
        date_fmt = workbook.add_format({"num_format": "yyyy-mm-dd"})
        highlight_fmt = workbook.add_format({"bg_color": "#FFF2CC", "bold": True})

        for sheet_name, frame in tables:
            frame_to_write = frame.copy()
            frame_to_write.to_excel(writer, sheet_name=sheet_name, index=False)
            ws = writer.sheets[sheet_name]
            ws.freeze_panes(1, 0)
            if len(frame_to_write.columns) > 0:
                ws.autofilter(0, 0, max(len(frame_to_write), 1), len(frame_to_write.columns) - 1)

            for col_idx, col in enumerate(frame_to_write.columns):
                ws.write(0, col_idx, col, header_fmt)
                width = _column_width(str(col), frame_to_write[col])
                cell_format = None
                if "比例%" in str(col) or str(col).endswith("比例"):
                    cell_format = percent_fmt
                elif str(col) == "Date":
                    cell_format = date_fmt
                elif pd.api.types.is_integer_dtype(frame_to_write[col]):
                    cell_format = integer_fmt
                elif pd.api.types.is_float_dtype(frame_to_write[col]):
                    cell_format = decimal_fmt
                ws.set_column(col_idx, col_idx, width, cell_format)

            if sheet_name == "05_参数组合统计" and not frame_to_write.empty and "推荐标签" in frame_to_write.columns:
                rec_col = frame_to_write.columns.get_loc("推荐标签")
                ws.conditional_format(1, rec_col, len(frame_to_write), rec_col, {
                    "type": "text", "criteria": "containing", "value": "方案", "format": highlight_fmt,
                })
            if sheet_name == "06_接近最佳组合" and not frame_to_write.empty:
                ws.set_tab_color("#70AD47")
            if sheet_name == "07_三个代表方案" and not frame_to_write.empty:
                ws.set_tab_color("#ED7D31")
            if sheet_name == "10_进水TP分组汇总" and not frame_to_write.empty:
                ws.set_tab_color("#5B9BD5")
            if sheet_name == "11_各组参数范围" and not frame_to_write.empty:
                ws.set_tab_color("#A5A5A5")
            if sheet_name == "12_各组参数组合" and not frame_to_write.empty:
                ws.set_tab_color("#4472C4")

    os.replace(tmp, OUTPUT_EXCEL)


# =============================================================================
# 5. 主计算
# =============================================================================

def main() -> None:
    start = time.time()
    cleanup_previous_outputs()

    log("开始四指标一级A参考目标反向筛选。")
    log(
        f"固定目标：TP <= {TP_TARGET:g}；TN <= {TN_TARGET:g}；"
        f"NH4+-N <= {NH4_TARGET:g}；COD <= {COD_TARGET:g} mg/L。"
    )
    log("正式输入仅使用十一个正向Excel中的‘第一轮粗扫描’记录。")

    try:
        input_files_df = check_input_files()

        canonical_keys: Optional[List[Tuple[float, float, float]]] = None
        canonical_key_to_pos: Dict[Tuple[float, float, float], int] = {}
        combo_do: Optional[np.ndarray] = None
        combo_hrt: Optional[np.ndarray] = None
        combo_feed: Optional[np.ndarray] = None
        combo_q_h: Optional[np.ndarray] = None
        combo_q_d: Optional[np.ndarray] = None

        tp_matrix: Optional[np.ndarray] = None
        tn_matrix: Optional[np.ndarray] = None
        nh4_matrix: Optional[np.ndarray] = None
        cod_matrix: Optional[np.ndarray] = None
        physical_matrix: Optional[np.ndarray] = None

        background_meta: List[Dict[str, Any]] = []
        seen_backgrounds: set[str] = set()
        source_file_stats: List[Dict[str, Any]] = []

        total_detail_rows = 0
        total_coarse_rows = 0
        total_supplement_rows = 0
        total_unknown_source_rows = 0
        total_physical_fail_rows = 0
        total_negative_flag_rows = 0
        total_nonfinite_prediction_rows = 0
        q_inconsistency_rows = 0
        order_realign_backgrounds = 0
        duplicate_combo_backgrounds = 0
        incomplete_combo_backgrounds = 0
        forward_flag_mismatch_rows = 0
        forward_flag_mismatch_examples: List[str] = []

        bg_index = 0

        for part_no, path in enumerate(FORWARD_FILES, start=1):
            log(f"读取正向预测第{part_no}部分：{path.name}")
            file_bg_count = 0
            file_detail_rows = 0
            file_coarse_rows = 0
            file_supplement_rows = 0
            detail_sheet_names_seen: List[str] = []

            for bg_id, raw_frame, detail_names in iter_background_candidate_frames(path):
                detail_sheet_names_seen = detail_names
                file_bg_count += 1
                bg_index += 1

                if bg_id in seen_backgrounds:
                    raise ValueError(f"十一个正向结果文件之间出现重复历史状态ID：{bg_id}")
                seen_backgrounds.add(bg_id)

                if bg_index > EXPECTED_BACKGROUND_N:
                    raise ValueError(
                        f"读取到的历史背景数已超过预期 {EXPECTED_BACKGROUND_N}；"
                        "如确有新数据，请显式修改 A2O_INVERSE_EXPECTED_BACKGROUND_N。"
                    )
                if bg_index == 1 or bg_index % PROGRESS_EVERY == 0 or bg_index == EXPECTED_BACKGROUND_N:
                    log(f"处理历史状态：{bg_index}/{EXPECTED_BACKGROUND_N}，当前={bg_id}")

                df = prepare_background_frame(raw_frame)
                total_detail_rows += len(df)
                file_detail_rows += len(df)

                source = df["网格来源"].astype(str).str.strip()
                coarse_mask = source.eq(COARSE_SOURCE)
                supplement_mask = source.eq(CURRENT_SUPPLEMENT_SOURCE)
                known_source_mask = coarse_mask | supplement_mask

                coarse = df.loc[coarse_mask].copy().reset_index(drop=True)
                supplement_n = int(supplement_mask.sum())
                unknown_n = int((~known_source_mask).sum())

                total_coarse_rows += len(coarse)
                total_supplement_rows += supplement_n
                total_unknown_source_rows += unknown_n
                file_coarse_rows += len(coarse)
                file_supplement_rows += supplement_n

                if STRICT_MODE and unknown_n > 0:
                    bad_values = sorted(set(source.loc[~known_source_mask].tolist()))
                    raise ValueError(f"{bg_id} 出现未知网格来源：{bad_values[:10]}")
                if coarse.empty:
                    raise ValueError(f"{bg_id} 没有‘{COARSE_SOURCE}’记录。")
                if coarse[["候选_DO", "候选_HRT", "候选_Feed"]].isna().any().any():
                    raise ValueError(f"{bg_id} 粗网格运行参数存在缺失。")

                duplicate_n = int(coarse[["候选_DO", "候选_HRT", "候选_Feed"]].duplicated().sum())
                if duplicate_n > 0:
                    duplicate_combo_backgrounds += 1
                    if STRICT_MODE:
                        raise ValueError(f"{bg_id} 粗网格存在 {duplicate_n} 条重复 DO×HRT×Feed 组合。")

                keys = [
                    combo_key(d, h, f)
                    for d, h, f in coarse[["候选_DO", "候选_HRT", "候选_Feed"]].to_numpy(dtype=float)
                ]

                if canonical_keys is None:
                    canonical_keys = keys
                    if len(canonical_keys) != EXPECTED_COMBINATION_N:
                        raise ValueError(
                            f"首个背景粗网格组合数={len(canonical_keys)}，预期={EXPECTED_COMBINATION_N}。"
                        )
                    if len(set(canonical_keys)) != len(canonical_keys):
                        raise ValueError("首个背景的标准粗网格存在重复组合。")

                    canonical_key_to_pos = {key: i for i, key in enumerate(canonical_keys)}
                    combo_n = len(canonical_keys)
                    combo_do = coarse["候选_DO"].to_numpy(dtype=float)
                    combo_hrt = coarse["候选_HRT"].to_numpy(dtype=float)
                    combo_feed = coarse["候选_Feed"].to_numpy(dtype=float)
                    combo_q_h = coarse["候选_Q_L_h"].to_numpy(dtype=float)
                    combo_q_d = coarse["候选_Q_L_d"].to_numpy(dtype=float)

                    shape = (combo_n, EXPECTED_BACKGROUND_N)
                    tp_matrix = np.full(shape, np.nan, dtype=np.float64)
                    tn_matrix = np.full(shape, np.nan, dtype=np.float64)
                    nh4_matrix = np.full(shape, np.nan, dtype=np.float64)
                    cod_matrix = np.full(shape, np.nan, dtype=np.float64)
                    physical_matrix = np.zeros(shape, dtype=bool)

                    expected_q_h = REACTOR_VOLUME_L / combo_hrt
                    expected_q_d = REACTOR_VOLUME_L * 24.0 / combo_hrt
                    bad_q = (
                        ~np.isfinite(combo_q_h) | ~np.isfinite(combo_q_d)
                        | (np.abs(combo_q_h - expected_q_h) > Q_CHECK_TOL)
                        | (np.abs(combo_q_d - expected_q_d) > Q_CHECK_TOL)
                    )
                    q_inconsistency_rows += int(np.sum(bad_q))
                    if STRICT_MODE and np.any(bad_q):
                        raise ValueError("首个背景候选 Q 与 57.6/HRT 换算不一致。")
                    order = np.arange(combo_n, dtype=int)
                else:
                    assert combo_do is not None and combo_hrt is not None and combo_feed is not None
                    assert combo_q_h is not None and combo_q_d is not None
                    assert tp_matrix is not None and tn_matrix is not None
                    assert nh4_matrix is not None and cod_matrix is not None and physical_matrix is not None

                    if len(keys) != len(canonical_keys):
                        incomplete_combo_backgrounds += 1
                        raise ValueError(f"{bg_id} 粗网格组合数={len(keys)}，标准组合数={len(canonical_keys)}。")
                    if keys == canonical_keys:
                        order = np.arange(len(keys), dtype=int)
                    else:
                        key_set = set(keys)
                        canonical_set = set(canonical_keys)
                        if key_set != canonical_set:
                            incomplete_combo_backgrounds += 1
                            raise ValueError(
                                f"{bg_id} 粗网格组合集合与标准网格不一致。"
                                f"缺失示例={list(canonical_set-key_set)[:5]}；新增示例={list(key_set-canonical_set)[:5]}"
                            )
                        order_realign_backgrounds += 1
                        order = np.asarray([canonical_key_to_pos[key] for key in keys], dtype=int)

                combo_n = len(canonical_keys)
                tp_raw = coarse["TP_预测"].to_numpy(dtype=float)
                tn_raw = coarse["TN_预测"].to_numpy(dtype=float)
                nh4_raw = coarse["NH4_预测"].to_numpy(dtype=float)
                cod_raw = coarse["COD_预测"].to_numpy(dtype=float)
                physical_raw = coarse["预测物理检查"].astype(str).str.strip().eq("通过").to_numpy(dtype=bool)

                finite_all_raw = np.isfinite(tp_raw) & np.isfinite(tn_raw) & np.isfinite(nh4_raw) & np.isfinite(cod_raw)
                nonfinite_n = int(np.sum(~finite_all_raw))
                total_nonfinite_prediction_rows += nonfinite_n

                # 反向程序独立重算正向文件中的五个一级A标签，仅作一致性核查。
                expected_tp_flag = np.isfinite(tp_raw) & (tp_raw >= -SATISFY_TOL) & (tp_raw <= TP_TARGET + SATISFY_TOL)
                expected_tn_flag = np.isfinite(tn_raw) & (tn_raw >= -SATISFY_TOL) & (tn_raw <= TN_TARGET + SATISFY_TOL)
                expected_nh4_flag = np.isfinite(nh4_raw) & (nh4_raw >= -SATISFY_TOL) & (nh4_raw <= NH4_TARGET + SATISFY_TOL)
                expected_cod_flag = np.isfinite(cod_raw) & (cod_raw >= -SATISFY_TOL) & (cod_raw <= COD_TARGET + SATISFY_TOL)
                expected_all_flag = physical_raw & finite_all_raw & expected_tp_flag & expected_tn_flag & expected_nh4_flag & expected_cod_flag

                expected_flags = [expected_tp_flag, expected_tn_flag, expected_nh4_flag, expected_cod_flag, expected_all_flag]
                for flag_col, expected_flag in zip(FORWARD_FLAG_COLUMNS, expected_flags):
                    actual_flag = coarse[flag_col].fillna("").astype(str).str.strip().eq("是").to_numpy(dtype=bool)
                    mismatch = actual_flag != expected_flag
                    mismatch_n = int(np.sum(mismatch))
                    forward_flag_mismatch_rows += mismatch_n
                    if mismatch_n and len(forward_flag_mismatch_examples) < 10:
                        bad_pos = np.flatnonzero(mismatch)[:3]
                        for pos in bad_pos:
                            forward_flag_mismatch_examples.append(f"{bg_id}/{flag_col}/行{int(pos)+1}")

                tp_c = np.full(combo_n, np.nan, dtype=float)
                tn_c = np.full(combo_n, np.nan, dtype=float)
                nh4_c = np.full(combo_n, np.nan, dtype=float)
                cod_c = np.full(combo_n, np.nan, dtype=float)
                physical_c = np.zeros(combo_n, dtype=bool)
                tp_c[order], tn_c[order], nh4_c[order], cod_c[order] = tp_raw, tn_raw, nh4_raw, cod_raw
                physical_c[order] = physical_raw

                col = bg_index - 1
                tp_matrix[:, col] = tp_c
                tn_matrix[:, col] = tn_c
                nh4_matrix[:, col] = nh4_c
                cod_matrix[:, col] = cod_c
                physical_matrix[:, col] = physical_c

                total_physical_fail_rows += int(np.sum(~physical_raw))
                negative_text = coarse["负预测指标"].fillna("").astype(str).str.strip()
                total_negative_flag_rows += int(negative_text.ne("").sum())

                q_h_raw = coarse["候选_Q_L_h"].to_numpy(dtype=float)
                q_d_raw = coarse["候选_Q_L_d"].to_numpy(dtype=float)
                hrt_raw = coarse["候选_HRT"].to_numpy(dtype=float)
                expected_q_h_raw = REACTOR_VOLUME_L / hrt_raw
                expected_q_d_raw = REACTOR_VOLUME_L * 24.0 / hrt_raw
                bad_q_raw = (
                    ~np.isfinite(q_h_raw) | ~np.isfinite(q_d_raw)
                    | (np.abs(q_h_raw - expected_q_h_raw) > Q_CHECK_TOL)
                    | (np.abs(q_d_raw - expected_q_d_raw) > Q_CHECK_TOL)
                )
                q_bad_n = int(np.sum(bad_q_raw))
                q_inconsistency_rows += q_bad_n
                if STRICT_MODE and q_bad_n > 0:
                    raise ValueError(f"{bg_id} 有 {q_bad_n} 条候选 Q 与 HRT 换算不一致。")

                first_row = df.iloc[0]
                background_meta.append({
                    "历史状态ID": bg_id,
                    "样本序号": safe_float(first_row.get("样本序号")),
                    "Date": excel_date_to_timestamp(first_row.get("Date")),
                    "来源文件": path.name,
                    "正式粗网格记录数": int(len(coarse)),
                    "当前工况补充记录数": supplement_n,
                })

            source_file_stats.append({
                "结果部分": f"第{part_no}部分", "输入文件": path.name,
                "明细副表数": len(detail_sheet_names_seen), "历史背景数": file_bg_count,
                "全部明细记录数": file_detail_rows, "正式粗网格记录数": file_coarse_rows,
                "当前工况补充记录数": file_supplement_rows,
            })

        if canonical_keys is None:
            raise RuntimeError("十一个正向预测文件均未读取到正式粗网格记录。")
        assert combo_do is not None and combo_hrt is not None and combo_feed is not None
        assert combo_q_h is not None and combo_q_d is not None
        assert tp_matrix is not None and tn_matrix is not None and nh4_matrix is not None and cod_matrix is not None
        assert physical_matrix is not None

        background_n = bg_index
        combination_n = len(canonical_keys)
        tp_matrix = tp_matrix[:, :background_n]
        tn_matrix = tn_matrix[:, :background_n]
        nh4_matrix = nh4_matrix[:, :background_n]
        cod_matrix = cod_matrix[:, :background_n]
        physical_matrix = physical_matrix[:, :background_n]

        expected_coarse_rows = background_n * combination_n
        if STRICT_MODE and background_n != EXPECTED_BACKGROUND_N:
            raise ValueError(f"共同历史背景数={background_n}，预期={EXPECTED_BACKGROUND_N}。")
        if STRICT_MODE and combination_n != EXPECTED_COMBINATION_N:
            raise ValueError(f"统一三参数组合数={combination_n}，预期={EXPECTED_COMBINATION_N}。")
        if STRICT_MODE and total_coarse_rows != expected_coarse_rows:
            raise ValueError(f"正式粗网格总数={total_coarse_rows}，理论={expected_coarse_rows}。")
        if STRICT_MODE and total_detail_rows != total_coarse_rows + total_supplement_rows + total_unknown_source_rows:
            raise ValueError("全部明细数与粗网格/补充/未知来源记录数之和不一致。")

        bg_ids = [row["历史状态ID"] for row in background_meta]
        if all(re.fullmatch(r"B\d{4}", x) for x in bg_ids):
            expected_ids = [f"B{i:04d}" for i in range(1, background_n + 1)]
            if STRICT_MODE and bg_ids != expected_ids:
                raise ValueError("历史背景ID没有按预期完整连续覆盖。")

        # 进水TP分组使用第1部分正向结果中的“03_历史状态”。
        # 只用于后续分组分析，不改写已经完成的正向预测。
        influent_tp_df = read_background_influent_tp(FORWARD_FILES[0])
        influent_tp_lookup = influent_tp_df.set_index("历史状态ID")
        missing_tp_backgrounds = [bg_id for bg_id in bg_ids if bg_id not in influent_tp_lookup.index]
        if missing_tp_backgrounds:
            raise ValueError(f"‘03_历史状态’缺少反向分析背景：{missing_tp_backgrounds[:10]}")
        extra_tp_backgrounds = [bg_id for bg_id in influent_tp_lookup.index.astype(str) if bg_id not in set(bg_ids)]
        if STRICT_MODE and extra_tp_backgrounds:
            raise ValueError(f"‘03_历史状态’存在反向分析未使用的额外背景：{extra_tp_backgrounds[:10]}")

        influent_tp_raw = np.asarray([
            safe_float(influent_tp_lookup.loc[bg_id, "进水TP原始值_mg_L"]) for bg_id in bg_ids
        ], dtype=float)
        influent_tp_values = np.asarray([
            safe_float(influent_tp_lookup.loc[bg_id, "进水TP分组使用值_mg_L"]) for bg_id in bg_ids
        ], dtype=float)
        if not np.all(np.isfinite(influent_tp_values)):
            raise ValueError("进水TP分组使用值仍存在缺失或非有限值。")

        tp_q1 = float(pd.Series(influent_tp_values).quantile(1.0 / 3.0, interpolation="linear"))
        tp_q2 = float(pd.Series(influent_tp_values).quantile(2.0 / 3.0, interpolation="linear"))
        tp_group_labels = np.empty(background_n, dtype=object)
        tp_group_labels[influent_tp_values <= tp_q1 + SATISFY_TOL] = "低进水TP"
        tp_group_labels[(influent_tp_values > tp_q1 + SATISFY_TOL) & (influent_tp_values <= tp_q2 + SATISFY_TOL)] = "中等进水TP"
        tp_group_labels[influent_tp_values > tp_q2 + SATISFY_TOL] = "高进水TP"

        # ---------------------------------------------------------------------
        # 四项一级A正式筛选
        # ---------------------------------------------------------------------
        finite_all = np.isfinite(tp_matrix) & np.isfinite(tn_matrix) & np.isfinite(nh4_matrix) & np.isfinite(cod_matrix)
        formal_valid = physical_matrix & finite_all
        tp_meet = formal_valid & (tp_matrix >= -SATISFY_TOL) & (tp_matrix <= TP_TARGET + SATISFY_TOL)
        tn_meet = formal_valid & (tn_matrix >= -SATISFY_TOL) & (tn_matrix <= TN_TARGET + SATISFY_TOL)
        nh4_meet = formal_valid & (nh4_matrix >= -SATISFY_TOL) & (nh4_matrix <= NH4_TARGET + SATISFY_TOL)
        cod_meet = formal_valid & (cod_matrix >= -SATISFY_TOL) & (cod_matrix <= COD_TARGET + SATISFY_TOL)
        all4_meet = tp_meet & tn_meet & nh4_meet & cod_meet

        physical_counts = np.sum(formal_valid, axis=1).astype(int)
        tp_counts = np.sum(tp_meet, axis=1).astype(int)
        tn_counts = np.sum(tn_meet, axis=1).astype(int)
        nh4_counts = np.sum(nh4_meet, axis=1).astype(int)
        cod_counts = np.sum(cod_meet, axis=1).astype(int)
        all4_counts = np.sum(all4_meet, axis=1).astype(int)

        n_valid = int(np.sum(formal_valid))
        n_tp = int(np.sum(tp_meet))
        n_tn = int(np.sum(tn_meet))
        n_nh4 = int(np.sum(nh4_meet))
        n_cod = int(np.sum(cod_meet))
        n_all4 = int(np.sum(all4_meet))
        screening_pct = 100.0 * n_all4 / total_coarse_rows if total_coarse_rows else np.nan

        # 从“历史背景”角度补充三个直观统计：
        # 1) 至少存在1组运行参数能够四项同时满足；
        # 2) 3780组参数都没有找到四项同时满足；
        # 3) 3780组参数全部四项同时满足。
        bg_all4_counts = np.sum(all4_meet, axis=0).astype(int)
        backgrounds_with_any_all4 = int(np.sum(bg_all4_counts > 0))
        backgrounds_with_no_all4 = int(np.sum(bg_all4_counts == 0))
        backgrounds_with_all_combos_all4 = int(np.sum(bg_all4_counts == combination_n))

        # ---------------------------------------------------------------------
        # 按进水TP三分位数分组，再分别比较3780组运行参数
        # ---------------------------------------------------------------------
        tp_group_order = ["低进水TP", "中等进水TP", "高进水TP"]
        tp_group_rule_text = {
            "低进水TP": f"Inf.TP <= {tp_q1:.6f} mg/L",
            "中等进水TP": f"{tp_q1:.6f} < Inf.TP <= {tp_q2:.6f} mg/L",
            "高进水TP": f"Inf.TP > {tp_q2:.6f} mg/L",
        }
        tp_group_summary_rows: List[Dict[str, Any]] = []
        tp_group_range_rows: List[Dict[str, Any]] = []
        tp_group_combo_frames: List[pd.DataFrame] = []
        tp_group_runtime: Dict[str, Dict[str, Any]] = {}

        for group_name in tp_group_order:
            bg_mask = tp_group_labels == group_name
            group_bg_n = int(np.sum(bg_mask))
            if group_bg_n <= 0:
                raise RuntimeError(f"{group_name}没有历史背景，无法进行分组统计。")

            group_tp_values = influent_tp_values[bg_mask]
            group_valid_counts = np.sum(formal_valid[:, bg_mask], axis=1).astype(int)
            group_tp_counts = np.sum(tp_meet[:, bg_mask], axis=1).astype(int)
            group_tn_counts = np.sum(tn_meet[:, bg_mask], axis=1).astype(int)
            group_nh4_counts = np.sum(nh4_meet[:, bg_mask], axis=1).astype(int)
            group_cod_counts = np.sum(cod_meet[:, bg_mask], axis=1).astype(int)
            group_all4_counts = np.sum(all4_meet[:, bg_mask], axis=1).astype(int)

            group_total_rows = group_bg_n * combination_n
            group_valid_rows = int(np.sum(formal_valid[:, bg_mask]))
            group_all4_rows = int(np.sum(all4_meet[:, bg_mask]))
            group_all4_pct = 100.0 * group_all4_rows / float(group_total_rows)
            group_best_count = int(np.max(group_all4_counts)) if len(group_all4_counts) else 0
            group_best_pct = 100.0 * group_best_count / float(group_bg_n)
            group_threshold = int(math.ceil(NEAR_BEST_RATIO * group_best_count - 1e-12)) if group_best_count > 0 else 0
            group_good_mask = (
                (group_all4_counts >= group_threshold) & (group_all4_counts > 0)
                if group_best_count > 0 else np.zeros(combination_n, dtype=bool)
            )
            group_best_mask = group_all4_counts == group_best_count if group_best_count > 0 else np.zeros(combination_n, dtype=bool)
            group_good_n = int(np.sum(group_good_mask))
            group_best_n = int(np.sum(group_best_mask))

            group_best_idx: Optional[int] = None
            if group_best_count > 0:
                candidates = np.flatnonzero(group_best_mask)
                order = np.lexsort((combo_feed[candidates], combo_hrt[candidates], combo_do[candidates]))
                group_best_idx = int(candidates[int(order[0])])

            tp_group_runtime[group_name] = {
                "背景数": group_bg_n,
                "总预测情景数": group_total_rows,
                "有效预测情景数": group_valid_rows,
                "达标预测情景数": group_all4_rows,
                "总体达标比例%": group_all4_pct,
                "最佳达标背景数": group_best_count,
                "最佳达标比例%": group_best_pct,
                "90%门槛背景数": group_threshold,
                "表现较好组合数": group_good_n,
                "并列最佳组合数": group_best_n,
                "代表最佳组合索引": group_best_idx,
            }

            tp_group_summary_rows.append({
                "进水TP水平": group_name,
                "分组规则": tp_group_rule_text[group_name],
                "本组实际进水TP最小值_mg_L": float(np.min(group_tp_values)),
                "本组实际进水TP最大值_mg_L": float(np.max(group_tp_values)),
                "历史背景数": group_bg_n,
                "运行参数组合数": combination_n,
                "全部预测情景数": group_total_rows,
                "通过物理和数值检查情景数": group_valid_rows,
                "四个出水指标同时达到要求情景数": group_all4_rows,
                "四个出水指标同时达到要求比例%": group_all4_pct,
                "本组最佳组合达标背景数": group_best_count,
                "本组最佳组合达标比例%": group_best_pct,
                "达到本组最佳90%以上所需最少达标背景数": group_threshold,
                "达到本组最佳90%以上的参数组合数": group_good_n,
                "并列最佳参数组合数": group_best_n,
                "代表最佳_DO_mg_L": combo_do[group_best_idx] if group_best_idx is not None else np.nan,
                "代表最佳_HRT_h": combo_hrt[group_best_idx] if group_best_idx is not None else np.nan,
                "代表最佳_Feed": combo_feed[group_best_idx] if group_best_idx is not None else np.nan,
                "说明": (
                    "本组没有找到四个出水指标同时达到参考要求的参数组合。"
                    if group_best_count <= 0
                    else "‘达到本组最佳90%以上’表示该参数组合的达标背景数不低于本组最佳组合的90%，不是要求90%的历史背景都达标。"
                ),
            })

            # 先汇总达到本组最佳表现90%以上的参数范围，再单独汇总并列最佳组合范围。
            for range_type, mask in [
                ("达到本组最佳90%以上的组合", group_good_mask),
                ("本组并列最佳组合", group_best_mask),
            ]:
                if np.any(mask):
                    for parameter, chinese_name, values, unit in [
                        ("Aerobic_DO", "好氧区DO", combo_do, "mg/L"),
                        ("HRT", "水力停留时间", combo_hrt, "h"),
                        ("Anaerobic_Feed_Fraction", "厌氧区进水分配比例", combo_feed, "比例"),
                    ]:
                        selected = values[mask]
                        tp_group_range_rows.append({
                            "进水TP水平": group_name,
                            "分组规则": tp_group_rule_text[group_name],
                            "范围类型": range_type,
                            "运行参数": parameter,
                            "中文含义": chinese_name,
                            "最小值": float(np.min(selected)),
                            "最大值": float(np.max(selected)),
                            "单位": unit,
                            "实际出现水平数": int(len(np.unique(selected))),
                            "对应参数组合数": int(np.sum(mask)),
                            "说明": "这些是筛选后参数组合实际出现的最小值和最大值，不表示三个参数区间内任意自由组合都能达到要求。",
                        })
                else:
                    tp_group_range_rows.append({
                        "进水TP水平": group_name,
                        "分组规则": tp_group_rule_text[group_name],
                        "范围类型": range_type,
                        "运行参数": "无",
                        "中文含义": "未找到符合该条件的参数组合",
                        "最小值": np.nan,
                        "最大值": np.nan,
                        "单位": "",
                        "实际出现水平数": 0,
                        "对应参数组合数": 0,
                        "说明": "当前扫描范围内没有可汇总的参数范围。",
                    })

            group_combo_df = pd.DataFrame({
                "进水TP水平": np.full(combination_n, group_name, dtype=object),
                "分组规则": np.full(combination_n, tp_group_rule_text[group_name], dtype=object),
                "本组历史背景数": np.full(combination_n, group_bg_n, dtype=int),
                "组合序号": np.arange(1, combination_n + 1, dtype=int),
                "候选组合ID": [recommendation_combo_id(d, h, f) for d, h, f in zip(combo_do, combo_hrt, combo_feed)],
                "候选_DO_mg_L": combo_do,
                "候选_HRT_h": combo_hrt,
                "候选_Feed": combo_feed,
                "候选_Feed_%": combo_feed * 100.0,
                "通过物理和数值检查背景数": group_valid_counts,
                "TP满足背景数": group_tp_counts,
                "TN满足背景数": group_tn_counts,
                "NH4满足背景数": group_nh4_counts,
                "COD满足背景数": group_cod_counts,
                "四个出水指标同时达到要求背景数": group_all4_counts,
                "四个出水指标同时达到要求比例%": 100.0 * group_all4_counts / float(group_bg_n),
                "相对本组最佳表现%": (
                    100.0 * group_all4_counts / float(group_best_count)
                    if group_best_count > 0 else np.full(combination_n, np.nan)
                ),
                "是否达到本组最佳90%以上": np.where(group_good_mask, "是", "否"),
                "是否本组并列最佳组合": np.where(group_best_mask, "是", "否"),
            })
            tp_group_combo_frames.append(group_combo_df)

        tp_group_summary_df = pd.DataFrame(tp_group_summary_rows)
        tp_group_range_df = pd.DataFrame(tp_group_range_rows)
        tp_group_combo_df = pd.concat(tp_group_combo_frames, ignore_index=True)
        tp_group_combo_df["进水TP水平"] = pd.Categorical(
            tp_group_combo_df["进水TP水平"], categories=tp_group_order, ordered=True
        )
        tp_group_combo_df = tp_group_combo_df.sort_values(
            ["进水TP水平", "候选_DO_mg_L", "候选_HRT_h", "候选_Feed"],
            kind="mergesort",
        ).reset_index(drop=True)
        tp_group_combo_df["进水TP水平"] = tp_group_combo_df["进水TP水平"].astype(str)

        # 每一项条件对最终结果的实际影响：先看其他三项，再加回该项。
        other3_tp = formal_valid & tn_meet & nh4_meet & cod_meet
        other3_tn = formal_valid & tp_meet & nh4_meet & cod_meet
        other3_nh4 = formal_valid & tp_meet & tn_meet & cod_meet
        other3_cod = formal_valid & tp_meet & tn_meet & nh4_meet
        condition_rows = []
        for name, limit, other3 in [
            ("TP", TP_TARGET, other3_tp), ("TN", TN_TARGET, other3_tn),
            ("NH4+-N", NH4_TARGET, other3_nh4), ("COD", COD_TARGET, other3_cod),
        ]:
            other3_n = int(np.sum(other3))
            extra_removed = other3_n - n_all4
            condition_rows.append({
                "指标": name,
                "一级A工程参考值_mg_L": limit,
                "暂时不要求这一项时_其他三项都满足记录数": other3_n,
                "加入这一项后_四项都满足记录数": n_all4,
                "加入这一项后减少的记录数": extra_removed,
                "是否实际减少了可选记录": "是" if extra_removed > 0 else "否",
                "说明": (
                    f"加入{name}条件后又减少 {extra_removed:,} 条记录。" if extra_removed > 0
                    else f"加入{name}条件后没有继续减少记录。"
                ),
            })
        condition_df = pd.DataFrame(condition_rows)

        # 四项组合级预测统计。
        def row_stats(matrix: np.ndarray) -> Tuple[np.ndarray, np.ndarray, np.ndarray]:
            return np.nanmin(matrix, axis=1), np.nanmedian(matrix, axis=1), np.nanmax(matrix, axis=1)

        tp_min, tp_med, tp_max = row_stats(tp_matrix)
        tn_min, tn_med, tn_max = row_stats(tn_matrix)
        nh4_min, nh4_med, nh4_max = row_stats(nh4_matrix)
        cod_min, cod_med, cod_max = row_stats(cod_matrix)

        any_all4_combo = all4_counts > 0
        combos_with_any = int(np.sum(any_all4_combo))
        combos_all_backgrounds = int(np.sum(all4_counts == background_n))

        if n_all4 > 0:
            do_min = float(np.min(combo_do[any_all4_combo])); do_max = float(np.max(combo_do[any_all4_combo]))
            hrt_min = float(np.min(combo_hrt[any_all4_combo])); hrt_max = float(np.max(combo_hrt[any_all4_combo]))
            feed_min = float(np.min(combo_feed[any_all4_combo])); feed_max = float(np.max(combo_feed[any_all4_combo]))
            q_h_min = float(np.min(combo_q_h[any_all4_combo])); q_h_max = float(np.max(combo_q_h[any_all4_combo]))
            q_d_min = float(np.min(combo_q_d[any_all4_combo])); q_d_max = float(np.max(combo_q_d[any_all4_combo]))
            do_level_n = int(len(np.unique(combo_do[any_all4_combo])))
            hrt_level_n = int(len(np.unique(combo_hrt[any_all4_combo])))
            feed_level_n = int(len(np.unique(combo_feed[any_all4_combo])))
        else:
            do_min = do_max = hrt_min = hrt_max = feed_min = feed_max = np.nan
            q_h_min = q_h_max = q_d_min = q_d_max = np.nan
            do_level_n = hrt_level_n = feed_level_n = 0

        # ---------------------------------------------------------------------
        # 三类代表组合：最佳适用 + 接近最佳表现的最低DO + 接近最佳表现的最短HRT
        # ---------------------------------------------------------------------
        best_idx: Optional[int] = None
        low_do_idx: Optional[int] = None
        short_hrt_idx: Optional[int] = None
        best_count = int(np.max(all4_counts)) if len(all4_counts) else 0
        threshold_count = int(math.ceil(NEAR_BEST_RATIO * best_count - 1e-12)) if best_count > 0 else 0
        near_best_mask = (all4_counts >= threshold_count) & (all4_counts > 0) if best_count > 0 else np.zeros_like(all4_counts, dtype=bool)
        near_best_idx = np.flatnonzero(near_best_mask)

        if best_count > 0:
            best_candidates = np.flatnonzero(all4_counts == best_count)
            best_order = np.lexsort((combo_feed[best_candidates], combo_hrt[best_candidates], combo_do[best_candidates]))
            best_idx = int(best_candidates[int(best_order[0])])

            low_order = np.lexsort((
                combo_feed[near_best_idx], combo_hrt[near_best_idx], -all4_counts[near_best_idx], combo_do[near_best_idx]
            ))
            low_do_idx = int(near_best_idx[int(low_order[0])])

            short_order = np.lexsort((
                combo_feed[near_best_idx], combo_do[near_best_idx], -all4_counts[near_best_idx], combo_hrt[near_best_idx]
            ))
            short_hrt_idx = int(near_best_idx[int(short_order[0])])

        recommendation_labels: Dict[int, List[str]] = {}
        if best_idx is not None: recommendation_labels.setdefault(best_idx, []).append("适用背景最多方案")
        if low_do_idx is not None: recommendation_labels.setdefault(low_do_idx, []).append("接近最佳表现的最低DO方案")
        if short_hrt_idx is not None: recommendation_labels.setdefault(short_hrt_idx, []).append("接近最佳表现的最短HRT方案")

        # ---------------------------------------------------------------------
        # 结果表
        # ---------------------------------------------------------------------
        settings_rows: List[Dict[str, Any]] = [
            {"类别": "版本", "项目": "脚本版本", "值": SCRIPT_VERSION},
            {"类别": "版本", "项目": "版本标识", "值": BUILD_FINGERPRINT},
            {"类别": "目标", "项目": "TP一级A工程参考值_mg_L", "值": TP_TARGET},
            {"类别": "目标", "项目": "TN一级A工程参考值_mg_L", "值": TN_TARGET},
            {"类别": "目标", "项目": "NH4+-N一级A工程参考值_mg_L", "值": NH4_TARGET},
            {"类别": "目标", "项目": "COD一级A工程参考值_mg_L", "值": COD_TARGET},
            {"类别": "目标", "项目": "筛选逻辑", "值": "四个出水指标必须同时达到参考要求"},
            {"类别": "输入", "项目": "正向输入文件数", "值": len(FORWARD_FILES)},
            {"类别": "输入", "项目": "正式反向数据来源", "值": "仅网格来源=第一轮粗扫描"},
            {"类别": "输入", "项目": "当前实际工况补充是否进入正式筛选", "值": "否"},
            {"类别": "筛选", "项目": "异常预测处理", "值": "预测物理/数值检查不通过的记录不进入正式筛选"},
            {"类别": "筛选", "项目": "正向一级A标签用途", "值": "只用于一致性核查；反向程序按原始预测值重新计算"},
            {"类别": "筛选", "项目": "历史接近程度等附加门槛", "值": "不使用"},
            {"类别": "筛选", "项目": "预测误差是否改写目标", "值": "否"},
            {"类别": "推荐", "项目": "接近最佳组合的保留规则", "值": "四项同时满足背景数达到最佳组合的90%以上"},
            {"类别": "推荐", "项目": "代表组合", "值": "适用背景最多、接近最佳表现的最低DO、接近最佳表现的最短HRT"},
            {"类别": "进水TP分组", "项目": "分组方法", "值": "按631个历史背景进水TP实际分布的三分位数分为低、中、高三个水平"},
            {"类别": "进水TP分组", "项目": "第一个分界值_mg_L", "值": tp_q1},
            {"类别": "进水TP分组", "项目": "第二个分界值_mg_L", "值": tp_q2},
            {"类别": "进水TP分组", "项目": "组内表现较好参数的保留规则", "值": "该组四个出水指标同时达到要求的背景数达到该组最佳组合的90%以上"},
            {"类别": "进水TP分组", "项目": "进水TP缺失处理", "值": "仅用于分组；按正向阶段相同的短内部线性插值规则补齐，不改写任何正向预测值"},
            {"类别": "解释", "项目": "代表方案预测范围说明", "值": "代表方案中的TP/TN/NH4/COD最小值、中位数和最大值均按该固定参数组合在全部历史背景下的预测统计，不只统计满足四项条件的背景"},
            {"类别": "解释", "项目": "NH4+-N结果使用边界", "值": "NH4+-N和其他三个指标一样参与筛选；但其最终机器学习模型对DO/HRT/Feed的运行响应很弱，且此前未超过roll3历史基线，因此不据此说明调整这三个运行参数能够改善实际NH4+-N"},
            {"类别": "解释", "项目": "参数范围含义", "值": "只是满足四项条件的记录实际出现过的范围，不表示范围内任意组合都满足"},
            {"类别": "解释", "项目": "预测性质", "值": "冻结模型的一步条件预测筛选，不是长期稳态模拟或实际运行保证"},
            {"类别": "输出", "项目": "结果Excel", "值": str(OUTPUT_EXCEL)},
            {"类别": "输出", "项目": "说明TXT", "值": str(OUTPUT_TEXT)},
        ]
        for _, row in input_files_df.iterrows():
            settings_rows.append({"类别": "输入文件", "项目": f"{row['结果部分']}_{Path(str(row['输入文件'])).name}", "值": row["SHA256"]})
        settings_df = pd.DataFrame(settings_rows)

        summary_df = pd.DataFrame([
            {"项目": "历史背景数", "数值": background_n, "单位/说明": "个"},
            {"项目": "统一三参数组合数", "数值": combination_n, "单位/说明": "个/背景"},
            {"项目": "正式粗网格记录数", "数值": total_coarse_rows, "单位/说明": "条"},
            {"项目": "当前实际工况补充记录数_不参与筛选", "数值": total_supplement_rows, "单位/说明": "条"},
            {"项目": "十一个正向文件全部明细记录数", "数值": total_detail_rows, "单位/说明": "条"},
            {"项目": "通过物理和数值检查的粗网格记录数", "数值": n_valid, "单位/说明": "条"},
            {"项目": "未通过物理或数值检查的粗网格记录数", "数值": total_coarse_rows - n_valid, "单位/说明": "条"},
            {"项目": "TP满足一级A参考值记录数", "数值": n_tp, "单位/说明": "条；仅统计正式有效记录"},
            {"项目": "TN满足一级A参考值记录数", "数值": n_tn, "单位/说明": "条；仅统计正式有效记录"},
            {"项目": "NH4+-N满足一级A参考值记录数", "数值": n_nh4, "单位/说明": "条；仅统计正式有效记录"},
            {"项目": "COD满足一级A参考值记录数", "数值": n_cod, "单位/说明": "条；仅统计正式有效记录"},
            {"项目": "四项同时满足记录数", "数值": n_all4, "单位/说明": "条"},
            {"项目": "四项同时满足比例", "数值": screening_pct, "单位/说明": "%（分母=全部正式粗网格）"},
            {"项目": "至少在1个背景下四项同时满足的参数组合数", "数值": combos_with_any, "单位/说明": "个"},
            {"项目": "在全部背景下四项都同时满足的参数组合数", "数值": combos_all_backgrounds, "单位/说明": "个"},
            {"项目": "至少有1组运行参数可四项同时满足的历史背景数", "数值": backgrounds_with_any_all4, "单位/说明": f"个 / {background_n}（{100.0*backgrounds_with_any_all4/background_n:.4f}%）"},
            {"项目": "没有任何扫描组合可四项同时满足的历史背景数", "数值": backgrounds_with_no_all4, "单位/说明": f"个 / {background_n}（{100.0*backgrounds_with_no_all4/background_n:.4f}%）"},
            {"项目": "全部扫描组合都可四项同时满足的历史背景数", "数值": backgrounds_with_all_combos_all4, "单位/说明": f"个 / {background_n}（{100.0*backgrounds_with_all_combos_all4/background_n:.4f}%）"},
            {"项目": "最佳组合可满足的历史背景数", "数值": best_count, "单位/说明": f"个 / {background_n}"},
            {"项目": "达到最佳组合90%以上表现所需的最少满足背景数", "数值": threshold_count, "单位/说明": f"个；ceil({best_count}×{NEAR_BEST_RATIO:.0%})"},
            {"项目": "达到最佳组合90%以上表现的参数组合数", "数值": int(np.sum(near_best_mask)), "单位/说明": "个"},
        ])

        if n_all4 > 0:
            range_df = pd.DataFrame([
                {"运行参数": "Aerobic_DO", "中文含义": "好氧区DO", "最小值": do_min, "最大值": do_max, "单位": "mg/L", "实际出现水平数": do_level_n, "说明": "四项同时满足记录中实际出现过的范围"},
                {"运行参数": "HRT", "中文含义": "水力停留时间", "最小值": hrt_min, "最大值": hrt_max, "单位": "h", "实际出现水平数": hrt_level_n, "说明": "四项同时满足记录中实际出现过的范围"},
                {"运行参数": "Anaerobic_Feed_Fraction", "中文含义": "厌氧区进水分配比例", "最小值": feed_min, "最大值": feed_max, "单位": "比例", "实际出现水平数": feed_level_n, "说明": "四项同时满足记录中实际出现过的范围；乘100可换算为%"},
                {"运行参数": "Q", "中文含义": "由HRT换算的进水流量", "最小值": q_h_min, "最大值": q_h_max, "单位": "L/h", "实际出现水平数": hrt_level_n, "说明": "Q=57.6/HRT，只用于辅助报告"},
                {"运行参数": "Q", "中文含义": "由HRT换算的日进水流量", "最小值": q_d_min, "最大值": q_d_max, "单位": "L/d", "实际出现水平数": hrt_level_n, "说明": "Q=57.6×24/HRT，只用于辅助报告"},
            ])
        else:
            range_df = pd.DataFrame([{"运行参数": "无", "中文含义": "没有筛选到四项同时满足的记录", "最小值": np.nan, "最大值": np.nan, "单位": "", "实际出现水平数": 0, "说明": "不输出虚构的参数范围"}])

        combo_df = pd.DataFrame({
            "组合序号": np.arange(1, combination_n + 1, dtype=int),
            "候选组合ID": [recommendation_combo_id(d, h, f) for d, h, f in zip(combo_do, combo_hrt, combo_feed)],
            "候选_DO_mg_L": combo_do, "候选_HRT_h": combo_hrt, "候选_Feed": combo_feed,
            "候选_Feed_%": combo_feed * 100.0, "候选_Q_L_h": combo_q_h, "候选_Q_L_d": combo_q_d,
            "全部背景数": np.full(combination_n, background_n, dtype=int),
            "通过物理和数值检查背景数": physical_counts,
            "TP满足背景数": tp_counts, "TN满足背景数": tn_counts, "NH4满足背景数": nh4_counts, "COD满足背景数": cod_counts,
            "四项同时满足背景数": all4_counts,
            "四项同时满足背景比例%": 100.0 * all4_counts / float(background_n),
            "相对最佳组合表现%": np.where(best_count > 0, 100.0 * all4_counts / float(best_count), np.nan),
            "是否达到最佳组合90%以上表现": np.where(near_best_mask, "是", "否"),
            "全部历史背景_TP预测最小值": tp_min, "全部历史背景_TP预测中位数": tp_med, "全部历史背景_TP预测最大值": tp_max,
            "全部历史背景_TN预测最小值": tn_min, "全部历史背景_TN预测中位数": tn_med, "全部历史背景_TN预测最大值": tn_max,
            "全部历史背景_NH4预测最小值": nh4_min, "全部历史背景_NH4预测中位数": nh4_med, "全部历史背景_NH4预测最大值": nh4_max,
            "全部历史背景_COD预测最小值": cod_min, "全部历史背景_COD预测中位数": cod_med, "全部历史背景_COD预测最大值": cod_max,
            "是否至少1个背景四项同时满足": np.where(any_all4_combo, "是", "否"),
            "推荐标签": ["；".join(recommendation_labels.get(i, [])) for i in range(combination_n)],
        })
        combo_df = combo_df.sort_values(["候选_DO_mg_L", "候选_HRT_h", "候选_Feed"], kind="mergesort").reset_index(drop=True)

        near_best_df = combo_df.loc[combo_df["是否达到最佳组合90%以上表现"].eq("是")].copy()
        if not near_best_df.empty:
            near_best_df = near_best_df.sort_values(
                ["四项同时满足背景数", "候选_DO_mg_L", "候选_HRT_h", "候选_Feed"],
                ascending=[False, True, True, True], kind="mergesort"
            ).reset_index(drop=True)

        recommendation_rows: List[Dict[str, Any]] = []
        strategies = [
            ("适用背景最多方案", best_idx, "先选择四项同时满足历史背景数最多的组合；并列时优先较低DO、较短HRT、较小Feed。"),
            ("接近最佳表现的最低DO方案", low_do_idx, "先保留四项同时满足背景数达到最佳组合90%以上的组合，再在其中优先选择最低DO。"),
            ("接近最佳表现的最短HRT方案", short_hrt_idx, "先保留四项同时满足背景数达到最佳组合90%以上的组合，再在其中优先选择最短HRT。"),
        ]
        for strategy, idx, rule in strategies:
            if idx is None:
                recommendation_rows.append({"推荐思路": strategy, "推荐状态": "无推荐", "选择规则": rule, "说明": "当前正向预测结果库中没有四项同时满足的正式粗网格记录。"})
                continue
            recommendation_rows.append({
                "推荐思路": strategy, "推荐状态": "已推荐", "选择规则": rule,
                "候选组合ID": recommendation_combo_id(combo_do[idx], combo_hrt[idx], combo_feed[idx]),
                "推荐_DO_mg_L": combo_do[idx], "推荐_HRT_h": combo_hrt[idx], "推荐_Feed": combo_feed[idx],
                "推荐_Feed_%": combo_feed[idx] * 100.0, "推荐_Q_L_h": combo_q_h[idx], "推荐_Q_L_d": combo_q_d[idx],
                "全部历史背景数": background_n, "通过物理和数值检查背景数": int(physical_counts[idx]),
                "四项同时满足背景数": int(all4_counts[idx]),
                "四项同时满足背景比例%": 100.0 * all4_counts[idx] / float(background_n),
                "相对最佳组合表现%": 100.0 * all4_counts[idx] / float(best_count) if best_count > 0 else np.nan,
                "达到最佳组合90%以上表现所需的最少满足背景数": threshold_count,
                "全部历史背景_TP预测最小值": tp_min[idx], "全部历史背景_TP预测中位数": tp_med[idx], "全部历史背景_TP预测最大值": tp_max[idx],
                "全部历史背景_TN预测最小值": tn_min[idx], "全部历史背景_TN预测中位数": tn_med[idx], "全部历史背景_TN预测最大值": tn_max[idx],
                "全部历史背景_NH4预测最小值": nh4_min[idx], "全部历史背景_NH4预测中位数": nh4_med[idx], "全部历史背景_NH4预测最大值": nh4_max[idx],
                "全部历史背景_COD预测最小值": cod_min[idx], "全部历史背景_COD预测中位数": cod_med[idx], "全部历史背景_COD预测最大值": cod_max[idx],
                "说明": "上述四项预测范围按该固定参数组合在全部历史背景下统计，并非只统计满足四项条件的背景；该组合是离散粗网格中的代表结果，不代表长期稳态或实际运行保证。",
            })
        recommendation_df = pd.DataFrame(recommendation_rows)

        bg_valid_counts = np.sum(formal_valid, axis=0).astype(int)
        bg_tp_counts = np.sum(tp_meet, axis=0).astype(int)
        bg_tn_counts = np.sum(tn_meet, axis=0).astype(int)
        bg_nh4_counts = np.sum(nh4_meet, axis=0).astype(int)
        bg_cod_counts = np.sum(cod_meet, axis=0).astype(int)
        background_df = pd.DataFrame(background_meta)
        background_df["进水TP原始值_mg_L"] = influent_tp_raw
        background_df["进水TP分组使用值_mg_L"] = influent_tp_values
        background_df["进水TP水平"] = tp_group_labels
        background_df["通过物理和数值检查组合数"] = bg_valid_counts
        background_df["TP满足组合数"] = bg_tp_counts
        background_df["TN满足组合数"] = bg_tn_counts
        background_df["NH4满足组合数"] = bg_nh4_counts
        background_df["COD满足组合数"] = bg_cod_counts
        background_df["四项同时满足组合数"] = bg_all4_counts
        background_df["四项同时满足组合比例%"] = 100.0 * bg_all4_counts / float(combination_n)

        check_rows: List[Dict[str, Any]] = []
        for row in source_file_stats:
            check_rows.append({
                "核查项目": f"{row['结果部分']}文件结构",
                "核查结果": "通过",
                "数值/说明": f"背景={row['历史背景数']}；明细={row['全部明细记录数']}；粗网格={row['正式粗网格记录数']}；补充={row['当前工况补充记录数']}；明细副表={row['明细副表数']}",
            })

        targets_fixed_ok = (
            math.isclose(TP_TARGET, 0.5) and math.isclose(TN_TARGET, 15.0)
            and math.isclose(NH4_TARGET, 5.0) and math.isclose(COD_TARGET, 50.0)
        )
        method_checks = [
            ("十一个正向Excel是否完整读取", len(FORWARD_FILES) == 11, f"文件数={len(FORWARD_FILES)}"),
            ("历史背景数", background_n == EXPECTED_BACKGROUND_N, f"实际={background_n}，预期={EXPECTED_BACKGROUND_N}"),
            ("统一组合数", combination_n == EXPECTED_COMBINATION_N, f"实际={combination_n}，预期={EXPECTED_COMBINATION_N}"),
            ("粗网格总数是否正确", total_coarse_rows == background_n * combination_n, f"实际={total_coarse_rows}，理论={background_n*combination_n}"),
            ("背景ID无重复", len(seen_backgrounds) == background_n, f"唯一ID={len(seen_backgrounds)}"),
            ("未知网格来源记录为0", total_unknown_source_rows == 0, f"未知来源={total_unknown_source_rows}"),
            ("重复组合背景数为0", duplicate_combo_backgrounds == 0, f"异常背景={duplicate_combo_backgrounds}"),
            ("组合不完整背景数为0", incomplete_combo_backgrounds == 0, f"异常背景={incomplete_combo_backgrounds}"),
            ("Q换算一致", q_inconsistency_rows == 0, f"不一致记录={q_inconsistency_rows}"),
            ("正向一级A标签与反向重新计算一致", forward_flag_mismatch_rows == 0, f"不一致单元={forward_flag_mismatch_rows}；示例={forward_flag_mismatch_examples}"),
            ("正式筛选只使用第一轮粗扫描", True, f"粗网格={total_coarse_rows}；当前实际工况补充={total_supplement_rows}，补充未进入筛选"),
            ("异常预测已排除", True, f"物理/数值检查未通过={total_coarse_rows-n_valid}；负预测字段非空={total_negative_flag_rows}；非有限预测={total_nonfinite_prediction_rows}"),
            ("四项一级A目标固定正确", targets_fixed_ok, f"TP={TP_TARGET}；TN={TN_TARGET}；NH4={NH4_TARGET}；COD={COD_TARGET}"),
            ("背景层面统计数量一致", backgrounds_with_any_all4 + backgrounds_with_no_all4 == background_n, f"有可行情景={backgrounds_with_any_all4}；无可行情景={backgrounds_with_no_all4}；总背景={background_n}"),
            ("接近最佳组合的90%规则已执行", True, f"最佳背景数={best_count}；90%门槛={threshold_count}；达到该水平的组合={int(np.sum(near_best_mask))}个"),
            ("进水TP分组覆盖全部历史背景", sum(int(tp_group_runtime[g]["背景数"]) for g in tp_group_order) == background_n, f"低/中/高={','.join(str(tp_group_runtime[g]['背景数']) for g in tp_group_order)}；总背景={background_n}"),
            ("进水TP分组后的预测情景数与粗网格一致", sum(int(tp_group_runtime[g]["总预测情景数"]) for g in tp_group_order) == total_coarse_rows, f"分组合计={sum(int(tp_group_runtime[g]['总预测情景数']) for g in tp_group_order)}；粗网格={total_coarse_rows}"),
            ("进水TP分组后的达标情景数与总体一致", sum(int(tp_group_runtime[g]["达标预测情景数"]) for g in tp_group_order) == n_all4, f"分组合计={sum(int(tp_group_runtime[g]['达标预测情景数']) for g in tp_group_order)}；总体={n_all4}"),
            ("进水TP分组没有未补齐缺失值", int(influent_tp_df.attrs.get("remaining_missing_n", 0)) == 0, f"原始缺失={int(influent_tp_df.attrs.get('original_missing_n', 0))}；插值后剩余={int(influent_tp_df.attrs.get('remaining_missing_n', 0))}"),
        ]
        for name, passed, detail in method_checks:
            check_rows.append({"核查项目": name, "核查结果": "通过" if passed else "失败", "数值/说明": detail})
        checks_df = pd.DataFrame(check_rows)
        if STRICT_MODE and checks_df["核查结果"].eq("失败").any():
            failed = checks_df.loc[checks_df["核查结果"].eq("失败"), "核查项目"].tolist()
            raise RuntimeError(f"严格模式方法检查失败：{failed}")

        write_result_excel(
            settings_df, summary_df, condition_df, range_df, combo_df,
            near_best_df, recommendation_df, background_df, checks_df,
            tp_group_summary_df, tp_group_range_df, tp_group_combo_df,
        )

        elapsed_minutes = (time.time() - start) / 60.0
        lines: List[str] = []
        lines.append("A²/O 四指标一级A参考目标反向筛选与运行参数推荐说明")
        lines.append("=" * 80)
        lines.append(f"运行时间：{now_str()}")
        lines.append(f"脚本版本：{SCRIPT_VERSION}")
        lines.append(f"版本标识：{BUILD_FINGERPRINT}")
        lines.append("")
        lines.append("一、筛选目标和输入")
        lines.append(f"1. TP：预测值 <= {TP_TARGET:g} mg/L。")
        lines.append(f"2. TN：预测值 <= {TN_TARGET:g} mg/L。")
        lines.append(f"3. NH4+-N：预测值 <= {NH4_TARGET:g} mg/L。")
        lines.append(f"4. COD：预测值 <= {COD_TARGET:g} mg/L。")
        lines.append("5. 四项必须同时满足；预测物理或数值检查不通过的记录不进入正式筛选。")
        lines.append("6. 正式反向输入只使用十一个正向结果文件中的‘第一轮粗扫描’记录；当前实际工况补充不参与推荐。")
        lines.append("")
        lines.append("二、数据规模")
        lines.append(f"1. 历史背景数：{background_n}。")
        lines.append(f"2. 每个背景统一运行参数组合数：{combination_n}。")
        lines.append(f"3. 正式粗网格记录数：{total_coarse_rows:,}。")
        lines.append(f"4. 当前实际工况补充记录数：{total_supplement_rows:,}（不参与正式筛选）。")
        lines.append(f"5. 通过物理和数值检查的正式粗网格记录数：{n_valid:,}。")
        lines.append(f"6. 未通过检查的正式粗网格记录数：{total_coarse_rows-n_valid:,}。")
        lines.append("")
        lines.append("三、四项筛选结果")
        lines.append(f"1. TP满足记录：{n_tp:,} 条。")
        lines.append(f"2. TN满足记录：{n_tn:,} 条。")
        lines.append(f"3. NH4+-N满足记录：{n_nh4:,} 条。")
        lines.append(f"4. COD满足记录：{n_cod:,} 条。")
        lines.append(f"5. 四项同时满足记录：{n_all4:,} 条，占全部正式粗网格 {screening_pct:.6f}%。")
        lines.append(f"6. 至少在1个历史背景下四项同时满足的三参数组合：{combos_with_any}/{combination_n}。")
        lines.append(f"7. 在全部{background_n}个历史背景下都四项同时满足的组合：{combos_all_backgrounds}。")
        lines.append(f"8. 从历史背景角度看，有 {backgrounds_with_any_all4}/{background_n} 个背景至少存在1组运行参数能够四项同时满足（{100.0*backgrounds_with_any_all4/background_n:.4f}%）。")
        lines.append(f"9. 有 {backgrounds_with_no_all4}/{background_n} 个背景在当前3780组扫描参数中没有找到四项同时满足的组合（{100.0*backgrounds_with_no_all4/background_n:.4f}%）。")
        lines.append(f"10. 有 {backgrounds_with_all_combos_all4}/{background_n} 个背景在3780组扫描参数下全部都能四项同时满足（{100.0*backgrounds_with_all_combos_all4/background_n:.4f}%）。")
        if n_tn == n_valid and n_nh4 == n_valid and n_cod == n_valid and n_all4 == n_tp:
            lines.append("11. 在当前有效预测中，TN、NH4+-N和COD都达到相应参考要求，四个指标能否同时达到要求主要取决于TP；因此后面进一步按进水TP水平分析运行参数。")
        else:
            lines.append("11. 四个指标均继续参与筛选；后面的进水TP分组分析用于进一步观察不同进水TP条件下运行参数的表现，不取消其他三个指标的达标要求。")
        lines.append("")
        lines.append("四、每一项条件对结果的实际影响")
        for i, row in condition_df.iterrows():
            lines.append(
                f"{i+1}. {row['指标']}：如果暂时不要求这一项，其他三项都满足的记录有 "
                f"{int(row['暂时不要求这一项时_其他三项都满足记录数']):,} 条；加入这一项后剩 "
                f"{int(row['加入这一项后_四项都满足记录数']):,} 条，因此另外减少 "
                f"{int(row['加入这一项后减少的记录数']):,} 条。"
            )
        lines.append("")
        lines.append("五、四项同时满足时实际出现的运行参数范围")
        if n_all4 <= 0:
            lines.append("当前正向预测结果库中没有找到四项同时满足的正式粗网格记录，因此不输出参数范围和推荐组合。")
        else:
            lines.append(f"1. DO：{do_min:g} ~ {do_max:g} mg/L。")
            lines.append(f"2. HRT：{hrt_min:g} ~ {hrt_max:g} h。")
            lines.append(f"3. 厌氧区进水分配比例：{feed_min:g} ~ {feed_max:g}（{feed_min*100:g}% ~ {feed_max*100:g}%）。")
            lines.append(f"4. 对应Q：{q_h_min:.6g} ~ {q_h_max:.6g} L/h（{q_d_min:.6g} ~ {q_d_max:.6g} L/d）。")
            lines.append("5. 这些只是四项同时满足记录实际出现过的最小值和最大值，不表示三个区间内任意自由组合都能满足四项要求。")
        lines.append("")
        lines.append("六、接近最佳表现的组合和三个代表方案")
        if best_count <= 0:
            lines.append("没有可推荐组合。")
        else:
            lines.append(f"1. 最佳组合最多可在 {best_count}/{background_n} 个历史背景下四项同时满足。")
            lines.append(f"2. 主分析采用90%规则：四项同时满足背景数达到最佳组合90%以上即可保留。因此至少需要 {threshold_count} 个背景四项同时满足，共有 {int(np.sum(near_best_mask))} 个组合达到这一水平。")
            for i, (strategy, idx, _) in enumerate(strategies, start=3):
                if idx is None:
                    lines.append(f"{i}. {strategy}：无。")
                    continue
                lines.append(
                    f"{i}. {strategy}：DO={combo_do[idx]:g} mg/L，HRT={combo_hrt[idx]:g} h，"
                    f"Feed={combo_feed[idx]:g}，Q={combo_q_h[idx]:.6g} L/h；"
                    f"四项同时满足 {all4_counts[idx]}/{background_n} 个背景（{100.0*all4_counts[idx]/background_n:.6f}%），"
                    f"相当于最佳组合表现的 {100.0*all4_counts[idx]/best_count:.6f}%。"
                )
        lines.append("")
        lines.append("七、按进水TP水平分别看运行参数")
        lines.append(
            f"1. 根据631个历史背景中进水TP的实际分布，按三分位数分为低、中、高三个水平；"
            f"分界值分别为 {tp_q1:.6f} mg/L 和 {tp_q2:.6f} mg/L。"
        )
        lines.append(
            "2. 这样分组是为了让三组历史背景数量接近，便于比较；这些分界值只代表本研究数据的分组位置，不是行业通用浓度界限。"
        )
        item_no = 3
        for group_name in tp_group_order:
            stat = tp_group_runtime[group_name]
            mask = tp_group_labels == group_name
            group_min = float(np.min(influent_tp_values[mask]))
            group_max = float(np.max(influent_tp_values[mask]))
            lines.append(
                f"{item_no}. {group_name}：进水TP实际范围 {group_min:.6f} ~ {group_max:.6f} mg/L，"
                f"历史背景 {stat['背景数']} 个；全部参数情景中四个出水指标同时达到要求的有 "
                f"{stat['达标预测情景数']:,}/{stat['总预测情景数']:,}（{stat['总体达标比例%']:.6f}%）。"
            )
            item_no += 1
            if stat["最佳达标背景数"] <= 0:
                lines.append(
                    f"{item_no}. {group_name}在当前DO、HRT和进水分配比例扫描范围内没有找到四个出水指标同时达到参考要求的参数组合，因此不输出该组运行参数范围。"
                )
                item_no += 1
                continue

            good_rows = tp_group_range_df.loc[
                (tp_group_range_df["进水TP水平"] == group_name)
                & (tp_group_range_df["范围类型"] == "达到本组最佳90%以上的组合")
                & (tp_group_range_df["运行参数"] != "无")
            ]
            best_rows = tp_group_range_df.loc[
                (tp_group_range_df["进水TP水平"] == group_name)
                & (tp_group_range_df["范围类型"] == "本组并列最佳组合")
                & (tp_group_range_df["运行参数"] != "无")
            ]

            def _get_range(rows: pd.DataFrame, parameter: str) -> Tuple[float, float]:
                part = rows.loc[rows["运行参数"] == parameter]
                if part.empty:
                    return np.nan, np.nan
                return float(part.iloc[0]["最小值"]), float(part.iloc[0]["最大值"])

            do_lo, do_hi = _get_range(good_rows, "Aerobic_DO")
            hrt_lo, hrt_hi = _get_range(good_rows, "HRT")
            feed_lo, feed_hi = _get_range(good_rows, "Anaerobic_Feed_Fraction")
            best_do_lo, best_do_hi = _get_range(best_rows, "Aerobic_DO")
            best_hrt_lo, best_hrt_hi = _get_range(best_rows, "HRT")
            best_feed_lo, best_feed_hi = _get_range(best_rows, "Anaerobic_Feed_Fraction")

            lines.append(
                f"{item_no}. {group_name}表现最好的固定参数组合可在 {stat['最佳达标背景数']}/{stat['背景数']} 个历史背景下达到要求"
                f"（{stat['最佳达标比例%']:.6f}%）；达到该组最佳表现90%以上的参数组合共有 {stat['表现较好组合数']} 个。"
            )
            item_no += 1
            lines.append(
                f"{item_no}. {group_name}达到本组最佳表现90%以上的组合中，DO实际范围为 {do_lo:g} ~ {do_hi:g} mg/L，"
                f"HRT为 {hrt_lo:g} ~ {hrt_hi:g} h，厌氧区进水分配比例为 {feed_lo:g} ~ {feed_hi:g}。"
            )
            item_no += 1
            lines.append(
                f"{item_no}. {group_name}并列最佳组合实际集中在：DO {best_do_lo:g} ~ {best_do_hi:g} mg/L，"
                f"HRT {best_hrt_lo:g} ~ {best_hrt_hi:g} h，厌氧区进水分配比例 {best_feed_lo:g} ~ {best_feed_hi:g}。"
            )
            item_no += 1
        lines.append(
            f"{item_no}. 上述参数范围只是筛选后组合实际出现的范围，不表示DO、HRT和进水分配比例在各自范围内可以任意自由组合。"
        )
        lines.append("")
        lines.append("八、解释边界")
        lines.append("1. 本结果来自已经冻结的正向预测数据库，不重新训练模型。")
        lines.append("2. 每条记录是固定进水、近期出水状态和装置结构后，仅改变当前DO/HRT/Feed得到的一步条件预测。")
        lines.append("3. 一级A数值在这里作为工程参考筛选条件，不直接等同于法规意义上的日均达标判定。")
        lines.append("4. 三个代表组合是离散粗网格中的模型筛选结果，不代表长期稳态或实际运行一定达到目标。")
        lines.append("5. 代表方案中TP、TN、NH4+-N和COD的预测最小值、中位数和最大值，均按该固定参数组合在全部历史背景下统计，并非只统计四项同时满足的背景。")
        lines.append("6. 90%表示‘这个参数组合的达标背景数达到最佳组合的90%以上’，不是‘90%的历史背景都能满足’。")
        lines.append("7. NH4+-N和其他三个指标一样参与筛选，但其最终机器学习模型对DO、HRT和Feed的运行响应很弱，且此前未超过roll3历史基线，因此不能据此说明调整这三个运行参数能够改善实际NH4+-N。")
        lines.append(f"8. 总耗时：{elapsed_minutes:.2f} 分钟。")
        lines.append("")
        lines.append("九、输出文件")
        lines.append(f"1. 结果Excel：{OUTPUT_EXCEL}")
        lines.append(f"2. 说明TXT：{OUTPUT_TEXT}")

        atomic_write_text(OUTPUT_TEXT, "\n".join(lines))

        log("四指标一级A反向筛选完成。")
        log(f"结果Excel：{OUTPUT_EXCEL}")
        log(f"说明TXT：{OUTPUT_TEXT}")
        log(f"历史背景={background_n}；组合={combination_n}；正式粗网格={total_coarse_rows:,}")
        log(f"四项同时满足={n_all4:,}；比例={screening_pct:.6f}%；最佳背景数={best_count}；90%门槛={threshold_count}")
        for strategy, idx, _ in strategies:
            if idx is not None:
                log(f"{strategy}：DO={combo_do[idx]:g}, HRT={combo_hrt[idx]:g}, Feed={combo_feed[idx]:g}, n={all4_counts[idx]}/{background_n}")
        log(f"进水TP三分位分界：Q1={tp_q1:.6f} mg/L；Q2={tp_q2:.6f} mg/L")
        for group_name in tp_group_order:
            stat = tp_group_runtime[group_name]
            log(
                f"{group_name}：背景={stat['背景数']}；总体达标比例={stat['总体达标比例%']:.6f}%；"
                f"最佳={stat['最佳达标背景数']}/{stat['背景数']}；达到本组最佳90%以上组合={stat['表现较好组合数']}"
            )

    except Exception as exc:
        tmp_excel = OUTPUT_EXCEL.with_name(OUTPUT_EXCEL.stem + ".part.xlsx")
        try:
            if tmp_excel.exists():
                tmp_excel.unlink()
        except Exception:
            pass

        failure_text = "\n".join([
            "A²/O 四指标一级A参考目标反向筛选运行失败",
            "=" * 80,
            f"时间：{now_str()}",
            f"脚本版本：{SCRIPT_VERSION}",
            f"输出目录：{OUTPUT_DIR}",
            f"固定目标：TP<={TP_TARGET:g}；TN<={TN_TARGET:g}；NH4+-N<={NH4_TARGET:g}；COD<={COD_TARGET:g} mg/L",
            "",
            f"错误类型：{type(exc).__name__}",
            f"错误信息：{exc}",
            "",
            "完整错误堆栈：",
            traceback.format_exc(),
            "",
            "说明：脚本失败后不会修改正向预测结果或模型，请根据错误信息核查十一个正向结果文件。",
        ])
        try:
            atomic_write_text(OUTPUT_TEXT, failure_text)
        except Exception:
            pass
        log(f"运行失败：{type(exc).__name__}: {exc}")
        raise


if __name__ == "__main__":
    main()
