# -*- coding: utf-8 -*-
"""
A²/O 四个单目标动态可用数据表构建与缺失值插补方法比较
============================================================

本脚本完成两项任务：
1. 从“异常值处理结果.xlsx”的“清洗后数据”工作表出发，分别为 TP、COD、
   NH4+-N 和 TN 构建单目标动态可用数据表；
2. 在四个动态可用数据表上，通过可复现的遮蔽实验比较均值填充、标准化 KNN
   插补和线性插值三种方法。

重要边界：
- 四个出水目标不进行插补；
- 只有 Inf.TP、Inf.COD、Inf.NH4+-N 和 Inf.TN 参与插补方法比较；
- 动态表中的真实输入缺失值保持为空，不输出可直接用于正式模型验证的全量插补表；
- 本阶段的插补实验仅用于离线比较插补策略；其中“线性插值”允许利用人工遮蔽点
  两侧的已知观测值恢复被遮蔽值，仅用于方法比较与分布诊断。
- 正式建模时仍须在每一次训练/验证划分内部重新拟合插补规则；时间顺序验证、
  滚动/扩展窗口验证中的验证样本严禁使用未来观测进行双向线性插值。

默认运行：
    python 02_动态可用表与缺失值插补比较.py

可选参数示例：
    python 02_动态可用表与缺失值插补比较.py \
        --input "${PROJECT_ROOT}/outputs/preprocessing/outlier/异常值处理结果.xlsx" \
        --output "${PROJECT_ROOT}/outputs/preprocessing/dynamic" \
        --repeats 30

依赖：
    pip install pandas numpy openpyxl scikit-learn
"""

from __future__ import annotations

import argparse
import math
import sys
import time
import traceback
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Dict, Iterable, List, Mapping, Optional, Sequence, Tuple

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sklearn.impute import KNNImputer


# =============================================================================
# 1. 默认路径与可调参数
# =============================================================================

PROJECT_ROOT = Path(__file__).resolve().parents[1]

DEFAULT_INPUT_FILE = Path(
    str(PROJECT_ROOT / "outputs/preprocessing/outlier/异常值处理结果.xlsx")
)
DEFAULT_OUTPUT_DIR = Path(
    str(PROJECT_ROOT / "outputs/preprocessing/dynamic")
)

PREFERRED_INPUT_SHEET = "清洗后数据"
OUTPUT_EXCEL_NAME = "动态可用数据与插补比较结果.xlsx"
OUTPUT_TXT_NAME = "动态可用数据与插补比较说明.txt"

MASTER_RANDOM_SEED = 20260713
DEFAULT_REPEATS = 30
RANDOM_MASK_RATIO = 0.10
BLOCK_MASK_RATIO = 0.10
ROW_MASK_RATIO = 0.10
BLOCK_MIN_LENGTH = 3
BLOCK_MAX_LENGTH = 7

KNN_N_NEIGHBORS = 5
KNN_WEIGHTS = "distance"
LINEAR_MAX_GAP = 7
LINEAR_FALLBACK = "median"  # 可选："mean" 或 "median"
NORMALIZATION_METHOD = "P95-P5"
EPS = 1e-12

ID_COLUMNS = ["样本序号", "Date"]
OPTIONAL_TRACKING_COLUMNS = ["Stage_ID"]

RAW_INPUT_COLUMNS = [
    "Aerobic_DO",
    "Anaerobic_Feed_Fraction",
    "Inf.TP",
    "Inf.COD",
    "Inf.NH4+-N",
    "Inf.TN",
    "HRT",
    "Anaerobic_Volume_Fraction",
    "Anoxic_Volume_Fraction",
]

IMPUTE_COLUMNS = ["Inf.TP", "Inf.COD", "Inf.NH4+-N", "Inf.TN"]

TARGETS: Mapping[str, Mapping[str, str]] = {
    "TP": {
        "target": "eff.TP",
        "sheet": "TP动态可用表",
        "display": "TP",
    },
    "COD": {
        "target": "eff.COD",
        "sheet": "COD动态可用表",
        "display": "COD",
    },
    "NH4-N": {
        "target": "eff.NH4+-N",
        "sheet": "氨氮动态可用表",
        "display": "NH4+-N",
    },
    "TN": {
        "target": "eff.TN",
        "sheet": "TN动态可用表",
        "display": "TN",
    },
}

METHODS = ["均值填充", "标准化KNN", "线性插值"]
MASK_SCENARIOS = ["随机遮蔽", "连续块遮蔽", "整行遮蔽"]


# =============================================================================
# 2. 数据结构
# =============================================================================


@dataclass(frozen=True)
class DynamicBuildInfo:
    target_key: str
    target_column: str
    original_rows: int
    target_valid_rows: int
    target_missing_rows_removed: int
    history_incomplete_rows_removed: int
    dynamic_rows: int
    first_sample: object
    last_sample: object
    first_date: object
    last_date: object


# =============================================================================
# 3. 基础工具函数
# =============================================================================


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="构建四个单目标动态可用表，并比较三种进水水质缺失值插补方法。"
    )
    parser.add_argument(
        "--input",
        type=Path,
        default=DEFAULT_INPUT_FILE,
        help="异常值处理结果 Excel 文件路径。",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=DEFAULT_OUTPUT_DIR,
        help="输出文件夹路径。",
    )
    parser.add_argument(
        "--sheet",
        type=str,
        default=PREFERRED_INPUT_SHEET,
        help="优先读取的工作表名称，默认“清洗后数据”。",
    )
    parser.add_argument(
        "--repeats",
        type=int,
        default=DEFAULT_REPEATS,
        help="每种遮蔽情形的重复次数，默认 30。",
    )
    parser.add_argument(
        "--seed",
        type=int,
        default=MASTER_RANDOM_SEED,
        help="总随机种子，默认 20260713。",
    )
    return parser.parse_args()


def require_columns(df: pd.DataFrame, columns: Sequence[str], context: str) -> None:
    missing = [c for c in columns if c not in df.columns]
    if missing:
        raise ValueError(f"{context}缺少必要字段：{missing}")


def choose_input_sheet(file_path: Path, preferred_sheet: str) -> str:
    """优先读取指定工作表；若不存在，则寻找包含全部必要字段的工作表。"""
    required = set(ID_COLUMNS + RAW_INPUT_COLUMNS + [v["target"] for v in TARGETS.values()])
    xls = pd.ExcelFile(file_path)

    if preferred_sheet in xls.sheet_names:
        header = pd.read_excel(file_path, sheet_name=preferred_sheet, nrows=0)
        if required.issubset(set(header.columns)):
            return preferred_sheet

    for sheet in xls.sheet_names:
        header = pd.read_excel(file_path, sheet_name=sheet, nrows=0)
        if required.issubset(set(header.columns)):
            return sheet

    raise ValueError(
        "未找到包含全部必要字段的工作表。已检查工作表："
        + "、".join(xls.sheet_names)
    )


def clean_column_names(df: pd.DataFrame) -> pd.DataFrame:
    result = df.copy()
    result.columns = [str(c).strip() for c in result.columns]
    return result


def convert_types(df: pd.DataFrame) -> pd.DataFrame:
    result = df.copy()
    result["Date"] = pd.to_datetime(result["Date"], errors="coerce")

    numeric_columns = [
        "样本序号",
        *RAW_INPUT_COLUMNS,
        *[v["target"] for v in TARGETS.values()],
    ]
    for col in numeric_columns:
        if col in result.columns:
            result[col] = pd.to_numeric(result[col], errors="coerce")
    return result


def sort_continuous_data(df: pd.DataFrame) -> Tuple[pd.DataFrame, str, List[str]]:
    """
    按连续运行顺序排序。

    首选样本序号；若样本序号缺失或重复，则优先使用 Date。函数返回排序后的数据、
    排序说明以及警告列表。
    """
    warnings: List[str] = []
    result = df.copy()

    sample_ok = (
        "样本序号" in result.columns
        and result["样本序号"].notna().all()
        and not result["样本序号"].duplicated().any()
    )
    date_ok = (
        "Date" in result.columns
        and result["Date"].notna().all()
        and not result["Date"].duplicated().any()
    )

    if sample_ok:
        sort_cols = ["样本序号"]
        if "Date" in result.columns:
            sort_cols.append("Date")
        result = result.sort_values(sort_cols, kind="mergesort").reset_index(drop=True)
        order_desc = "按样本序号升序排序，Date 用于并列核查"
        if not date_ok:
            warnings.append("Date 存在缺失或重复，但样本序号完整且唯一，历史顺序以样本序号为准。")
    elif date_ok:
        result = result.sort_values(["Date"], kind="mergesort").reset_index(drop=True)
        order_desc = "样本序号存在缺失或重复，改按 Date 升序排序"
        warnings.append("样本序号不完整或不唯一，历史特征顺序改由 Date 决定。")
    else:
        raise ValueError(
            "样本序号和 Date 均无法提供明确、唯一的连续运行顺序，请先核查数据。"
        )

    return result, order_desc, warnings


def safe_value(value: object) -> str:
    if pd.isna(value):
        return "无"
    if isinstance(value, pd.Timestamp):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, (np.floating, float)):
        return f"{float(value):.6g}"
    return str(value)


def format_float(value: object, digits: int = 6) -> str:
    if value is None or pd.isna(value):
        return "无"
    return f"{float(value):.{digits}f}".rstrip("0").rstrip(".")


# =============================================================================
# 4. 四个单目标动态可用数据表
# =============================================================================


def build_dynamic_table(
    sorted_df: pd.DataFrame,
    target_key: str,
    target_column: str,
) -> Tuple[pd.DataFrame, DynamicBuildInfo]:
    """
    先仅剔除对应目标为空的样本，再基于“当前样本之前最近的有效目标观测”构造：
    lag1、lag2 和 roll3。最后剔除无法形成完整历史状态的样本。
    """
    original_rows = len(sorted_df)

    target_base = sorted_df.loc[sorted_df[target_column].notna()].copy()
    target_base = target_base.reset_index(drop=True)
    target_valid_rows = len(target_base)

    lag1_col = f"{target_column}_lag1"
    lag2_col = f"{target_column}_lag2"
    roll3_col = f"{target_column}_roll3"

    target_base[lag1_col] = target_base[target_column].shift(1)
    target_base[lag2_col] = target_base[target_column].shift(2)
    # 只使用当前样本之前最近三个有效目标观测值，不包含当前目标值。
    target_base[roll3_col] = target_base[target_column].shift(1).rolling(
        window=3, min_periods=3
    ).mean()

    history_columns = [lag1_col, lag2_col, roll3_col]
    dynamic = target_base.dropna(subset=history_columns).copy().reset_index(drop=True)

    tracking_cols = [c for c in ID_COLUMNS + OPTIONAL_TRACKING_COLUMNS if c in dynamic.columns]
    output_columns = tracking_cols + RAW_INPUT_COLUMNS + [target_column] + history_columns
    dynamic = dynamic[output_columns].copy()

    info = DynamicBuildInfo(
        target_key=target_key,
        target_column=target_column,
        original_rows=original_rows,
        target_valid_rows=target_valid_rows,
        target_missing_rows_removed=original_rows - target_valid_rows,
        history_incomplete_rows_removed=target_valid_rows - len(dynamic),
        dynamic_rows=len(dynamic),
        first_sample=dynamic["样本序号"].iloc[0] if len(dynamic) else np.nan,
        last_sample=dynamic["样本序号"].iloc[-1] if len(dynamic) else np.nan,
        first_date=dynamic["Date"].iloc[0] if len(dynamic) else pd.NaT,
        last_date=dynamic["Date"].iloc[-1] if len(dynamic) else pd.NaT,
    )
    return dynamic, info


def build_missing_summary(dynamic_tables: Mapping[str, pd.DataFrame]) -> pd.DataFrame:
    rows: List[Dict[str, object]] = []
    for target_key, table in dynamic_tables.items():
        for variable in IMPUTE_COLUMNS:
            missing_count = int(table[variable].isna().sum())
            rows.append(
                {
                    "目标指标": TARGETS[target_key]["display"],
                    "变量": variable,
                    "动态样本数": len(table),
                    "非缺失数": int(table[variable].notna().sum()),
                    "缺失数": missing_count,
                    "缺失比例(%)": missing_count / len(table) * 100 if len(table) else np.nan,
                }
            )
    return pd.DataFrame(rows)


# =============================================================================
# 5. 三种插补方法
# =============================================================================


def mean_impute(data: pd.DataFrame) -> pd.DataFrame:
    result = data.copy().astype(float)
    for col in result.columns:
        mean_value = result[col].mean(skipna=True)
        if pd.isna(mean_value):
            raise ValueError(f"变量 {col} 全部缺失，无法执行均值填充。")
        result[col] = result[col].fillna(mean_value)
    return result


def standardized_knn_impute(
    data: pd.DataFrame,
    n_neighbors: int = KNN_N_NEIGHBORS,
    weights: str = KNN_WEIGHTS,
) -> pd.DataFrame:
    """仅使用四个进水水质变量，先按列标准化，再执行 KNN 插补。"""
    arr = data.to_numpy(dtype=float)
    means = np.nanmean(arr, axis=0)
    stds = np.nanstd(arr, axis=0, ddof=0)

    if np.isnan(means).any():
        bad_cols = [data.columns[i] for i, flag in enumerate(np.isnan(means)) if flag]
        raise ValueError(f"以下变量全部缺失，无法执行 KNN 插补：{bad_cols}")

    stds[~np.isfinite(stds) | (stds < EPS)] = 1.0
    standardized = (arr - means) / stds

    # 邻居数不能超过至少含一个观测值的样本数。
    usable_rows = int(np.sum(np.isfinite(standardized).any(axis=1)))
    effective_neighbors = max(1, min(n_neighbors, usable_rows))

    imputer = KNNImputer(
        n_neighbors=effective_neighbors,
        weights=weights,
        metric="nan_euclidean",
    )
    imputed_z = imputer.fit_transform(standardized)

    # 某些 sklearn 版本对全空行或特殊模式可能仍产生 NaN，统一用标准化后的 0
    # 兜底，相当于使用对应变量均值。该情况会在整行遮蔽情景中体现 KNN 的局限。
    imputed_z = np.where(np.isfinite(imputed_z), imputed_z, 0.0)
    imputed = imputed_z * stds + means

    return pd.DataFrame(imputed, index=data.index, columns=data.columns)


def _interpolate_short_gaps_array(values: np.ndarray, max_gap: int) -> np.ndarray:
    arr = values.astype(float, copy=True)
    n = len(arr)
    i = 0
    while i < n:
        if np.isfinite(arr[i]):
            i += 1
            continue
        start = i
        while i < n and not np.isfinite(arr[i]):
            i += 1
        end = i - 1
        gap_length = end - start + 1

        has_left = start > 0 and np.isfinite(arr[start - 1])
        has_right = i < n and np.isfinite(arr[i])
        if gap_length <= max_gap and has_left and has_right:
            left = arr[start - 1]
            right = arr[i]
            step = (right - left) / (gap_length + 1)
            for k in range(gap_length):
                arr[start + k] = left + step * (k + 1)
    return arr


def offline_bidirectional_linear_impute(
    data: pd.DataFrame,
    max_gap: int = LINEAR_MAX_GAP,
    fallback: str = LINEAR_FALLBACK,
) -> pd.DataFrame:
    """
    离线双向线性插值，仅用于本脚本的人工遮蔽恢复实验和分布保持诊断。

    只对两侧均有有效观测且连续缺口长度不超过 max_gap 的区间进行线性插值；
    端点缺失、过长缺口或其他未填补位置，使用该列非缺失值的均值或中位数兜底。

    重要：该函数会使用缺口右侧观测值，因此不得直接用于正式的时间顺序验证、
    滚动/扩展窗口验证或任何需要模拟“仅使用预测时点之前信息”的验证流程。
    正式建模阶段应在对应训练/验证划分内部实现无未来信息的插补逻辑。
    """
    if fallback not in {"mean", "median"}:
        raise ValueError("LINEAR_FALLBACK 只能为 'mean' 或 'median'。")

    result = pd.DataFrame(index=data.index)
    for col in data.columns:
        values = data[col].to_numpy(dtype=float)
        filled = _interpolate_short_gaps_array(values, max_gap=max_gap)

        observed = values[np.isfinite(values)]
        if len(observed) == 0:
            raise ValueError(f"变量 {col} 全部缺失，无法执行线性插值。")
        fallback_value = (
            float(np.nanmedian(observed)) if fallback == "median" else float(np.nanmean(observed))
        )
        filled[~np.isfinite(filled)] = fallback_value
        result[col] = filled
    return result


def apply_imputation_method(method: str, data: pd.DataFrame) -> pd.DataFrame:
    if method == "均值填充":
        return mean_impute(data)
    if method == "标准化KNN":
        return standardized_knn_impute(data)
    if method == "线性插值":
        # 仅用于本脚本离线遮蔽实验/分布诊断；正式时间验证不可直接复用。
        return offline_bidirectional_linear_impute(data)
    raise ValueError(f"未知插补方法：{method}")


# =============================================================================
# 6. 可复现遮蔽实验
# =============================================================================


def create_rng(master_seed: int, target_idx: int, scenario_idx: int, repeat: int) -> Tuple[np.random.Generator, int]:
    seed_sequence = np.random.SeedSequence(
        [int(master_seed), int(target_idx), int(scenario_idx), int(repeat)]
    )
    child_seed = int(seed_sequence.generate_state(1, dtype=np.uint32)[0])
    return np.random.default_rng(child_seed), child_seed


def random_cell_mask(data: pd.DataFrame, rng: np.random.Generator, ratio: float) -> pd.DataFrame:
    mask = pd.DataFrame(False, index=data.index, columns=data.columns)
    for col in data.columns:
        candidates = np.flatnonzero(data[col].notna().to_numpy())
        if len(candidates) == 0:
            continue
        n_mask = max(1, int(round(len(candidates) * ratio)))
        n_mask = min(n_mask, len(candidates))
        selected = rng.choice(candidates, size=n_mask, replace=False)
        mask.iloc[selected, mask.columns.get_loc(col)] = True
    return mask


def _consecutive_true_runs(observed: np.ndarray) -> List[Tuple[int, int]]:
    """返回 observed=True 的连续区间，格式为 [start, end)，仅按原始行位置判断连续。"""
    runs: List[Tuple[int, int]] = []
    n = len(observed)
    i = 0
    while i < n:
        if not observed[i]:
            i += 1
            continue
        start = i
        while i < n and observed[i]:
            i += 1
        runs.append((start, i))
    return runs


def contiguous_block_mask(
    data: pd.DataFrame,
    rng: np.random.Generator,
    ratio: float,
    min_length: int,
    max_length: int,
) -> pd.DataFrame:
    """
    构造严格连续的人工块状缺失。

    仅从“原本连续且非缺失”的观测区间内抽取完整连续块，每个块长度均在
    min_length–max_length 之间；不同块不得重叠。不会再用随机散点补足目标遮蔽数。

    因为必须保证块严格连续，当剩余目标数量小于 min_length，或数据中已经没有
    合适的连续非缺失区间时，实际遮蔽比例允许略低于设定 ratio。
    """
    if min_length < 1:
        raise ValueError("BLOCK_MIN_LENGTH 必须至少为 1。")
    if max_length < min_length:
        raise ValueError("BLOCK_MAX_LENGTH 不能小于 BLOCK_MIN_LENGTH。")
    if not (0 < ratio < 1):
        raise ValueError("连续块遮蔽比例 ratio 必须位于 (0, 1) 内。")

    mask = pd.DataFrame(False, index=data.index, columns=data.columns)

    for col_idx, col in enumerate(data.columns):
        observed = data[col].notna().to_numpy(dtype=bool)
        observed_count = int(observed.sum())
        if observed_count < min_length:
            continue

        target_count = max(1, int(round(observed_count * ratio)))
        target_count = min(target_count, observed_count)

        current_count = 0
        while current_count < target_count:
            remaining = target_count - current_count
            if remaining < min_length:
                break

            available = observed & ~mask[col].to_numpy(dtype=bool)
            runs = _consecutive_true_runs(available)

            candidates: List[Tuple[int, int]] = []
            max_allowed_length = min(max_length, remaining)

            for run_start, run_end in runs:
                run_length = run_end - run_start
                if run_length < min_length:
                    continue

                longest = min(max_allowed_length, run_length)
                for length in range(min_length, longest + 1):
                    for start in range(run_start, run_end - length + 1):
                        candidates.append((start, length))

            if not candidates:
                break

            chosen_idx = int(rng.integers(0, len(candidates)))
            start, length = candidates[chosen_idx]
            mask.iloc[start:start + length, col_idx] = True
            current_count += length

    return mask


def whole_row_mask(data: pd.DataFrame, rng: np.random.Generator, ratio: float) -> pd.DataFrame:
    mask = pd.DataFrame(False, index=data.index, columns=data.columns)
    complete_rows = np.flatnonzero(data.notna().all(axis=1).to_numpy())
    if len(complete_rows) == 0:
        return mask

    n_rows_to_mask = max(1, int(round(len(complete_rows) * ratio)))
    n_rows_to_mask = min(n_rows_to_mask, len(complete_rows))
    selected = rng.choice(complete_rows, size=n_rows_to_mask, replace=False)
    mask.iloc[selected, :] = True
    return mask


def create_mask(
    scenario: str,
    data: pd.DataFrame,
    rng: np.random.Generator,
) -> pd.DataFrame:
    if scenario == "随机遮蔽":
        return random_cell_mask(data, rng, RANDOM_MASK_RATIO)
    if scenario == "连续块遮蔽":
        return contiguous_block_mask(
            data,
            rng,
            ratio=BLOCK_MASK_RATIO,
            min_length=BLOCK_MIN_LENGTH,
            max_length=BLOCK_MAX_LENGTH,
        )
    if scenario == "整行遮蔽":
        return whole_row_mask(data, rng, ROW_MASK_RATIO)
    raise ValueError(f"未知遮蔽情形：{scenario}")


def normalization_scale(series: pd.Series) -> Tuple[float, str]:
    values = series.dropna().to_numpy(dtype=float)
    if len(values) == 0:
        return 1.0, "无有效值，使用1"

    if NORMALIZATION_METHOD == "P95-P5":
        scale = float(np.nanpercentile(values, 95) - np.nanpercentile(values, 5))
        desc = "P95-P5"
    elif NORMALIZATION_METHOD == "标准差":
        scale = float(np.nanstd(values, ddof=1))
        desc = "标准差"
    elif NORMALIZATION_METHOD == "极差":
        scale = float(np.nanmax(values) - np.nanmin(values))
        desc = "最大值-最小值"
    else:
        raise ValueError(f"未知归一化方式：{NORMALIZATION_METHOD}")

    if not np.isfinite(scale) or scale < EPS:
        fallback = float(np.nanstd(values, ddof=1)) if len(values) > 1 else 0.0
        if np.isfinite(fallback) and fallback >= EPS:
            return fallback, f"{desc}过小，改用标准差"
        return 1.0, f"{desc}和标准差均过小，使用1"
    return scale, desc


def compute_metrics(
    true_values: np.ndarray,
    predicted_values: np.ndarray,
    scale: float,
) -> Dict[str, float]:
    valid = np.isfinite(true_values) & np.isfinite(predicted_values)
    if not valid.any():
        return {"RMSE": np.nan, "MAE": np.nan, "NRMSE": np.nan, "NMAE": np.nan}

    errors = predicted_values[valid] - true_values[valid]
    rmse = float(np.sqrt(np.mean(errors ** 2)))
    mae = float(np.mean(np.abs(errors)))
    return {
        "RMSE": rmse,
        "MAE": mae,
        "NRMSE": rmse / scale,
        "NMAE": mae / scale,
    }


def run_masking_experiments(
    dynamic_tables: Mapping[str, pd.DataFrame],
    repeats: int,
    master_seed: int,
) -> Tuple[pd.DataFrame, pd.DataFrame]:
    """返回逐变量重复实验结果和遮蔽规模记录。"""
    metric_rows: List[Dict[str, object]] = []
    mask_rows: List[Dict[str, object]] = []

    target_keys = list(TARGETS.keys())
    for target_idx, target_key in enumerate(target_keys):
        table = dynamic_tables[target_key]
        water = table[IMPUTE_COLUMNS].astype(float).copy()
        scales = {col: normalization_scale(water[col]) for col in IMPUTE_COLUMNS}

        print(f"\n[{target_idx + 1}/{len(target_keys)}] 正在评价 {TARGETS[target_key]['display']} 动态表……")

        for scenario_idx, scenario in enumerate(MASK_SCENARIOS):
            print(f"  - {scenario}：{repeats} 次重复")
            for repeat in range(1, repeats + 1):
                rng, child_seed = create_rng(master_seed, target_idx, scenario_idx, repeat)
                mask = create_mask(scenario, water, rng)
                total_masked = int(mask.to_numpy().sum())

                if total_masked == 0:
                    mask_rows.append(
                        {
                            "目标指标": TARGETS[target_key]["display"],
                            "遮蔽情形": scenario,
                            "重复序号": repeat,
                            "子随机种子": child_seed,
                            "遮蔽单元格数": 0,
                            "遮蔽行数": 0,
                            "状态": "未找到可遮蔽位置",
                        }
                    )
                    continue

                masked_data = water.mask(mask)
                mask_rows.append(
                    {
                        "目标指标": TARGETS[target_key]["display"],
                        "遮蔽情形": scenario,
                        "重复序号": repeat,
                        "子随机种子": child_seed,
                        "遮蔽单元格数": total_masked,
                        "遮蔽行数": int(mask.any(axis=1).sum()),
                        "状态": "成功",
                    }
                )

                # 三种方法使用完全相同的遮蔽位置。
                for method in METHODS:
                    try:
                        imputed = apply_imputation_method(method, masked_data)
                        method_status = "成功"
                        error_message = ""
                    except Exception as exc:  # 保留失败记录，避免静默丢失。
                        imputed = pd.DataFrame(np.nan, index=water.index, columns=water.columns)
                        method_status = "失败"
                        error_message = str(exc)

                    for variable in IMPUTE_COLUMNS:
                        variable_mask = mask[variable].to_numpy(dtype=bool)
                        true_values = water.loc[variable_mask, variable].to_numpy(dtype=float)
                        predicted_values = imputed.loc[variable_mask, variable].to_numpy(dtype=float)
                        scale, scale_desc = scales[variable]
                        metrics = compute_metrics(true_values, predicted_values, scale)

                        metric_rows.append(
                            {
                                "目标指标": TARGETS[target_key]["display"],
                                "遮蔽情形": scenario,
                                "重复序号": repeat,
                                "子随机种子": child_seed,
                                "插补方法": method,
                                "变量": variable,
                                "该变量遮蔽数": int(variable_mask.sum()),
                                "归一化尺度": scale,
                                "尺度定义": scale_desc,
                                **metrics,
                                "执行状态": method_status,
                                "错误信息": error_message,
                            }
                        )

    return pd.DataFrame(metric_rows), pd.DataFrame(mask_rows)


# =============================================================================
# 7. 结果汇总与推荐
# =============================================================================


def summarize_variable_metrics(repeat_results: pd.DataFrame) -> pd.DataFrame:
    valid = repeat_results.loc[repeat_results["执行状态"] == "成功"].copy()
    group_cols = ["目标指标", "遮蔽情形", "插补方法", "变量"]
    grouped = valid.groupby(group_cols, dropna=False)

    summary = grouped.agg(
        成功重复数=("重复序号", "nunique"),
        平均遮蔽数=("该变量遮蔽数", "mean"),
        RMSE均值=("RMSE", "mean"),
        RMSE标准差=("RMSE", "std"),
        MAE均值=("MAE", "mean"),
        MAE标准差=("MAE", "std"),
        NRMSE均值=("NRMSE", "mean"),
        NRMSE标准差=("NRMSE", "std"),
        NMAE均值=("NMAE", "mean"),
        NMAE标准差=("NMAE", "std"),
    ).reset_index()

    quantiles = grouped[["NRMSE", "NMAE"]].quantile([0.025, 0.5, 0.975]).unstack()
    quantiles.columns = [f"{metric}_{int(q * 1000) / 10:g}分位" for metric, q in quantiles.columns]
    quantiles = quantiles.reset_index()
    return summary.merge(quantiles, on=group_cols, how="left")


def build_target_scenario_summary(repeat_results: pd.DataFrame) -> pd.DataFrame:
    """
    先在每次重复内对四个变量等权平均，再按目标×遮蔽情形×方法汇总。
    这样不会因某个变量遮蔽数量略多而获得更高权重。
    """
    valid = repeat_results.loc[repeat_results["执行状态"] == "成功"].copy()
    per_repeat = (
        valid.groupby(
            ["目标指标", "遮蔽情形", "插补方法", "重复序号", "子随机种子"],
            dropna=False,
        )[["NRMSE", "NMAE"]]
        .mean()
        .reset_index()
    )
    per_repeat["归一综合误差"] = (per_repeat["NRMSE"] + per_repeat["NMAE"]) / 2

    grouped = per_repeat.groupby(["目标指标", "遮蔽情形", "插补方法"], dropna=False)
    summary = grouped.agg(
        成功重复数=("重复序号", "nunique"),
        NRMSE均值=("NRMSE", "mean"),
        NRMSE标准差=("NRMSE", "std"),
        NMAE均值=("NMAE", "mean"),
        NMAE标准差=("NMAE", "std"),
        归一综合误差均值=("归一综合误差", "mean"),
        归一综合误差标准差=("归一综合误差", "std"),
        归一综合误差中位数=("归一综合误差", "median"),
    ).reset_index()

    ci = grouped["归一综合误差"].quantile([0.025, 0.975]).unstack().reset_index()
    ci = ci.rename(columns={0.025: "归一综合误差2.5分位", 0.975: "归一综合误差97.5分位"})
    summary = summary.merge(ci, on=["目标指标", "遮蔽情形", "插补方法"], how="left")

    summary["情形内排名"] = summary.groupby(["目标指标", "遮蔽情形"])[
        "归一综合误差均值"
    ].rank(method="min", ascending=True)
    return summary.sort_values(
        ["目标指标", "遮蔽情形", "情形内排名", "插补方法"]
    ).reset_index(drop=True)


def distribution_preservation_check(dynamic_tables: Mapping[str, pd.DataFrame]) -> pd.DataFrame:
    rows: List[Dict[str, object]] = []

    for target_key, table in dynamic_tables.items():
        water = table[IMPUTE_COLUMNS].astype(float).copy()
        for method in METHODS:
            imputed = apply_imputation_method(method, water)
            for variable in IMPUTE_COLUMNS:
                original = water[variable].dropna().to_numpy(dtype=float)
                completed = imputed[variable].to_numpy(dtype=float)
                original_missing = int(water[variable].isna().sum())

                def stats(values: np.ndarray) -> Dict[str, float]:
                    return {
                        "均值": float(np.mean(values)),
                        "标准差": float(np.std(values, ddof=1)) if len(values) > 1 else 0.0,
                        "P5": float(np.percentile(values, 5)),
                        "Q1": float(np.percentile(values, 25)),
                        "中位数": float(np.percentile(values, 50)),
                        "Q3": float(np.percentile(values, 75)),
                        "P95": float(np.percentile(values, 95)),
                    }

                before = stats(original)
                after = stats(completed)
                iqr_before = before["Q3"] - before["Q1"]
                iqr_after = after["Q3"] - after["Q1"]

                mean_shift = abs(after["均值"] - before["均值"]) / max(abs(before["均值"]), EPS) * 100
                std_shift = abs(after["标准差"] - before["标准差"]) / max(abs(before["标准差"]), EPS) * 100
                iqr_shift = abs(iqr_after - iqr_before) / max(abs(iqr_before), EPS) * 100
                distribution_score = float(np.mean([mean_shift, std_shift, iqr_shift]))

                rows.append(
                    {
                        "目标指标": TARGETS[target_key]["display"],
                        "插补方法": method,
                        "变量": variable,
                        "动态样本数": len(water),
                        "实际缺失数": original_missing,
                        "实际缺失比例(%)": original_missing / len(water) * 100 if len(water) else np.nan,
                        "原始均值": before["均值"],
                        "插补后均值": after["均值"],
                        "均值相对变化(%)": mean_shift,
                        "原始标准差": before["标准差"],
                        "插补后标准差": after["标准差"],
                        "标准差相对变化(%)": std_shift,
                        "原始P5": before["P5"],
                        "插补后P5": after["P5"],
                        "原始Q1": before["Q1"],
                        "插补后Q1": after["Q1"],
                        "原始中位数": before["中位数"],
                        "插补后中位数": after["中位数"],
                        "原始Q3": before["Q3"],
                        "插补后Q3": after["Q3"],
                        "原始P95": before["P95"],
                        "插补后P95": after["P95"],
                        "IQR相对变化(%)": iqr_shift,
                        "分布偏移指标": distribution_score,
                    }
                )
    return pd.DataFrame(rows)


def build_overall_ranking(
    target_scenario_summary: pd.DataFrame,
    distribution_check: pd.DataFrame,
) -> pd.DataFrame:
    # 各目标、各遮蔽情形等权平均。
    error_summary = (
        target_scenario_summary.groupby("插补方法", as_index=False)
        .agg(
            总体NRMSE=("NRMSE均值", "mean"),
            总体NMAE=("NMAE均值", "mean"),
            总体归一综合误差=("归一综合误差均值", "mean"),
            跨情形标准差=("归一综合误差均值", "std"),
        )
    )

    structured = target_scenario_summary.loc[
        target_scenario_summary["遮蔽情形"].isin(["连续块遮蔽", "整行遮蔽"])
    ]
    structured_summary = (
        structured.groupby("插补方法", as_index=False)["归一综合误差均值"]
        .mean()
        .rename(columns={"归一综合误差均值": "结构缺失综合误差"})
    )

    distribution_summary = (
        distribution_check.groupby("插补方法", as_index=False)
        .agg(
            平均分布偏移=("分布偏移指标", "mean"),
            最大分布偏移=("分布偏移指标", "max"),
        )
    )

    result = error_summary.merge(structured_summary, on="插补方法", how="left")
    result = result.merge(distribution_summary, on="插补方法", how="left")

    result["总体误差排名"] = result["总体归一综合误差"].rank(method="min", ascending=True)
    result["结构缺失排名"] = result["结构缺失综合误差"].rank(method="min", ascending=True)
    result["分布保持排名"] = result["平均分布偏移"].rank(method="min", ascending=True)

    # 不设置人为加权总分，采用透明的逐级排序：
    # 总体归一误差 → 结构缺失误差 → 分布偏移。
    result = result.sort_values(
        ["总体归一综合误差", "结构缺失综合误差", "平均分布偏移", "插补方法"],
        ascending=[True, True, True, True],
    ).reset_index(drop=True)
    result["综合推荐顺序"] = np.arange(1, len(result) + 1)
    result["是否推荐"] = np.where(result["综合推荐顺序"] == 1, "推荐", "备选")

    ordered_cols = [
        "综合推荐顺序",
        "插补方法",
        "是否推荐",
        "总体NRMSE",
        "总体NMAE",
        "总体归一综合误差",
        "跨情形标准差",
        "结构缺失综合误差",
        "平均分布偏移",
        "最大分布偏移",
        "总体误差排名",
        "结构缺失排名",
        "分布保持排名",
    ]
    return result[ordered_cols]


def build_parameter_table(
    input_file: Path,
    input_sheet: str,
    output_dir: Path,
    repeats: int,
    master_seed: int,
    order_desc: str,
) -> pd.DataFrame:
    rows = [
        ("输入文件", str(input_file)),
        ("输入工作表", input_sheet),
        ("输出文件夹", str(output_dir)),
        ("连续运行排序规则", order_desc),
        ("总随机种子", master_seed),
        ("每种遮蔽情形重复次数", repeats),
        ("随机遮蔽比例", RANDOM_MASK_RATIO),
        ("连续块遮蔽目标比例", BLOCK_MASK_RATIO),
        ("整行遮蔽完整行比例", ROW_MASK_RATIO),
        ("连续块最短长度", BLOCK_MIN_LENGTH),
        ("连续块最长长度", BLOCK_MAX_LENGTH),
        ("连续块生成规则", "仅从原本连续非缺失区间抽取3–7个连续样本块；不使用随机散点补足"),
        ("KNN邻居数", KNN_N_NEIGHBORS),
        ("KNN权重", KNN_WEIGHTS),
        ("线性插值最大连续缺口", LINEAR_MAX_GAP),
        ("线性插值兜底方式", LINEAR_FALLBACK),
        ("线性插值用途限制", "仅用于离线人工遮蔽恢复和分布诊断；正式时间验证不得使用未来观测双向插值"),
        ("NRMSE/NMAE归一化尺度", NORMALIZATION_METHOD),
        ("参与插补比较的变量", "、".join(IMPUTE_COLUMNS)),
        ("插补方法", "、".join(METHODS)),
        ("遮蔽情形", "、".join(MASK_SCENARIOS)),
        ("正式建模限制", "本结果仅用于选择策略；正式验证须在训练/验证划分内部重新拟合插补规则"),
    ]
    return pd.DataFrame(rows, columns=["参数", "设置"])


# =============================================================================
# 8. Excel 与 TXT 输出
# =============================================================================


def write_excel(
    output_path: Path,
    dynamic_tables: Mapping[str, pd.DataFrame],
    build_summary: pd.DataFrame,
    missing_summary: pd.DataFrame,
    overall_ranking: pd.DataFrame,
    target_scenario_summary: pd.DataFrame,
    variable_summary: pd.DataFrame,
    repeat_results: pd.DataFrame,
    mask_summary: pd.DataFrame,
    distribution_check: pd.DataFrame,
    parameter_table: pd.DataFrame,
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(output_path, engine="openpyxl", datetime_format="yyyy-mm-dd") as writer:
        build_summary.to_excel(writer, sheet_name="样本构建汇总", index=False)
        for target_key, table in dynamic_tables.items():
            table.to_excel(writer, sheet_name=TARGETS[target_key]["sheet"], index=False)
        missing_summary.to_excel(writer, sheet_name="动态表缺失统计", index=False)
        overall_ranking.to_excel(writer, sheet_name="插补方法总评", index=False)
        target_scenario_summary.to_excel(writer, sheet_name="目标情形汇总", index=False)
        variable_summary.to_excel(writer, sheet_name="变量误差汇总", index=False)
        distribution_check.to_excel(writer, sheet_name="分布保持检查", index=False)
        mask_summary.to_excel(writer, sheet_name="遮蔽规模记录", index=False)
        repeat_results.to_excel(writer, sheet_name="重复实验结果", index=False)
        parameter_table.to_excel(writer, sheet_name="实验参数", index=False)

    style_excel(output_path)


def style_excel(output_path: Path) -> None:
    wb = load_workbook(output_path)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    sub_fill = PatternFill("solid", fgColor="D9EAF7")
    recommended_fill = PatternFill("solid", fgColor="E2F0D9")
    thin_gray = Side(style="thin", color="D9E1F2")
    border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.sheet_view.showGridLines = False

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border

        ws.row_dimensions[1].height = 30

        # 数据区样式和日期/数字格式。
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(vertical="center", wrap_text=False)
                if isinstance(cell.value, datetime):
                    cell.number_format = "yyyy-mm-dd"
                elif isinstance(cell.value, float):
                    cell.number_format = "0.000000"

        # 合理列宽，避免极端自动扩展。
        for col_idx in range(1, ws.max_column + 1):
            letter = get_column_letter(col_idx)
            max_len = 0
            for cell in ws[letter]:
                value = "" if cell.value is None else str(cell.value)
                max_len = max(max_len, len(value))
            width = min(max(max_len + 2, 10), 36)
            ws.column_dimensions[letter].width = width

        # 动态表中 Date 格式和固定数值精度。
        header_map = {cell.value: cell.column for cell in ws[1]}
        if "Date" in header_map:
            col_letter = get_column_letter(header_map["Date"])
            for cell in ws[col_letter][1:]:
                cell.number_format = "yyyy-mm-dd"
        if "样本序号" in header_map:
            col_letter = get_column_letter(header_map["样本序号"])
            for cell in ws[col_letter][1:]:
                cell.number_format = "0"

        # 总评表突出推荐方法，并添加误差色阶。
        if ws.title == "插补方法总评":
            headers = {cell.value: cell.column for cell in ws[1]}
            if "是否推荐" in headers:
                flag_col = headers["是否推荐"]
                for r in range(2, ws.max_row + 1):
                    if ws.cell(r, flag_col).value == "推荐":
                        for c in range(1, ws.max_column + 1):
                            ws.cell(r, c).fill = recommended_fill
                            ws.cell(r, c).font = Font(bold=True)
            for metric in ["总体归一综合误差", "结构缺失综合误差", "平均分布偏移"]:
                if metric in headers and ws.max_row >= 3:
                    col_letter = get_column_letter(headers[metric])
                    ws.conditional_formatting.add(
                        f"{col_letter}2:{col_letter}{ws.max_row}",
                        ColorScaleRule(
                            start_type="min", start_color="63BE7B",
                            mid_type="percentile", mid_value=50, mid_color="FFEB84",
                            end_type="max", end_color="F8696B",
                        ),
                    )

        if ws.title in {"样本构建汇总", "动态表缺失统计", "实验参数"}:
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    if cell.column == 1:
                        cell.fill = sub_fill

    wb.save(output_path)


def build_txt_report(
    output_path: Path,
    elapsed_seconds: float,
    input_file: Path,
    input_sheet: str,
    excel_output: Path,
    sorted_df: pd.DataFrame,
    order_desc: str,
    warnings: Sequence[str],
    build_infos: Sequence[DynamicBuildInfo],
    missing_summary: pd.DataFrame,
    overall_ranking: pd.DataFrame,
    target_scenario_summary: pd.DataFrame,
    distribution_check: pd.DataFrame,
    repeats: int,
    master_seed: int,
) -> None:
    lines: List[str] = []
    add = lines.append

    add("A²/O 动态可用数据表与缺失值插补比较结果说明")
    add("=" * 76)
    add("")
    add("一、任务与文件信息")
    add(f"1. 执行时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    add(f"2. 运行耗时：{elapsed_seconds:.2f} 秒")
    add(f"3. 输入文件：{input_file}")
    add(f"4. 输入工作表：{input_sheet}")
    add(f"5. 输出 Excel：{excel_output}")
    add(f"6. 输出说明：{output_path}")
    add(f"7. 原始数据行数：{len(sorted_df)}")
    add(f"8. 连续运行排序规则：{order_desc}")
    add("")

    add("二、四个单目标动态可用数据表构建规则")
    add("1. 从异常值处理后的“清洗后数据”出发，不重新执行异常值筛查。")
    add("2. 每个目标独立构建基础表，只剔除该目标自身为空的样本。")
    add("3. 历史状态按连续运行顺序构造，使用当前样本之前最近的有效目标观测。")
    add("4. lag1 为最近 1 个有效目标值，lag2 为最近第 2 个有效目标值。")
    add("5. roll3 为当前样本之前最近 3 个有效目标值的平均，不包含当前目标值。")
    add("6. 仅删除无法形成完整 lag1、lag2 和 roll3 的样本，不机械删除前三个日期。")
    add("7. 四个进水水质变量中的真实缺失继续保留，不在动态表中提前补全。")
    add("8. 四个动态表将作为后续 F0、F1、F2 和 F3 的同样本基础。")
    add("")

    add("三、动态可用样本构建结果")
    for idx, info in enumerate(build_infos, start=1):
        add(
            f"{idx}. {TARGETS[info.target_key]['display']}：原始 {info.original_rows} 行；"
            f"目标有效 {info.target_valid_rows} 行；因目标缺失剔除 {info.target_missing_rows_removed} 行；"
            f"因历史状态不完整剔除 {info.history_incomplete_rows_removed} 行；"
            f"最终动态可用 {info.dynamic_rows} 行；样本范围 {safe_value(info.first_sample)}–{safe_value(info.last_sample)}；"
            f"日期范围 {safe_value(info.first_date)}–{safe_value(info.last_date)}。"
        )
    add("")

    add("四、动态表中四个进水变量的真实缺失情况")
    for target in missing_summary["目标指标"].drop_duplicates():
        add(f"【{target}】")
        sub = missing_summary.loc[missing_summary["目标指标"] == target]
        for _, row in sub.iterrows():
            add(
                f"- {row['变量']}：缺失 {int(row['缺失数'])}/{int(row['动态样本数'])}，"
                f"缺失比例 {row['缺失比例(%)']:.2f}%。"
            )
    add("")

    add("五、插补方法比较设计")
    add("1. 参与比较的变量：Inf.TP、Inf.COD、Inf.NH4+-N、Inf.TN。")
    add("2. 不对 Aerobic_DO、Anaerobic_Feed_Fraction、HRT 或容积分数进行数据驱动插补。")
    add("3. 不对 eff.TP、eff.COD、eff.NH4+-N 或 eff.TN 目标值进行插补。")
    add(f"4. 总随机种子：{master_seed}；每种目标×遮蔽情形重复 {repeats} 次。")
    add(f"5. 随机遮蔽：每个进水变量随机遮蔽约 {RANDOM_MASK_RATIO * 100:.0f}% 的已有观测值。")
    add(
        f"6. 连续块遮蔽：目标为每个变量约 {BLOCK_MASK_RATIO * 100:.0f}% 的已有观测值，"
        f"仅从原本连续非缺失区间抽取长度 {BLOCK_MIN_LENGTH}–{BLOCK_MAX_LENGTH} 的完整连续块；"
        "不同块不重叠，也不使用随机散点补足。为保证严格连续，实际比例允许略低于目标比例。"
    )
    add(f"7. 整行遮蔽：从四个进水变量均有观测的样本中遮蔽约 {ROW_MASK_RATIO * 100:.0f}% 的完整行。")
    add("8. 三种方法在每次重复中使用完全相同的遮蔽位置。")
    add("9. 均值填充：使用当前目标动态表中、遮蔽后仍可见的该变量均值。")
    add(
        f"10. 标准化 KNN：四个进水变量先标准化，K={KNN_N_NEIGHBORS}，"
        f"权重为 {KNN_WEIGHTS}，再转换回原量纲。"
    )
    add(
        f"11. 线性插值：本阶段仅作为离线人工遮蔽恢复方法；两侧均有观测且连续缺口不超过 "
        f"{LINEAR_MAX_GAP} 个样本时进行双向线性插值，其余位置使用 {LINEAR_FALLBACK} 兜底。"
    )
    add("12. 上述双向线性插值不得直接用于正式时间顺序或滚动/扩展窗口验证；正式验证必须使用预测时点之前可获得的信息。")
    add(
        f"13. NRMSE 和 NMAE 的归一化尺度采用 {NORMALIZATION_METHOD}；"
        "先按变量计算，再在每次重复内对四个变量等权平均。"
    )
    add("")

    add("六、插补方法总体比较")
    for _, row in overall_ranking.iterrows():
        add(
            f"{int(row['综合推荐顺序'])}. {row['插补方法']}（{row['是否推荐']}）："
            f"总体 NRMSE={format_float(row['总体NRMSE'])}，"
            f"总体 NMAE={format_float(row['总体NMAE'])}，"
            f"总体归一综合误差={format_float(row['总体归一综合误差'])}，"
            f"结构缺失综合误差={format_float(row['结构缺失综合误差'])}，"
            f"平均分布偏移={format_float(row['平均分布偏移'])}%。"
        )
    recommended = overall_ranking.iloc[0]["插补方法"] if len(overall_ranking) else "无法确定"
    add(f"综合推荐方法：{recommended}。推荐顺序按总体归一误差、结构缺失误差和分布偏移逐级确定，不使用人为加权总分。")
    add("")

    add("七、各目标和缺失情形下的最佳方法")
    best_rows = target_scenario_summary.loc[target_scenario_summary["情形内排名"] == 1]
    for _, row in best_rows.iterrows():
        add(
            f"- {row['目标指标']} / {row['遮蔽情形']}：{row['插补方法']}，"
            f"归一综合误差均值={format_float(row['归一综合误差均值'])}，"
            f"标准差={format_float(row['归一综合误差标准差'])}，"
            f"95%经验区间=[{format_float(row['归一综合误差2.5分位'])}, "
            f"{format_float(row['归一综合误差97.5分位'])}]。"
        )
    add("")

    add("八、实际缺失值补全后的分布保持检查")
    dist_summary = (
        distribution_check.groupby("插补方法", as_index=False)
        .agg(平均分布偏移=("分布偏移指标", "mean"), 最大分布偏移=("分布偏移指标", "max"))
        .sort_values("平均分布偏移")
    )
    for _, row in dist_summary.iterrows():
        add(
            f"- {row['插补方法']}：平均分布偏移 {row['平均分布偏移']:.3f}%，"
            f"最大分布偏移 {row['最大分布偏移']:.3f}%。"
        )
    add("分布偏移指标为均值、标准差和 IQR 相对变化绝对值的平均，仅用于辅助判断，不替代遮蔽误差。")
    add("")

    add("九、Excel 工作表说明")
    add("1. 样本构建汇总：四个目标从原始数据到动态可用数据的样本数量变化。")
    add("2. TP动态可用表、COD动态可用表、氨氮动态可用表、TN动态可用表：保留真实输入缺失的单目标建模基础表。")
    add("3. 动态表缺失统计：四个进水变量在各目标动态表中的缺失数量和比例。")
    add("4. 插补方法总评：三种方法的总体误差、结构缺失误差、分布偏移和推荐顺序。")
    add("5. 目标情形汇总：各目标、各遮蔽情形和各方法的重复实验汇总及情形内排名。")
    add("6. 变量误差汇总：逐目标、逐变量、逐遮蔽情形的 RMSE、MAE、NRMSE 和 NMAE。")
    add("7. 分布保持检查：三种方法处理真实缺失后，各变量描述统计的变化。")
    add("8. 遮蔽规模记录：每次重复的子随机种子、遮蔽单元格数和遮蔽行数。")
    add("9. 重复实验结果：每一次重复、每个变量和每种方法的完整误差记录。")
    add("10. 实验参数：路径、随机种子、重复次数、遮蔽比例和插补参数。")
    add("")

    add("十、重要解释边界与后续使用")
    add("1. 动态可用表中的真实进水缺失值仍保持为空，这是有意设计，不是脚本遗漏。")
    add("2. 本阶段对实际缺失值的补全仅用于分布保持诊断，没有输出为正式建模数据。")
    add("3. 正式 Bootstrap .632+、时间顺序验证和滚动/扩展窗口验证中，插补器必须仅在训练集或 in-bag 样本内拟合。")
    add("4. OOB、验证集和未来样本不能参与训练阶段均值、KNN 参数、标准化参数或插值兜底值的确定。")
    add("5. 本脚本中的离线双向线性插值函数仅用于人工遮蔽恢复和分布诊断；时间顺序和滚动/扩展窗口验证中的验证样本不能使用未来观测进行双向线性插值。")
    add("6. 本次推荐结果应与不同目标、不同缺失情形和分布保持结果联合解释，不能只依据某一个变量或某一次重复。")
    add("7. 下一步可根据推荐方法，在验证内部构造 F0、F1、F2 和 F3 四组特征表并开展模型训练。")

    if warnings:
        add("")
        add("十一、运行警告")
        for idx, warning in enumerate(warnings, start=1):
            add(f"{idx}. {warning}")

    output_path.write_text("\n".join(lines), encoding="utf-8-sig")


# =============================================================================
# 9. 主程序
# =============================================================================


def main() -> int:
    args = parse_args()
    start_time = time.perf_counter()

    input_file: Path = args.input.expanduser().resolve()
    output_dir: Path = args.output.expanduser().resolve()
    repeats: int = int(args.repeats)
    master_seed: int = int(args.seed)

    if repeats < 2:
        raise ValueError("重复次数至少应为 2，以便计算重复实验标准差。")
    if not input_file.exists():
        raise FileNotFoundError(f"输入文件不存在：{input_file}")

    output_dir.mkdir(parents=True, exist_ok=True)
    excel_output = output_dir / OUTPUT_EXCEL_NAME
    txt_output = output_dir / OUTPUT_TXT_NAME

    input_sheet = choose_input_sheet(input_file, args.sheet)
    raw = pd.read_excel(input_file, sheet_name=input_sheet)
    raw = clean_column_names(raw)

    required_columns = ID_COLUMNS + RAW_INPUT_COLUMNS + [v["target"] for v in TARGETS.values()]
    require_columns(raw, required_columns, "输入数据")

    typed = convert_types(raw)
    sorted_df, order_desc, warnings = sort_continuous_data(typed)

    if sorted_df["Date"].isna().any():
        warnings.append(f"Date 有 {int(sorted_df['Date'].isna().sum())} 个无法解析值。")

    dynamic_tables: Dict[str, pd.DataFrame] = {}
    build_infos: List[DynamicBuildInfo] = []
    build_rows: List[Dict[str, object]] = []

    for target_key, meta in TARGETS.items():
        dynamic, info = build_dynamic_table(sorted_df, target_key, meta["target"])
        if len(dynamic) == 0:
            raise ValueError(f"{meta['display']} 无法形成动态可用样本。")
        dynamic_tables[target_key] = dynamic
        build_infos.append(info)
        build_rows.append(
            {
                "目标指标": meta["display"],
                "目标字段": meta["target"],
                "原始数据行数": info.original_rows,
                "目标有效行数": info.target_valid_rows,
                "因目标缺失剔除行数": info.target_missing_rows_removed,
                "因历史状态不完整剔除行数": info.history_incomplete_rows_removed,
                "动态可用行数": info.dynamic_rows,
                "首个动态样本序号": info.first_sample,
                "末个动态样本序号": info.last_sample,
                "首个动态日期": info.first_date,
                "末个动态日期": info.last_date,
                "历史特征": f"{meta['target']}_lag1、{meta['target']}_lag2、{meta['target']}_roll3",
            }
        )

    build_summary = pd.DataFrame(build_rows)
    missing_summary = build_missing_summary(dynamic_tables)

    print("已完成四个单目标动态可用数据表构建。")
    print(build_summary[["目标指标", "目标有效行数", "动态可用行数"]].to_string(index=False))

    repeat_results, mask_summary = run_masking_experiments(
        dynamic_tables=dynamic_tables,
        repeats=repeats,
        master_seed=master_seed,
    )

    if repeat_results.empty:
        raise RuntimeError("遮蔽实验未生成有效结果。")

    variable_summary = summarize_variable_metrics(repeat_results)
    target_scenario_summary = build_target_scenario_summary(repeat_results)
    distribution_check = distribution_preservation_check(dynamic_tables)
    overall_ranking = build_overall_ranking(target_scenario_summary, distribution_check)
    parameter_table = build_parameter_table(
        input_file=input_file,
        input_sheet=input_sheet,
        output_dir=output_dir,
        repeats=repeats,
        master_seed=master_seed,
        order_desc=order_desc,
    )

    write_excel(
        output_path=excel_output,
        dynamic_tables=dynamic_tables,
        build_summary=build_summary,
        missing_summary=missing_summary,
        overall_ranking=overall_ranking,
        target_scenario_summary=target_scenario_summary,
        variable_summary=variable_summary,
        repeat_results=repeat_results,
        mask_summary=mask_summary,
        distribution_check=distribution_check,
        parameter_table=parameter_table,
    )

    elapsed = time.perf_counter() - start_time
    build_txt_report(
        output_path=txt_output,
        elapsed_seconds=elapsed,
        input_file=input_file,
        input_sheet=input_sheet,
        excel_output=excel_output,
        sorted_df=sorted_df,
        order_desc=order_desc,
        warnings=warnings,
        build_infos=build_infos,
        missing_summary=missing_summary,
        overall_ranking=overall_ranking,
        target_scenario_summary=target_scenario_summary,
        distribution_check=distribution_check,
        repeats=repeats,
        master_seed=master_seed,
    )

    print("\n处理完成。")
    print(f"Excel 输出：{excel_output}")
    print(f"文字说明：{txt_output}")
    print("\n插补方法综合推荐顺序：")
    print(
        overall_ranking[
            ["综合推荐顺序", "插补方法", "总体归一综合误差", "结构缺失综合误差", "平均分布偏移"]
        ].to_string(index=False)
    )
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception as exc:
        print("\n脚本运行失败：", exc, file=sys.stderr)
        traceback.print_exc()
        sys.exit(1)